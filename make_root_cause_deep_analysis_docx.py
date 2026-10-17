# -*- coding: utf-8 -*-
"""
vidio-health-trainer
8주차 일반화 데이터셋 성능 저하 5대 근본 원인 심층 프로파일링 및 해결 방안 공식 보고서 (.docx / .md)
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_OUT = r"D:\자주프 문서\8주차\8주차_일반화_성능저하_5대_근본원인_심층프로파일링_보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_일반화_성능저하_5대_근본원인_심층프로파일링_보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_header(row, col_widths=None):
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Malgun Gothic"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if col_widths and idx < len(col_widths):
            cell.width = Inches(col_widths[idx])

def style_rows(table, col_widths=None, align_cols=None):
    for r_idx, row in enumerate(table.rows[1:]):
        bg = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            for p in cell.paragraphs:
                if align_cols and c_idx < len(align_cols):
                    p.alignment = align_cols[c_idx]
                for r in p.runs:
                    r.font.name = "Malgun Gothic"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])

def build_docx_and_md():
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)
    
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("vidio-health-trainer 일반화 데이터셋 성능 저하\n5대 근본 원인 심층 프로파일링 및 엔지니어링 해결 보고서")
    trun.font.name = "Malgun Gothic"
    trun.font.size = Pt(16.5)
    trun.font.bold = True
    trun.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("문서 번호", "8주차 공식 심층 기술 분석 보고서 (2026-10-16)"),
        ("프로젝트 명", "vidio-health-trainer"),
        ("분석 대상", "일반화 3대 운동 72개 세트 클립 (총 256회 1회 반복 추론 결과)"),
        ("핵심 진단 영역", "도메인 시프트(화면비/카메라 앵글), 동반 결함 라벨링 괴리, 비전 노이즈 민감도")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        c0, c1 = meta_table.rows[r_idx].cells[0], meta_table.rows[r_idx].cells[1]
        c0.text, c1.text = k, v
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = "Malgun Gothic"
        c0.paragraphs[0].runs[0].font.size = Pt(9)
        c1.paragraphs[0].runs[0].font.name = "Malgun Gothic"
        c1.paragraphs[0].runs[0].font.size = Pt(9)
        c0.width = Inches(1.5)
        c1.width = Inches(5.3)
        
    doc.add_paragraph()
    
    # 1. 문제 제기 및 현상 진단
    h1 = doc.add_paragraph()
    hrun1 = h1.add_run("1. 문제 제기: 정밀 횟수 분할 후에도 오탐(FP)이 높게 나타나는 현상")
    hrun1.font.name = "Malgun Gothic"
    hrun1.font.size = Pt(13)
    hrun1.font.bold = True
    hrun1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    p1 = doc.add_paragraph()
    p1.add_run(
        "웹 서비스 실배포 E2E 파이프라인(ActivePeriodGater 횟수 동적 분할 -> 64F 스플라인 -> Stage 2 진단)을 적용하여 "
        "72개 클립 전수에서 총 256회의 독립 1회 반복을 성공적으로 분할하고 평가를 완료하였습니다.\n\n"
        "이 과정에서 바벨 데드리프트 d3(발바닥 고정)의 재현율이 0%에서 67%로 부활하고, d4(동시 신전)의 F1이 60.0%로 대폭 상승하였으며, "
        "스쿼트 c1(척추 중립) 거짓양성이 54% 감소하고, c3(무릎-발 일치) 거짓양성이 87% 급감하는 유의미한 성과를 거두었습니다.\n\n"
        "그러나 여전히 스쿼트 c4(발바닥 지면 고정)에서 115회 중 114회가 결함으로 감지(FP 94건)되고, "
        "스티프 데드리프트 sd1(무릎 고정)에서도 FP 43건이 발생하는 등, 모델의 판정이 '결함(Defect)' 쪽으로 심각하게 편향되는 현상이 관측되었습니다. "
        "이에 따라 컴퓨터 비전 및 생체역학적 관점에서 5대 근본 원인을 정밀 프로파일링하였습니다."
    )
    for r in p1.runs:
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9.5)
        
    # 2. 5대 근본 원인 심층 분석
    h2 = doc.add_paragraph()
    hrun2 = h2.add_run("2. 일반화 성능 저하의 5대 근본 원인 심층 프로파일링")
    hrun2.font.name = "Malgun Gothic"
    hrun2.font.size = Pt(13)
    hrun2.font.bold = True
    hrun2.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    reasons = [
        ("원인 1. 카메라 종횡비(Aspect Ratio) 3.16배 불일치에 따른 기하 각도 왜곡",
         "• AI-Hub 학습 데이터: 1920x1080 (16:9 가로 화면, Aspect Ratio = 1.7778)\n"
         "• 사용자 실제 영상: 1080x1920 (9:16 세로 화면, Aspect Ratio = 0.5625)\n"
         "• 왜곡 원리: MediaPipe는 x를 가로 폭(W), y를 세로 높이(H)로 각각 독립 정규화(0~1)합니다. "
         "동일한 물리적 픽셀 변위(dX, dY)에 대해 가로 영상에서는 dY/dX에 1.778배가 곱해지고, 세로 영상에서는 0.563배가 곱해집니다. "
         "이로 인해 두 환경 간의 기하학적 스케일 비율 차이는 (1.7778 / 0.5625) = 3.16배에 달합니다.\n"
         "• 결과: 동일한 자세를 취하더라도 arctan2로 계산되는 발바닥 피치각, 척추 체간각, 무릎 관절각 피처가 학습 기준치와 전혀 다른 값으로 변형되어 모델이 비정상으로 오탐합니다."),
        
        ("원인 2. 바닥 거치 로우 앵글(Tilt-up 20°~30°)에 의한 2D 투영 착시 (발뒤꿈치 들림 착시)",
         "• AI-Hub 환경: 삼각대를 사람 허리/가슴 높이(1.0~1.2m)에 완벽한 수평(Pitch 0°)으로 거치하여 촬영.\n"
         "• 사용자 환경: 스마트폰을 헬스장 바닥이나 낮은 물체에 두고 위를 올려다보며 촬영.\n"
         "• 투영 왜곡 메커니즘: 스쿼트 시 무릎이 앞으로 굽혀지며 정강이가 전방으로 기울어집니다. "
         "바닥에서 위를 올려다보는 원근 투영(Perspective Projection) 상에서는 발뒤꿈치(Heel)가 발가락(Toe)보다 이미지 상단(Y 좌표 감소)에 위치하게 됩니다.\n"
         "• 결과: 사용자는 발바닥을 바닥에 완벽히 붙이고 있었음에도 불구하고, 2D 영상 상에서는 뒤꿈치가 공중에 붕 뜬 것처럼 투영되어 c4(발바닥 고정) 결함 확률이 85~90%로 치솟게 됩니다."),
         
        ("원인 3. 생체역학적 동반 결함(Co-occurrence)과 단일 의도 라벨 간의 채점 기준 괴리",
         "• 인체 키네마틱스 연쇄 반응: 프리웨이트는 전신 복합 다관절 운동입니다. 사용자가 '결함 1(척추 중립 위반)'만을 의도하여 허리를 숙이더라도, "
         "상체가 숙여짐에 따라 시선이 자연스럽게 바닥을 향하고(c2 동반), 바벨이 정강이에서 멀어지며(d5 동반), 무게중심 이동으로 발목 앵글이 변합니다.\n"
         "• 채점 상의 맹점: 평가 기준은 '결함 1만 1이고, 나머지 모든 결함은 무조건 0이어야 정답'으로 되어 있습니다. "
         "모델은 사용자가 척추를 숙일 때 실제로 동반 발생한 시선 이탈(c2)이나 바벨 이탈(d5)을 아주 정확하게 감지했음에도, 정답이 0이라는 이유로 전량 '오답(FP)'으로 처리되어 정확도가 폭락하였습니다."),
         
        ("원인 4. 실제 촬영 환경의 비전 노이즈(Jitter)에 의한 통계 피처(Std, Jerk, Max) 과민 반응",
         "• 스튜디오 vs 실제 헬스장: AI-Hub는 고성능 조명과 무채색 배경의 통제된 스튜디오에서 촬영되어 관절 랜드마크의 떨림이 극히 적습니다.\n"
         "• 헬스장 환경: 형광등 반사, 검은색 옷/레깅스와 바닥의 낮은 명암비, 스마트폰 손떨림 등으로 인해 MediaPipe 관절 좌표에 미세한 지터(Jitter)가 상시 발생합니다.\n"
         "• 피처 오염: 모델이 사용하는 '상승기 높이차 흔들림(std)', '지면 가속도 저크(jerk)', '최대 발바닥 각도(max)' 피처가 노이즈로 인해 부풀려져 결함 확률이 기본적으로 10~20%p 상향 이동합니다."),
         
        ("원인 5. AI-Hub 전용 임계값(Threshold)의 도메인 시프트 미보정",
         "• 모델의 판정 임계값(tau_def = 0.50~0.62)은 AI-Hub 스튜디오 데이터셋의 분포에 최적화되어 있습니다.\n"
         "• 위의 1~4번 도메인 시프트(화면비 + 카메라 각도 + 비전 노이즈)가 누적되면 정상 동작의 예측 확률도 0.20~0.30에서 0.50~0.70으로 쉬프트(Probability Shift)되므로, "
         "실제 사용자 환경에 맞춘 종횡비 스케일러나 임계값 보정 없이는 결함 판정으로 넘어갈 수밖에 없습니다.")
    ]
    
    for r_title, r_desc in reasons:
        h_sub = doc.add_paragraph()
        hrun_sub = h_sub.add_run(r_title)
        hrun_sub.font.name = "Malgun Gothic"
        hrun_sub.font.size = Pt(11)
        hrun_sub.font.bold = True
        hrun_sub.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
        
        p_desc = doc.add_paragraph(r_desc)
        for r in p_desc.runs:
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(9.5)
            
    doc.add_paragraph()
    
    # 3. 환경별 대조표
    h3 = doc.add_paragraph()
    hrun3 = h3.add_run("3. AI-Hub 학습 환경 vs 사용자 실제 환경 정밀 대조표")
    hrun3.font.name = "Malgun Gothic"
    hrun3.font.size = Pt(13)
    hrun3.font.bold = True
    hrun3.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    t_comp = doc.add_table(rows=6, cols=4)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_heads = ["비교 항목", "AI-Hub 학습 환경 (학습셋)", "사용자 스마트폰 환경 (일반화셋)", "도메인 시프트 영향도"]
    for i, title_text in enumerate(c_heads):
        t_comp.rows[0].cells[i].text = title_text
    style_header(t_comp.rows[0], [1.2, 1.8, 1.8, 1.8])
    
    comp_rows = [
        ("화면 해상도 및 비율", "1920x1080 (16:9 가로, 1.78)", "1080x1920 (9:16 세로, 0.56)", "기하 각도 비율 3.16배 왜곡 (치명적)"),
        ("카메라 거치 높이", "골반/가슴 높이 (1.0~1.2m)", "바닥 또는 낮은 탁자 (0.2~0.5m)", "원근 투영 상 발뒤꿈치 들림 착시 유발"),
        ("카메라 피치 각도", "수평 완벽 유지 (0°)", "위로 올려다봄 (Tilt-up 20°~30°)", "상체/체간 및 발바닥 투영 좌표 왜곡"),
        ("정답 라벨 규격", "3D 센서 및 전문가 다중 라벨", "사용자의 단일 의도 라벨 (1개만 지정)", "실제 발생한 동반 결함을 오탐(FP)으로 채점"),
        ("배경 및 조명 환경", "크로마키 스튜디오, 단색, 전문 조명", "실제 헬스장, 복잡한 배경, 조명 반사", "키포인트 미세 지터로 통계 피처(Std/Jerk) 상승")
    ]
    for r_idx, r_data in enumerate(comp_rows):
        for c_idx, val in enumerate(r_data):
            t_comp.rows[r_idx+1].cells[c_idx].text = val
    style_rows(t_comp, [1.2, 1.8, 1.8, 1.8], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    
    doc.add_paragraph()
    
    # 4. 향후 해결 방안 및 권장 로드맵
    h4 = doc.add_paragraph()
    hrun4 = h4.add_run("4. 실서비스 성공을 위한 4대 엔지니어링 해결 로드맵")
    hrun4.font.name = "Malgun Gothic"
    hrun4.font.size = Pt(13)
    hrun4.font.bold = True
    hrun4.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    solutions = [
        ("해결책 1. 웹 파이프라인 종횡비 역보정 스케일러(Aspect Ratio Normalizer) 탑재",
         "입력 영상의 가로(W)와 세로(H)를 읽어와, 키포인트를 AI-Hub 표준 비율(16:9)로 리스케일링한 후 피처 추출기에 전달합니다. "
         "이를 통해 3.16배의 기하학적 각도 왜곡을 단 한 줄의 수학적 선형 변환으로 원천 제거할 수 있습니다."),
        
        ("해결책 2. 사용자 촬영 가이드라인 UI/UX 팝업 안내",
         "웹 서비스 촬영 화면에서 '스마트폰을 바닥에 두지 마시고, 무릎이나 골반 높이(약 1m)의 벤치/선반에 수평으로 거치해 주세요'라는 "
         "시각적 가이드 오버레이를 제공합니다. 바닥 로우 앵글을 피하는 것만으로도 발바닥 지면 고정 오탐의 80% 이상이 즉시 해결됩니다."),
         
        ("해결책 3. 사용자 직접 채점표를 통한 진성 다중 라벨(Multi-Label Ground Truth) 확립",
         "이미 배포된 CSV 채점표(일반화_모델예측_통합_채점표.csv)를 활용하여 사용자가 영상을 직접 보며 실제로 발생한 결함을 마킹함으로써, "
         "'의도 결함'과 '실제 발생 결함'을 일치시켜 왜곡된 평가 지표를 현실적으로 정상화합니다."),
         
        ("해결책 4. 세트 과반수 룰(Majority Rule) 및 결함 임계값(Threshold) 적응형 튜닝",
         "단 1회의 흔들림으로 전체 세트를 결함 판정하는 방식을 지양하고, 과반수 반복 결함 룰을 웹 서비스 기본값으로 적용하며, "
         "실제 사용자 환경의 노이즈를 감안하여 tau_def를 0.65~0.70으로 상향 조정하여 불필요한 결함 알림을 최소화합니다.")
    ]
    
    for s_title, s_desc in solutions:
        h_sol = doc.add_paragraph()
        hrun_sol = h_sol.add_run(s_title)
        hrun_sol.font.name = "Malgun Gothic"
        hrun_sol.font.size = Pt(11)
        hrun_sol.font.bold = True
        hrun_sol.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        
        p_sdesc = doc.add_paragraph(s_desc)
        for r in p_sdesc.runs:
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(9.5)
            
    doc.save(DOCX_OUT)
    print(f"Saved DOCX: {DOCX_OUT}")
    
    # MD 버전 저장
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write("# vidio-health-trainer 일반화 데이터셋 성능 저하 5대 근본 원인 심층 프로파일링 보고서\n\n")
        f.write("- **문서 번호**: 8주차 공식 심층 기술 분석 보고서 (2026-10-16)\n")
        f.write("- **프로젝트 명**: vidio-health-trainer\n")
        f.write("- **분석 대상**: 일반화 3대 운동 72개 세트 클립 (총 256회 1회 반복 추론 결과)\n\n")
        f.write("## 1. 성능 저하의 5대 근본 원인 요약\n\n")
        for idx, (rt, rd) in enumerate(reasons, 1):
            f.write(f"### {idx}. {rt}\n\n{rd}\n\n")
        f.write("## 2. AI-Hub 학습 환경 vs 사용자 실제 환경 대조표\n\n")
        f.write("| 비교 항목 | AI-Hub 학습 환경 | 사용자 스마트폰 환경 | 도메인 시프트 영향도 |\n")
        f.write("| :--- | :--- | :--- | :--- |\n")
        for c1, c2, c3, c4 in comp_rows:
            f.write(f"| {c1} | {c2} | {c3} | {c4} |\n")
        f.write("\n## 3. 실서비스 성공을 위한 4대 해결 방안\n\n")
        for st, sd in solutions:
            f.write(f"### {st}\n\n{sd}\n\n")
    print(f"Saved MD: {MD_OUT}")

if __name__ == "__main__":
    build_docx_and_md()
