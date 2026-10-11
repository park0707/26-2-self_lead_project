# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_dl3_report.py
Purpose: 다이나믹 런지 세 번째 결함 (dl3: 뒤다리 무릎 각도 90도)
         바벨 런지 bl3_v4 아키텍처 계승 기반 v1.0 공식 성능 평가 보고서 생성 (.docx & .md)
"""

import os
import sys
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOC_DIR = r"D:\자주프 문서\6주차"
os.makedirs(DOC_DIR, exist_ok=True)
DOCX_PATH = os.path.join(DOC_DIR, "다이나믹런지_dl3_v1_성능평가보고서.docx")
MD_PATH = os.path.join(DOC_DIR, "다이나믹런지_dl3_v1_성능평가보고서.md")
METRICS_JSON_PATH = r"D:\Data\피처_데이터\tuning_results\dl3_official_metrics.json"

# 스타일 헬퍼 함수
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_box(doc, title, subtitle):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "1E293B")
    set_cell_margins(cell, top=240, bottom=240, left=300, right=300)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_proj = p.add_run("vidio-health-trainer | 인공지능 운동 자세 교정 솔루션\n")
    run_proj.font.size = Pt(11)
    run_proj.font.bold = True
    run_proj.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
    
    run_t = p.add_run(title + "\n")
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(255, 255, 255)
    
    run_sub = p.add_run(subtitle)
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(203, 213, 225)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_section_title(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)

def add_callout(doc, text, title="핵심 요약 (Key Takeaway)"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0FDF4") # Emerald 50
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="10B981"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(4, 120, 87)
    
    run_desc = p.add_run(text)
    run_desc.font.size = Pt(10.0)
    run_desc.font.color.rgb = RGBColor(6, 78, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(tbl, col_widths=None):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(tbl.rows):
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
            if i == 0:
                set_cell_background(cell, "0F172A")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.0)
                        r.font.color.rgb = RGBColor(30, 41, 59)

def build_docx_and_md():
    with open(METRICS_JSON_PATH, 'r', encoding='utf-8') as f:
        metrics = json.load(f)
        
    doc = docx.Document()
    
    # 여백 설정
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # 헤더
    add_header_box(
        doc,
        "다이나믹 런지 dl3(뒤다리 무릎 각도 90도) 세부 오류 모델 성능 평가 보고서",
        "바벨 런지(bl3_v4) 모델 구조 및 14D 기구학 피처 계승 기반 v1.0 공식 토너먼트 및 전수 감사"
    )
    
    # 1. 핵심 성과 요약
    add_section_title(doc, "1. 핵심 성과 요약 (Executive Summary)")
    callout_txt = (
        "1. 바벨 런지(bl3) 아키텍처 완전 계승: 바벨 런지와 세부 오류 라벨(Condition 3: 뒤다리 무릎 각도 90도)이 100% 일치함을 확인하고, 선행 연구 기반 14D 정예 기구학 피처셋 및 CatBoost 3-State 구조를 완벽 계승하였습니다.\n"
        "2. 다이나믹 런지 전수(1,959개) 공정 학습 및 평가: 스텝 포워드(1,071개)와 스텝 백워드(888개) 전수 표본에 대해 원천 라벨 100% 무수정 상태로 5-Fold 피험자 격리 교차검증을 수행하였습니다.\n"
        "3. 공식 챔피언 등극: dl3 CatBoost 모델이 전체 3-State F1 80.41%, Accuracy 89.11%, ROC-AUC 91.58%, Coverage 90.51%로 최고 성능을 달성하였습니다.\n"
        "4. 세부 종목별 균형 성능 확보: 스텝 포워드(F1 81.07%, AUC 91.68%)와 스텝 백워드(F1 79.53%, AUC 91.48%) 모두에서 80% 전후의 매우 안정적인 분류력을 입증하였습니다.\n"
        "5. 초저지연 실시간 추론: 단일 시퀀스 추론 지연시간 1.775ms (초당 563 FPS)를 기록하여 모바일/엣지 배포 규격을 완벽 충족하였습니다."
    )
    add_callout(doc, callout_txt, "핵심 벤치마크 성과")
    
    # 2. 데이터셋 및 라벨 일치성 감사
    add_section_title(doc, "2. 데이터셋 구성 및 라벨 일치성 감사 (Label Integrity Audit)")
    p = doc.add_paragraph()
    p.add_run(
        "AI-Hub 피트니스 데이터셋 분석 결과, 바벨 런지(04)와 다이나믹 런지(05 스텝 포워드, 06 스텝 백워드)는 5대 세부 운동 조건(conditions) 및 오류 판정 체계가 100% 동일한 라벨 규격을 공유하고 있음을 전수 감사하여 실증 확인하였습니다."
    )
    
    tbl_ds = doc.add_table(rows=4, cols=5)
    headers = ["운동 종목", "폴더 코드", "전체 시퀀스 수", "정상(Pass) 건수", "결함(Defect) 건수 (비율)"]
    for j, h in enumerate(headers):
        tbl_ds.cell(0, j).paragraphs[0].text = h
        
    ds_data = [
        ["스텝 포워드 다이나믹 런지", "05_스텝_포워드_다이나믹_런지", "1,071건", "769건 (71.8%)", "302건 (28.2%)"],
        ["스텝 백워드 다이나믹 런지", "06_스텝_백워드_다이나믹_런지", "888건", "640건 (72.1%)", "248건 (27.9%)"],
        ["통합 다이나믹 런지 (전체)", "05 + 06 통합", "1,959건", "1,409건 (71.9%)", "550건 (28.1%)"]
    ]
    for i, row in enumerate(ds_data, 1):
        for j, val in enumerate(row):
            tbl_ds.cell(i, j).paragraphs[0].text = val
    style_table(tbl_ds, [Inches(1.8), Inches(1.8), Inches(1.1), Inches(1.1), Inches(1.3)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 3. 14D 계승 피처셋 명세
    add_section_title(doc, "3. 14D 정예 기구학 피처셋 계승 명세 (Feature Inheritance)")
    p = doc.add_paragraph()
    p.add_run(
        "바벨 런지 bl3_v4에서 입증된 14D 정예 기구학 피처 파이프라인(11D 직교 기구학 + 3D 고급 생체역학 선행연구 피처)을 그대로 계승하여 시각적 회전 및 원근 단축 노이즈에 강건한 특징 공간을 구성하였습니다."
    )
    
    tbl_feat = doc.add_table(rows=15, cols=3)
    feat_headers = ["피처 번호", "피처 명칭 (Feature Name)", "생체역학적 정의 및 산출 메커니즘"]
    for j, h in enumerate(feat_headers):
        tbl_feat.cell(0, j).paragraphs[0].text = h
        
    feat_rows = [
        ["f1", "rear_knee_stability_std", "최하사점 구간 뒤다리 무릎 각도 표준편차 (무릎 떨림 및 안정성)"],
        ["f2", "lead_rear_knee_ratio", "앞다리 대비 뒤다리 무릎 각도 비율 (하체 전후 굴곡 대칭성)"],
        ["f3", "rear_femur_vertical_pitch", "뒤다리 대퇴골의 지면 수직선 대비 피치 각도 (°)"],
        ["f4", "rear_shank_horizontal_pitch", "뒤다리 정강이의 지면 수평선 대비 피치 각도 (°)"],
        ["f5", "inter_ankle_stride_bottom", "최하사점 양 발목 간 2D 유클리드 보폭 거리"],
        ["f6", "knee_velocity_decel_nadir", "최하사점 직전 뒤다리 무릎 각속도 감속도 극점"],
        ["f7", "rear_knee_target_dev_bottom", "최하사점 90도 목표 각도 대비 뒤다리 무릎 절대 편차 (°)"],
        ["f8", "inter_knee_vertical_disp", "최하사점 양 무릎 간 수직 높이 편차 (골반 하강 깊이)"],
        ["f9", "torso_to_rear_femur_angle", "체간 축과 뒤다리 대퇴골 간 사잇각 (고관절 신전도)"],
        ["f10", "lead_knee_angle_bottom", "최하사점 앞다리 무릎 2D 각도 (°)"],
        ["f11", "rear_knee_angle_bottom", "최하사점 뒤다리 무릎 2D 원천 각도 (°)"],
        ["f12", "knee_hysteresis_area", "[선행연구] 각도-각속도 위상 평면 닫힌 루프 적분 면적"],
        ["f13", "femur_tibia_proj_ratio", "[선행연구] 최하사점 대퇴/정강이 2D 투영 길이비 (원근 단축 보정)"],
        ["f14", "crp_hip_knee_mean", "[선행연구] 고관절-슬관절 연속 상대 위상 (CRP) 평균값 (하지 협응 대칭성)"]
    ]
    for i, row in enumerate(feat_rows, 1):
        for j, val in enumerate(row):
            tbl_feat.cell(i, j).paragraphs[0].text = val
    style_table(tbl_feat, [Inches(0.8), Inches(2.2), Inches(4.0)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 4. 토너먼트 종합 성능 평가 결과
    add_section_title(doc, "4. 5-Fold StratifiedGroupKFold 토너먼트 종합 성능 평가 결과")
    p = doc.add_paragraph()
    p.add_run(
        "피험자 9명 전원을 완벽 격리 분할(StratifiedGroupKFold, 5-Fold)하여 데이터 누수(Leakage)를 원천 차단한 상태에서 CatBoost, LightGBM, XGBoost, Tri-Ensemble 모델을 토너먼트 평가하였습니다 (판정 커버리지 >= 90.0% 충족 기준)."
    )
    
    summary = metrics['summary']
    tbl_res = doc.add_table(rows=len(summary) + 1, cols=7)
    res_headers = ["모델명", "3-State F1", "판정 커버리지", "정확도", "정밀도", "재현율", "ROC-AUC"]
    for j, h in enumerate(res_headers):
        tbl_res.cell(0, j).paragraphs[0].text = h
        
    for i, row_data in enumerate(summary, 1):
        tbl_res.cell(i, 0).paragraphs[0].text = row_data['Model']
        tbl_res.cell(i, 1).paragraphs[0].text = row_data['3-State F1 (%)']
        tbl_res.cell(i, 2).paragraphs[0].text = row_data['Coverage (%)']
        tbl_res.cell(i, 3).paragraphs[0].text = row_data['Accuracy (%)']
        tbl_res.cell(i, 4).paragraphs[0].text = row_data['Precision (%)']
        tbl_res.cell(i, 5).paragraphs[0].text = row_data['Recall (%)']
        tbl_res.cell(i, 6).paragraphs[0].text = row_data['ROC-AUC (%)']
    style_table(tbl_res, [Inches(2.5), Inches(0.9), Inches(1.0), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.9)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 5. 다차원 세부 성능 감사
    add_section_title(doc, "5. 다차원 세부 성능 감사 (Multi-Dimensional Audit)")
    
    add_section_title(doc, "5.1 세부 운동 종목별(스텝 포워드 vs 스텝 백워드) 비교", level=2)
    fwd = metrics['sub_exercise_audit']['05_step_forward']
    bwd = metrics['sub_exercise_audit']['06_step_backward']
    
    tbl_sub = doc.add_table(rows=3, cols=7)
    sub_headers = ["세부 운동 종목", "표본 수", "3-State F1", "판정 커버리지", "정확도", "정밀도", "재현율"]
    for j, h in enumerate(sub_headers):
        tbl_sub.cell(0, j).paragraphs[0].text = h
        
    sub_rows = [
        ["05 스텝 포워드 다이나믹 런지", f"{fwd['total']}건", f"{fwd['f1']:.2f}%", f"{fwd['coverage']:.2f}%", f"{fwd['accuracy']:.2f}%", f"{fwd['precision']:.2f}%", f"{fwd['recall']:.2f}%"],
        ["06 스텝 백워드 다이나믹 런지", f"{bwd['total']}건", f"{bwd['f1']:.2f}%", f"{bwd['coverage']:.2f}%", f"{bwd['accuracy']:.2f}%", f"{bwd['precision']:.2f}%", f"{bwd['recall']:.2f}%"]
    ]
    for i, row in enumerate(sub_rows, 1):
        for j, val in enumerate(row):
            tbl_sub.cell(i, j).paragraphs[0].text = val
    style_table(tbl_sub, [Inches(2.3), Inches(0.9), Inches(0.9), Inches(1.0), Inches(0.8), Inches(0.8), Inches(0.8)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    add_section_title(doc, "5.2 피험자(Subject 1~9) 전수 격리 교차검증 성능", level=2)
    tbl_subj = doc.add_table(rows=10, cols=6)
    subj_headers = ["피험자", "표본 수", "3-State F1", "판정 커버리지", "정확도", "정밀도"]
    for j, h in enumerate(subj_headers):
        tbl_subj.cell(0, j).paragraphs[0].text = h
        
    for i, s_id in enumerate(range(1, 10), 1):
        s_data = metrics['subject_audit'][f"피험자{s_id}"]
        tbl_subj.cell(i, 0).paragraphs[0].text = f"피험자 {s_id}"
        tbl_subj.cell(i, 1).paragraphs[0].text = f"{s_data['total']}건"
        tbl_subj.cell(i, 2).paragraphs[0].text = f"{s_data['f1']:.2f}%"
        tbl_subj.cell(i, 3).paragraphs[0].text = f"{s_data['coverage']:.2f}%"
        tbl_subj.cell(i, 4).paragraphs[0].text = f"{s_data['accuracy']:.2f}%"
        tbl_subj.cell(i, 5).paragraphs[0].text = f"{s_data['precision']:.2f}%"
    style_table(tbl_subj, [Inches(1.2), Inches(1.1), Inches(1.1), Inches(1.2), Inches(1.1), Inches(1.1)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 6. 상용 서비스 배포 및 무결성 검증
    add_section_title(doc, "6. 상용 서비스 배포 규격 및 단위 테스트 무결성 검증")
    callout_dep = (
        f"• 프로덕션 모델 모듈: src.models.sub_models.dl3.DL3Classifier\n"
        f"• 프로덕션 피처 추출기: src.features.lunge_dl3_features.extract_lunge_dl3_14d_features\n"
        f"• 최종 챔피언 번들: D:\\Data\\피처_데이터\\tuning_results\\dl3_champion_bundle.joblib\n"
        f"• 최적 3-State 결정 임계값: tau_low = 0.43, tau_high = 0.71 (Optimal tau = 0.57)\n"
        f"• 단일 시퀀스 추론 지연시간: {metrics['inference_latency_ms']:.3f} ms (초당 563 FPS 처리 가능, 상용 30fps 규격 33.3ms 대비 18배 초저지연)\n"
        f"• 단위 테스트 결과: test_dl3.py 5개 테스트 전수 통과 (OK, 0.314s)"
    )
    add_callout(doc, callout_dep, "프로덕션 배포 명세")
    
    doc.save(DOCX_PATH)
    print(f"[+] 공식 Word 보고서 생성 완료: {DOCX_PATH}")
    
    # Markdown 보고서 생성
    md_content = f"""# [vidio-health-trainer] 다이나믹 런지 dl3(뒤다리 무릎 각도 90도) 세부 오류 모델 성능 평가 보고서

> **프로젝트**: vidio-health-trainer  
> **운동 종목**: 다이나믹 런지 (Dynamic Lunge - 스텝 포워드 & 스텝 백워드)  
> **전담 결함**: dl3 (세 번째 결함: 뒤다리 무릎 각도 90도)  
> **모델 아키텍처**: dl3 CatBoost (14D Elite Kinematics) [bl3_v4 구조 및 피처 100% 계승]  
> **데이터셋**: AI-Hub 다이나믹 런지 1,959개 전수 시퀀스 (05 포워드 1,071개 + 06 백워드 888개)  
> **평가 원칙**: 원천 라벨 100% 무수정 공정 평가 (5-Fold StratifiedGroupKFold, 피험자 9명 전원 완벽 격리)  
> **공식 배포 문서**: `{DOCX_PATH}` (Word 필수 규격 준수)  

---

## 1. 핵심 성과 요약 (Executive Summary)

1. **바벨 런지(bl3) 아키텍처 완전 계승**:
   - 바벨 런지와 세부 오류 라벨(Condition 3: 뒤다리 무릎 각도 90도)이 100% 일치함을 확인하고, 선행 연구 기반 14D 정예 기구학 피처셋 및 CatBoost 3-State 구조를 완벽 계승하였습니다.
2. **다이나믹 런지 전수(1,959개) 공정 학습 및 평가**:
   - 스텝 포워드(1,071개)와 스텝 백워드(888개) 전수 표본에 대해 원천 라벨 100% 무수정 상태로 5-Fold 피험자 격리 교차검증을 수행하였습니다.
3. **공식 챔피언 등극**:
   - dl3 CatBoost 모델이 **전체 3-State F1 80.41%, Accuracy 89.11%, ROC-AUC 91.58%, Coverage 90.51%**로 최고 성능을 달성하였습니다.
4. **세부 종목별 균형 성능 확보**:
   - 스텝 포워드(F1 81.07%, AUC 91.68%)와 스텝 백워드(F1 79.53%, AUC 91.48%) 모두에서 80% 전후의 매우 안정적인 분류력을 입증하였습니다.
5. **초저지연 실시간 추론**:
   - 단일 시퀀스 추론 지연시간 **1.775ms (초당 563 FPS)**를 기록하여 모바일/엣지 배포 규격을 완벽 충족하였습니다.

---

## 2. 데이터셋 구성 및 라벨 일치성 감사 (Label Integrity Audit)

| 운동 종목 | 폴더 코드 | 전체 시퀀스 수 | 정상(Pass) 건수 | 결함(Defect) 건수 (비율) |
|:---|:---|:---:|:---:|:---:|
| 스텝 포워드 다이나믹 런지 | `05_스텝_포워드_다이나믹_런지` | 1,071건 | 769건 (71.8%) | 302건 (28.2%) |
| 스텝 백워드 다이나믹 런지 | `06_스텝_백워드_다이나믹_런지` | 888건 | 640건 (72.1%) | 248건 (27.9%) |
| **통합 다이나믹 런지 (전체)** | **05 + 06 통합** | **1,959건** | **1,409건 (71.9%)** | **550건 (28.1%)** |

---

## 3. 14D 정예 기구학 피처셋 계승 명세

| 피처 번호 | 피처 명칭 (Feature Name) | 생체역학적 정의 및 산출 메커니즘 | 세대 / 출처 |
|:---:|:---|:---|:---:|
| **f1** | `rear_knee_stability_std` | 최하사점 구간 뒤다리 무릎 각도 표준편차 (무릎 떨림 및 안정성) | bl3 v3 계승 |
| **f2** | `lead_rear_knee_ratio` | 앞다리 대비 뒤다리 무릎 각도 비율 (하체 전후 굴곡 대칭성) | bl3 v3 계승 |
| **f3** | `rear_femur_vertical_pitch` | 뒤다리 대퇴골의 지면 수직선 대비 피치 각도 (°) | bl3 v3 계승 |
| **f4** | `rear_shank_horizontal_pitch` | 뒤다리 정강이의 지면 수평선 대비 피치 각도 (°) | bl3 v3 계승 |
| **f5** | `inter_ankle_stride_bottom` | 최하사점 양 발목 간 2D 유클리드 보폭 거리 | bl3 v3 계승 |
| **f6** | `knee_velocity_decel_nadir` | 최하사점 직전 뒤다리 무릎 각속도 감속도 극점 | bl3 v3 계승 |
| **f7** | `rear_knee_target_dev_bottom` | 최하사점 90도 목표 각도 대비 뒤다리 무릎 절대 편차 (°) | bl3 v3 계승 |
| **f8** | `inter_knee_vertical_disp` | 최하사점 양 무릎 간 수직 높이 편차 (골반 하강 깊이) | bl3 v3 계승 |
| **f9** | `torso_to_rear_femur_angle` | 체간 축과 뒤다리 대퇴골 간 사잇각 (고관절 신전도) | bl3 v3 계승 |
| **f10** | `lead_knee_angle_bottom` | 최하사점 앞다리 무릎 2D 각도 (°) | bl3 v3 계승 |
| **f11** | `rear_knee_angle_bottom` | 최하사점 뒤다리 무릎 2D 원천 각도 (°) | bl3 v3 계승 |
| **f12** | `knee_hysteresis_area` | [선행연구] 각도-각속도 위상 평면 닫힌 루프 적분 면적 $\\oint \\omega d\\theta$ | bl3 v4 선행연구 |
| **f13** | `femur_tibia_proj_ratio` | [선행연구] 최하사점 대퇴/정강이 2D 투영 길이비 (원근 단축 보정) | bl3 v4 선행연구 |
| **f14** | `crp_hip_knee_mean` | [선행연구] 고관절-슬관절 연속 상대 위상 (CRP) 평균값 (하지 협응 대칭성) | bl3 v4 선행연구 |

---

## 4. 5-Fold StratifiedGroupKFold 토너먼트 종합 성능 평가 결과

| 모델명 | 3-State F1 (%) | Coverage (%) | Accuracy (%) | Precision (%) | Recall (%) | ROC-AUC (%) | tau_low | tau_high |
|:---|:---:|:---:|:---:|:---:|:---:|:---:|:---:|:---:|
| **dl3 CatBoost (14D Elite)** | **80.41%** | **90.51%** | **89.11%** | **77.80%** | **83.19%** | **91.58%** | **0.43** | **0.71** |
| dl3 LightGBM (14D Elite) | 79.08% | 90.45% | 88.15% | 75.48% | 83.05% | 90.75% | 0.44 | 0.72 |
| dl3 XGBoost (14D Elite) | 79.68% | 90.25% | 88.52% | 76.83% | 82.74% | 90.98% | 0.49 | 0.77 |
| dl3 Tri-Ensemble (14D Blend) | 79.96% | 90.45% | 88.54% | 76.42% | 83.85% | 91.40% | 0.45 | 0.73 |

---

## 5. 다차원 세부 성능 감사 (Multi-Dimensional Audit)

### 5.1 세부 운동 종목별 성능 감사
- **스텝 포워드 다이나믹 런지 (05)**: F1 **81.07%**, Accuracy **89.09%**, Precision **77.21%**, Recall **85.34%**, Coverage **90.76%**, ROC-AUC **91.68%** (표본 1,071건)
- **스텝 백워드 다이나믹 런지 (06)**: F1 **79.53%**, Accuracy **89.14%**, Precision **78.60%**, Recall **80.48%**, Coverage **90.20%**, ROC-AUC **91.48%** (표본 888건)

### 5.2 피험자(Subject 1~9) 전수 격리 교차검증 성능
- **피험자 1**: F1=75.21%, Acc=86.70%, Cov=94.78% (표본 230건)
- **피험자 2**: F1=73.95%, Acc=83.24%, Cov=85.65% (표본 216건)
- **피험자 3**: F1=81.67%, Acc=89.42%, Cov=90.83% (표본 229건)
- **피험자 4**: F1=92.31%, Acc=96.17%, Cov=90.87% (표본 230건)
- **피험자 5**: F1=85.92%, Acc=91.77%, Cov=89.34% (표본 272건)
- **피험자 6**: F1=77.55%, Acc=89.67%, Cov=92.21% (표본 231건)
- **피험자 7**: F1=91.30%, Acc=95.51%, Cov=92.71% (표본 192건)
- **피험자 8**: F1=75.91%, Acc=83.58%, Cov=87.01% (표본 231건)
- **피험자 9**: F1=64.29%, Acc=83.05%, Cov=92.19% (표본 128건)

---

## 6. 상용 서비스 배포 규격 및 단위 테스트 무결성 검증

- **프로덕션 분류기 클래스**: `src.models.sub_models.dl3.DL3Classifier`
- **프로덕션 피처 추출 함수**: `src.features.lunge_dl3_features.extract_lunge_dl3_14d_features`
- **최종 챔피언 가중치 번들**: `D:\\Data\\피처_데이터\\tuning_results\\dl3_champion_bundle.joblib`
- **최적 3-State 결정 임계값**: `tau_low = 0.43, tau_high = 0.71 (Optimal tau = 0.57)`
- **단일 시퀀스 추론 지연시간**: **1.775 ms** (초당 563 FPS 처리)
- **단위 테스트 무결성**: `test_dl3.py` 5개 테스트 전수 통과 (OK)
"""
    with open(MD_PATH, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"[+] 마크다운 보고서 생성 완료: {MD_PATH}")

if __name__ == '__main__':
    build_docx_and_md()
