import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_mediapipe_report(docx_path: str = "d:/자주프 문서/1주차/Mediapipe 기록.docx"):
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc = docx.Document()

    # 1인치 여백 설정
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(24, 43, 73)
        p.paragraph_format.space_after = Pt(6)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(100, 110, 125)
        p.paragraph_format.space_after = Pt(20)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 52, 96)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 86, 160)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def add_p(text, bold=False, italic=False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(10)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_bullet(text, bold_prefix="", level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Malgun Gothic'
            run_b.font.size = Pt(9.5)
            run_b.bold = True
            run_b.font.color.rgb = RGBColor(40, 50, 60)
        run_t = p.add_run(text)
        run_t.font.name = 'Malgun Gothic'
        run_t.font.size = Pt(9.5)
        return p

    def style_table(table, col_widths, headers, data, align_center_cols=[]):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_row = table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.text = text
            shading = parse_xml(r'<w:shd {} w:fill="182B49"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Malgun Gothic'
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, row_data in enumerate(data):
            row = table.rows[row_idx + 1]
            fill_color = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = str(val)
                shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
                cell._tc.get_or_add_tcPr().append(shading)
                p = cell.paragraphs[0]
                if col_idx in align_center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Malgun Gothic'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(40, 40, 40)

        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    # 제목 영역
    add_title("2차원 복합 분할 (Subject × View Dual Split) 및 MediaPipe 기술 보고서")
    add_subtitle("피험자 독립 분할과 5대 시점 홀드아웃의 동시 교차 분할 및 미사용 데이터 전수 분석")

    # 1. 개요 및 연구 배경
    add_h1("1. 개요 및 연구 배경 (Overview & Motivation)")
    add_p("본 기술 문서는 모델의 실전 배포 환경을 가장 엄격하게 평가하기 위해, 피험자 단위 독립 분할(Subject Split)과 5대 카메라 시점 홀드아웃(View Split)을 동시에 교차 적용한 2차원 복합 분할(Dual Hold-out Matrix)의 설계 사양과 구축 결과를 기록합니다.")

    # 2. 2차원 복합 분할 (Dual Split Matrix) 사양
    add_h1("2. 2차원 복합 분할 (Dual Hold-out Matrix) 매트릭스 사양")
    add_p("피험자 그룹(Train 6명, Valid 2명, Test 1명)과 카메라 시점(Train 3개, Valid 1개, Test 1개)을 교차하여 완벽히 누수가 차단된 5대 데이터셋을 생성하였습니다.")

    dual_headers = ["데이터셋 파일명", "피험자 구성 (Subject)", "카메라 시점 (View)", "총 샘플 수 ($N$)", "Shape (A / B)", "연구 평가 목적 및 의의"]
    dual_data = [
        ["Dual_Train.npz", "피험자 1, 2, 3, 4, 5, 6", "view1, 2, 5 (정면, 대각, 후면)", "19,932개", "(19932, 97)\n(19932, 64, 69)", "순수 학습 본체 세트 (새로운 사람/각도 차단)"],
        ["Dual_Valid.npz", "피험자 7, 8", "view4 (후방 대각선)", "2,195개", "(2195, 97)\n(2195, 64, 69)", "인물과 각도가 모두 분리된 하이퍼파라미터 튜닝"],
        ["Dual_Test_Strict.npz", "피험자 9", "view3 (완전 측면 90°)", "798개", "(798, 97)\n(798, 64, 69)", "★ 복합 홀드아웃 (미학습 인물 + 미학습 각도 실전 성능)"],
        ["Dual_Test_Unseen_Subj.npz", "피험자 9", "view1, 2, 5 (학습 시점)", "2,394개", "(2394, 97)\n(2394, 64, 69)", "인물 일반화 평가 (미학습 인물, 학습 각도)"],
        ["Dual_Test_Unseen_View.npz", "피험자 1, 2, 3, 4, 5, 6", "view3 (완전 측면 90°)", "6,644개", "(6644, 97)\n(6644, 64, 69)", "시점 일반화 평가 (학습 인물, 미학습 각도)"]
    ]
    t_dual = doc.add_table(rows=len(dual_data)+1, cols=6)
    style_table(t_dual, [Inches(1.5), Inches(1.2), Inches(1.3), Inches(0.9), Inches(1.1), Inches(1.5)], dual_headers, dual_data, align_center_cols=[1, 2, 3, 4])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. 미사용 데이터 전수 분석
    add_h1("3. 로컬 데이터 전수 조사 및 미사용 데이터 현황")
    add_p("현재 워크스페이스 내에 보관되어 있으나 2D 분류 파이프라인에서 미사용 상태인 원본 데이터 자산 목록입니다.")
    
    unused_headers = ["데이터 자산 구분", "보관 디렉토리 경로", "데이터 성격 및 규모", "현재 미사용 사유 및 향후 확장 방안"]
    unused_data = [
        ["AI Hub 3D 라벨 데이터", "Data/01_raw_data/labels_json/*/3D/", "9대 운동 전체 시퀀스 3D 월드 좌표 JSON", "현재 2D 15-Joint 분류에 집중함 (향후 3D 관절 모멘트/부하 계산 2단계 연구에 활용 가능)"],
        ["UI-PRMD Vicon 3D 모캡", "Data/06_external_benchmarks/01_UI-PRMD/Vicon_모션캡처/", "고정밀 광학식 3D 마커 및 관절 각도 TXT 파일", "Kinect 2D 투영 데이터 우선 활용 (향후 모션캡처 정밀도 비교 연구 시 활용 가능)"],
        ["REHAB24-6 3D 및 마커", "Data/06_external_benchmarks/02_REHAB24-6/3D_*, 2D_마커_*", "26개 관절 3D 좌표 및 인체 표면 마커 NPY 파일", "2D 26개 관절 좌표 우선 활용 (향후 생체역학 정밀 역운동학 모델링 시 활용 가능)"]
    ]
    t_unused = doc.add_table(rows=len(unused_data)+1, cols=4)
    style_table(t_unused, [Inches(1.5), Inches(2.2), Inches(1.8), Inches(2.0)], unused_headers, unused_data, align_center_cols=[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. MediaPipe 추출 결과 및 Delta 피처
    add_h1("4. MediaPipe 추출 결과 및 Delta/Concat 피처")
    add_p("MediaPipe Pose를 통해 6,047개 시퀀스에 대해 추출을 완료하였으며, AI Hub 라벨과의 오차 및 결합 피처를 생성하였습니다.")
    add_bullet("Left/Right Shoulder, Elbow, Wrist, Hip, Knee, Ankle (12개)", "1) 1:1 직접 매핑: ")
    add_bullet("Pelvis(골반중점), Neck(어깨중점), Torso(체간중점) (3개)", "2) 파생 중점 관절: ")
    add_bullet("Format A (97D), Format B (64, 69), Delta (97D), Concat (194D)", "3) 피처 세트: ")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. 저장 경로 안내
    add_h1("5. 최종 데이터셋 저장 경로")
    add_bullet("d:/Data/피처_데이터/Dual_Splits/Dual_Train.npz (19,932개)", "Dual Train Set: ")
    add_bullet("d:/Data/피처_데이터/Dual_Splits/Dual_Valid.npz (2,195개)", "Dual Valid Set: ")
    add_bullet("d:/Data/피처_데이터/Dual_Splits/Dual_Test_Strict.npz (798개 - 복합 홀드아웃)", "Dual Test Strict: ")
    add_bullet("d:/Data/피처_데이터/Dual_Splits/Dual_Test_Unseen_Subj.npz (2,394개)", "Test Unseen Subject: ")
    add_bullet("d:/Data/피처_데이터/Dual_Splits/Dual_Test_Unseen_View.npz (6,644개)", "Test Unseen View: ")

    doc.save(docx_path)
    print(f"[+] Word 문서 갱신 완료: {docx_path}")

if __name__ == '__main__':
    create_mediapipe_report()
