import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_pipeline_plan(docx_path: str = "d:/자주프 문서/1주차/모델_파이프라인.docx"):
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(24, 43, 73)
        p.paragraph_format.space_after = Pt(6)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(100, 110, 125)
        p.paragraph_format.space_after = Pt(18)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 52, 96)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 86, 160)
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)

    def add_p(text, bold=False, italic=False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Malgun Gothic'
        run.font.size = Pt(9.5)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        return p

    def add_bullet(text, bold_prefix="", level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Malgun Gothic'
            run_b.font.size = Pt(9.0)
            run_b.bold = True
            run_b.font.color.rgb = RGBColor(40, 50, 60)
        run_t = p.add_run(text)
        run_t.font.name = 'Malgun Gothic'
        run_t.font.size = Pt(9.0)
        return p

    def style_table(table, col_widths, headers, data, align_center_cols=[]):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_row = table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.text = text
            shading = parse_xml(r'<w:shd {} w:fill="182B49"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Malgun Gothic'
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, row_data in enumerate(data):
            row = table.rows[row_idx + 1]
            fill_color = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = str(val)
                shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
                cell._tc.get_or_add_tcPr().append(shading)
                p = cell.paragraphs[0]
                if col_idx in align_center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Malgun Gothic'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(40, 40, 40)

        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    # 제목 영역
    add_title("3단계 계층형 모델 파이프라인 전 단계 실증 벤치마크 보고서")
    add_subtitle("1단계(9종 분류) ──> 2단계(정상/오류 판정) ──> 3단계(세부 결함 진단) 전 단계 구현 및 비교 평가")

    # 1. 아키텍처 개요
    add_h1("1. 파이프라인 아키텍처 개요 (System Architecture)")
    add_p("본 연구의 핵심 파이프라인은 '① 무슨 운동인가? -> ② 잘하고 있는가? -> ③ 무엇이 틀렸는가?'를 단계별로 진단하는 계층형 캐스케이드 구조입니다.")
    add_bullet("src/models/first_model/ : 9종 운동 대분류 (Multiclass Classification)", "1단계: ")
    add_bullet("src/models/second_models/ : 정상 vs 오류 이진 판정 (Binary Quality Assessment)", "2단계: ")
    add_bullet("src/models/third_models/ : 세부 감점 조건 다중 라벨 진단 (Multi-Label Defect Diagnosis)", "3단계: ")

    # 2. 1단계 실증 벤치마크 결과
    add_h1("2. 1단계 9종 운동 종목 분류 실증 벤치마크 결과 (Stage 1)")
    bench1_headers = ["모델명 (Model Architecture)", "입력 규격", "Valid F1", "Strict Test F1 (홀드아웃)", "Subj Test F1 (인물일반화)", "추론 속도"]
    bench1_data = [
        ["1D-CNN + Bi-GRU (Temporal)", "Format B (64, 69)", "92.12%", "58.04% (1위)", "96.53% (1위)", "441.7 FPS"],
        ["XGBoost (Hist Gradient)", "Format A (97D)", "90.02%", "53.89% (2위)", "92.90%", "150,652 FPS"],
        ["MLP (Dense Neural Net)", "Format A (97D)", "89.75%", "51.90%", "92.21%", "1,017,742 FPS"],
        ["LightGBM (Fast GBDT)", "Format A (97D)", "89.87%", "45.16%", "92.60%", "155,373 FPS"],
        ["Random Forest (200 Trees)", "Format A (97D)", "86.93%", "40.86%", "91.77%", "31,059 FPS"]
    ]
    t1 = doc.add_table(rows=len(bench1_data)+1, cols=6)
    style_table(t1, [Inches(1.8), Inches(1.1), Inches(0.9), Inches(1.2), Inches(1.2), Inches(0.9)], bench1_headers, bench1_data, align_center_cols=[1, 2, 3, 4, 5])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. 2단계 실증 벤치마크 결과
    add_h1("3. 2단계 정상/오류 판별 실증 벤치마크 결과 (Stage 2)")
    add_p("정상 5.5% vs 오류 94.5%의 극심한 불균형 조건에서 비용 민감(Cost-Sensitive) 학습 및 Focal Loss를 적용한 결과입니다.")
    bench2_headers = ["모델명 (Model Architecture)", "입력 규격", "Valid Macro-F1", "Strict Test F1 (홀드아웃)", "Subj Test F1 (인물일반화)", "PR-AUC (Subj)"]
    bench2_data = [
        ["1D-CNN + Bi-GRU (Focal)", "Format B (64, 69)", "66.69% (1위)", "53.83% (1위)", "64.79%", "0.334"],
        ["Cost-Sensitive XGBoost", "Format A (97D)", "65.98%", "49.60%", "66.13% (1위)", "0.414 (1위)"],
        ["Balanced Random Forest", "Format A (97D)", "57.13%", "48.62%", "56.87%", "0.310"],
        ["Cost-Sensitive LightGBM", "Format A (97D)", "48.59%", "48.62%", "48.62%", "0.124"]
    ]
    t2 = doc.add_table(rows=len(bench2_data)+1, cols=6)
    style_table(t2, [Inches(1.8), Inches(1.1), Inches(1.0), Inches(1.2), Inches(1.2), Inches(0.9)], bench2_headers, bench2_data, align_center_cols=[1, 2, 3, 4, 5])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. 3단계 실증 벤치마크 결과
    add_h1("4. 3단계 세부 결함(4대 감점 조건) 다중 라벨 진단 실증 벤치마크 (Stage 3)")
    add_p("01_바벨_스쿼트의 4가지 세부 결함('척추 중립', '고개 정면', '무릎 방향 일치', '지면 고정')을 동시 진단한 결과입니다.")
    bench3_headers = ["모델명 (Model Architecture)", "입력 규격", "Valid Micro-F1", "Strict Test F1 (홀드아웃)", "Subj Test F1 (인물일반화)", "Hamming Loss"]
    bench3_data = [
        ["1D-CNN + Bi-GRU (Multi-Head)", "Format B (64, 69)", "75.82%", "60.98% (1위)", "72.67%", "0.4648"],
        ["Multi-Output Random Forest", "Format A (97D)", "76.16% (1위)", "60.42%", "74.36%", "0.4453 (1위)"],
        ["Multi-Output LightGBM", "Format A (97D)", "75.93%", "51.36%", "75.27% (1위)", "0.4883"],
        ["Multi-Output XGBoost", "Format A (97D)", "75.27%", "51.18%", "73.24%", "0.4844"]
    ]
    t3 = doc.add_table(rows=len(bench3_data)+1, cols=6)
    style_table(t3, [Inches(1.8), Inches(1.1), Inches(1.0), Inches(1.2), Inches(1.2), Inches(0.9)], bench3_headers, bench3_data, align_center_cols=[1, 2, 3, 4, 5])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 5. 소스 코드 구조 요약
    add_h1("5. 구축 완료된 전체 소스 코드 패키지 구조")
    add_bullet("src/models/first_model/ : 1단계 9종 운동 분류 (RF, XGB, LGBM, MLP, Bi-GRU, evaluate_all)", "1단계 패키지: ")
    add_bullet("src/models/second_models/ : 2단계 정상/오류 판정 (Balanced RF, Cost-XGB, Cost-LGBM, Focal Bi-GRU)", "2단계 패키지: ")
    add_bullet("src/models/third_models/ : 3단계 세부 결함 진단 (Multi-Output RF, XGB, LGBM, Multi-Head Bi-GRU)", "3단계 패키지: ")

    doc.save(docx_path)
    print(f"[+] Word 문서 갱신 완료: {docx_path}")

if __name__ == '__main__':
    create_pipeline_plan()
