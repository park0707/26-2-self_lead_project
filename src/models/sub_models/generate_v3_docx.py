import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('vidio-health-trainer', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('굿모닝 척추의 중립 오류 모델 성능 평가 보고서 (v3)', level=1)

# Summary
doc.add_heading('1. 개요 및 목적', level=2)
doc.add_paragraph('본 문서는 굿모닝 척추의 중립 오류(Error 3) 판별에 대해, 가장 우수한 피처 조합을 보였던 v2_a (MediaPipe + SpinePose) 아키텍처를 기반으로 판단 보류율(Reject Rate)을 약 10% 수준으로 조정한 신규 모델(v3)의 성능을 평가한 보고서입니다.')

# Architecture
doc.add_heading('2. 모델 아키텍처 및 튜닝 전략', level=2)
doc.add_paragraph(' 1) 피처 세트: 45D (MediaPipe 33관절 기반 Stiff Spine & Kinematics + SpinePose 9관절 정밀 척추)', style='List Bullet')
doc.add_paragraph(' 2) 3-State 판단 보류 조건 수정: Coverage(판별률) 하한선을 기존 85%에서 90%로 상향 조정 (보류율 10% 이하 강제)', style='List Bullet')

# Performance
doc.add_heading('3. 핵심 성능 지표 (5-Fold 교차 검증)', level=2)

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Metric'
hdr_cells[1].text = 'v2_a (보류율 15%)'
hdr_cells[2].text = 'v3 (보류율 10% 강제)'

data = [
    ('Coverage', '85.20%', '90.78%'),
    ('F1 Score', '90.13%', '88.41%'),
    ('Accuracy', '90.45%', '87.90%'),
    ('Precision', '93.58%', '86.44%'),
    ('Recall', '86.93%', '90.47%'),
    ('ROC-AUC', '93.07%', '93.07%'),
]

for item in data:
    row_cells = table.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]

doc.add_heading('4. 상세 분석 및 결론', level=2)
doc.add_paragraph('판단 보류율(Reject Rate)을 15%에서 약 10%로 줄임에 따라 판별 구간(Coverage)이 90.78%로 확대되었습니다. 이에 따라 더 많은 데이터를 강제로 판별하게 되어 전체적인 정밀도(Precision)는 93.58%에서 86.44%로 약 7%p 가량 하락하였으며, F1 Score 역시 90.13%에서 88.41%로 소폭 하락했습니다.')
doc.add_paragraph('하지만 Coverage 100%였던 v1 모델(F1 85.16%)과 비교할 때, 여전히 약 10%의 애매한 데이터를 걸러냄으로써 3%p 이상의 F1 상승(85.16% -> 88.41%)을 확보하고 있습니다. 따라서 서비스 제공 시 "사용자 피드백 응답률(Coverage 90%)"과 "피드백의 정확성(F1 88%)" 사이의 비즈니스적 밸런스를 맞추기에 가장 적합한 최종 타협 모델로 평가됩니다.')

doc.save(r'D:\자주프 문서\gm3_v3_performance_report.docx')
print('Report saved.')
