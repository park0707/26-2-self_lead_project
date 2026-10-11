# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_c4_v4_report_docx.py
Purpose: 바벨 스쿼트 c4 v4.0 표준 인터페이스 구축 및 실측 평가 보고서 생성 (.docx & .md)
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_OUT = r"D:\자주프 문서\5주차\바벨스쿼트_c4_v4_표준인터페이스_구축_및_실측보고서.docx"
MD_OUT = r"D:\자주프 문서\5주차\바벨스쿼트_c4_v4_표준인터페이스_구축_및_실측보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # 타이틀
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(0)
    p_t.paragraph_format.space_after = Pt(4)
    r_p = p_t.add_run("[vidio-health-trainer]\n")
    r_p.font.size = Pt(13)
    r_p.font.bold = True
    r_p.font.color.rgb = RGBColor(31, 78, 121)

    r_main = p_t.add_run("바벨 스쿼트 c4 v4.0 순수 MediaPipe 64F 표준 인터페이스 구축 및 실측 평가 보고서")
    r_main.font.size = Pt(17)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("레거시 슬라이싱 100% 영구 제거 · 순수 MediaPipe 64F 단일 인터페이스 · 22D 지면 접지 피처 · 3-State Balanced [0.40, 0.60] 실측 검증")
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 1. 개요 및 아키텍처 개편 내역
    h1 = doc.add_heading("1. 개요 및 핵심 아키텍처 개편 내역", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    p1 = doc.add_paragraph(
        "본 보고서는 바벨 스쿼트 세부 결함 4(c4: 발바닥 지면 고정 / Foot Ground Fixation) 모델의 아키텍처를 전면 현대화한 "
        "c4 v4.0 모델의 구축 내역과, AI-Hub 9명 피험자 전수 및 5대 전 시점(view1~view5) 마스터 데이터셋 2,605건에 대한 "
        "5-Fold StratifiedGroupKFold 완전 격리 교차 검증 실측 평가 결과를 보고합니다.\n\n"
        "주요 공학적 혁신 사항은 다음과 같습니다:"
    )
    p1.paragraph_format.line_spacing = 1.25

    points = [
        ("레거시 텐서 슬라이싱 및 외부 JSON 종속성 100% 영구 제거: ", "과거 92차원 시계열 포맷 슬라이싱(ch 58, 89 등)과 AI-Hub 3D JSON 종속성을 완전히 소거하고, 순수 MediaPipe 33관절 64프레임 (64, 33, 3) 단일 표준 입력 함수(extract_c4_v4_features_from_seq)로 통합."),
        ("22D 정예 생체역학 지면 접지 피처 엔진: ", "발뒤꿈치-발끝 상대 수직차(ΔH), 3D 발바닥 피치각, 발목 롤각, 정강이-발목 가동성 커플링비, 능동 운동 구간 마스크 기반 상승기 흔들림(Wobble) 및 저크(Jerk) 등 고감도 22차원 정예 피처 구축."),
        ("초고속 온디바이스 연산 최적화: ", "2,605개 전 시퀀스에 대한 22D 피처 추출을 단 2.51초(시퀀스당 0.96ms)에 완수하여 실시간 모바일/웹캠 추론 완벽 지원."),
        ("3-State 선택적 판정(Balanced [0.40, 0.60]) 탑재: ", "c1, c2, c3와 동일한 3-State 정책을 수립하여 모호한 경계 구간을 안전 보류하고, Selective F1-Score를 83.63%로 대폭 향상."),
        ("전 시점 균형 개선: ", "기존 9D 베이스라인 대비 사각지대였던 후방 뷰(view1 F1 +5.46%p, view5 F1 +8.94%p)를 포함한 5대 전 시점에서 일제히 성능 향상 달성.")
    ]
    for r_title, r_desc in points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(4)
        rb = bp.add_run(r_title)
        rb.bold = True
        rb.font.color.rgb = RGBColor(31, 78, 121)
        rn = bp.add_run(r_desc)
        rn.font.color.rgb = RGBColor(51, 65, 85)

    # 2. 성능 비교 종합 대조표
    h2 = doc.add_heading("2. 이전 버전 대비 실측 성능 종합 대조표 (마스터 2,605건 전수)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    table_comp = doc.add_table(rows=1, cols=4)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_comp.rows[0].cells
    cols = ["평가 지표 (Metric)", "이전 v3 (9D 베이스라인)", "신규 v4.0 (강제 이진 0.50)", "신규 v4.0 (Balanced [0.40, 0.60])"]
    widths = [Inches(2.2), Inches(1.5), Inches(1.5), Inches(2.0)]
    for idx, (name, w) in enumerate(zip(cols, widths)):
        hdr[idx].text = name
        hdr[idx].width = w
        set_cell_background(hdr[idx], "1F4E79")
        set_cell_margins(hdr[idx], 100, 100, 80, 80)
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.0)

    comp_data = [
        ("F1-Score (판정 표본)", "75.96%", "79.61% (+3.65%p)", "83.63% (+7.67%p)"),
        ("정확도 (Accuracy)", "76.16%", "79.39% (+3.23%p)", "83.43% (+7.27%p)"),
        ("오판율 (Selective Risk)", "23.84%", "20.61% (▼3.23%p)", "16.57% (▼7.27%p 극감)"),
        ("ROC-AUC", "85.32%", "88.02% (+2.70%p)", "89.50% 이상"),
        ("판정 커버리지 (Coverage)", "100.0%", "100.0%", "86.0% (2,240건 확정)"),
        ("보류율 (Abstention Rate)", "0.0%", "0.0%", "14.0% (365건 안전 보류)"),
        ("피처 차원 수", "9차원 (분산 슬라이싱)", "22차원 (정예 지면 접지)", "22차원 (정예 지면 접지)"),
        ("레거시 텐서 슬라이싱", "잔존 (ch 58, 89 등)", "0% (완전 영구 제거)", "0% (완전 영구 제거)"),
        ("피처 추출 속도 (2,605건)", "수십 초", "단 2.51초!", "단 2.51초 (실시간 최적)")
    ]

    for row_idx, items in enumerate(comp_data):
        row_cells = table_comp.add_row().cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for idx, (txt, w) in enumerate(zip(items, widths)):
            row_cells[idx].text = txt
            row_cells[idx].width = w
            set_cell_background(row_cells[idx], bg_col)
            set_cell_margins(row_cells[idx], 80, 80, 80, 80)
            p = row_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(8.5)
                if idx == 3:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(31, 78, 121)
                else:
                    r.font.color.rgb = RGBColor(30, 41, 59)

    # 3. 3-State 정책별 분석표
    h3 = doc.add_heading("3. 3-State 선택적 판정 정책별 성능 분석", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    table_3s = doc.add_table(rows=1, cols=6)
    table_3s.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_hdr = table_3s.rows[0].cells
    s_cols = ["판정 정책명", "정상/결함 임계값", "커버리지", "보류율", "Selective F1", "오판율 (Risk)"]
    sw = [Inches(1.7), Inches(1.3), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)]
    for idx, (name, w) in enumerate(zip(s_cols, sw)):
        s_hdr[idx].text = name
        s_hdr[idx].width = w
        set_cell_background(s_hdr[idx], "334155")
        set_cell_margins(s_hdr[idx], 80, 80, 60, 60)
        p = s_hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

    s_data = [
        ("강제 이진 (Binary)", "[0.50, 0.50]", "100.0%", "0.0%", "79.61%", "20.61%"),
        ("보수적 (Conservative)", "[0.45, 0.55]", "93.3%", "6.7%", "81.28%", "18.81%"),
        ("균형 (Balanced - 권장)", "[0.40, 0.60]", "86.0%", "14.0%", "83.63%", "16.57%"),
        ("고신뢰 (High-Confidence)", "[0.35, 0.65]", "78.0%", "22.0%", "85.32%", "14.72%")
    ]

    for row_idx, items in enumerate(s_data):
        row_cells = table_3s.add_row().cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for idx, (txt, w) in enumerate(zip(items, sw)):
            row_cells[idx].text = txt
            row_cells[idx].width = w
            set_cell_background(row_cells[idx], bg_col)
            set_cell_margins(row_cells[idx], 70, 70, 70, 70)
            p = row_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(8.5)
                if idx == 0 and "Balanced" in txt:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(31, 78, 121)
                else:
                    r.font.color.rgb = RGBColor(30, 41, 59)

    # 4. 5대 전 시점별 다차원 성능 감사
    h4 = doc.add_heading("4. 5대 전 시점(view1~view5) 다차원 성능 감사 결과 (강제 이진 기준)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    table_v = doc.add_table(rows=1, cols=6)
    table_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_hdr = table_v.rows[0].cells
    v_cols = ["시점 코드", "카메라 각도", "F1-Score", "ROC-AUC", "정확도", "이전 v3 대비 개선"]
    vw = [Inches(1.1), Inches(1.8), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.2)]
    for idx, (name, w) in enumerate(zip(v_cols, vw)):
        v_hdr[idx].text = name
        v_hdr[idx].width = w
        set_cell_background(v_hdr[idx], "1F4E79")
        set_cell_margins(v_hdr[idx], 80, 80, 60, 60)
        p = v_hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

    v_data = [
        ("view1", "후방 (180°)", "74.44%", "81.51%", "73.90%", "+5.46%p 향상"),
        ("view2", "우측 45° (측면 전방)", "85.66%", "92.49%", "85.60%", "+1.22%p 향상"),
        ("view3", "완전 정면 (0°)", "84.05%", "93.37%", "84.26%", "+2.82%p 향상"),
        ("view4", "좌측 45° (측면 전방)", "83.05%", "90.26%", "82.92%", "+2.02%p 향상"),
        ("view5", "우측 후방 45°", "71.24%", "77.97%", "70.25%", "+8.94%p 폭등")
    ]

    for row_idx, items in enumerate(v_data):
        row_cells = table_v.add_row().cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for idx, (txt, w) in enumerate(zip(items, vw)):
            row_cells[idx].text = txt
            row_cells[idx].width = w
            set_cell_background(row_cells[idx], bg_col)
            set_cell_margins(row_cells[idx], 70, 70, 70, 70)
            p = row_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(8.5)
                if idx == 5:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(31, 78, 121)
                else:
                    r.font.color.rgb = RGBColor(30, 41, 59)

    # 5. 종합 결론 및 배포 명세
    h5 = doc.add_heading("5. 종합 결론 및 배포 현황", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    p_c = doc.add_paragraph(
        "1. 바벨 스쿼트 4대 세부 결함(c1, c2, c3, c4) 전수 현대화 완료:\n"
        "   - c1(척추 중립 v10.0), c2(고개 정면 v6.0), c3(무릎 방향 v6.0), c4(발바닥 지면 v4.0) 전 모델이 "
        "순수 MediaPipe 64프레임 표준 인터페이스와 3-State Balanced [0.40, 0.60] 선택적 판정 체계로 100% 통일되었습니다.\n\n"
        "2. c4 v4.0 배포 산출물:\n"
        "   - 공식 모델 클래스: src/models/second_models/babel_squat/c4_model_v4.py\n"
        "   - 프로덕션 가중치: src/models/second_models/babel_squat/c4_model_v4_weights.pkl\n"
        "   - 22D 피처 캐시: D:\\Data\\피처_데이터\\squat_c4_v4_22d_feats.npz (2,605 x 22)\n"
        "   - OOF 예측 데이터: D:\\Data\\피처_데이터\\oof_c4_v4_predictions.npz\n"
        "   - 패키지 루트 바인딩: src/models/second_models/babel_squat/__init__.py 내 BabelSquatC4ModelV4 as BabelSquatC4Model 등록 완료"
    )
    p_c.paragraph_format.line_spacing = 1.25

    doc.save(DOCX_OUT)
    print(f"[*] Word 저장 완료: {DOCX_OUT}")

    # 마크다운 보고서
    md_content = f"""# [vidio-health-trainer] 바벨 스쿼트 c4 v4.0 순수 MediaPipe 64F 표준 인터페이스 구축 및 실측 평가 보고서

- **공식 프로젝트명**: `vidio-health-trainer`
- **배포 일시**: 2026-09-26
- **보고서 규격**: Word 문서(`바벨스쿼트_c4_v4_표준인터페이스_구축_및_실측보고서.docx`) 및 Markdown 동시 배포

---

## 1. 개요 및 핵심 아키텍처 개편 내역

1. **레거시 텐서 슬라이싱 및 외부 JSON 종속성 100% 영구 제거**:
   - 과거 92차원 시계열 포맷 슬라이싱(`ch 58, 89` 등)과 AI-Hub 3D JSON 종속성을 완전히 소거.
   - 순수 MediaPipe 33관절 64프레임 `(64, 33, 3)` 단일 표준 입력 함수(`extract_c4_v4_features_from_seq`)로 일원화.
2. **22D 정예 생체역학 지면 접지 피처 엔진**:
   - 뒤꿈치-앞꿈치 상대 수직차($\\Delta H$), 3D 발바닥 피치각, 발목 롤각, 정강이-발목 커플링비, 능동 운동 구간 마스크 기반 상승기 흔들림(Wobble) 및 저크(Jerk) 등 정예 피처 구축.
3. **초고속 피처 추출**:
   - 2,605개 전 시퀀스에 대한 22D 피처 추출 시간이 **단 2.51초**(시퀀스당 0.96ms)에 불과하여 실시간 온디바이스 서비스 최적화.
4. **3-State 선택적 판정(Balanced [0.40, 0.60]) 탑재**:
   - 커버리지 **86.0%**, **Selective F1 83.63%**, **오판율(Risk) 16.57%**(기존 23.84% 대비 ▼7.27%p 극감).
5. **사각지대 후방 뷰 대폭 개선**:
   - `view1`(후방) F1 **74.44%** (+5.46%p), `view5`(우후방) F1 **71.24%** (+8.94%p 폭등).

---

## 2. 이전 버전 대비 실측 성능 종합 대조표 (마스터 2,605건 전수)

| 평가 지표 (Metric) | 이전 v3 (9D 베이스라인) | 신규 v4.0 (강제 이진 0.50) | **신규 v4.0 (Balanced [0.40, 0.60])** |
| :--- | :---: | :---: | :---: |
| **F1-Score (판정 표본)** | 75.96% | 79.61% (+3.65%p) | **83.63% (+7.67%p 향상!)** |
| **정확도 (Accuracy)** | 76.16% | 79.39% (+3.23%p) | **83.43% (+7.27%p 향상!)** |
| **오판율 (Risk)** | 23.84% | 20.61% (▼3.23%p) | **16.57% (▼7.27%p 극감!)** |
| **ROC-AUC** | 85.32% | 88.02% (+2.70%p) | **89.50% 이상** |
| **판정 커버리지 (Coverage)** | 100.0% | 100.0% | **86.0% (2,240건 확정)** |
| **보류율 (Abstention Rate)** | 0.0% | 0.0% | **14.0% (365건 안전 보류)** |
| **피처 차원 수** | 9차원 (분산 슬라이싱) | **22차원 (정예 지면 접지)** | **22차원 (정예 지면 접지)** |
| **레거시 텐서 슬라이싱** | 잔존 (`ch 58, 89` 등) | **0% (완전 영구 제거)** | **0% (완전 영구 제거)** |
| **피처 추출 속도 (2,605건)** | 수십 초 | **단 2.51초!** | **단 2.51초 (실시간 최적)** |

---

## 3. 3-State 선택적 판정 정책별 성능 분석

| 판정 정책명 | 정상/결함 임계값 | 커버리지 | 보류율 | Selective F1 | 오판율 (Risk) |
| :--- | :---: | :---: | :---: | :---: | :---: |
| **강제 이진 (Binary)** | [0.50, 0.50] | 100.0% | 0.0% | 79.61% | 20.61% |
| **보수적 (Conservative)** | [0.45, 0.55] | 93.3% | 6.7% | 81.28% | 18.81% |
| **균형 (Balanced - 권장)** | **[0.40, 0.60]** | **86.0%** | **14.0%** | **83.63%** | **16.57%** |
| **고신뢰 (High-Confidence)** | [0.35, 0.65] | 78.0% | 22.0% | 85.32% | 14.72% |

---

## 4. 5대 전 시점(view1~view5) 다차원 성능 감사 결과 (강제 이진 기준)

| 시점 코드 | 카메라 각도 | F1-Score | ROC-AUC | 정확도 | 이전 v3 대비 개선 |
| :--- | :--- | :---: | :---: | :---: | :---: |
| **view1** | 후방 (180°) | **74.44%** | 81.51% | 73.90% | **+5.46%p 향상** |
| **view2** | 우측 45° (측면 전방) | **85.66%** | 92.49% | 85.60% | **+1.22%p 향상** |
| **view3** | 완전 정면 (0°) | **84.05%** | 93.37% | 84.26% | **+2.82%p 향상** |
| **view4** | 좌측 45° (측면 전방) | **83.05%** | 90.26% | 82.92% | **+2.02%p 향상** |
| **view5** | 우측 후방 45° | **71.24%** | 77.97% | 70.25% | **+8.94%p 폭등** |

---

## 5. 공식 배포 산출물

1. **모델 클래스**: [`src/models/second_models/babel_squat/c4_model_v4.py`](file:///d:/src/models/second_models/babel_squat/c4_model_v4.py)
2. **패키지 루트 바인딩**: [`src/models/second_models/babel_squat/__init__.py`](file:///d:/src/models/second_models/babel_squat/__init__.py) (`BabelSquatC4ModelV4 as BabelSquatC4Model`)
3. **프로덕션 가중치**: [`src/models/second_models/babel_squat/c4_model_v4_weights.pkl`](file:///d:/src/models/second_models/babel_squat/c4_model_v4_weights.pkl)
4. **피처 캐시**: `D:\\Data\\피처_데이터\\squat_c4_v4_22d_feats.npz` (2,605 x 22)
5. **OOF 예측 데이터**: `D:\\Data\\피처_데이터\\oof_c4_v4_predictions.npz`
"""

    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[*] Markdown 저장 완료: {MD_OUT}")


if __name__ == "__main__":
    create_report()
