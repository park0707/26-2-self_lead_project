# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_external_ai_background_briefing_docx.py
Purpose: 외부 AI 기술 자문용 바벨 스쿼트 통합 진단 엔진 v9.0 배경지식 명세서 DOCX 생성
Targets:
  1) D:\자주프 문서\vidio-health-trainer_외부AI_자문용_바벨스쿼트_통합진단_배경지식_및_아키텍처_명세서.docx
  2) D:\자주프 문서\5주차\vidio-health-trainer_외부AI_자문용_바벨스쿼트_통합진단_배경지식_및_아키텍처_명세서.docx
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_docx(target_path):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc = Document()

    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    style_normal = doc.styles['Normal']
    style_normal.font.name = '맑은 고딕'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 표지 타이틀
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("vidio-health-trainer\n바벨 스쿼트 AI 통합 진단 엔진 v9.0 상세 아키텍처 및\n외부 AI 기술 자문 배경지식 명세서")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Technical Specification & Baseline Briefing for External AI Advisors: Enhancing Clean F1 & 4-Defect EMR")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_meta = p_meta.add_run("작성일자: 2026. 09. 27 | 시스템 버전: v9.0 Master Production | 프로젝트 공식 명칭: vidio-health-trainer")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph()

    # 1. 문서의 목적 및 외부 AI 자문 핵심 목표
    doc.add_heading("1. 문서의 목적 및 외부 AI 자문 핵심 목표 (Executive Mission)", level=1)
    p_goal_intro = doc.add_paragraph(
        "본 명세서는 vidio-health-trainer 프로젝트의 바벨 스쿼트 AI 진단 시스템에 대해 외부 최신 LLM 및 멀티모달 AI 연구원(Advisor)에게 현재 시스템의 전처리 파이프라인, 4대 서브모델(Layer 1), 다층 앙상블 및 동적 라우팅(Layer 2~4), 데이터 구조를 완전히 공개하고, [정상 운동자 분류 F1-Score (Clean F1)] 및 [4대 결함 동시 적중률 (4-Defect EMR)]을 추가로 비약시키기 위한 최신 AI 기술 솔루션을 도출하기 위해 작성되었습니다."
    )
    p_goal_intro.paragraph_format.line_spacing = 1.2

    # 핵심 목표 테이블
    t_goal = doc.add_table(rows=4, cols=3)
    t_goal.alignment = WD_TABLE_ALIGNMENT.CENTER
    g_headers = ["핵심 최우선 과제", "현재 v9.0 달성 수준 (N=2,605)", "차세대 v10.0 목표치"]
    for j, h in enumerate(g_headers):
        cell = t_goal.cell(0, j)
        cell.text = h
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, 120, 120, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "맑은 고딕"
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    g_rows = [
        ("4-Defect EMR (완전 일치율)\n4개 결함 여부를 동시 100% 적중", "71.24% (선택적 판정 적용 시)\n(100% 전수 기준: 68.98% / 1,797건 적중)", "80.00% 이상 돌파\n(전수 커버리지 75% 이상)"),
        ("Clean F1-Score (정상 판정)\n무결점 정상 운동자 종합 적중 F1", "67.63%\n(Precision: 60.27%, Recall: 77.05%)", "75.00% 이상 달성\n(Precision 70% 이상)"),
        ("RETAKE (재촬영 권고율)\n판정 불확실/노이즈 시 안전 보류", "3.19% (83건 / 2,605건)\n(데드밴드 [0.32, 0.38], 마진 Δ < 0.05)", "4.00% 이하 유지\n(사용자 경험 저해 최소화)")
    ]
    for i, row in enumerate(g_rows):
        for j, val in enumerate(row):
            cell = t_goal.cell(i+1, j)
            cell.text = val
            set_cell_margins(cell, 80, 80, 80, 80)
            if i % 2 == 1:
                set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "맑은 고딕"
                r.font.size = Pt(8.5)

    doc.add_paragraph()

    # 2. 데이터셋 규격 및 전처리 표준 파이프라인
    doc.add_heading("2. 데이터셋 규격 및 전처리 표준 파이프라인 (Data Pipeline Invariants)", level=1)
    p_pipe = doc.add_paragraph(
        "■ 1) 데이터셋 및 피험자 격리 원칙 (Subject Isolation):\n"
        "• AI-Hub 헬스 피트니스 바벨 스쿼트 골든셋 전수 2,605건 사용.\n"
        "• 피험자 1번부터 9번까지 9명 전수를 학습 및 검증에 반영하며, 평가 시 StratifiedGroupKFold (n_splits=5, groups=subject_id)를 적용하여 학습 세트와 검증 세트 간 피험자 데이터가 절대 섞이지 않도록 완전 격리.\n\n"
        "■ 2) 공식 5대 카메라 시점 체계 (Official 5-View Taxonomy):\n"
        "• 시점 A (우측 후방 45도 / 내부 코드 view5): 521건 (고개 및 안면 랜드마크 비가시 영역 발생)\n"
        "• 시점 B (우측 전방 45도 / 내부 코드 view2): 521건 (전신 3차원 깊이 정보 양호)\n"
        "• 시점 C (정면 0도 / 내부 코드 view3): 521건 (좌우 대칭성 및 관상면 무릎 외반 최적 관측)\n"
        "• 시점 D (좌측 전방 45도 / 내부 코드 view4): 521건 (전신 3차원 깊이 정보 양호)\n"
        "• 시점 E (좌측 후방 45도 / 내부 코드 view1): 521건 (고개 및 안면 랜드마크 비가시 영역 발생)\n"
        "※ 5대 시점 전수 균등 포함 (각 20% 점유)하여 특정 앵글에 대한 과적합을 방지함.\n\n"
        "■ 3) 64프레임 3차 스플라인 표준 전처리 (64-Frame Cubic Spline Normalization):\n"
        "• 입력 영상 파이프라인: 원본 연속 영상 -> 단일 횟수(1 Repetition) 시작-최하사점-종료 구간 엄밀 분할 -> 64프레임으로 3차 스플라인(Cubic Spline) 보간.\n"
        "• 기본 입력 텐서 규격: shape: (64, 33, 3) (MediaPipe Pose 33개 랜드마크의 3D 공간 좌표 X, Y, Z).\n"
        "• 보조 텐서 규격 (c1 척추 모델 전용): shape: (64, 9, 3) (SpinePose 9대 분절: 경추 3점, 흉추 3점, 요추 3점)."
    )
    p_pipe.paragraph_format.line_spacing = 1.2

    doc.add_paragraph()

    # 3. 4대 최신 세부 결함 모델 세부 명세
    doc.add_heading("3. 4대 최신 챔피언 세부 결함 모델 세부 명세 (Layer 1 Submodels)", level=1)
    t_sub = doc.add_table(rows=5, cols=5)
    t_sub.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["결함 코드", "공식 명칭", "최신 버전", "피처 및 모델 아키텍처", "현재 F1 / Acc / AUC"]
    for j, h in enumerate(s_headers):
        cell = t_sub.cell(0, j)
        cell.text = h
        set_cell_background(cell, "2B6CB0")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "맑은 고딕"
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    s_rows = [
        ("c1", "척추 중립 위반\n(요추 과전만/후만)", "v11.1 Tri-Ensemble", "• 38D LightGBM (체간 기저계 투영)\n• 27D DietGBM (경량 도메인 피처)\n• 24D Dilated TCN (시계열 텐서)\n※ 0.10:0.40:0.50 가중 앙상블", "F1: 95.76%\nAcc: 95.89%\nAUC: 0.9882"),
        ("c2", "시선 및 고개 정면\n(경추 과신전/굴곡)", "v7.1 시점적응형", "• 19D CatBoost + LightGBM\n• 코-귀 관통 벡터, 경추 각도\n• 후방(시점 A, E) 시선 사각지대 판정 보류/면제 규칙 내장", "F1: 83.43%\nAcc: 82.76%\nAUC: 0.9124"),
        ("c3", "무릎-발 방향 일치\n(무릎 외반/내측붕괴)", "v6.0 3D 관상면", "• 22D CatBoost\n• 최하사점 절대 무릎/발목 비율(KA Ratio)\n• 3D 대퇴-경골 관상각, 비대칭 외반 피크", "F1: 94.04%\nAcc: 94.17%\nAUC: 0.9781"),
        ("c4", "발바닥 완전 접지\n(뒤꿈치/앞꿈치 들림)", "v5.1 동역학 위상", "• 25D CatBoost + Dilated TCN\n• 앞꿈치/뒤꿈치 Y변위, 착지 면적\n• 떨림(Wobble), 저크(Jerk), FFT 0차(F0 Rate)", "F1: 87.66%\nAcc: 87.72%\nAUC: 0.9415")
    ]
    for i, row in enumerate(s_rows):
        for j, val in enumerate(row):
            cell = t_sub.cell(i+1, j)
            cell.text = val
            set_cell_margins(cell, 80, 80, 80, 80)
            if i % 2 == 1:
                set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2, 4] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "맑은 고딕"
                r.font.size = Pt(8.5)

    doc.add_paragraph()

    # 4. 상위 통합 진단 마스터 엔진 v9.0 아키텍처
    doc.add_heading("4. 상위 통합 진단 마스터 엔진 v9.0 아키텍처 (Layer 1.5 ~ Layer 4)", level=1)
    p_arch = doc.add_paragraph(
        "■ Layer 1.5: 5대 시점 적응형 온도 스케일링 (Temperature Scaling)\n"
        "• 목적: 트리 기반 모델 및 신경망의 출력 확률 과확신(Over-confidence) 및 불확실성 왜곡 교정.\n"
        "• 수식: p_cal = 1 / (1 + exp(-logit / T))\n"
        "• 적용 파라미터: T_c1 = 0.456, T_c3 = 1.112, T_c4 = 1.001\n"
        "• c2 모델 시점별 적응형 온도: 시점 A(0.893), 시점 B(0.898), 시점 C(0.963), 시점 D(0.890), 시점 E(0.903)\n\n"
        "■ Layer 3: 도메인 보정 피처 추출기 (해결책 4 & 해결책 5)\n"
        "• [해결책 4] 12D 관측성 품질 게이트 (Quality Gate):\n"
        "  - 체간 길이 상대표준편차(std/mean), 대퇴골 상대표준편차, 경골 상대표준편차 (신체 분절 뼈대 길이 안정성)\n"
        "  - 관절 중심 떨림(Wobble), 급격한 가속도 변화(Jerk), 3D Z축 깊이 분산, 관절 가려짐(Occlusion) 비율 등 12개 지표.\n"
        "  - 종합 품질 점수 Q_overall 산출 및 이상 샘플 필터링.\n"
        "• [해결책 5] 9D 3단계 위상 운동사슬 피처 (Kinetic Chain Interaction):\n"
        "  - 스쿼트 전 구간을 하강기(Eccentric), 최하사점(Nadir), 상승기(Concentric)의 3위상으로 분할.\n"
        "  - 위상별 [발바닥-무릎], [발바닥-체간], [무릎-체간] 분절 간 각속도 상관계수 및 위상차 9차원 벡터.\n"
        "  - 발목의 흔들림이 무릎 외반과 요추 굴곡으로 이어지는 보상 작용(Compensatory movement) 포착.\n\n"
        "■ Layer 2: 39D View-Conditioned Mixture-of-Experts (MoE) 동적 라우팅 (해결책 6)\n"
        "• 입력 피처 39D 결합 벡터:\n"
        "  - [0~4] 5대 시점 One-Hot 인코딩 (5D)\n"
        "  - [5~8] Layer 1.5 교정 결함 확률 [p_c1, p_c2, p_c3, p_c4] (4D)\n"
        "  - [9~12] 모델별 확신도 마진 및 엔트로피 (4D)\n"
        "  - [13~17] c4 발동역학 핵심 피처 (Wobble, Y-Delta, Toe-Lift, Jerk, F0-Rate) (5D)\n"
        "  - [18~29] Layer 3 12D 관측성 품질 게이트 피처 (12D)\n"
        "  - [30~38] Layer 3 9D 3단계 위상 운동사슬 피처 (9D)\n"
        "• 3대 전문가(Sub-Experts) 구조:\n"
        "  - Expert 1 (Front/Oblique Specialist, Linear(39->48->16)): 정면 및 45도 전방 시점 판정 특화\n"
        "  - Expert 2 (Rear Specialist, Linear(39->48->16)): 후방 45도 시점 및 시선 사각지대 판정 특화\n"
        "  - Expert 3 (Kinetic Chain Specialist, Linear(39->48->16)): 발-무릎-체간 다관절 위상 협응 특화\n"
        "• Gating Network: 39D -> 48 -> 3 (Softmax 출력 g1, g2, g3)\n"
        "• 출력: 16-상태 결합 확률 분포 P_16 = g1*E1(x) + g2*E2(x) + g3*E3(x) (16-Way: 2^4 결함 조합)\n"
        "• 로드 밸런싱 손실 (Load-Balancing Loss): L_balance = 3 * sum(P_i * f_i) 적용으로 게이트 붕괴 원천 차단 (실제 부하 분산: 30.5% : 34.0% : 35.5% 달성)\n\n"
        "■ Layer 4: MoE 16-상태 결합 비대칭 선택적 판정 (Quality-Gated Selective Prediction)\n"
        "• 정상 확률 정의: P(Clean) = P(State 0) = P((0, 0, 0, 0))\n"
        "• 데드밴드 판정: P(Clean) in [0.32, 0.38] 또는 Top-1과 Top-2 결합 상태 간 마진 ΔP < 0.05인 경우, 억지 추론을 배제하고 안전하게 재촬영 권고(RETAKE)로 분류.\n"
        "• 결과: EMR 71.24% 달성 및 Clean Recall 77.05% 확보, RETAKE율 3.19%로 극소화."
    )
    p_arch.paragraph_format.line_spacing = 1.2

    doc.add_paragraph()

    # 5. 7단계 마일스톤 진화 궤적
    doc.add_heading("5. 7단계 마일스톤 진화 궤적 (Evolution Milestones)", level=1)
    t_m = doc.add_table(rows=8, cols=6)
    t_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_headers = ["버전", "핵심 도입 기술", "4-Defect EMR", "Clean F1", "Clean Recall", "RETAKE율"]
    for j, h in enumerate(m_headers):
        cell = t_m.cell(0, j)
        cell.text = h
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "맑은 고딕"
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    m_rows = [
        ("v6.0", "단순 독립 4모델 Argmax 결합", "15.90%", "34.43%", "45.00%", "0.00%"),
        ("v8.0", "온도 보정 + 후방 적응형 휴리스틱", "48.02%", "55.40%", "58.12%", "0.00%"),
        ("v8.2", "16-상태 결합 신경망 (18D Joint Meta)", "52.17%", "58.73%", "62.45%", "0.00%"),
        ("v8.3", "비대칭 비용 데드밴드 선택적 판정", "56.40%", "62.45%", "68.20%", "4.85%"),
        ("v8.4", "12D 관측성 품질 게이트 (해결책 4)", "59.94%", "65.13%", "72.04%", "3.80%"),
        ("v8.5", "3단계 위상 운동사슬 39D (해결책 5)", "66.08%", "66.43%", "74.80%", "3.55%"),
        ("v9.0 Master", "39D View-Conditioned MoE 동적 라우팅", "71.24% (전수 68.98%)", "67.63%", "77.05%", "3.19%")
    ]
    for i, row in enumerate(m_rows):
        for j, val in enumerate(row):
            cell = t_m.cell(i+1, j)
            cell.text = val
            set_cell_margins(cell, 80, 80, 80, 80)
            if i % 2 == 1:
                set_cell_background(cell, "F7FAFC")
            if i == 6:
                set_cell_background(cell, "EBF8FF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2, 3, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "맑은 고딕"
                r.font.size = Pt(8.5)
                if i == 6:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    doc.add_paragraph()

    # 6. 현재 남아있는 기술적 병목 및 문제점 분석
    doc.add_heading("6. 현재 남아있는 기술적 병목 및 문제점 분석 (Current Bottlenecks)", level=1)
    p_bot = doc.add_paragraph(
        "■ 1) 정상 판정의 정밀도(Precision 60.27%) 대비 재현율(Recall 77.05%) 불균형:\n"
        "• 재현율은 77%에 도달했으나, 1개 이상의 미세 결함이 있는 샘플을 '정상'으로 오분류하는 위양성(False Positive)이 여전히 존재하여 Precision이 60%에 머무름.\n"
        "• 원인: 데이터셋 내 '전체 무결점 정상' 샘플의 희소성(약 10~15%) 및 4개 부위 중 1개 부위만 미세하게 결함인 경계 샘플의 모호성.\n\n"
        "■ 2) c2 (시선 및 고개) 모델의 성능 한계 (F1 83.43%):\n"
        "• c1(95.76%), c3(94.04%), c4(87.66%)에 비해 c2가 가장 낮은 F1(83.43%)을 기록하며 전체 4-Defect EMR의 주요 발목을 잡고 있음.\n"
        "• 후방 45도(시점 A, E)에서 얼굴 전면이 가려지므로 척추-경추 각도만으로 추정해야 하는 본질적 가시성 한계 존재.\n\n"
        "■ 3) 16-상태 결합 공간에서의 희귀 복합 결합(Rare Multi-Defect) 불균형:\n"
        "• 16개 결합 상태 중 단일 결함이나 2개 결함은 샘플이 충분하나, 3개 결함 동시 발생((1,1,1,0) 등)이나 4개 결함 전수 발생((1,1,1,1)) 등 희귀 상태는 학습 표본이 극소수임.\n\n"
        "■ 4) 뼈대 길이 고정(Skeletal Length Invariance) 전처리의 부재:\n"
        "• MediaPipe 3D 좌표는 카메라 원근감과 피험자 체형에 따라 프레임마다 뼈대 길이(대퇴골, 경골, 체간)가 미세하게 변동함.\n"
        "• 현재는 이를 Layer 3에서 12D 상대표준편차로 검출만 하고 있을 뿐, 좌표 자체를 고정된 스켈레톤 길이로 리타겟팅(Normalization)하지는 않고 있음."
    )
    p_bot.paragraph_format.line_spacing = 1.2

    doc.add_paragraph()

    # 7. 외부 AI 자문 요청 질문서
    doc.add_heading("7. 외부 AI 자문 요청 질문서 (Specific Inquiries for External AI)", level=1)
    p_inq = doc.add_paragraph(
        "외부 AI 연구원께서는 위의 배경과 기술 명세를 바탕으로, 다음 6대 질문에 대해 [구체적인 수학적 원리, 네트워크 아키텍처, 손실 함수 수식, 파이썬 구현 가이드]를 포함하여 실질적인 해결책을 제안해 주십시오.\n\n"
        "■ [질문 1: Clean Precision 70%+ 돌파를 위한 손실 함수 및 분류 전략]\n"
        "• 정상(Clean) 클래스의 정밀도(Precision)를 현재 60.27%에서 70% 이상으로 끌어올리고 FP를 억제하기 위한 최신 비용 민감 학습(Cost-Sensitive / Asymmetric Focal Loss) 또는 정상 상태 중심 Contrastive Learning(대조 학습) 설계 방안은 무엇인가?\n\n"
        "■ [질문 2: c2 시선/고개 모델의 후방 뷰 극복을 위한 경추-흉추 운동학 연동]\n"
        "• 후방 45도 시점에서 얼굴 특징점 결측을 보상하기 위해, 흉추(Thoracic Spine)와 견갑대(Scapula)의 3D 회전 및 각속도로부터 경추(Cervical) 각도를 간접 추정하는 생체역학적 커플링(Biomechanical Coupling) 모델링 기법은 무엇인가?\n\n"
        "■ [질문 3: 4-Defect EMR 80%+ 달성을 위한 16-상태 결합 모델 고도화]\n"
        "• 현재의 MoE(3개 전문가) 구조를 넘어, 4대 결함 간의 조건부 의존성(Conditional Dependency: 예, c4 발목 불안정이 c3 무릎 외반을 유발하고 c1 요추 보상으로 이어지는 인과 그래프)을 직접 모델링하는 Graph Neural Network (GNN) 또는 Directed Acyclic Graph (DAG) 베이지안 신경망 적용 방안은?\n\n"
        "■ [질문 4: 뼈대 길이 고정(Skeletal Length Normalization) 알고리즘]\n"
        "• 64프레임 (64, 33, 3) 3D 좌표에 대해, 피험자 기립 시의 안정된 세그먼트 길이를 기준으로 프레임별 관절 방향 벡터는 유지하면서 뼈 길이를 완벽히 일정하게 고정하는 스켈레톤 정규화(Kinematic Retargeting) 수식 및 파이프라인 삽입 위치(서브모델 입력 전 vs 통합 피처 추출 전)는?\n\n"
        "■ [질문 5: 시공간 통합 신경망 (Spatio-Temporal Transformer / ST-GCN)]\n"
        "• 현재 수작업 피처(Hand-crafted 39D) + MoE 방식과 원천 시퀀스 (64, 33, 3)를 직접 입력받는 Spatial-Temporal Graph Convolutional Network (ST-GCN) 또는 Pose-Transformer의 하이브리드 결합 가능성 및 최적 융합(Late Fusion vs Intermediate Feature Fusion) 전략은?\n\n"
        "■ [질문 6: 4대 세부 모델의 신뢰도 보정(Calibration) 고도화]\n"
        "• 현재 스칼라 온도 스케일링(Temperature Scaling)을 넘어, 피처 벡터(39D)의 공간적 위치에 따라 온도가 국소적으로 변하는 가우시안 프로세스 온도 스케일링(GP-Temperature) 또는 매트릭스 스케일링(Matrix Scaling) 도입이 판정 불확실성 감소에 기여할 수 있는가?"
    )
    p_inq.paragraph_format.line_spacing = 1.2

    doc.save(target_path)
    print(f"DOCX successfully saved: {target_path}")

if __name__ == "__main__":
    OUT_DOCX_1 = r"D:\자주프 문서\vidio-health-trainer_외부AI_자문용_바벨스쿼트_통합진단_배경지식_및_아키텍처_명세서.docx"
    OUT_DOCX_2 = r"D:\자주프 문서\5주차\vidio-health-trainer_외부AI_자문용_바벨스쿼트_통합진단_배경지식_및_아키텍처_명세서.docx"

    create_docx(OUT_DOCX_1)
    os.makedirs(os.path.dirname(OUT_DOCX_2), exist_ok=True)
    shutil.copyfile(OUT_DOCX_1, OUT_DOCX_2)
    print(f"Copied to: {OUT_DOCX_2}")
