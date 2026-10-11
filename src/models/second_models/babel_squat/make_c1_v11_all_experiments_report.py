# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_c1_v11_all_experiments_report.py
Purpose: c1(척추 중립 위반) v11 전 실험(E0, E1, E2) 및 챔피언 모델 확정 정식 Word 보고서(.docx) 생성
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DOCX_PATH = r"D:\자주프 문서\5주차\vidio-health-trainer_c1_v11_전체실험_E0_E1_E2_최종검증보고서.docx"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_callout(doc, text_p_list, title="핵심 연구 성과 요약 (Executive Summary)", border_hex="1B365D", bg_hex="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    run_t = p0.add_run(f"■ {title}")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    for item in text_p_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(item)
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    title_p = doc.add_paragraph()
    t_run = title_p.add_run("vidio-health-trainer")
    t_run.font.name = "맑은 고딕"
    t_run.font.size = Pt(12)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    sub_p = doc.add_paragraph()
    s_run = sub_p.add_run("바벨 스쿼트 c1 (척추 중립 위반) v11 전 실험(E0·E1·E2) 실측 및 최종 챔피언 확정 보고서")
    s_run.font.name = "맑은 고딕"
    s_run.font.size = Pt(18)
    s_run.font.bold = True
    s_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    m_run = meta_p.add_run("작성일: 2026-09-26 | 대상 결함: c1 척추의 중립 위반 | 평가 데이터: 마스터 2,605건 전수 (5대 시점 x 9명 피험자 격리)")
    m_run.font.name = "맑은 고딕"
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    add_callout(
        doc,
        [
            "외부 AI가 제안한 3대 핵심 실험(E0, E1, E2)을 2,605개 마스터 전수 대상 5-Fold 피험자 완전 격리로 전수 실측 완수.",
            "실험 E0 (Nadir 국소 윈도우 피처): 하강 최하사점 신호 희석을 해결하여 정면(View 3) F1을 74.30%에서 82.57%로 +8.27%p 폭등시킴.",
            "실험 E1 (경량 Dilated TCN + GBDT Fusion): 64F 시계열 텐서를 직접 읽는 TCN과 GBDT의 최적 Late Fusion(TCN 55% + GBDT 45%)을 통해 F1 84.24%(+4.51%p 대도약)를 달성하고, 5대 시점 전수 83% 이상으로 폭등시킴.",
            "실험 E2 (View-Conditioned Normalization 감사): 시점별 피처 정규화 시 절대 회전각 신호 파괴로 F1 79.20%로 열화됨을 확인하여, 외부 AI 권고 기준에 따라 즉시 채택 중단(Aborted) 확정.",
            "최종 챔피언 확정: 전 시점 무결성과 최고 일반화 성능을 입증한 c1_v11_b (E1 TCN-GBDT Fusion, F1 84.24%)를 공식 프로덕션 모델로 배포 완료."
        ]
    )

    # 1. 3대 실험 요약 및 설계
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. 외부 AI 제안 기반 c1 3대 실험 로드맵 개요")
    r1.font.name = "맑은 고딕"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl_intro = doc.add_table(rows=4, cols=4)
    tbl_intro.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_intro.autofit = False
    col_w_i = [Inches(1.0), Inches(1.3), Inches(3.0), Inches(1.2)]
    headers_i = ["실험 코드", "모델 버전", "핵심 기술 메커니즘", "최종 실측 판정"]
    rows_i = [
        ["E0", "c1_v11_a_diet", "Nadir ±8F 국소 윈도우 피처 10D 추출 및 다중공선성 제거 (27D)", "성공 (F1 80.52%, 정면 +8.27%p)"],
        ["E1", "c1_v11_b", "64F 관절 시계열 텐서(64, 24) 직접 입력 경량 TCN + GBDT Fusion", "★ 대성공 (F1 84.24%, 5대시점 전수 83%+)"],
        ["E2", "c1_v11_c", "View-Conditioned Normalization (시점별 피처 평균/분산 강제 정규화)", "중단 (F1 79.20% 열화, 기각)"]
    ]

    for c_idx, h_text in enumerate(headers_i):
        cell = tbl_intro.cell(0, c_idx)
        cell.width = col_w_i[c_idx]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows_i, 1):
        bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = tbl_intro.cell(r_idx, c_idx)
            cell.width = col_w_i[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if "★ 대성공" in val:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x66, 0x00)
            elif "중단" in val:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # 2. 전수 실측 성능 직접 대조표
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("2. 모델 세대별 실측 성능 종합 대조 (N=2,605 전수, 5-Fold 피험자 완전 격리)")
    r2.font.name = "맑은 고딕"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl_perf = doc.add_table(rows=6, cols=6)
    tbl_perf.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_perf.autofit = False
    col_w_p = [Inches(1.5), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)]
    headers_p = ["지표 (Metric)", "v10 (28D 베이스)", "E0 (v11_a_diet)", "E1 TCN 단독", "E1 Fusion (v11_b)", "E2 View-Norm"]
    rows_p = [
        ["F1-Score (결함)", "79.73%", "80.52%", "82.04%", "84.24%", "79.20%"],
        ["ROC-AUC", "88.33%", "88.13%", "89.97%", "91.59%", "87.45%"],
        ["Precision (정밀도)", "80.00%", "81.88%", "79.13%", "84.81%", "80.00%"],
        ["Recall (재현율)", "79.46%", "79.20%", "85.18%", "83.67%", "78.41%"],
        ["Accuracy (정확도)", "80.38%", "81.54%", "82.03%", "84.91%", "80.15%"]
    ]

    for c_idx, h_text in enumerate(headers_p):
        cell = tbl_perf.cell(0, c_idx)
        cell.width = col_w_p[c_idx]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows_p, 1):
        bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = tbl_perf.cell(r_idx, c_idx)
            cell.width = col_w_p[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if c_idx == 4:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x44, 0x88)

    # 3. 5대 시점별 세부 감사
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("3. 5대 시점별 세부 성능 감사 (Multi-View Performance Audit)")
    r3.font.name = "맑은 고딕"
    r3.font.size = Pt(13)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl_v = doc.add_table(rows=6, cols=6)
    tbl_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_v.autofit = False
    col_w_v = [Inches(1.7), Inches(0.9), Inches(1.0), Inches(1.0), Inches(1.0), Inches(0.9)]
    headers_v = ["시점 (카메라 앵글)", "v10 F1", "E0 (diet) F1", "E1 TCN F1", "★ E1 Fusion", "E2 View-Norm"]
    rows_v = [
        ["view1 (완전 후방 180°)", "76.12%", "76.11%", "82.51%", "83.27%", "74.49%"],
        ["view2 (우측 45도 측면)", "83.21%", "82.89%", "80.52%", "83.69%", "80.32%"],
        ["view3 (완전 정면 0°)", "74.30%", "82.57%", "81.19%", "83.40%", "82.16%"],
        ["view4 (좌측 45도 측면)", "81.85%", "81.40%", "81.13%", "83.37%", "81.04%"],
        ["view5 (우측 후방 45도)", "83.18%", "79.67%", "84.85%", "87.43%", "77.87%"]
    ]

    for c_idx, h_text in enumerate(headers_v):
        cell = tbl_v.cell(0, c_idx)
        cell.width = col_w_v[c_idx]
        set_cell_background(cell, "2C3E50")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows_v, 1):
        bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = tbl_v.cell(r_idx, c_idx)
            cell.width = col_w_v[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if c_idx == 4:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # 4. 결론
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r4 = h4.add_run("4. 최종 판정 및 프로덕션 배포 결론")
    r4.font.name = "맑은 고딕"
    r4.font.size = Pt(13)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.line_spacing = 1.25
    p_end.add_run(
        "1. [E0 검증 성공]: Nadir 국소 윈도우 피처는 정면 시점의 극심한 버트윙크 희석 문제를 완벽히 해결하여 F1 82.57%를 기록했습니다.\n"
        "2. [E1 챔피언 등극]: 경량 Dilated TCN(시계열 24D)과 GBDT(생체역학 27D)의 Late Fusion 앙상블은 상호보완적 시너지를 발휘하여 "
        "전체 F1 84.24%를 달성했으며, 과거 가장 취약했던 후방(83.27%)과 정면(83.40%)을 포함해 5대 시점 전수에서 83% 이상을 고르게 확보했습니다.\n"
        "3. [E2 중단 확정]: View-Conditioned Normalization은 시점 간 절대 자세 각도 신호를 훼손하여 F1 79.20%로 열화되었으므로 채택을 공식 중단했습니다.\n"
        "4. [최종 배포]: vidio-health-trainer c1의 공식 프로덕션 모델로 c1_v11_b (TCN 55% + GBDT 45% Fusion)를 최종 확정 배포합니다."
    )

    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)
    doc.save(OUT_DOCX_PATH)
    print(f"[*] 정식 종합 Word 보고서 생성 완료: {OUT_DOCX_PATH}")


if __name__ == '__main__':
    main()
