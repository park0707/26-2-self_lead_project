# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_v9_vs_v10_comparison_report_docx.py
Purpose: v9.0 vs v10.0 해결책 1 공식 성능 비교 및 검증 DOCX 보고서 생성
Target: D:\자주프 문서\5주차\vidio-health-trainer_v9_vs_v10_해결책1_성능비교_및_검증보고서.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="1A365D", bg_color="F7FAFC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.9)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}\n")
    run_t.bold = True
    run_t.font.name = '맑은 고딕'
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = '맑은 고딕'
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def style_heading(p, text, level=1):
    p.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 6))
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = '맑은 고딕'
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def format_table(table, col_widths, headers, data, header_bg="1A365D", alt_bg="F7FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_row = table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(heading)
        run.bold = True
        run.font.name = '맑은 고딕'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_col = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = '맑은 고딕'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            
            if c_idx in [1, 2, 3, 4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E0"/>
        <w:left w:val="none"/>
        <w:bottom w:val="single" w:sz="12" w:space="0" w:color="1A365D"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def create_report(output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    style_normal = doc.styles['Normal']
    style_normal.font.name = '맑은 고딕'
    style_normal.font.size = Pt(9.5)
    style_normal.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("vidio-health-trainer\nv9.0 vs v10.0 해결책 1 (False-Clean Asymmetric Loss) 공식 성능 비교 및 검증 보고서")
    r_title.bold = True
    r_title.font.name = '맑은 고딕'
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Benchmark & Multi-View Audit of False-Clean Reverse-ASL: Elevating Clean Precision to 75.41% (+7.74%p)")
    r_sub.font.name = '맑은 고딕'
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_meta = p_meta.add_run("작성일자: 2026. 09. 27 | 시스템 버전: v10.0 Master Production | 프로젝트 공식 명칭: vidio-health-trainer")
    r_meta.font.name = '맑은 고딕'
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph()

    # 1. 개요 및 실험 목적
    style_heading(doc.add_paragraph(), "1. 실험 배경 및 해결책 1 도입 목적", level=1)
    doc.add_paragraph(
        "본 보고서는 외부 AI 6대 기술 자문 중 가장 파급력이 크고 즉각 도입 가능한 [해결책 1: False-Clean 비대칭 손실(Reverse-ASL) 및 안전 임계값 제어]을 vidio-health-trainer v9.0 마스터 베이스라인에 정밀 적용하고, 기존 v9.0 모델과의 전수 100% 강제 판정 및 선택적 판정(Selective Prediction) 성능을 전수 대조 평가한 결과를 기술합니다."
    )

    add_callout(
        doc,
        "해결책 1의 핵심 메커니즘 및 도입 가치",
        "1. Reverse-ASL 손실 함수: 결함 샘플(2,470건)이 Clean으로 오판되는 치명적 임상 오류(False Clean)에 대해 p^2 및 lambda_fc=3.0의 비대칭 지수적 패널티를 부과하여 모델의 경계선 과신을 원천 억제함.\n"
        "2. 완전한 베이스라인 계승: 기존 c1~c4 GBDT 서브모델 챔피언 및 39D 피처셋(품질 12D + 운동사슬 9D + 시점 5D + 물리 5D + 보정확률 8D)을 100% 보존하면서 MoE 메타 신경망 손실만 교체함.\n"
        "3. 정밀 제약 임계값: Clean Precision 70%+ 돌파를 목표로 타깃 임계값을 제어하여 Recall 손실 없이 Precision을 극대화함.",
        border_color="2B6CB0", bg_color="EBF8FF"
    )

    # 2. 종합 비교 결과
    style_heading(doc.add_paragraph(), "2. v9.0 vs v10.0 종합 성능 비교 (2,605건 전수)", level=1)

    t_all = doc.add_table(rows=5, cols=7)
    headers_all = ["평가 모드 / 버전", "Clean F1", "Precision", "Recall", "4-Defect EMR", "False Clean", "RETAKE율"]
    widths_all = [Inches(2.1), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.9), Inches(0.8), Inches(0.7)]
    data_all = [
        ["v9.0 베이스라인 (전수 강제판정)", "67.16%", "67.67%", "66.67%", "69.40%", "43건", "0.00%"],
        ["v10.0 해결책 1 (전수 강제판정)", "71.60%", "75.41%", "68.15%", "69.60%", "30건", "0.00%"],
        ["v9.0 베이스라인 (선택적 판정)", "70.12%", "68.22%", "72.13%", "71.71%", "-", "9.90%"],
        ["v10.0 해결책 1 (선택적 판정)", "74.29%", "75.83%", "72.80%", "71.78%", "-", "9.94%"]
    ]
    format_table(t_all, widths_all, headers_all, data_all, header_bg="1A365D")
    doc.add_paragraph()

    # 3. 핵심 성과 분석
    style_heading(doc.add_paragraph(), "3. 핵심 지표 개선 성과 및 정밀 분석", level=1)
    doc.add_paragraph(
        "■ Clean Precision 75.41% 달성 (+7.74%p 상승, 목표 70%+ 돌파):\n"
        "  - 전수 100% 강제 판정 기준 기존 67.67%에서 75.41%로 +7.74%p 대폭 상승하였습니다.\n"
        "  - 선택적 판정 기준에서도 기존 68.22%에서 75.83%로 +7.61%p 상승하여 완벽하게 75% 벽을 돌파하였습니다.\n\n"
        "■ False Clean 건수 43건 → 30건 (-13건, -30.2% 감소):\n"
        "  - 실제 결함이 존재하는 운동인데 시스템이 정상이라고 잘못 판정한 치명적 오분류가 43건에서 30건으로 30.2% 대폭 축소되었습니다.\n\n"
        "■ Clean F1 71.60% 달성 (+4.43%p 상승) 및 Recall 동반 상승 (+1.48%p):\n"
        "  - 통상 Precision을 올리면 Recall이 급락하지만, v10.0에서는 Reverse-ASL이 경계선 샘플을 정교하게 정돈하면서 Recall도 66.67%에서 68.15%로 동반 상승하였습니다.\n\n"
        "■ RETAKE율 9.94% 완벽 제어 (v9.0의 9.90% 수준 유지):\n"
        "  - 데드밴드를 [0.32, 0.38]로 정밀 일치시켜 재촬영 요구 비율을 9.94%로 유지하면서 Coverage 90.1%를 안정적으로 확보하였습니다."
    )

    # 4. 5대 카메라 시점별 전수 대조 감사
    style_heading(doc.add_paragraph(), "4. 5대 카메라 시점별 세부 Clean F1 및 정밀도 감사 (전수 강제판정)", level=1)

    t_view = doc.add_table(rows=6, cols=6)
    headers_view = ["카메라 시점 명칭", "v9.0 F1", "v10.0 F1", "Δ F1 (향상도)", "v10.0 Precision", "False Clean 감소"]
    widths_view = [Inches(1.8), Inches(0.9), Inches(0.9), Inches(1.1), Inches(1.1), Inches(1.1)]
    data_view = [
        ["시점 A (우측 후방 45도 / view5)", "51.85%", "61.54%", "+9.69%p", "64.00% (+12.15%p)", "13건 → 9건 (-4건)"],
        ["시점 B (우측 전방 45도 / view2)", "71.70%", "77.78%", "+6.08%p", "77.78% (+4.70%p)", "7건 → 6건 (-1건)"],
        ["시점 C (정면 0도 / view3)",       "81.97%", "84.75%", "+2.78%p", "78.12% (+4.60%p)", "9건 → 7건 (-2건)"],
        ["시점 D (좌측 전방 45도 / view4)", "73.08%", "73.47%", "+0.39%p", "81.82% (+5.82%p)", "6건 → 4건 (-2건)"],
        ["시점 E (좌측 후방 45도 / view1)", "54.17%", "55.81%", "+1.65%p", "75.00% (+13.10%p)", "8건 → 4건 (-4건)"]
    ]
    format_table(t_view, widths_view, headers_view, data_view, header_bg="2B6CB0")
    doc.add_paragraph()

    add_callout(
        doc,
        "5대 시점 전수 성능 감사 합격 (시점 편향 제로 검증)",
        "공식 규정(Rule 2, Rule 4)에 따라 5대 전 시점의 성능을 감사한 결과, 5개 시점 모두에서 F1과 Precision이 단 1건의 퇴행(Regression) 없이 일제히 상승하였습니다.\n"
        "특히 취약 구간이던 후방 시점(시점 A F1 +9.69%p, 시점 E Precision +13.10%p)의 거짓 정상(False Clean)이 절반 가까이 제거되어 전 시점 안정성이 극대화되었습니다.",
        border_color="2B6CB0", bg_color="EBF8FF"
    )

    # 5. 4대 세부 결함 모델 무결성 감사
    style_heading(doc.add_paragraph(), "5. 4대 세부 결함(c1~c4) 모델 성능 무결성 검증", level=1)
    
    t_def = doc.add_table(rows=5, cols=5)
    headers_def = ["세부 결함 진단 모듈", "v9.0 F1", "v10.0 F1", "변화량 (Δ F1)", "검증 판정"]
    widths_def = [Inches(2.2), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.4)]
    data_def = [
        ["c1: 척추의 중립 위반", "95.76%", "95.64%", "-0.12%p", "안정 유지 (합격)"],
        ["c2: 고개 정면 및 시선 위반", "83.97%", "84.00%", "+0.02%p", "미세 향상 (합격)"],
        ["c3: 발과 무릎의 방향 일치 위반", "94.09%", "93.41%", "-0.68%p", "정상 범위 (합격)"],
        ["c4: 발바닥 지면 고정 위반", "87.85%", "87.70%", "-0.15%p", "안정 유지 (합격)"]
    ]
    format_table(t_def, widths_def, headers_def, data_def, header_bg="1A365D")
    doc.add_paragraph()

    # 6. 결론 및 향후 계획
    style_heading(doc.add_paragraph(), "6. 결론 및 향후 v10.0 고도화 계획", level=1)
    doc.add_paragraph(
        "1. 해결책 1(Reverse-ASL)의 완벽한 효과 실증:\n"
        "   - Clean Precision을 75.41%까지 끌어올리고 False Clean을 30.2% 제거함으로써, 외부 AI가 지목한 가장 큰 병목이 성공적으로 해소되었습니다.\n"
        "   - 통합 엔진 파일(babel_squat_integrated_v10_0.py), 메타 파라미터(joint_meta_16way_v10_0_meta.json), 모델 가중치(joint_meta_16way_v10_0_weights.pth)가 성공적으로 확정되었습니다.\n\n"
        "2. 다음 고도화 단계:\n"
        "   - Phase 1의 [해결책 4: Exact 16-State Energy Head(CRF)]와 [해결책 3: c2 불확실성 분산 헤드 및 흉곽 좌표계]를 연계하여 4-Defect EMR 75%+ 및 Clean F1 75%+를 달성할 예정입니다."
    )

    doc.save(output_path)
    print(f"Report successfully saved: {output_path}")

if __name__ == '__main__':
    target = r"D:\자주프 문서\5주차\vidio-health-trainer_v9_vs_v10_해결책1_성능비교_및_검증보고서.docx"
    create_report(target)
