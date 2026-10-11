# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_external_ai_briefing_v2.py
Purpose: 외부 AI 재조사 의뢰용 바벨 스쿼트 세부 오류 모델(c1~c4) 및 통합 정상/오류 분류 모델 상세 명세서 생성
         - PDF 및 DOCX 동시 생성
Destination:
  - D:\자주프 문서\외부_AI_자문용_바벨스쿼트_세부오류_및_통합분류_모델구조_명세서.pdf
  - D:\자주프 문서\외부_AI_자문용_바벨스쿼트_세부오류_및_통합분류_모델구조_명세서.docx
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

DEST_DIR = r"D:\자주프 문서"
PDF_PATH = os.path.join(DEST_DIR, "외부_AI_자문용_바벨스쿼트_세부오류_및_통합분류_모델구조_명세서.pdf")
DOCX_PATH = os.path.join(DEST_DIR, "외부_AI_자문용_바벨스쿼트_세부오류_및_통합분류_모델구조_명세서.docx")

# ---------------------------------------------------------------------------
# ReportLab Canvas (헤더 / 푸터 / 페이지 번호)
# ---------------------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("MalgunGothic", 8)
        self.setFillColor(colors.HexColor("#6B7280"))
        
        # 상단 헤더
        self.drawString(54, 800, "vidio-health-trainer | AI Biomechanics & System Architecture Technical Briefing")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 794, 541, 794)

        # 하단 푸터
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(541, 35, page_str)
        self.drawString(54, 35, "CONFIDENTIAL - FOR EXTERNAL AI RESEARCH & CONSULTATION ONLY")
        self.line(54, 45, 541, 45)
        self.restoreState()


# ---------------------------------------------------------------------------
# 1. PDF 생성 함수
# ---------------------------------------------------------------------------
def generate_pdf():
    font_path = "C:/Windows/Fonts/malgun.ttf"
    font_bold_path = "C:/Windows/Fonts/malgunbd.ttf"
    pdfmetrics.registerFont(TTFont("MalgunGothic", font_path))
    pdfmetrics.registerFont(TTFont("MalgunGothic-Bold", font_bold_path))

    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        fontName='MalgunGothic-Bold',
        fontSize=17,
        leading=22,
        textColor=colors.HexColor("#0F2D59"),
        spaceAfter=4
    )
    sub_title_style = ParagraphStyle(
        'DocSubTitle',
        fontName='MalgunGothic-Bold',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#2B6CB0"),
        spaceAfter=6
    )
    desc_box_style = ParagraphStyle(
        'DescBox',
        fontName='MalgunGothic',
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#2D3748")
    )
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        fontName='MalgunGothic-Bold',
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#0F2D59"),
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'Heading2_Custom',
        fontName='MalgunGothic-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1A365D"),
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        fontName='MalgunGothic',
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#374151"),
        spaceAfter=5
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        fontName='MalgunGothic',
        fontSize=7.5,
        leading=10.5,
        textColor=colors.HexColor("#1F2937")
    )
    table_header_style = ParagraphStyle(
        'TableHeader',
        fontName='MalgunGothic-Bold',
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.white
    )
    q_title_style = ParagraphStyle(
        'QTitle',
        fontName='MalgunGothic-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#0056B3"),
        spaceAfter=2
    )

    story = []

    # 헤더
    story.append(Paragraph("vidio-health-trainer | AI Biomechanics & System Architecture Technical Briefing", sub_title_style))
    story.append(Paragraph("바벨 스쿼트 세부 오류 모델(c1~c4) 및 통합 정상/오류 분류 모델 상세 명세서<br/>[외부 AI 재조사 의뢰용: 정상 F1 & 4개 오류 라벨 완전 일치율(EMR) 60%대 정체 돌파 자문서]", title_style))
    story.append(Spacer(1, 4))

    # 요약 박스
    summary_text = (
        "<b>[의뢰 배경 및 자문 목표]</b><br/>"
        "본 문서는 <b>vidio-health-trainer</b> 프로젝트에서 바벨 스쿼트 3차원 골격 시계열 데이터를 바탕으로 "
        "4대 세부 결함(c1 척추, c2 고개/시선, c3 무릎, c4 발바닥) 및 전신 정상/오류를 진단하는 통합 AI 엔진 구축 중 발생한 "
        "<b>'정상 F1-Score 및 4대 라벨 완전 일치율(EMR)의 60%대 정체 병목'</b>을 해결하기 위해 외부 AI에게 최적의 해결책을 자문받고자 작성되었습니다.<br/>"
        "외부 AI의 1차 자문 제안 7가지(소프트 리스크 퓨전, 관측 가능 정상 게이팅, 사선 왜곡 보정, 발목 롤 특이점 보정, 5대 시점 동적 임계치, 3-State 선택적 판정 등)를 "
        "성공적으로 적용하여 전반적인 성능이 대폭 향상되었으나, <b>(1) 정상 F1-Score (~68%)</b> 및 <b>(2) 4개 라벨 완전 일치율 EMR (~64%)</b>이 "
        "여전히 60%대에 머무르고 있습니다. 이를 80% 이상으로 극대화하기 위한 모델 구조, 손실 함수, 결합 기법의 재조사를 의뢰합니다."
    )
    summary_table = Table([[Paragraph(summary_text, desc_box_style)]], colWidths=[487])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F0F4F8")),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 7),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#1A365D")),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 6))

    # 섹션 1
    story.append(Paragraph("1. 데이터셋 환경 명세 및 극심한 클래스 불균형 현황", h1_style))
    p1 = (
        "• <b>데이터셋 출처 및 규모</b>: AI-Hub 피트니스 데이터셋 바벨 스쿼트 전수 N=2,605건 시퀀스<br/>"
        "• <b>피험자 전수 반영</b>: 피험자 1번~9번 전수 (9명) 반영 및 5-Fold StratifiedGroupKFold 피험자 완전 격리 검증<br/>"
        "• <b>공식 5대 시점 전수 반영</b>: 시점 A(우측 후방 45°), 시점 B(우측 전방 45°), 시점 C(정면 0°), 시점 D(좌측 전방 45°), 시점 E(좌측 후방 45°)<br/>"
        "• <b>전처리 표준</b>: 원천 영상 전 시퀀스 무삭제 64프레임 3차 스플라인(Cubic Spline) 보간 완료 (횟수 분할 파편화 100% 영구 배제)<br/>"
        "• <b>골든 정답 라벨 (Golden Ground Truth)</b>: AI-Hub 원천 라벨의 72.8% 노이즈를 전문가 57선 체크리스트로 정제한 무결성 라벨셋<br/>"
        "• <b>★ 극심한 정상/결함 클래스 불균형</b>:<br/>"
        "  - <b>전신 완전 정상 (Clean / c1~c4 모두 0)</b>: <b>135건 (5.18%)</b><br/>"
        "  - <b>전신 결함 존재 (Defect / c1~c4 중 1개 이상 1)</b>: <b>2,470건 (94.82%) [불균형 비율 1:18.3]</b><br/>"
        "  - 각 결함별 발생률: c1(척추) 46.96%, c2(시선) 53.73%, c3(무릎) 49.69%, c4(발바닥) 49.74% (개별 결함은 ~50% 균형)"
    )
    story.append(Paragraph(p1, body_style))

    # 섹션 2
    story.append(Paragraph("2. 4대 세부 오류 모델(c1~c4)의 구성, 입력, 피처셋 및 처리 파이프라인", h1_style))
    story.append(Paragraph(
        "모든 세부 오류 모델은 실제 웹캠/카메라 서비스 배포를 위해 <b>AI-Hub 3D JSON을 100% 완전 배제</b>하고, "
        "<b>순수 MediaPipe 33관절 64프레임 직접 계측 파이프라인</b>으로 구축되었습니다.", body_style
    ))

    t_sub_data = [
        [
            Paragraph("결함 코드", table_header_style),
            Paragraph("결함 명칭 및 물리적 의미", table_header_style),
            Paragraph("입력 관절 및 처리 파이프라인", table_header_style),
            Paragraph("추출 피처셋 명세", table_header_style),
            Paragraph("단독 모델 및 실측 성능", table_header_style)
        ],
        [
            Paragraph("<b>c1</b>", table_cell_style),
            Paragraph("<b>척추의 중립 위반</b><br/>(요추 굴곡/과신전,<br/>버트윙크, 굿모닝)", table_cell_style),
            Paragraph("MediaPipe 33관절 64F<br/>+ 체간 로컬 정규직교계<br/>+ Nadir ±8F 국소 윈도우", table_cell_style),
            Paragraph("<b>27D 정예(diet) 생체역학</b>:<br/>L1-S1 2차 곡률 계수 poly_a, L1-L5 Cobb 전만각, C1-T8 힌지각, 골반 대비 L5 후방말림, Nadir ±8F 요추-골반 상대각 국소 SD 등 16D + 11D MP 직접계측<br/>+ <b>24D 시계열 텐서(64x24)</b>", table_cell_style),
            Paragraph("<b>Tri-Ensemble</b>:<br/>CatBoost(10%) + LightGBM(40%) + Dilated TCN(50%)<br/><b>F1 80.52%</b><br/>Acc 81.54%<br/>AUC 88.94%", table_cell_style)
        ],
        [
            Paragraph("<b>c2</b>", table_cell_style),
            Paragraph("<b>고개 정면 미유지</b><br/>(경추 과굴곡/신전,<br/>시선 바닥/천장 이탈)", table_cell_style),
            Paragraph("MediaPipe 33관절 64F<br/>(코 #0, 양귀 #7/8,<br/>어깨 #11/12, 골반 #23/24)", table_cell_style),
            Paragraph("<b>19D 생체역학 피처셋</b>:<br/>체간 로컬 시상면 경추 굴곡각 8D + 단일 귀(Visible Ear) 적응형 안면 지향 피처 4D + 전/후방 법선 Z 성분 2D 등", table_cell_style),
            Paragraph("CatBoost(60%) + LightGBM(40%)<br/><b>이진 F1 76.53%</b><br/>선택적 F1 83.90%<br/>(Coverage 63.8%)", table_cell_style)
        ],
        [
            Paragraph("<b>c3</b>", table_cell_style),
            Paragraph("<b>발-무릎 방향 불일치</b><br/>(무릎 모임/외반 Valgus,<br/>좌우 비대칭 붕괴)", table_cell_style),
            Paragraph("MediaPipe 33관절 64F<br/>(무릎 #25/26, 발목 #27/28,<br/>골반 #23/24)", table_cell_style),
            Paragraph("<b>22D 생체역학 피처셋</b>:<br/>골반 관상면(Coronal) 투영, 절대 무릎-발목 너비 비율(K/A Ratio 6D), 좌우 독립 외반합/동시붕괴 6D, 대퇴-경골 3D 관상각 4D, 무릎-발끝 정렬 오차 4D 등", table_cell_style),
            Paragraph("CatBoost(70%) + LightGBM(30%)<br/><b>F1 91.45%</b> 🏆<br/>Acc 91.21%<br/>AUC 96.50%", table_cell_style)
        ],
        [
            Paragraph("<b>c4</b>", table_cell_style),
            Paragraph("<b>발바닥 지면 미고정</b><br/>(뒤꿈치/앞꿈치 들림,<br/>발목 롤 회내/회외)", table_cell_style),
            Paragraph("MediaPipe 33관절 64F<br/>(발목 #27/28, 뒤꿈치 #29/30,<br/>앞꿈치 #31/32)", table_cell_style),
            Paragraph("<b>25D 생체역학 피처셋</b>:<br/>최하사점 발바닥 경사각 Pitch, 수직 높이차 ydelta, 앞/뒤꿈치 들림량, 정강이-발목 커플링비, F0 접촉 상태 전이 빈도 등", table_cell_style),
            Paragraph("CatBoost(60%) + LightGBM(30%) + RF(10%)<br/><b>F1 79.46%</b><br/>Acc 79.42%<br/>AUC 87.73%", table_cell_style)
        ]
    ]

    t_sub = Table(t_sub_data, colWidths=[35, 95, 100, 162, 95])
    t_sub.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F2D59")),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (0, 2), (-1, 2), colors.white),
        ('BACKGROUND', (0, 3), (-1, 3), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (0, 4), (-1, 4), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_sub)
    story.append(Spacer(1, 6))

    # 섹션 3
    story.append(Paragraph("3. 바벨 스쿼트 통합 정상/오류 분류 모델의 아키텍처 및 처리 방식", h1_style))
    p3 = (
        "통합 정상/오류 분류 모델(통합 진단 엔진)은 4대 세부 모델의 예측 결과와 신체 전체의 동역학적 특성을 융합하여 "
        "<b>전신 완전 정상(Clean / State 0)</b>, <b>자세 결함(Defect / State 1)</b>, <b>재촬영 권고(Retake / State 2)</b>의 3-State 의사결정을 수행합니다.<br/>"
        "• <b>입력 데이터</b>: 4대 서브모델의 연속 사후확률 벡터 [p_c1, p_c2, p_c3, p_c4], 불확실성 마진 |p - 0.5|, 공식 5대 시점 원핫 벡터(view1~5), 상체/하체 상호작용 로짓<br/>"
        "• <b>지금까지 적용된 7대 개선 조치</b>:<br/>"
        "  1) <b>Hard-AND 탈피</b>: 독립 이진 판정 후 AND 연산(오차 누적으로 Clean F1 30%대 폭락)을 폐기하고 소프트 결합 도입<br/>"
        "  2) <b>Soft-Risk Fusion</b>: 결함 미발생 결합확률 P(Clean) = Π (1 - p_k) 도입<br/>"
        "  3) <b>Observable Clean (v7.1)</b>: 후방(시점 A, E)에서 얼굴 미관측 시 c2 불확실성을 감안하여 3대 결함(c1, c3, c4) 안전 시 정상 판정 구제<br/>"
        "  4) <b>Oblique Adaptive Soft Fusion (v7.3)</b>: 사선(시점 B, D) 45° 원근 왜곡 완화 가중치 적용<br/>"
        "  5) <b>Heel-Toe Roll Singularity Compensation (v7.4)</b>: 발목 회내/회외 특이점 보정<br/>"
        "  6) <b>5-View Dynamic Threshold Optimization (v7.5)</b>: 공식 5대 시점별(A~E) 독립 동적 임계치 최적화<br/>"
        "  7) <b>3-State Selective Classification</b>: 모호 구간 강제 판정을 지양하고 7%대 최소 보류율로 안전성 확보<br/>"
        "• <b>현재 달성된 성능과 한계</b>: 7가지 조치 적용으로 Clean F1이 30%대에서 <b>68.42%</b>로 대폭 상승했으나, "
        "<b>정상 F1 (~68%)</b>과 <b>4개 라벨 완전 일치율(Exact Match Ratio, EMR: ~64%)</b>이 여전히 60%대에 머물러 서비스 상용화 목표(80% 이상)에 미달함."
    )
    story.append(Paragraph(p3, body_style))

    # 섹션 4
    story.append(Paragraph("4. 왜 여전히 60%대에 머물고 있는가? 근본 원인 진단", h1_style))
    p4 = (
        "• <b>원인 1: 다중 레이블 오차 곱셈 법칙 (Multi-Label Error Multiplication)</b><br/>"
        "  개별 모델의 정확도가 80%~91%로 우수하더라도, 4개 라벨을 동시에 모두 정확히 맞출 확률(EMR)은 "
        "  독립 가정 하에 0.80 × 0.85 × 0.91 × 0.80 ≈ 49.5%에 불과합니다. 현재 64.2%는 독립 곱보다는 높으나, "
        "  라벨 간 상관관계(Label Dependency)를 명시적으로 최적화하는 결합 헤드가 부재하여 발생하는 구조적 한계입니다.<br/>"
        "• <b>원인 2: 극단적 정상 희소성(5.18%)에 기인한 오탐(False Positive) 치명타</b><br/>"
        "  전체 2,605개 샘플 중 완전 정상이 135개에 불과하므로, 2,470개의 결함 샘플 중 단 3%(~74개)만 '정상'으로 잘못 판정되어도 "
        "  정상 Precision이 60%대로 급락하여 정상 F1-Score가 60%대에 묶이게 됩니다.<br/>"
        "• <b>원인 3: 운동사슬(Kinetic Chain) 연쇄 붕괴의 미반영</b><br/>"
        "  스쿼트 동작은 '발목 가동성 부족(c4) → 무릎 안쪽 모임(c3) → 골반 보상 및 척추 말림(c1)'의 운동학적 연쇄성을 가집니다. "
        "  현재의 파이프라인은 4개 모델을 개별 학습한 뒤 사후 결합(Late Fusion)하므로 신체 분절 간 인과적 전이를 학습하지 못합니다.<br/>"
        "• <b>원인 4: 단안 카메라 시점별 깊이 정보 비대칭성</b><br/>"
        "  정면(시점 C)에서는 Z축 전후 깊이 변화(척추 굴곡, 발 들림)가 가려지고, 후방(시점 A, E)에서는 얼굴이 보이지 않습니다."
    )
    story.append(Paragraph(p4, body_style))

    # 섹션 5
    story.append(Paragraph("5. 외부 AI에게 요청하는 4대 핵심 재조사 의뢰 과제", h1_style))
    story.append(Paragraph(
        "위 시스템 구성 및 진단 결과를 바탕으로, 다음 4대 핵심 질문에 대해 학술 문헌(CVPR, ICCV, NeurIPS, KDD 등) 및 "
        "오픈소스 SOTA 피트니스/의료 AI 사례를 조사하여 구체적인 수학적 수식과 알고리즘 아키텍처를 제안해 주십시오.", body_style
    ))

    q_box_data = [
        [
            Paragraph("과제 1: [정상 F1-Score 80% 돌파] 극단적 불균형(5% 정상 vs 95% 결함) 극복 기법", q_title_style),
            Paragraph(
                "• <b>질문</b>: 4개 결함 확률을 소프트 퓨전할 때, 정상 Precision을 훼손하지 않으면서 정상 Recall을 끌어올려 "
                "<b>Clean F1-Score를 현재 68%에서 80% 이상으로 달성하기 위한 최적의 결합 기법</b>은 무엇인가?<br/>"
                "• <b>검토 요청 분야</b>: Cost-Sensitive Meta-Classifier, Multi-Label Conformal Prediction, "
                "Focal Loss 기반 2단계 메타러너, 최적 베이지안 판정 경계(Bayes Decision Rule under Asymmetric Cost) 수식.", body_style
            )
        ],
        [
            Paragraph("과제 2: [4개 오류 라벨 완전 일치율(EMR) 80% 돌파] 레이블 의존성 모델링", q_title_style),
            Paragraph(
                "• <b>질문</b>: 독립 분류기의 오차 누적을 극복하고, <b>4개 이진 라벨 [c1, c2, c3, c4]을 동시에 모두 정확히 맞추는 비율(EMR)을 "
                "현재 64%에서 80% 이상으로 끌어올리기 위한 SOTA 구조</b>는 무엇인가?<br/>"
                "• <b>검토 요청 분야</b>: Classifier Chains (순서 최적화: c3→c4→c1→c2 등), Probabilistic Graphical Models (CRF/Markov Random Field), "
                "Multi-Task Joint Representation Head, Structured SVM / Energy-Based Models.", body_style
            )
        ],
        [
            Paragraph("과제 3: [운동사슬 연쇄 반영] 관절 간 시공간 상호작용 피처 결합 방안", q_title_style),
            Paragraph(
                "• <b>질문</b>: 4대 서브모델이 각각 독립된 피처를 추출하는 현재 방식을 보완하여, "
                "<b>발목(c4)-무릎(c3)-척추(c1) 간의 상호작용(Kinetic Chain Interaction)을 포착하면서도 과적합을 방지하는 최적의 중간 표현(Intermediate Representation)</b>은 무엇인가?<br/>"
                "• <b>검토 요청 분야</b>: Spatio-Temporal Graph (ST-GCN) 가중치 전이, Cross-Attention 융합 블록, 피처 레벨 앙상블 규격.", body_style
            )
        ],
        [
            Paragraph("과제 4: [5대 시점 비대칭성 극복] 시점 조건부 메타 앙상블(View-Conditioned MoE) 수식", q_title_style),
            Paragraph(
                "• <b>질문</b>: 공식 5대 시점(후방, 사선, 정면)의 비대칭 관측성을 처리할 때, 단순 임계치 튜닝을 넘어 "
                "<b>카메라 각도(시점 원핫 또는 법선 Z)에 따라 서브모델의 신뢰도 가중치를 동적으로 배분하는 게이팅 네트워크(Gating Network)의 최적 수식과 학습 전략</b>은 무엇인가?", body_style
            )
        ]
    ]

    for q_head, q_body in q_box_data:
        q_table = Table([[q_head], [q_body]], colWidths=[487])
        q_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 5),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 2),
            ('TOPPADDING', (0, 1), (-1, 1), 2),
            ('BOTTOMPADDING', (0, 1), (-1, 1), 5),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('LINEBELOW', (0, 0), (-1, 0), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        story.append(q_table)
        story.append(Spacer(1, 4))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[+] PDF 브리핑 문서 생성 완료: {PDF_PATH}")


# ---------------------------------------------------------------------------
# 2. DOCX 생성 함수
# ---------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1A365D"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_header_cell(row, idx, text, width, align=WD_ALIGN_PARAGRAPH.CENTER):
    c = row.cells[idx]
    c.width = width
    set_cell_background(c, "1A365D")
    set_cell_margins(c, top=140, bottom=140)
    p = c.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

def add_data_cell(row, idx, text, width, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, color_rgb=None, bg_hex=None):
    c = row.cells[idx]
    c.width = width
    if bg_hex:
        set_cell_background(c, bg_hex)
    set_cell_margins(c, top=90, bottom=90)
    p = c.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(8.5)
    r.font.bold = bold
    if color_rgb:
        r.font.color.rgb = color_rgb

def generate_docx():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("vidio-health-trainer | AI Biomechanics & System Architecture Technical Briefing")
    r_sub.font.name = "맑은 고딕"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("바벨 스쿼트 세부 오류 모델(c1~c4) 및 통합 정상/오류 분류 모델 상세 명세서\n[외부 AI 재조사 의뢰용: 정상 F1 & 4개 라벨 완전 일치율(EMR) 60%대 정체 돌파 자문서]")
    r_title.font.name = "맑은 고딕"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("프로젝트: vidio-health-trainer | 대상 운동: 바벨 스쿼트 | 데이터셋: N=2,605 64F 스플라인 | 목적: 외부 AI 기술 자문")
    r_meta.font.name = "맑은 고딕"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    doc.add_paragraph()

    # 요약 박스
    tbl_box = doc.add_table(rows=1, cols=1)
    tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_box = tbl_box.rows[0].cells[0]
    c_box.width = Inches(6.5)
    set_cell_background(c_box, "F0F4F8")
    set_cell_margins(c_box, top=140, bottom=140, left=180, right=180)
    tcPr = c_box._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1A365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p_b = c_box.paragraphs[0]
    r_bt = p_b.add_run("■ [의뢰 배경 및 자문 목표]\n")
    r_bt.font.name = "맑은 고딕"
    r_bt.font.size = Pt(10)
    r_bt.font.bold = True
    r_bt.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    r_bc = p_b.add_run(
        "본 문서는 vidio-health-trainer 프로젝트에서 바벨 스쿼트 3차원 골격 시계열 데이터를 바탕으로 "
        "4대 세부 결함(c1 척추, c2 고개/시선, c3 무릎, c4 발바닥) 및 전신 정상/오류를 진단하는 통합 AI 엔진 구축 중 발생한 "
        "'정상 F1-Score 및 4대 라벨 완전 일치율(EMR)의 60%대 정체 병목'을 해결하기 위해 외부 AI에게 최적의 해결책을 자문받고자 작성되었습니다.\n"
        "외부 AI의 1차 자문 제안 7가지(소프트 리스크 퓨전, 관측 가능 정상 게이팅, 사선 왜곡 보정, 발목 롤 특이점 보정, 5대 시점 동적 임계치, 3-State 선택적 판정 등)를 "
        "성공적으로 적용하여 전반적인 성능이 대폭 향상되었으나, (1) 정상 F1-Score (~68%) 및 (2) 4개 라벨 완전 일치율 EMR (~64%)이 "
        "여전히 60%대에 머무르고 있습니다. 이를 80% 이상으로 극대화하기 위한 모델 구조, 손실 함수, 결합 기법의 재조사를 의뢰합니다."
    )
    r_bc.font.name = "맑은 고딕"
    r_bc.font.size = Pt(9)
    doc.add_paragraph()

    # 섹션 1
    h1 = doc.add_heading("1. 데이터셋 환경 명세 및 극심한 클래스 불균형 현황", level=1)
    h1.runs[0].font.name = "맑은 고딕"
    h1.runs[0].font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)
    p_s1 = doc.add_paragraph()
    p_s1.add_run(
        "• 데이터셋 규모: AI-Hub 피트니스 데이터셋 바벨 스쿼트 전수 N=2,605건 시퀀스\n"
        "• 피험자 전수 반영: 피험자 1번~9번 전수(9명) 반영 및 5-Fold StratifiedGroupKFold 피험자 완전 격리 검증 (데이터 누수 0%)\n"
        "• 공식 5대 시점 전수 반영: 시점 A(우측 후방 45°), 시점 B(우측 전방 45°), 시점 C(정면 0°), 시점 D(좌측 전방 45°), 시점 E(좌측 후방 45°)\n"
        "• 전처리 표준: 원천 영상 전 시퀀스 무삭제 64프레임 3차 스플라인(Cubic Spline) 보간 완료 (횟수 분할 파편화 100% 영구 배제)\n"
        "• 골든 정답 라벨: AI-Hub 원천 라벨의 72.8% 노이즈를 전문가 57선 체크리스트로 정제한 무결성 라벨셋 (squat_golden_ground_truth.npz)\n"
        "• ★ 극심한 정상/결함 클래스 불균형:\n"
        "  - 전신 완전 정상 (Clean / c1~c4 모두 0): 135건 (5.18%)\n"
        "  - 전신 결함 존재 (Defect / c1~c4 중 1개 이상 1): 2,470건 (94.82%) [불균형 비율 1:18.3]\n"
        "  - 각 결함별 발생률: c1 46.96%, c2 53.73%, c3 49.69%, c4 49.74% (개별 결함은 ~50% 균형)"
    )

    # 섹션 2
    h2 = doc.add_heading("2. 4대 세부 오류 모델(c1~c4)의 구성, 입력, 피처셋 및 처리 파이프라인", level=1)
    h2.runs[0].font.name = "맑은 고딕"
    h2.runs[0].font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    t_sub = doc.add_table(rows=5, cols=5)
    t_sub.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sub)
    w_sub = [Inches(0.9), Inches(1.3), Inches(1.3), Inches(1.8), Inches(1.2)]
    for i, title in enumerate(["결함 코드", "결함 명칭 및 의미", "입력 및 처리 파이프라인", "추출 피처셋 명세", "단독 모델 및 성능"]):
        add_header_cell(t_sub.rows[0], i, title, w_sub[i])

    rows_sub = [
        ("c1", "척추의 중립 위반\n(요추 굴곡/과신전, 버트윙크, 굿모닝)", "MediaPipe 33관절 64F\n+ 체간 로컬 정규직교계\n+ Nadir ±8F 국소 윈도우", "27D 정예 생체역학\n(L1-S1 곡률 poly_a, Cobb각, C1-T8 힌지각, L5 후방말림 등 16D + 11D MP 직접계측)\n+ 24D 시계열 텐서(64x24)", "Tri-Ensemble\n(CatBoost+LGBM+TCN)\nF1 80.52%\nAcc 81.54%\nAUC 88.94%"),
        ("c2", "고개 정면 미유지\n(경추 과굴곡/신전, 시선 바닥/천장 이탈)", "MediaPipe 33관절 64F\n(코, 양귀, 양어깨, 양골반)\n※ AI-Hub 3D JSON 배제", "19D 생체역학 피처셋\n(체간 로컬 시상면 경추 굴곡각 8D + 단일 귀 적응형 안면 지향 4D + 법선 Z 2D)", "CatBoost+LGBM\n이진 F1 76.53%\n선택적 F1 83.90%\n(Coverage 63.8%)"),
        ("c3", "발-무릎 방향 불일치\n(무릎 모임 Valgus, 좌우 비대칭 붕괴)", "MediaPipe 33관절 64F\n(양 무릎, 양 발목, 양 골반)\n※ Format B 슬라이싱 제거", "22D 생체역학 피처셋\n(골반 관상면 투영 K/A 비율 6D, 외반합 6D, 대퇴-경골 관상각 4D, 무릎-발끝 정렬오차 4D)", "CatBoost+LGBM\nF1 91.45% 🏆\nAcc 91.21%\nAUC 96.50%"),
        ("c4", "발바닥 지면 미고정\n(뒤꿈치/앞꿈치 들림, 발목 롤 회내/회외)", "MediaPipe 33관절 64F\n(양 발목, 양 뒤꿈치, 양 앞꿈치)", "25D 생체역학 피처셋\n(최하사점 발바닥 경사각 Pitch, 수직차 ydelta, 들림량, 커플링비, F0 접촉전이 빈도)", "CatBoost+LGBM+RF\nF1 79.46%\nAcc 79.42%\nAUC 87.73%")
    ]
    for r_idx, r_data in enumerate(rows_sub, 1):
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, val in enumerate(r_data):
            bold = (c_idx == 0)
            add_data_cell(t_sub.rows[r_idx], c_idx, val, w_sub[c_idx], bold=bold, bg_hex=bg)

    doc.add_paragraph()

    # 섹션 3
    h3 = doc.add_heading("3. 바벨 스쿼트 통합 정상/오류 분류 모델의 아키텍처 및 처리 방식", level=1)
    h3.runs[0].font.name = "맑은 고딕"
    h3.runs[0].font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)
    p_s3 = doc.add_paragraph()
    p_s3.add_run(
        "• 입력 데이터: 4대 서브모델의 사후확률 벡터 [p_c1, p_c2, p_c3, p_c4], 불확실성 마진 |p - 0.5|, 공식 5대 시점 원핫 벡터(view1~5), 상체/하체 상호작용 로짓\n"
        "• 지금까지 적용된 7대 개선 조치:\n"
        "  1) Hard-AND 탈피: 각 모델 0.50 이진 판정 후 AND 연산(Clean F1 30%대 폭락) 폐기 및 소프트 결합 도입\n"
        "  2) Soft-Risk Fusion: 결함 미발생 결합확률 P(Clean) = Π (1 - p_k) 도입\n"
        "  3) Observable Clean (v7.1): 후방(시점 A, E) 시선 가림 고려, 3대 결함(c1, c3, c4) 안전 시 정상 인정\n"
        "  4) Oblique Adaptive Soft Fusion (v7.3): 사선(시점 B, D) 45° 원근 왜곡 완화 가중치 적용\n"
        "  5) Heel-Toe Roll Singularity Compensation (v7.4): 발목 롤 특이점 보정\n"
        "  6) 5-View Dynamic Threshold Optimization (v7.5): 공식 5대 시점별(A~E) 독립 동적 임계치 최적화\n"
        "  7) 3-State Selective Classification: State 0(CLEAN), State 1(DEFECT), State 2(RETAKE 보류)\n"
        "• 현재 실측 성능: 7가지 조치 적용으로 Clean F1이 30%대에서 68.42%로 대폭 상승했으나, 정상 F1 (~68%)과 4개 라벨 완전 일치율(Exact Match Ratio, EMR: ~64%)이 여전히 60%대에 머물러 서비스 상용화 목표(80% 이상)에 미달함."
    )

    # 섹션 4
    h4 = doc.add_heading("4. 왜 여전히 60%대에 머물고 있는가? 근본 원인 진단", level=1)
    h4.runs[0].font.name = "맑은 고딕"
    h4.runs[0].font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)
    p_s4 = doc.add_paragraph()
    p_s4.add_run(
        "• 원인 1: 다중 레이블 오차 곱셈 법칙 (Multi-Label Error Multiplication)\n"
        "  개별 모델의 정확도가 80%~91%라 하더라도, 4개 라벨을 동시에 모두 정확히 맞출 확률(EMR)은 독립 가정 하에 0.80 × 0.85 × 0.91 × 0.80 ≈ 49.5%에 불과합니다. 현재 64.2%는 독립 곱보다는 높으나, 라벨 간 상관관계를 명시적으로 최적화하는 결합 헤드가 부재하여 발생하는 구조적 한계입니다.\n"
        "• 원인 2: 극단적 정상 희소성(5.18%)에 기인한 오탐(False Positive) 치명타\n"
        "  전체 2,605개 샘플 중 정상이 135개뿐이므로, 2,470개의 결함 샘플 중 단 3%(~74개)만 '정상'으로 오분류되어도 정상 Precision이 60%대로 급락하여 정상 F1-Score가 60%대에 묶입니다.\n"
        "• 원인 3: 운동사슬(Kinetic Chain) 연쇄 붕괴의 미반영\n"
        "  스쿼트는 '발목 가동성 부족(c4) → 무릎 외반(c3) → 척추 말림(c1)'의 운동학적 연쇄성을 가지나, 현재의 독립 앙상블은 인과적 전이를 학습하지 못합니다.\n"
        "• 원인 4: 단안 카메라 시점별 깊이 정보 비대칭성\n"
        "  정면(시점 C)에서는 Z축 전후 깊이 변화(척추 굴곡, 발 들림)가 가려지고, 후방(시점 A, E)에서는 얼굴이 보이지 않습니다."
    )

    # 섹션 5
    h5 = doc.add_heading("5. 외부 AI에게 요청하는 4대 핵심 재조사 의뢰 과제", level=1)
    h5.runs[0].font.name = "맑은 고딕"
    h5.runs[0].font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    q_data = [
        ("과제 1: [정상 F1-Score 80% 돌파] 극단적 불균형(5% 정상 vs 95% 결함) 극복 기법",
         "• 질문: 4개 결함 확률을 소프트 퓨전할 때, 정상 Precision을 훼손하지 않으면서 정상 Recall을 끌어올려 Clean F1-Score를 현재 68%에서 80% 이상으로 달성하기 위한 최적의 결합 기법은 무엇인가?\n"
         "• 검토 요청 분야: Cost-Sensitive Meta-Classifier, Multi-Label Conformal Prediction, Focal Loss 기반 2단계 메타러너, 최적 비대칭 베이지안 판정 경계 수식."),
        ("과제 2: [4개 오류 라벨 완전 일치율(EMR) 80% 돌파] 레이블 의존성 모델링",
         "• 질문: 독립 분류기의 오차 누적을 극복하고, 4개 이진 라벨 [c1, c2, c3, c4]을 동시에 모두 정확히 맞추는 비율(EMR)을 현재 64%에서 80% 이상으로 끌어올리기 위한 SOTA 구조는 무엇인가?\n"
         "• 검토 요청 분야: Classifier Chains (최적 순서: c3→c4→c1→c2 등), Probabilistic Graphical Models (CRF/Markov Random Field), Multi-Task Joint Representation Head, Structured Output SVM."),
        ("과제 3: [운동사슬 연쇄 반영] 관절 간 시공간 상호작용 피처 결합 방안",
         "• 질문: 4대 서브모델이 각각 독립된 피처를 추출하는 현재 방식을 보완하여, 발목(c4)-무릎(c3)-척추(c1) 간의 상호작용(Kinetic Chain Interaction)을 포착하면서도 과적합을 방지하는 최적의 중간 표현(Intermediate Representation)은 무엇인가?\n"
         "• 검토 요청 분야: Spatio-Temporal Graph (ST-GCN) 가중치 전이, Cross-Attention 융합 블록, 피처 레벨 앙상블 규격."),
        ("과제 4: [5대 시점 비대칭성 극복] 시점 조건부 메타 앙상블(View-Conditioned MoE) 수식",
         "• 질문: 공식 5대 시점(후방, 사선, 정면)의 비대칭 관측성을 처리할 때, 단순 임계치 튜닝을 넘어 카메라 각도(시점 원핫 또는 법선 Z)에 따라 서브모델의 신뢰도 가중치를 동적으로 배분하는 게이팅 네트워크(Gating Network)의 최적 수식과 학습 전략은 무엇인가?")
    ]

    for q_title, q_body in q_data:
        p_q = doc.add_paragraph()
        r_qt = p_q.add_run(q_title + "\n")
        r_qt.font.name = "맑은 고딕"
        r_qt.font.size = Pt(10)
        r_qt.font.bold = True
        r_qt.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
        r_qb = p_q.add_run(q_body)
        r_qb.font.name = "맑은 고딕"
        r_qb.font.size = Pt(9)
        doc.add_paragraph()

    os.makedirs(DEST_DIR, exist_ok=True)
    doc.save(DOCX_PATH)
    print(f"[+] DOCX 브리핑 문서 생성 완료: {DOCX_PATH}")

if __name__ == '__main__':
    generate_pdf()
    generate_docx()
