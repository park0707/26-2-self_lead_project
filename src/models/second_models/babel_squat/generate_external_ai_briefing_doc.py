# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Module: generate_external_ai_briefing_doc.py
Purpose: 외부 AI 연구·컨설팅용 베이스 지식 문서 (.docx 및 .pdf) 동시 생성
Destination:
  - D:\자주프 문서\외부ai 용 2단계 베이스 지식.docx
  - D:\자주프 문서\외부ai 용 2단계 베이스 지식.pdf
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

DEST_DIR = r"D:\자주프 문서"
DOCX_PATH = os.path.join(DEST_DIR, "외부ai 용 2단계 베이스 지식.docx")
PDF_PATH = os.path.join(DEST_DIR, "외부ai 용 2단계 베이스 지식.pdf")


# ---------------------------------------------------------------------------
# 1. DOCX 생성 유틸리티 및 본문 작성
# ---------------------------------------------------------------------------
def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_docx():
    doc = docx.Document()

    # 페이지 여백 설정 (상하 20mm, 좌우 20mm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 기본 스타일 폰트 설정
    normal_style = doc.styles['Normal']
    normal_style.font.name = '맑은 고딕'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 1. 헤더 타이틀 블록
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("vidio-health-trainer | AI Biomechanics & System Architecture Briefing\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)

    run_title = title_p.add_run("외부 AI 연구·자문용 베이스 지식 및 아키텍처 진단서")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    # 부제 박스
    callout_p = doc.add_paragraph()
    callout_p.paragraph_format.space_before = Pt(4)
    callout_p.paragraph_format.space_after = Pt(16)
    r_desc = callout_p.add_run(
        "본 문서는 'vidio-health-trainer' 프로젝트에서 바벨 스쿼트 3D 스켈레톤 시계열 데이터를 바탕으로 "
        "4대 운동 결함 진단 및 전신 정상/오류 분류를 수행하는 과정에서 직면한 '세부 모델 결합 시 오차 누적 병목' 및 "
        "'세부 결함 성능 상한선' 문제를 외부 AI에게 정확히 전달하고, 최적의 차세대 아키텍처 자문을 구하기 위해 작성된 공식 기술 배경서입니다."
    )
    r_desc.font.size = Pt(9.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # -----------------------------------------------------------------------
    # 섹션 1: 프로젝트 개요 및 환경 명세
    # -----------------------------------------------------------------------
    h1 = doc.add_heading("1. 프로젝트 개요 및 데이터 환경 명세", level=1)
    h1.style.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    doc.add_paragraph(
        "• 프로젝트 명: vidio-health-trainer\n"
        "• 대상 운동: 바벨 스쿼트 (Barbell Squat) — 3차원 골격 궤적 기반 운동 역학 분석\n"
        "• 입력 데이터: 64프레임 3차원 관절 좌표계 (MediaPipe Pose 33 관절점 + DFKI SpinePose 9대 척추 분절: S1~C1)\n"
        "• 데이터셋 규모: AI-Hub 피트니스 데이터셋 N=1,906건 시퀀스 (단일 반복 동작 기준 64프레임 3차 스플라인 보간 완료)\n"
        "• 카메라 시점 전수 사용: 5대 시점 전수 반영 (view1: 후방, view2: 우측 45°, view3: 완전 정면, view4: 좌측 45°, view5: 우측 후방 45°)\n"
        "• 검증 표준 프로토콜: 피험자 9명 전수 격리 5-Fold StratifiedGroupKFold 교차검증 (동일 피험자 train/val 동시 등장 원천 차단)"
    )

    # -----------------------------------------------------------------------
    # 섹션 2: 4대 세부 결함(Defects) 정의 및 단독 모델 실측 성적표
    # -----------------------------------------------------------------------
    h2 = doc.add_heading("2. 4대 세부 결함(Defects) 정의 및 개별 전문 모델 실측 성능", level=1)
    h2.style.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    doc.add_paragraph(
        "본 시스템은 바벨 스쿼트 수행 중 발생하는 인체역학적 핵심 결함을 4가지 세부 결함 클래스로 분류합니다. "
        "모든 단독 모델은 피험자 격리 5-Fold OOF(Out-Of-Fold) 방식으로 실측 검증되었습니다."
    )

    # 표: 4대 결함 성능표
    table1 = doc.add_table(rows=5, cols=6)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.style = 'Table Grid'

    headers1 = ["결함 코드", "결함 명칭 및 생체역학적 정의", "적용 모델 아키텍처", "Precision", "Recall", "F1-Score"]
    for i, h in enumerate(headers1):
        cell = table1.cell(0, i)
        cell.text = h
        set_cell_shading(cell, "0F2D59")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9)

    data1 = [
        ("c1", "척추의 중립 위반 (Spine Neutrality)\n- 요추 굴곡(Flexion), 신전, 버트윙크, 굿모닝", "45D Deep-Diet Hybrid\n(28D 기하학 + 17D Conv-BiGRU-Attn)", "82.14%", "77.42%", "79.71%\n(Clean 80.34%)"),
        ("c2", "고개 정면 미유지 (Head Alignment)\n- 시선 바닥/천장 과도 주시, 경추 꺾임", "13D 정예 기하학 앙상블\n(Gated Face/Gaze Vector)", "89.24%", "82.17%", "85.56%"),
        ("c3", "발과 무릎 방향 불일치 (Knee Valgus)\n- 무릎 모임(외반) 및 벌어짐(내반)", "17D Deep-Temporal Attention Net\n(무릎-발 정렬 64프레임 동역학)", "93.88%", "91.82%", "92.84%"),
        ("c4", "발바닥 지면 미고정 (Foot Grounding)\n- 뒤꿈치/앞꿈치 들림, 발목 안팎 뒤틀림", "9D 지면 접지 역학 피처셋\n(발목 회내/회외 3D 변위)", "92.16%", "86.37%", "89.17%")
    ]

    for row_idx, row_data in enumerate(data1, start=1):
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = table1.cell(row_idx, col_idx)
            cell.text = text
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            if col_idx in [0, 3, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)
                if col_idx == 0 or (col_idx == 5 and "92.84" in text):
                    r.font.bold = True

    # -----------------------------------------------------------------------
    # 섹션 3: 현재 통합 엔진 아키텍처 (계층적 투-트랙 캐스케이드)
    # -----------------------------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    h3 = doc.add_heading("3. 현재의 시스템 아키텍처: 계층적 투-트랙 캐스케이드 (Two-Track Cascade)", level=1)
    h3.style.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    doc.add_paragraph(
        "우리 시스템은 추론 효율성과 오탐 억제를 위해 전체 판별을 상체(Upper)와 하체(Lower)로 분할한 계층적 투-트랙 구조를 채택하고 있습니다.\n\n"
        "1) 1차 Tier-1 게이트키퍼 필터링:\n"
        "   • 상체 게이트키퍼 (M_upper, 42D): P_upper < 0.55 이면 상체 정상 확정 -> c1=0, c2=0 강제 할당 및 추론 스킵\n"
        "   • 하체 게이트키퍼 (M_lower, 26D): P_lower < 0.45 이면 하체 정상 확정 -> c3=0, c4=0 강제 할당 및 추론 스킵\n\n"
        "2) 2차 Tier-2 세부 전문 모델 추론:\n"
        "   • 게이트키퍼를 통과한 결함 의심 시퀀스에 한해서만 c1, c2, c3, c4 개별 모델을 가동하여 결함 여부를 0 또는 1로 판정\n\n"
        "3) 최종 전신 정상(Clean) vs 전신 오류(Defect) 판정 규칙 (Bottom-up 결합):\n"
        "   • 전신 정상 (Clean): 4개 결함이 모두 0인 경우 (c1=0 AND c2=0 AND c3=0 AND c4=0)\n"
        "   • 전신 오류 (Defect): 4개 결함 중 단 하나라도 1인 경우 (c1=1 OR c2=1 OR c3=1 OR c4=1)"
    )

    # -----------------------------------------------------------------------
    # 섹션 4: 현재 직면한 2대 본질적 병목 및 실측 딜레마
    # -----------------------------------------------------------------------
    h4 = doc.add_heading("4. 현재 직면한 2대 본질적 병목과 실측 딜레마", level=1)
    h4.style.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    doc.add_paragraph(
        "현재 시스템을 공식 5-Fold 실측 평가한 결과, 사용자님께서 정확히 통찰하신 바와 같이 "
        "두 가지 거대한 수학적·구조적 병목에 직면해 있습니다."
    )

    # 딜레마 1
    p_d1 = doc.add_paragraph()
    p_d1.add_run("병목 1. 오차의 승수 누적 (Multiplicative Error Stacking) 현상\n").bold = True
    p_d1.add_run(
        "개별 세부 결함 모델들의 성능(F1 80%~93%)은 매우 훌륭함에도 불구하고, 이를 OR/AND 결합으로 묶어 전신 정상을 판정하는 순간 "
        "치명적인 불균형이 발생합니다.\n"
        "• 전신 오류(Defect) 탐지: F1 97.23% / Precision 97.84% (극도로 뛰어남)\n"
        "• 전신 정상(Clean) 탐지: F1 63.97% / Recall 69.60% (심각하게 저조함)\n"
        "• 4대 결함 100% 완전 동시 적중률 (Exact Match Ratio): 61.80% (1,906건 중 1,178건만 4개 결함을 완벽 일치)\n\n"
        "수학적 원인: '전신 정상'은 P(c1=0) * P(c2=0) * P(c3=0) * P(c4=0)의 곱연산 구조를 가집니다. "
        "4개 모델 중 단 하나라도 정상 데이터를 결함(1)으로 오진(False Positive)하면, 그 동작은 즉시 '전신 오류'로 뒤집힙니다. "
        "이로 인해 데이터셋 내 존재하는 진짜 정상(Clean 125건) 중 약 30.4%가 억울하게 오류자로 판정되는 문제가 고착화되었습니다."
    )

    # 딜레마 2
    p_d2 = doc.add_paragraph()
    p_d2.add_run("\n병목 2. 세부 결함 모델의 성능 상한선 도달 (Bayes Error Rate의 벽)\n").bold = True
    p_d2.add_run(
        "특히 척추 중립 위반(c1)의 경우, 최신 컴퓨터 비전 및 딥러닝 기법(CTR-GCN 동적 토폴로지 신경망, SupCon 대조학습, "
        "피크 굴곡 각도 다중 태스크 회귀)을 전면 적용했음에도 F1 79.7% 수준에서 정체되었습니다.\n"
        "• 원인 분석: 모델 알고리즘의 결함이 아니라 원천 데이터셋의 라벨 모호성 때문입니다. "
        "데이터 검증 결과, 라벨은 '정상(0)'으로 되어 있으나 실제로는 척추가 40° 이상 심각하게 굴곡된 '가짜 정상(Level 4 모호성)' 데이터가 "
        "최소 52건 이상 원천 데이터셋에 포함되어 있습니다. "
        "정답지 자체에 노이즈가 존재하므로 지도학습 신경망이 이론적 베이즈 에러율(Bayes Error Rate) 한계에 도달한 상태입니다."
    )

    # -----------------------------------------------------------------------
    # 섹션 5: 외부 AI에게 의뢰하는 핵심 자문 및 연구 과제 (Research Questions)
    # -----------------------------------------------------------------------
    h5 = doc.add_heading("5. 외부 AI에게 의뢰하는 핵심 연구 및 설계 질문", level=1)
    h5.style.font.color.rgb = RGBColor(0x0F, 0x2D, 0x59)

    doc.add_paragraph(
        "외부 AI 연구원께서는 위의 배경과 정량 실측 수치를 바탕으로 아래 3가지 핵심 질문에 대해 "
        "수학적 이론, 최신 학술 연구(CVPR, ICCV, NeurIPS, 스포츠 AI 논문), 그리고 실무 아키텍처 관점에서 구체적인 해법을 제시해 주시기 바랍니다."
    )

    # Q1
    q1_box = doc.add_paragraph()
    q1_box.paragraph_format.left_indent = Inches(0.2)
    q1_r1 = q1_box.add_run("[핵심 질문 1] 세부 오류 모델의 결합(Bottom-up)으로 전신 정상을 판별하려 한 시도 자체가 구조적 결함인가?\n")
    q1_r1.bold = True
    q1_r1.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
    q1_box.add_run(
        "• 멀티레이블 분류(Multi-label Classification) 및 이상 탐지(Anomaly Detection) 관점에서, "
        "개별 결함 모델들의 판정 결과(c1~c4)를 논리곱(AND)으로 결합하여 정상을 판정하는 방식은 본질적으로 오차 누적이 불가피한 실패 설계입니까?\n"
        "• 산업계 및 학계에서 다중 결함 검출과 정상 판별을 동시에 요구받을 때 일반적으로 채택하는 표준 설계 패턴은 무엇입니까?"
    )

    # Q2
    q2_box = doc.add_paragraph()
    q2_box.paragraph_format.left_indent = Inches(0.2)
    q2_r2 = q2_box.add_run("[핵심 질문 2] 시도 자체는 맞지만, 성능 향상을 위한 결합 방식의 '본질'을 놓친 것인가?\n")
    q2_r2.bold = True
    q2_r2.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
    q2_box.add_run(
        "• 만약 Bottom-up 접근을 유지해야 한다면, 각 모델의 위양성(False Positive)이 곱연산으로 전이되는 것을 차단하는 "
        "고도화된 확률 융합 기법(예: 플랫 스케일링/아이소토닉 확률 보정, 베이지안 네트워크 결합, Copula 기반 상관관계 모델링, 신뢰도 마진 기반 동적 임계값) 중 "
        "가장 효과적인 해결책은 무엇입니까?"
    )

    # Q3
    q3_box = doc.add_paragraph()
    q3_box.paragraph_format.left_indent = Inches(0.2)
    q3_r3 = q3_box.add_run("[핵심 질문 3] 차세대 아키텍처로서 어떤 방향성이 가장 타당한가?\n")
    q3_r3.bold = True
    q3_r3.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
    q3_box.add_run(
        "우리가 검토할 수 있는 다음 대안들 중 어떤 아키텍처를 우선적으로 추진해야 합니까?\n"
        "  1) 대안 A (Top-down 2단계 분리): 1단계에서 전신 정상/오류를 독립적으로 직접 판별하는 '전신 정상 전용 모델(One-Class Classifier or Binary Gatekeeper)'을 구축하고, 2단계에서 오류로 판정된 건에 대해서만 c1~c4 세부 진단을 수행하는 구조\n"
        "  2) 대안 B (End-to-End Multi-Task Learning): 64프레임 전신 스켈레톤을 단일 딥러닝 백본에 입력하고, 전신 정상 여부 헤드(Binary Head)와 4대 결함 헤드(Multi-label Head)를 상호 공유 표현으로 동시 학습시키는 단일 통합 신경망\n"
        "  3) 대안 C (Hierarchical Soft Probability Fusion): 기존 분할 모델 구조를 유지하되, 각 모델의 사후 확률을 0/1 이진화하지 않고 다차원 로지스틱 회귀나 메타 러너(Stacking Classifier)로 전신 정상을 최종 종합 추론하는 구조"
    )

    doc.save(DOCX_PATH)
    print(f"[*] Word 문서 (.docx) 생성 완료: {DOCX_PATH}")


# ---------------------------------------------------------------------------
# 2. PDF 생성 유틸리티 및 본문 작성 (ReportLab)
# ---------------------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("MalgunGothic", 8)
        self.setFillColor(colors.HexColor("#6B7280"))
        # 상단 헤더
        self.drawString(54, 800, "vidio-health-trainer | AI Biomechanics & System Architecture Briefing")
        self.setStrokeColor(colors.HexColor("#E5E7EB"))
        self.setLineWidth(0.5)
        self.line(54, 794, 541, 794)

        # 하단 푸터
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(541, 35, page_str)
        self.drawString(54, 35, "CONFIDENTIAL - FOR EXTERNAL AI RESEARCH & CONSULTATION ONLY")
        self.line(54, 45, 541, 45)
        self.restoreState()


def create_pdf():
    # 폰트 등록
    font_path = "C:/Windows/Fonts/malgun.ttf"
    font_bold_path = "C:/Windows/Fonts/malgunbd.ttf"
    pdfmetrics.registerFont(TTFont("MalgunGothic", font_path))
    pdfmetrics.registerFont(TTFont("MalgunGothic-Bold", font_bold_path))

    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        fontName='MalgunGothic-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#111827"),
        spaceAfter=6
    )
    sub_title_style = ParagraphStyle(
        'DocSubTitle',
        fontName='MalgunGothic-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#0056B3"),
        spaceAfter=4
    )
    desc_box_style = ParagraphStyle(
        'DescBox',
        fontName='MalgunGothic',
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#4B5563")
    )
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        fontName='MalgunGothic-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0F2D59"),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        fontName='MalgunGothic',
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#374151"),
        spaceAfter=6
    )
    bold_body_style = ParagraphStyle(
        'BoldBody_Custom',
        fontName='MalgunGothic-Bold',
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#111827"),
        spaceAfter=4
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        fontName='MalgunGothic',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1F2937")
    )
    table_header_style = ParagraphStyle(
        'TableHeader',
        fontName='MalgunGothic-Bold',
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.white
    )
    q_title_style = ParagraphStyle(
        'QTitle',
        fontName='MalgunGothic-Bold',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#0056B3"),
        spaceAfter=4
    )

    story = []

    # 헤더
    story.append(Spacer(1, 10))
    story.append(Paragraph("vidio-health-trainer | AI Biomechanics & System Architecture Briefing", sub_title_style))
    story.append(Paragraph("외부 AI 연구·자문용 베이스 지식 및 아키텍처 진단서", title_style))

    # 요약 박스
    summary_text = (
        "<b>[문서 목적]</b> 본 문서는 <b>vidio-health-trainer</b> 프로젝트에서 바벨 스쿼트 3D 스켈레톤 시계열 데이터를 바탕으로 "
        "4대 운동 결함 진단 및 전신 정상/오류 분류를 수행하는 과정에서 직면한 <b>'세부 모델 결합 시 오차 누적 병목'</b> 및 "
        "<b>'세부 결함 성능 상한선'</b> 문제를 외부 AI에게 정확히 전달하고, 최적의 차세대 아키텍처 자문을 구하기 위해 작성된 공식 기술 배경서입니다."
    )
    summary_table = Table([[Paragraph(summary_text, desc_box_style)]], colWidths=[487])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F3F4F6")),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 10))

    # 섹션 1
    story.append(Paragraph("1. 프로젝트 개요 및 데이터 환경 명세", h1_style))
    p1_text = (
        "• <b>프로젝트 명</b>: vidio-health-trainer<br/>"
        "• <b>대상 운동</b>: 바벨 스쿼트 (Barbell Squat) — 3차원 골격 궤적 기반 운동 역학 분석<br/>"
        "• <b>입력 데이터</b>: 64프레임 3차원 관절 좌표계 (MediaPipe Pose 33 관절점 + DFKI SpinePose 9대 척추 분절: S1~C1)<br/>"
        "• <b>데이터셋 규모</b>: AI-Hub 피트니스 데이터셋 N=1,906건 시퀀스 (단일 반복 동작 기준 64프레임 3차 스플라인 보간 완료)<br/>"
        "• <b>카메라 시점 전수 사용</b>: 5대 시점 전수 반영 (후방, 우측45°, 완전정면, 좌측45°, 우측후방45°)<br/>"
        "• <b>검증 표준 프로토콜</b>: 피험자 9명 전수 격리 5-Fold StratifiedGroupKFold 교차검증 (Data Leakage 원천 차단)"
    )
    story.append(Paragraph(p1_text, body_style))

    # 섹션 2
    story.append(Paragraph("2. 4대 세부 결함(Defects) 정의 및 개별 전문 모델 실측 성능", h1_style))
    story.append(Paragraph(
        "각 세부 결함 모델은 생체역학적 도메인 지식과 최신 딥러닝 아키텍처(Conv-BiGRU-Attention)가 융합된 전문 모델이며, "
        "모두 피험자 격리 5-Fold 실측 검증되었습니다.", body_style
    ))

    t_data = [
        [
            Paragraph("결함 코드", table_header_style),
            Paragraph("결함 명칭 및 생체역학적 정의", table_header_style),
            Paragraph("적용 모델 아키텍처", table_header_style),
            Paragraph("Precision", table_header_style),
            Paragraph("Recall", table_header_style),
            Paragraph("F1-Score", table_header_style)
        ],
        [
            Paragraph("<b>c1</b>", table_cell_style),
            Paragraph("<b>척추 중립 위반</b> (Spine Neutrality)<br/>- 요추 굴곡/신전, 버트윙크, 굿모닝", table_cell_style),
            Paragraph("45D Deep-Diet Hybrid<br/>(28D 기하학 + 17D 시계열 딥러닝)", table_cell_style),
            Paragraph("82.14%", table_cell_style),
            Paragraph("77.42%", table_cell_style),
            Paragraph("<b>79.71%</b><br/>(Clean 80.34%)", table_cell_style)
        ],
        [
            Paragraph("<b>c2</b>", table_cell_style),
            Paragraph("<b>고개 정면 미유지</b> (Head Alignment)<br/>- 시선 바닥/천장 과도 주시, 경추 꺾임", table_cell_style),
            Paragraph("13D 정예 기하학 앙상블<br/>(Gated Face/Gaze Vector)", table_cell_style),
            Paragraph("89.24%", table_cell_style),
            Paragraph("82.17%", table_cell_style),
            Paragraph("<b>85.56%</b>", table_cell_style)
        ],
        [
            Paragraph("<b>c3</b>", table_cell_style),
            Paragraph("<b>발과 무릎 불일치</b> (Knee Valgus)<br/>- 무릎 모임(외반) 및 벌어짐(내반)", table_cell_style),
            Paragraph("17D Deep-Temporal Net<br/>(무릎-발 정렬 64프레임 동역학)", table_cell_style),
            Paragraph("93.88%", table_cell_style),
            Paragraph("91.82%", table_cell_style),
            Paragraph("<b>92.84%</b> 🏆", table_cell_style)
        ],
        [
            Paragraph("<b>c4</b>", table_cell_style),
            Paragraph("<b>발바닥 지면 미고정</b> (Foot Grounding)<br/>- 뒤꿈치/앞꿈치 들림, 발목 회내/회외", table_cell_style),
            Paragraph("9D 지면 접지 역학 피처셋<br/>(발목 3D 변위 및 압력 중심)", table_cell_style),
            Paragraph("92.16%", table_cell_style),
            Paragraph("86.37%", table_cell_style),
            Paragraph("<b>89.17%</b>", table_cell_style)
        ]
    ]

    t1 = Table(t_data, colWidths=[40, 150, 147, 50, 50, 50])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F2D59")),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (0, 2), (-1, 2), colors.white),
        ('BACKGROUND', (0, 3), (-1, 3), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (0, 4), (-1, 4), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t1)
    story.append(Spacer(1, 10))

    # 섹션 3
    story.append(Paragraph("3. 현재의 시스템 아키텍처: 계층적 투-트랙 캐스케이드 (Two-Track Cascade)", h1_style))
    p3_text = (
        "• <b>Tier-1 게이트키퍼 필터링</b>:<br/>"
        "  - 상체 게이트키퍼 (42D): P_upper &lt; 0.55 이면 상체 정상 확정 &rarr; c1=0, c2=0 강제 할당 및 추론 스킵<br/>"
        "  - 하체 게이트키퍼 (26D): P_lower &lt; 0.45 이면 하체 정상 확정 &rarr; c3=0, c4=0 강제 할당 및 추론 스킵<br/>"
        "• <b>Tier-2 세부 전문 모델 추론</b>: 게이트키퍼 통과 시퀀스에 한해 c1~c4 개별 모델 가동<br/>"
        "• <b>전신 판정 규칙 (Bottom-up 결합)</b>:<br/>"
        "  - 전신 정상 (Clean): (c1 == 0) &amp; (c2 == 0) &amp; (c3 == 0) &amp; (c4 == 0)<br/>"
        "  - 전신 오류 (Defect): (c1 == 1) | (c2 == 1) | (c3 == 1) | (c4 == 1)"
    )
    story.append(Paragraph(p3_text, body_style))

    # 섹션 4
    story.append(Paragraph("4. 현재 직면한 2대 본질적 병목과 실측 딜레마", h1_style))
    d1_box = (
        "<b>[병목 1] 오차의 승수 누적 (Multiplicative Error Stacking) 현상</b><br/>"
        "개별 세부 결함 모델들의 성능(F1 80%~93%)은 매우 훌륭함에도 불구하고, 이를 OR/AND 결합으로 묶어 전신 정상을 판정하는 순간 치명적인 불균형이 발생합니다.<br/>"
        "• <b>전신 오류(Defect) 탐지</b>: <b>F1 97.23% / Precision 97.84%</b> (극도로 뛰어남)<br/>"
        "• <b>전신 정상(Clean) 탐지</b>: <b>F1 63.97% / Recall 69.60%</b> (심각하게 저조함)<br/>"
        "• <b>4대 결함 100% 완전 동시 적중률 (Exact Match Ratio)</b>: <b>61.80%</b> (1,906건 중 1,178건만 일치)<br/>"
        "<i>수학적 원인: 전신 정상이 되려면 4개 모델이 모두 0이어야 하므로, 4개 중 단 하나라도 위양성(False Positive)을 범하면 세트 전체가 오류로 뒤집힙니다. "
        "진짜 정상(125건) 중 약 30.4%가 억울하게 오류로 오진되는 문제가 발생합니다.</i>"
    )
    t_d1 = Table([[Paragraph(d1_box, body_style)]], colWidths=[487])
    t_d1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FEF2F2")),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#F87171")),
    ]))
    story.append(t_d1)
    story.append(Spacer(1, 6))

    d2_box = (
        "<b>[병목 2] 세부 결함 모델의 성능 상한선 도달 (Bayes Error Rate의 벽)</b><br/>"
        "척추 중립 위반(c1)의 경우, 최신 CTR-GCN 동적 토폴로지, SupCon 대조학습, 피크 각도 다중 태스크 회귀를 적용했음에도 F1 79.7%에서 정체되었습니다.<br/>"
        "• <b>원인 분석</b>: 모델 알고리즘의 한계가 아니라 <b>AI-Hub 원천 라벨 자체의 노이즈</b> 때문입니다. "
        "라벨은 '정상(0)'으로 되어 있으나 실제로는 척추가 40° 이상 심각하게 굴곡된 '가짜 정상(Level 4 노이즈)'이 원천 데이터셋에 52건 존재합니다. "
        "정답지 자체에 모호성이 있어 지도학습 알고리즘이 이론적 베이즈 에러율 한계에 부딪혔습니다."
    )
    t_d2 = Table([[Paragraph(d2_box, body_style)]], colWidths=[487])
    t_d2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FFFBEB")),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#FBBF24")),
    ]))
    story.append(t_d2)
    story.append(Spacer(1, 8))

    # 섹션 5
    story.append(Paragraph("5. 외부 AI에게 의뢰하는 핵심 연구 및 설계 질문", h1_style))
    story.append(Paragraph(
        "외부 AI 연구원께서는 위의 배경과 정량 실측 수치를 바탕으로 아래 3가지 질문에 대해 "
        "수학적 이론, 최신 학술 논문, 그리고 실무 아키텍처 관점에서 구체적인 해법을 제시해 주시기 바랍니다.", body_style
    ))

    q_boxes = [
        (
            "[핵심 질문 1] 세부 오류 모델의 결합(Bottom-up)으로 전신 정상을 판별하려 한 시도 자체가 구조적 결함인가?",
            "• 멀티레이블 분류(Multi-label Classification) 및 이상 탐지(Anomaly Detection) 관점에서, "
            "개별 결함 모델(c1~c4)을 논리곱(AND)으로 결합하여 정상을 판정하는 방식은 본질적으로 오차 누적이 불가피한 실패 설계입니까?<br/>"
            "• 산업계 및 학계에서 다중 결함 검출과 정상 판별을 동시에 요구받을 때 일반적으로 채택하는 표준 설계 패턴은 무엇입니까?"
        ),
        (
            "[핵심 질문 2] 시도 자체는 맞지만, 성능 향상을 위한 결합 방식의 '본질'을 놓친 것인가?",
            "• 만약 Bottom-up 접근을 유지해야 한다면, 각 모델의 위양성(False Positive)이 곱연산으로 전이되는 것을 차단하는 "
            "고도화된 확률 융합 기법(예: 플랫 스케일링/아이소토닉 확률 보정, 베이지안 네트워크 결합, Copula 기반 상관관계 모델링, 신뢰도 마진 기반 동적 임계값) 중 "
            "가장 효과적인 해결책은 무엇입니까?"
        ),
        (
            "[핵심 질문 3] 차세대 아키텍처로서 어떤 방향성이 가장 타당한가?",
            "우리가 검토할 수 있는 다음 대안들 중 어떤 아키텍처를 우선적으로 추진해야 합니까?<br/>"
            "<b>1) 대안 A (Top-down 2단계 분리)</b>: 1단계에서 전신 정상/오류를 독립적으로 직접 판별하는 '전신 정상 전용 모델(One-Class Classifier or Binary Gatekeeper)'을 구축하고, 2단계에서 오류로 판정된 건에 대해서만 c1~c4 세부 진단을 수행하는 구조<br/>"
            "<b>2) 대안 B (End-to-End Multi-Task Learning)</b>: 64프레임 전신 스켈레톤을 단일 딥러닝 백본에 입력하고, 전신 정상 여부 헤드(Binary Head)와 4대 결함 헤드(Multi-label Head)를 상호 공유 표현으로 동시 학습시키는 단일 통합 신경망<br/>"
            "<b>3) 대안 C (Hierarchical Soft Probability Fusion)</b>: 기존 분할 모델 구조를 유지하되, 각 모델의 사후 확률을 0/1 이진화하지 않고 다차원 로지스틱 회귀나 메타 러너(Stacking Classifier)로 전신 정상을 최종 종합 추론하는 구조"
        )
    ]

    for q_t, q_b in q_boxes:
        q_content = [
            Paragraph(q_t, q_title_style),
            Paragraph(q_b, body_style)
        ]
        q_table = Table([[q_content]], colWidths=[487])
        q_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#93C5FD")),
        ]))
        story.append(q_table)
        story.append(Spacer(1, 6))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[*] PDF 문서 (.pdf) 생성 완료: {PDF_PATH}")


def main():
    os.makedirs(DEST_DIR, exist_ok=True)
    create_docx()
    create_pdf()
    print("=" * 80)
    print("  [완료] 외부 AI 연구용 베이스 지식 문서 (.docx 및 .pdf) 동시 배포 완료")
    print(f"  - Word: {DOCX_PATH}")
    print(f"  - PDF : {PDF_PATH}")
    print("=" * 80)


if __name__ == '__main__':
    main()
