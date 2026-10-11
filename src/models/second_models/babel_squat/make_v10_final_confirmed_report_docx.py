# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_v10_final_confirmed_report_docx.py
Purpose: 바벨스쿼트 AI 통합 진단 엔진 v10.0 공식 마스터 프로덕션 챔피언 최종 확정 보고서 생성
Target: D:\자주프 문서\5주차\vidio-health-trainer_바벨스쿼트_통합진단_v10_0_공식확정보고서.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="1A365D", bg_color="F7FAFC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.9)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}\n")
    run_t.bold = True
    run_t.font.name = '맑은 고딕'
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = '맑은 고딕'
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def style_heading(p, text, level=1):
    p.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 6))
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = '맑은 고딕'
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def format_table(table, col_widths, headers, data, header_bg="1A365D", alt_bg="F7FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_row = table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(heading)
        run.bold = True
        run.font.name = '맑은 고딕'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_col = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = '맑은 고딕'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            
            if c_idx in [1, 2, 3, 4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E0"/>
        <w:left w:val="none"/>
        <w:bottom w:val="single" w:sz="12" w:space="0" w:color="1A365D"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def create_report(output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    style_normal = doc.styles['Normal']
    style_normal.font.name = '맑은 고딕'
    style_normal.font.size = Pt(9.5)
    style_normal.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("vidio-health-trainer\n바벨 스쿼트 AI 통합 진단 엔진 v10.0 공식 마스터 프로덕션 챔피언 최종 확정 보고서")
    r_title.bold = True
    r_title.font.name = '맑은 고딕'
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Official Production Confirmation of v10.0 Master Engine: Selective Clean F1 74.29% & Precision 75.83%")
    r_sub.font.name = '맑은 고딕'
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_meta = p_meta.add_run("공식 확정일: 2026. 09. 27 | 시스템 버전: v10.0 Master Production | 프로젝트 공식 명칭: vidio-health-trainer")
    r_meta.font.name = '맑은 고딕'
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph()

    # 1. 확정 선언
    style_heading(doc.add_paragraph(), "1. 공식 승격 및 마스터 챔피언 확정 선언 (Executive Declaration)", level=1)
    doc.add_paragraph(
        "vidio-health-trainer 프로젝트의 바벨 스쿼트 4대 결함 및 완전 정상 다중진단 파이프라인에서, [해결책 1: False-Clean Asymmetric Loss (Reverse-ASL)]을 성공적으로 장착한 'v10.0 선택적 판정 모델'이 Clean F1 74.29%, Clean Precision 75.83%, 4-Defect EMR 71.78%, RETAKE율 9.94%(커버리지 90.06%)를 기록하며 사용자의 최종 승인을 받아 [공식 마스터 프로덕션 챔피언 v10.0]으로 정식 승격 및 확정되었음을 선언합니다."
    )

    add_callout(
        doc,
        "v10.0 마스터 프로덕션 공식 확정 핵심 사양 (Master Spec)",
        "■ 버전 명칭: vidio-health-trainer 바벨 스쿼트 통합 엔진 v10.0 (BabelSquatIntegratedEngineV10_0)\n"
        "■ 마스터 가중치: joint_meta_16way_v10_0_weights.pth (73.5 KB)\n"
        "■ 메타 파라미터: joint_meta_16way_v10_0_meta.json (tau_clean=0.38, deadband=[0.32, 0.38], eta=0.10, T_meta=1.20)\n"
        "■ 통합 엔진 소스: D:\\src\\models\\second_models\\babel_squat\\babel_squat_integrated_v10_0.py\n"
        "■ 공식 판정 성능 (선택적 모드): Clean F1 74.29% | Precision 75.83% | Recall 72.80% | 4-EMR 71.78% | RETAKE 9.94%",
        border_color="1A365D", bg_color="F7FAFC"
    )

    # 2. 전수 성능 대조표
    style_heading(doc.add_paragraph(), "2. v9.0 베이스라인 대비 공식 성능 향상 총괄표 (2,605건 전수)", level=1)

    t_all = doc.add_table(rows=5, cols=7)
    headers_all = ["평가 모드 / 버전", "Clean F1", "Precision", "Recall", "4-Defect EMR", "False Clean", "RETAKE율"]
    widths_all = [Inches(2.1), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.9), Inches(0.8), Inches(0.7)]
    data_all = [
        ["v9.0 베이스라인 (선택적 판정)", "70.12%", "68.22%", "72.13%", "71.71%", "-", "9.90%"],
        ["v10.0 공식 확정 (선택적 판정)", "74.29%", "75.83%", "72.80%", "71.78%", "-", "9.94%"],
        ["v9.0 베이스라인 (전수 100% 강제)", "67.16%", "67.67%", "66.67%", "69.40%", "43건", "0.00%"],
        ["v10.0 공식 확정 (전수 100% 강제)", "71.60%", "75.41%", "68.15%", "69.60%", "30건", "0.00%"]
    ]
    format_table(t_all, widths_all, headers_all, data_all, header_bg="1A365D")
    doc.add_paragraph()

    # 3. 5대 시점별 전수 감사
    style_heading(doc.add_paragraph(), "3. 공식 5대 카메라 시점별 전수 성능 감사 (퇴행 제로 검증)", level=1)

    t_view = doc.add_table(rows=6, cols=6)
    headers_view = ["카메라 시점 규격", "v9.0 F1", "v10.0 F1", "향상도 (Δ F1)", "v10.0 Precision", "False Clean 감소"]
    widths_view = [Inches(1.8), Inches(0.9), Inches(0.9), Inches(1.1), Inches(1.1), Inches(1.1)]
    data_view = [
        ["시점 A (우측 후방 45° / view5)", "51.85%", "61.54%", "+9.69%p ▲", "64.00% (+12.15%p)", "13건 → 9건 (-4건)"],
        ["시점 B (우측 전방 45° / view2)", "71.70%", "77.78%", "+6.08%p ▲", "77.78% (+4.70%p)", "7건 → 6건 (-1건)"],
        ["시점 C (정면 0° / view3)",       "81.97%", "84.75%", "+2.78%p ▲", "78.12% (+4.60%p)", "9건 → 7건 (-2건)"],
        ["시점 D (좌측 전방 45° / view4)", "73.08%", "73.47%", "+0.39%p ▲", "81.82% (+5.82%p)", "6건 → 4건 (-2건)"],
        ["시점 E (좌측 후방 45° / view1)", "54.17%", "55.81%", "+1.65%p ▲", "75.00% (+13.10%p)", "8건 → 4건 (-4건)"]
    ]
    format_table(t_view, widths_view, headers_view, data_view, header_bg="2B6CB0")
    doc.add_paragraph()

    # 4. 아키텍처 및 4-Layer 파이프라인
    style_heading(doc.add_paragraph(), "4. v10.0 마스터 아키텍처 및 4-Layer 진단 파이프라인", level=1)
    doc.add_paragraph(
        "■ Layer 1 [부위별 3-State 진단]:\n"
        "  - c1 (척추 중립 위반): v11.1 (F1 95.64%)\n"
        "  - c2 (고개 정면/시선 위반): v7.1 (F1 84.00%)\n"
        "  - c3 (발과 무릎 방향 일치 위반): v6.0 (F1 93.41%)\n"
        "  - c4 (발바닥 지면 고정 위반): v5.1 (F1 87.70%)\n\n"
        "■ Layer 2 [39D MoE 메타 신경망 + Reverse-ASL]:\n"
        "  - 입력 특성: 4D 보정확률 + 4D 마진 + 5D 시점원핫 + 5D c4물리 + 12D 관측품질 + 9D 운동사슬 = 총 39D\n"
        "  - 전문가 구성: Front/Oblique 특화, Rear 특화, Kinetic Chain 특화의 3대 서브전문가 동적 라우팅\n"
        "  - 최적화 손실: Reverse-ASL (lambda_fc=3.0) 적용으로 False Clean 페널티 지수적 부과\n\n"
        "■ Layer 3 [물리적 가속도/저크 필터링]:\n"
        "  - Wobble(0.008), Jerk(220.0), F0-Rate(0.35) 한계치를 통한 이상치 정상 판정 원천 차단\n\n"
        "■ Layer 4 [신체 품질 게이트 및 선택적 판정]:\n"
        "  - 관측 품질 미달(q_overall < 0.05) 및 확률 데드밴드 [0.32, 0.38] 구간의 경계선 샘플을 RETAKE로 안전 유도"
    )

    # 5. 서명 및 배포 승인
    style_heading(doc.add_paragraph(), "5. 시스템 배포 상태 및 파일 무결성 확인", level=1)
    doc.add_paragraph(
        "1. 소스 코드 배포: D:\\src\\models\\second_models\\babel_squat\\babel_squat_integrated_v10_0.py 정상 생성 및 import 검증 완료\n"
        "2. 가중치 배포: joint_meta_16way_v10_0_weights.pth (SHA/크기 정상 확인)\n"
        "3. 메타데이터 배포: joint_meta_16way_v10_0_meta.json 상태 'OFFICIAL_CONFIRMED' 갱신 완료\n"
        "4. 문서화 배포: 5주차 폴더 내 공식 확정 DOCX 보고서 배포 완료"
    )

    doc.save(output_path)
    print(f"Report successfully saved: {output_path}")

if __name__ == '__main__':
    target = r"D:\자주프 문서\5주차\vidio-health-trainer_바벨스쿼트_통합진단_v10_0_공식확정보고서.docx"
    create_report(target)
