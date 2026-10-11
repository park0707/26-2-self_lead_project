# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_c1_v11_final_report.py
Purpose: c1 v11_final (E0 + E1 융합 챔피언) 정식 Word 보고서(.docx) 생성
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DOCX_PATH = r"D:\자주프 문서\5주차\vidio-health-trainer_c1_v11_final_융합챔피언_최종검증보고서.docx"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_callout(doc, text_p_list, title="핵심 연구 성과 요약 (Executive Summary)", border_hex="1B365D", bg_hex="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    run_t = p0.add_run(f"■ {title}")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    for item in text_p_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(item)
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    title_p = doc.add_paragraph()
    t_run = title_p.add_run("vidio-health-trainer")
    t_run.font.name = "맑은 고딕"
    t_run.font.size = Pt(12)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    sub_p = doc.add_paragraph()
    s_run = sub_p.add_run("바벨 스쿼트 c1 v11_final (E0 + E1 궁극의 융합 챔피언) 실측 검증 보고서")
    s_run.font.name = "맑은 고딕"
    s_run.font.size = Pt(18)
    s_run.font.bold = True
    s_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    m_run = meta_p.add_run("작성일: 2026-09-26 | 대상: 척추 중립 위반(c1) | 모델: v11_final Tri-Ensemble | 데이터: 마스터 2,605건 전수")
    m_run.font.name = "맑은 고딕"
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    add_callout(
        doc,
        [
            "사용자의 직관적 가설 'E0와 E1이 각각 성능을 향상시키니 두 모델을 결합하면 더 높은 성능에 도달할 것'이 5-Fold 실측 결과 완벽히 적중함.",
            "v11_a(E0 38D GBDT 10%) + v11_a_diet(E0 27D GBDT 40%) + v11_b(E1 24D TCN 50%)의 3-Way Tri-Ensemble 아키텍처 구축.",
            "전체 F1 84.47%(기존 v10 79.73% 대비 +4.74%p 상승), 전체 정확도 85.22%(+4.84%p 상승)로 역대 최고 성능 공식 경신.",
            "5대 시점 전수에서 F1 83.3% ~ 87.4%를 기록하며, 최약 시점이었던 후방(84.12%, +8.00%p)과 정면(83.51%, +9.21%p)을 완벽 정복함.",
            "c1_model_v11_final을 vidio-health-trainer의 공식 최종 프로덕션 챔피언 모델로 배포 완료."
        ]
    )

    # 1. 모델 세대별 실측 성능 종합 비교표
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. c1 모델 세대별 5-Fold 실측 성능 비교 (N=2,605 전수, 피험자 완전 격리)")
    r1.font.name = "맑은 고딕"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl = doc.add_table(rows=6, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    col_w = [Inches(1.5), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)]
    headers = ["지표 (Metric)", "v10 (28D 베이스)", "v11_a_diet (27D)", "v11_b (TCN단독)", "★ v11_final", "v10 대비 증감"]
    rows = [
        ["F1-Score (결함)", "79.73%", "80.52%", "82.04%", "84.47%", "+4.74%p"],
        ["Accuracy (정확도)", "80.38%", "81.54%", "82.03%", "85.22%", "+4.84%p"],
        ["Precision (정밀도)", "80.00%", "81.88%", "79.13%", "85.54%", "+5.54%p"],
        ["Recall (재현율)", "79.46%", "79.20%", "85.18%", "83.43%", "+3.97%p"],
        ["ROC-AUC", "88.33%", "88.13%", "89.97%", "91.58%", "+3.25%p"]
    ]

    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.width = col_w[c_idx]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows, 1):
        bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = tbl.cell(r_idx, c_idx)
            cell.width = col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if c_idx == 4:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x44, 0x88)
            elif c_idx == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    # 2. 5대 시점별 세부 성능 감사
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("2. 5대 시점별 세부 성능 감사 (Multi-View Audit)")
    r2.font.name = "맑은 고딕"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl_v = doc.add_table(rows=6, cols=6)
    tbl_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_v.autofit = False
    col_w_v = [Inches(1.8), Inches(0.8), Inches(1.0), Inches(1.1), Inches(1.0), Inches(0.8)]
    headers_v = ["시점 (카메라 앵글)", "표본수", "v10.0 F1", "★ v11_final F1", "v10 대비 증감", "정확도"]
    rows_v = [
        ["view1 (완전 후방 180°)", "521", "76.12%", "84.12%", "+8.00%p (폭등)", "85.22%"],
        ["view2 (우측 45도 측면)", "521", "83.21%", "84.02%", "+0.81%p", "84.45%"],
        ["view3 (완전 정면 0°)", "521", "74.30%", "83.51%", "+9.21%p (폭등)", "84.84%"],
        ["view4 (좌측 45도 측면)", "521", "81.85%", "83.30%", "+1.45%p", "83.69%"],
        ["view5 (우측 후방 45도)", "521", "83.18%", "87.37%", "+4.19%p", "87.91%"]
    ]

    for c_idx, h_text in enumerate(headers_v):
        cell = tbl_v.cell(0, c_idx)
        cell.width = col_w_v[c_idx]
        set_cell_background(cell, "2C3E50")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows_v, 1):
        bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = tbl_v.cell(r_idx, c_idx)
            cell.width = col_w_v[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if c_idx == 3:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x44, 0x88)
            elif "폭등" in val:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # 3. 결론
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("3. 결론 및 배포 안내")
    r3.font.name = "맑은 고딕"
    r3.font.size = Pt(13)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.line_spacing = 1.25
    p_end.add_run(
        "사용자의 제안에 따라 E0(v11_a 국소 피처)와 E1(v11_b 시계열 신경망)을 융합한 c1_v11_final은 "
        "F1 84.47%, 정확도 85.22%를 달성하여 프로젝트 역사상 가장 뛰어난 일반화 성능을 기록했습니다.\n"
        "특히 정면(+9.21%p)과 후방(+8.00%p)의 병목을 완벽히 해소하여 모든 카메라 각도에서 83.3% 이상의 균일한 성능을 보장합니다."
    )

    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)
    doc.save(OUT_DOCX_PATH)
    print(f"[*] 정식 최종 Word 보고서 생성 완료: {OUT_DOCX_PATH}")


if __name__ == '__main__':
    main()
