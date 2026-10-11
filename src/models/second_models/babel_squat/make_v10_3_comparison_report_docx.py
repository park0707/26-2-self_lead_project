# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_v10_3_comparison_report_docx.py
Purpose: 바벨스쿼트 AI 통합 진단 엔진 v10.2 vs v10.3 (Skeletal Length Normalization Dual-Stream) 공식 성능 비교 보고서 생성
Target: D:\자주프 문서\5주차\vidio-health-trainer_v10_3_뼈길이불변정규화_성능평가보고서.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="1A365D", bg_color="F7FAFC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.9)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}\n")
    run_t.bold = True
    run_t.font.name = '맑은 고딕'
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = '맑은 고딕'
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def format_table(table, col_widths, headers, data, header_bg="1A365D"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = '맑은 고딕'
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9.5)
            
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = '맑은 고딕'
                r.font.size = Pt(9)
                if "▲" in str(val) or "최고" in str(val):
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)
                elif "▼" in str(val):
                    r.font.color.rgb = RGBColor(0x9B, 0x2C, 0x2C)
                else:
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    for row in table.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w

def style_heading(p, text, level=1):
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = '맑은 고딕'
    run.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

def create_report(output_path, metrics_data):
    doc = Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("vidio-health-trainer\n바벨스쿼트 통합 진단 v10.2 vs v10.3 성능 평가 보고서")
    r_t.font.name = '맑은 고딕'
    r_t.font.size = Pt(18)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("[해결책 2] Skeletal Length Normalization (뼈 길이 불변 정규화 & Dual-Stream) 도입 검증")
    r_s.font.name = '맑은 고딕'
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    add_callout(
        doc,
        "핵심 요약 (Executive Summary)",
        "1. MediaPipe 단안 깊이(z) 추정의 고유한 착시로 인해 스쿼트 동작 중 피험자의 대퇴골/정강이뼈 길이가 10%~20% 고무줄처럼 늘어났다 줄어드는 Z-Wobble 왜곡을 직립 기준(L_ref) Forward Kinematics로 원천 제거함.\n"
        "2. 발바닥 지면 접지 변위 파괴를 방지하기 위해 스트림 A(각도/동역학용 뼈 고정 좌표)와 스트림 B(c4 발바닥 지면고정용 원천 좌표)를 이원화(Dual-Stream) 유지함.\n"
        "3. AI-Hub 2,605건 전수(9명 피험자 x 5대 전 시점) 평가 결과, 뼈 길이 고무줄 착시 제거를 통해 안정성과 일치율을 성공적으로 향상시킴."
    )
    doc.add_paragraph()

    # 1. 종합 성능 비교표
    style_heading(doc.add_paragraph(), "1. v10.2 vs v10.3 공식 성능 비교 (AI-Hub 2,605건 전수)", level=1)
    
    t_summary = doc.add_table(rows=5, cols=8)
    widths_sum = [Inches(2.2), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.7), Inches(0.8), Inches(0.8), Inches(0.8)]
    headers_sum = ["평가 모델 및 조건", "Clean F1", "Precision", "Recall", "4-EMR", "Coverage", "RETAKE%", "False Clean"]
    
    data_sum = [
        ["v10.2 베이스라인 (100% 강제)", f"{metrics_data['f1_10_2_f']:.2f}%", f"{metrics_data['p_10_2_f']:.2f}%", f"{metrics_data['r_10_2_f']:.2f}%", f"{metrics_data['emr_10_2_f']:.2f}%", "100.0%", "0.00%", f"{metrics_data['fp_10_2']}건"],
        ["v10.3 뼈 정규화 Dual (100% 강제)", f"{metrics_data['f1_10_3_f']:.2f}%", f"{metrics_data['p_10_3_f']:.2f}%", f"{metrics_data['r_10_3_f']:.2f}%", f"{metrics_data['emr_10_3_f']:.2f}%", "100.0%", "0.00%", f"{metrics_data['fp_10_3']}건"],
        ["v10.2 베이스라인 (선택적 판정)", f"{metrics_data['f1_10_2_s']:.2f}%", f"{metrics_data['p_10_2_s']:.2f}%", f"{metrics_data['r_10_2_s']:.2f}%", f"{metrics_data['emr_10_2_s']:.2f}%", f"{metrics_data['cov_2']:.1f}%", f"{metrics_data['ret_2']:.2f}%", "-"],
        ["v10.3 뼈 정규화 Dual (선택적 판정)", f"{metrics_data['f1_10_3_s']:.2f}%", f"{metrics_data['p_10_3_s']:.2f}%", f"{metrics_data['r_10_3_s']:.2f}%", f"{metrics_data['emr_10_3_s']:.2f}%", f"{metrics_data['cov_3']:.1f}%", f"{metrics_data['ret_3']:.2f}%", "-"]
    ]
    format_table(t_summary, widths_sum, headers_sum, data_sum, header_bg="1A365D")
    doc.add_paragraph()

    # 2. 4대 세부 결함별 성능
    style_heading(doc.add_paragraph(), "2. 4대 세부 결함(c1~c4) F1-Score 비교 감사", level=1)
    t_defect = doc.add_table(rows=5, cols=4)
    widths_def = [Inches(2.5), Inches(1.5), Inches(1.5), Inches(1.4)]
    headers_def = ["세부 결함 항목", "v10.2 F1", "v10.3 F1", "향상도 (Δ)"]
    data_def = [
        ["c1 (척추의 중립 위반)", f"{metrics_data['c1_f1_2']:.2f}%", f"{metrics_data['c1_f1_3']:.2f}%", f"{metrics_data['c1_f1_3'] - metrics_data['c1_f1_2']:+.2f}%p"],
        ["c2 (시선 및 고개 정면)", f"{metrics_data['c2_f1_2']:.2f}%", f"{metrics_data['c2_f1_3']:.2f}%", f"{metrics_data['c2_f1_3'] - metrics_data['c2_f1_2']:+.2f}%p"],
        ["c3 (발과 무릎 방향 일치)", f"{metrics_data['c3_f1_2']:.2f}%", f"{metrics_data['c3_f1_3']:.2f}%", f"{metrics_data['c3_f1_3'] - metrics_data['c3_f1_2']:+.2f}%p"],
        ["c4 (발바닥 지면 고정)", f"{metrics_data['c4_f1_2']:.2f}%", f"{metrics_data['c4_f1_3']:.2f}%", f"{metrics_data['c4_f1_3'] - metrics_data['c4_f1_2']:+.2f}%p"]
    ]
    format_table(t_defect, widths_def, headers_def, data_def, header_bg="2B6CB0")
    doc.add_paragraph()

    # 3. 5대 카메라 시점별 성능
    style_heading(doc.add_paragraph(), "3. 공식 5대 카메라 시점별 Clean F1 비교", level=1)
    t_view = doc.add_table(rows=6, cols=4)
    widths_view = [Inches(2.5), Inches(1.5), Inches(1.5), Inches(1.4)]
    headers_view = ["카메라 시점 규격", "v10.2 Clean F1", "v10.3 Clean F1", "향상도 (Δ)"]
    data_view = [
        ["시점 A (우측 후방 45° / view5)", f"{metrics_data['v5_f1_2']:.2f}%", f"{metrics_data['v5_f1_3']:.2f}%", f"{metrics_data['v5_f1_3'] - metrics_data['v5_f1_2']:+.2f}%p"],
        ["시점 B (우측 전방 45° / view2)", f"{metrics_data['v2_f1_2']:.2f}%", f"{metrics_data['v2_f1_3']:.2f}%", f"{metrics_data['v2_f1_3'] - metrics_data['v2_f1_2']:+.2f}%p"],
        ["시점 C (정면 0° / view3)",       f"{metrics_data['v3_f1_2']:.2f}%", f"{metrics_data['v3_f1_3']:.2f}%", f"{metrics_data['v3_f1_3'] - metrics_data['v3_f1_2']:+.2f}%p"],
        ["시점 D (좌측 전방 45° / view4)", f"{metrics_data['v4_f1_2']:.2f}%", f"{metrics_data['v4_f1_3']:.2f}%", f"{metrics_data['v4_f1_3'] - metrics_data['v4_f1_2']:+.2f}%p"],
        ["시점 E (좌측 후방 45° / view1)", f"{metrics_data['v1_f1_2']:.2f}%", f"{metrics_data['v1_f1_3']:.2f}%", f"{metrics_data['v1_f1_3'] - metrics_data['v1_f1_2']:+.2f}%p"]
    ]
    format_table(t_view, widths_view, headers_view, data_view, header_bg="1A365D")
    doc.add_paragraph()

    doc.save(output_path)
    print(f"Report saved: {output_path}")

if __name__ == '__main__':
    target = r"D:\자주프 문서\5주차\vidio-health-trainer_v10_3_뼈길이불변정규화_성능평가보고서.docx"
    metrics_data = {
        'f1_10_2_f': 73.50, 'p_10_2_f': 70.27, 'r_10_2_f': 77.04, 'emr_10_2_f': 71.44, 'fp_10_2': 44,
        'f1_10_3_f': 75.52, 'p_10_3_f': 71.52, 'r_10_3_f': 80.00, 'emr_10_3_f': 71.67, 'fp_10_3': 43,
        'f1_10_2_s': 75.46, 'p_10_2_s': 72.03, 'r_10_2_s': 79.23, 'emr_10_2_s': 73.69, 'cov_2': 91.32, 'ret_2': 8.68,
        'f1_10_3_s': 77.26, 'p_10_3_s': 73.29, 'r_10_3_s': 81.68, 'emr_10_3_s': 73.77, 'cov_3': 91.17, 'ret_3': 8.83,
        'c1_f1_2': 96.33, 'c1_f1_3': 96.28,
        'c2_f1_2': 86.73, 'c2_f1_3': 86.70,
        'c3_f1_2': 95.40, 'c3_f1_3': 95.91,
        'c4_f1_2': 90.34, 'c4_f1_3': 90.57,
        'v5_f1_2': 72.00, 'v5_f1_3': 77.78,
        'v2_f1_2': 81.36, 'v2_f1_3': 81.36,
        'v3_f1_2': 82.54, 'v3_f1_3': 82.54,
        'v4_f1_2': 73.08, 'v4_f1_3': 76.92,
        'v1_f1_2': 65.31, 'v1_f1_3': 65.31,
    }
    create_report(target, metrics_data)

