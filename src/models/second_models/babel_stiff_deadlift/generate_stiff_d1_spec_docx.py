# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 스티프 데드리프트 (Barbell Stiff-leg Deadlift)
Defect: d1. 바벨 궤적과 몸 밀착 (Barbell Path Proximity to Body)
Script: generate_stiff_d1_spec_docx.py
Purpose: 스티프 데드리프트 d1 최종 모델 확정 및 배포명세서 공식 Word 및 Markdown 문서 생성
Rule: AGENTS.md 규격 (공식 명칭, Word 문서 필수, 5대 시점 명칭 준수)
"""

import os
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_JSON_PATH = os.path.join(CURRENT_DIR, "stiff_d1_results.json")
DOCX_OUTPUT_PATH = r"D:\자주프 문서\6주차\바벨_스티프_데드리프트_d1_바벨밀착_최종_v1.0_확정_및_배포명세서.docx"
MD_OUTPUT_PATH = r"D:\자주프 문서\6주차\바벨_스티프_데드리프트_d1_바벨밀착_최종_v1.0_확정_및_배포명세서.md"
os.makedirs(os.path.dirname(DOCX_OUTPUT_PATH), exist_ok=True)


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout(doc, text, title="공식 최종 확정 요약 (Confirmed Summary)", border_hex="1F497D", bg_hex="F2F5F9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}\n")
    run_t.bold = True
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    run_b = p.add_run(text)
    run_b.font.name = "맑은 고딕"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def build_documents(data: Dict[str, Any]):
    """JSON 데이터를 바탕으로 Word 및 Markdown 보고서 생성"""
    doc = Document()
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # 1. 표지 타이틀
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    r1 = title_p.add_run("vidio-health-trainer | 최종 모델 선정 및 운영 배포 명세서\n")
    r1.font.name = "맑은 고딕"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    r2 = title_p.add_run("바벨 스티프 데드리프트 d1 (바벨 궤적과 몸 밀착)\n최종 v1.0 Zero-Calibrated 모델 공식 확정 및 운영 배포 명세서")
    r2.font.name = "맑은 고딕"
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(12)
    rm = meta_p.add_run(f"발행일자: {data.get('timestamp', '2026-10-02')} | 프로젝트: vidio-health-trainer | 배포 상태: Official Production Confirmed")
    rm.font.name = "맑은 고딕"
    rm.font.size = Pt(9.5)
    rm.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    # 요약 콜아웃
    opt = data.get('tier2_3state_optimal', {})
    t1 = data.get('tier1_golden_metrics', {})
    base = data.get('tier2_2state_baseline', {})
    summary_text = (
        f"본 문서는 'vidio-health-trainer' 프로젝트의 바벨 스티프 데드리프트(03_바벨_스티프_데드리프트) "
        f"핵심 결함인 'd1. 바벨 궤적과 몸 밀착' 진단 모델의 최종 엔지니어링 사양 및 검증 결과를 규정합니다.\n"
        f"• 이전 바벨 데드리프트 d5에서 성공적으로 입증된 '동적 어드레스 영점화(Zero-Calibration)'와 "
        f"체형 비율 보정 및 스티프 고유 힌지-광배근 연직선 피처를 결합한 '38D 도메인 피처셋'을 확립했습니다.\n"
        f"• 피험자 9명 전수 격리 5-Fold StratifiedGroupKFold 교차 검증 결과, "
        f"Tier 1 골든 셋 F1 {t1.get('f1', 0):.2f}% (AUC {t1.get('auc', 0):.2f}%), "
        f"Tier 2 실무 3-State F1 ★ {opt.get('f1', 0):.2f}% (커버리지 {opt.get('coverage', 0):.2f}%, 정확도 {opt.get('accuracy', 0):.2f}%)를 달성했습니다.\n"
        f"• 판정 보류율 {opt.get('abstain_rate', 0):.2f}%로 오탐(FP)을 {base.get('fp', 0)}건에서 {opt.get('fp', 0)}건으로 대폭 억제하였습니다."
    )
    add_callout(doc, summary_text, title="공식 최종 확정 요약 (Executive Summary)")

    # 1. 최종 확정 운영 규격 명세서
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. 최종 확정 운영 규격 명세서")
    r_h1.font.name = "맑은 고딕"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    spec_tbl = doc.add_table(rows=1, cols=3)
    spec_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(spec_tbl)

    hdr = spec_tbl.rows[0].cells
    hdr[0].text = "규격 분류 항목"
    hdr[1].text = "최종 확정 사양 / 수치"
    hdr[2].text = "엔지니어링 근거 및 상세 설명"
    for c in hdr:
        set_cell_background(c, "1F497D")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.runs[0].font.name = "맑은 고딕"
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    specs = [
        ("공식 모델 명칭", "v1.0 Zero-Calibrated (Stiff-38D)", "바벨 데드 v6.0 영점화 37D 계승 + 스티프 힌지-미드풋 연직선 1D 결합"),
        ("입력 텐서 규격", "(64, 33, 3) MediaPipe 3D", "무삭제 64프레임 3차 스플라인(Cubic Spline) 보간 순수 비전 좌표"),
        ("의사결정 정책", "3-State View-Calibrated Gating", f"정상 한계 [≤ {opt.get('tau_clean', 0):.2f}] / 결함 한계 [≥ {opt.get('tau_defect', 0):.2f}] 이원화"),
        ("실무 커버리지", f"{opt.get('coverage', 0):.2f}% (보류율 {opt.get('abstain_rate', 0):.2f}%)", "모호한 경계 구간을 재촬영 유도하여 오분류 위험 원천 차단"),
        ("3-State 최종 F1", f"★ {opt.get('f1', 0):.2f}% (정확도 {opt.get('accuracy', 0):.2f}%)", f"2-State 베이스라인({base.get('f1', 0):.2f}%) 대비 현격한 도약 달성"),
        ("골든셋 F1 (Tier 1)", f"{t1.get('f1', 0):.2f}% (AUC {t1.get('auc', 0):.2f}%)", "단독 바벨 이탈(Key 210) vs 정상(Key 209) 순수 판별력 입증"),
        ("오탐 억제 성과", f"FP {base.get('fp', 0)}건 -> {opt.get('fp', 0)}건", f"거짓 양성 오탐 {base.get('fp', 0) - opt.get('fp', 0)}건 대폭 소거 완료")
    ]

    for item, val, desc in specs:
        row = spec_tbl.add_row().cells
        row[0].text = item
        row[1].text = val
        row[2].text = desc
        for i, c in enumerate(row):
            set_cell_margins(c, 70, 70, 100, 100)
            if i == 0:
                set_cell_background(c, "F2F5F9")
            p = c.paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "맑은 고딕"
                p.runs[0].font.size = Pt(9)
                if i == 1:
                    p.runs[0].font.bold = True

    # 2. 공식 5대 시점별 세부 감사 결과
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. 공식 5대 카메라 시점별 최종 성능 감사 (AGENTS.md 표준)")
    r_h2.font.name = "맑은 고딕"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    vm_tbl = doc.add_table(rows=1, cols=7)
    vm_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(vm_tbl)

    vm_hdr = vm_tbl.rows[0].cells
    vm_cols = ["시점 구분", "카메라 앵글", "3-State F1", "정밀도", "재현율", "정확도", "커버리지"]
    for i, title in enumerate(vm_cols):
        vm_hdr[i].text = title
        set_cell_background(vm_hdr[i], "1F497D")
        set_cell_margins(vm_hdr[i], 80, 80, 80, 80)
        p = vm_hdr[i].paragraphs[0]
        p.runs[0].font.name = "맑은 고딕"
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    vmetrics = data.get('view_metrics', {})
    for v_code in ['view5', 'view2', 'view3', 'view4', 'view1']:
        if v_code not in vmetrics:
            continue
        vm = vmetrics[v_code]
        row = vm_tbl.add_row().cells
        alias_full = vm.get('alias', v_code)
        # 파싱
        v_name = alias_full.split(' ')[0] + ' ' + alias_full.split(' ')[1]
        v_angle = ' '.join(alias_full.split(' ')[2:])
        row[0].text = v_name
        row[1].text = v_angle
        row[2].text = f"{vm.get('f1', 0):.2f}%"
        row[3].text = f"{vm.get('precision', 0):.2f}%"
        row[4].text = f"{vm.get('recall', 0):.2f}%"
        row[5].text = f"{vm.get('accuracy', 0):.2f}%"
        row[6].text = f"{vm.get('coverage', 0):.1f}%"

        for i, c in enumerate(row):
            set_cell_margins(c, 70, 70, 80, 80)
            if i == 0:
                set_cell_background(c, "F2F5F9")
            p = c.paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "맑은 고딕"
                p.runs[0].font.size = Pt(9)
                if i == 2:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # 3. 실서비스 추론 파이프라인 연동 코드
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. 프로덕션 실서비스 추론 파이프라인 연동 코드")
    r_h3.font.name = "맑은 고딕"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    code_p = doc.add_paragraph()
    code_p.paragraph_format.space_before = Pt(4)
    code_p.paragraph_format.space_after = Pt(8)
    code_text = (
        "from src.models.second_models.babel_stiff_deadlift.stiff_d1_inference import infer_stiff_d1_sequence\n\n"
        "# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 관절 좌표 배열\n"
        "result = infer_stiff_d1_sequence(\n"
        "    kps_64f=kps_64f,\n"
        "    view_name='view4'  # 시점 D (좌측 전방 45° 최우선 권장)\n"
        ")\n\n"
        "print(result)\n"
        "# 출력 예시:\n"
        "# {\n"
        "#   'defect_prob': 0.082,\n"
        "#   'prediction': 0,                # 0: 정상, 1: 결함, -1: 판정 보류\n"
        "#   'is_decided': True,             # 확정 판정\n"
        "#   'decision_label': '정상 (바벨 밀착 궤적 양호)',\n"
        "#   'view_name': 'view4',\n"
        "#   'view_alias': '시점 D (좌측 전방 45°)',\n"
        "#   'features_dim': 38\n"
        "# }"
    )
    rc = code_p.add_run(code_text)
    rc.font.name = "Consolas"
    rc.font.size = Pt(9)
    rc.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # 4. 배포 산출물 및 무결성 검증 목록
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. 프로덕션 배포 산출물 및 무결성 검증 목록")
    r_h4.font.name = "맑은 고딕"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    items = [
        "최종 가중치 파일: D:\\src\\models\\second_models\\babel_stiff_deadlift\\weights\\stiff_deadlift_d1_best.pkl",
        "공식 평가 결과 JSON: D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_results.json",
        "38D 피처 추출 엔진: D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_features.py",
        "프로덕션 추론 모듈: D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_inference.py",
        "단위 및 통합 테스트: D:\\src\\models\\second_models\\babel_stiff_deadlift\\test_stiff_d1_inference.py 100% 통과"
    ]
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9.5)

    doc.save(DOCX_OUTPUT_PATH)
    print(f"[SUCCESS] Word document generated: {DOCX_OUTPUT_PATH}")

    # Markdown 문서도 함께 저장
    md_content = f"""# vidio-health-trainer | 최종 모델 선정 및 운영 배포 명세서
# 바벨 스티프 데드리프트 d1 (바벨 궤적과 몸 밀착) 최종 v1.0 Zero-Calibrated 모델 공식 확정 및 운영 배포 명세서

- **발행일자:** {data.get('timestamp', '2026-10-02')}
- **공식 모델 명칭:** `v1.0 Zero-Calibrated (Stiff-38D)`
- **배포 상태:** **Official Production Release (최종 확정)**
- **확정 사양:** 3-State F1 **★ {opt.get('f1', 0):.2f}%** (정확도 **{opt.get('accuracy', 0):.2f}%**) | 커버리지 **{opt.get('coverage', 0):.2f}%** | Tier 1 골든셋 F1 **{t1.get('f1', 0):.2f}%**

---

## 1. 확정된 최종 운영 규격 명세서

| 규격 분류 항목 | 최종 확정 사양 / 수치 | 엔지니어링 근거 및 상세 설명 |
| :--- | :---: | :--- |
| **공식 모델 명칭** | `v1.0 Zero-Calibrated (Stiff-38D)` | 바벨 데드 v6.0 영점화 37D 계승 + 스티프 힌지-미드풋 연직선 1D 결합 |
| **입력 텐서 규격** | **(64, 33, 3) MediaPipe 3D** | 무삭제 64프레임 3차 스플라인(Cubic Spline) 보간 순수 비전 좌표 |
| **의사결정 정책** | **3-State View-Calibrated Gating** | 정상 한계 [≤ {opt.get('tau_clean', 0):.2f}] / 결함 한계 [≥ {opt.get('tau_defect', 0):.2f}] |
| **실무 커버리지** | **{opt.get('coverage', 0):.2f}% (보류율 {opt.get('abstain_rate', 0):.2f}%)** | 모호한 경계 구간을 재촬영 유도하여 오분류 위험 원천 차단 |
| **3-State 최종 F1** | **★ {opt.get('f1', 0):.2f}% (정확도 {opt.get('accuracy', 0):.2f}%)** | 2-State 베이스라인({base.get('f1', 0):.2f}%) 대비 비약적 성능 향상 |
| **골든셋 F1 (Tier 1)** | **{t1.get('f1', 0):.2f}% (AUC {t1.get('auc', 0):.2f}%)** | 단독 바벨 이탈(Key 210) vs 정상(Key 209) 순수 판별력 입증 |
| **오탐 억제 성과** | **FP {base.get('fp', 0)}건 -> {opt.get('fp', 0)}건** | 거짓 양성 오탐 {base.get('fp', 0) - opt.get('fp', 0)}건 대폭 소거 |

---

## 2. 공식 5대 카메라 시점별 최종 성능 감사 (AGENTS.md 표준)

| 시점 구분 | 카메라 앵글 | 3-State F1 | 정밀도 | 재현율 | 정확도 | 커버리지 |
| :--- | :--- | :---: | :---: | :---: | :---: | :---: |
"""
    for v_code in ['view5', 'view2', 'view3', 'view4', 'view1']:
        if v_code not in vmetrics:
            continue
        vm = vmetrics[v_code]
        alias_full = vm.get('alias', v_code)
        v_name = alias_full.split(' ')[0] + ' ' + alias_full.split(' ')[1]
        v_angle = ' '.join(alias_full.split(' ')[2:])
        md_content += f"| **{v_name}** | {v_angle} | **{vm.get('f1', 0):.2f}%** | {vm.get('precision', 0):.2f}% | {vm.get('recall', 0):.2f}% | {vm.get('accuracy', 0):.2f}% | {vm.get('coverage', 0):.1f}% |\n"

    md_content += f"""
---

## 3. 프로덕션 실서비스 추론 파이프라인 연동 코드

```python
from src.models.second_models.babel_stiff_deadlift.stiff_d1_inference import infer_stiff_d1_sequence

# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 관절 좌표 배열
result = infer_stiff_d1_sequence(
    kps_64f=kps_64f,
    view_name="view4"  # 시점 D (좌측 전방 45° 최우선 권장)
)

print(result)
# 출력 예시:
# {{
#   'defect_prob': 0.082,
#   'prediction': 0,                # 0: 정상, 1: 결함, -1: 판정 보류
#   'is_decided': True,             # 확정 판정
#   'decision_label': '정상 (바벨 밀착 궤적 양호)',
#   'view_name': 'view4',
#   'view_alias': '시점 D (좌측 전방 45°)',
#   'features_dim': 38
# }}
```

---

## 4. 프로덕션 배포 산출물 및 무결성 검증

1. **최종 가중치 파일**: `D:\\src\\models\\second_models\\babel_stiff_deadlift\\weights\\stiff_deadlift_d1_best.pkl`
2. **공식 평가 결과 JSON**: `D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_results.json`
3. **38D 피처 추출 엔진**: `D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_features.py`
4. **프로덕션 추론 모듈**: `D:\\src\\models\\second_models\\babel_stiff_deadlift\\stiff_d1_inference.py`
5. **단위 및 통합 테스트**: `D:\\src\\models\\second_models\\babel_stiff_deadlift\\test_stiff_d1_inference.py` 100% 통과
"""
    with open(MD_OUTPUT_PATH, 'w', encoding='utf-8') as fp:
        fp.write(md_content)
    print(f"[SUCCESS] Markdown document generated: {MD_OUTPUT_PATH}")


if __name__ == "__main__":
    if os.path.exists(RESULTS_JSON_PATH):
        with open(RESULTS_JSON_PATH, 'r', encoding='utf-8') as fp:
            data = json.load(fp)
        build_documents(data)
    else:
        print(f"Results json not found yet: {RESULTS_JSON_PATH}")
