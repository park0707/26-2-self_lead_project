# -*- coding: utf-8 -*-
"""
vidio-health-trainer / 바벨 스티프 데드리프트 sd1
Report: sd1 v6 구성요소 전수 절제(Ablation) 및 v7 최적화 / v6 vs v7 성능 비교 보고서
Output: D:/자주프 문서/6주차/vidio-health-trainer_스티프데드리프트_sd1_v7_절제최적화_성능비교보고서.docx (+pdf)
"""
import os, json
import numpy as np
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

CUR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = r"D:\자주프 문서\6주차"
BASE = "vidio-health-trainer_스티프데드리프트_sd1_v7_절제최적화_성능비교보고서"
DOCX_PATH, PDF_PATH = os.path.join(OUT_DIR, BASE + ".docx"), os.path.join(OUT_DIR, BASE + ".pdf")
NAVY = RGBColor(27, 54, 93)

S3 = json.load(open(os.path.join(CUR, 'v7_stage3_candidates.json'), encoding='cp949'))
S2 = json.load(open(os.path.join(CUR, 'v7_stage2_ablation.json'), encoding='cp949'))
V6 = json.load(open(os.path.join(CUR, 'stiff_deadlift_knee_v6_benchmark_results.json'), encoding='utf-8'))
V7 = json.load(open(os.path.join(CUR, 'stiff_deadlift_knee_v7_benchmark_results.json'), encoding='utf-8'))


def shade(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hexcolor}"/>'))


def run(p, text, bold=False, size=9.5, color=None):
    r = p.add_run(text); r.bold = bold; r.font.name = "맑은 고딕"; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), "맑은 고딕")
    if color is not None: r.font.color.rgb = color
    return r


def heading(doc, text):
    h = doc.add_heading(level=1); r = h.add_run(text); r.font.name = "맑은 고딕"; r.font.color.rgb = NAVY
    r._element.rPr.rFonts.set(qn('w:eastAsia'), "맑은 고딕")


def para(doc, text, bold=False, size=9.5, after=5):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(after)
    run(p, text, bold=bold, size=size); return p


def callout(doc, title, text, alert=False):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, "FFF3CD" if alert else "F0F4F8")
    p = c.paragraphs[0]
    run(p, f"[{title}] ", bold=True, color=RGBColor(180, 40, 40) if alert else NAVY)
    run(p, text, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(doc, headers, rows, widths=None, hl_rows=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, "1B365D"); p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, h, bold=True, size=8.5, color=RGBColor(255, 255, 255))
    for r_i, row in enumerate(rows, 1):
        for c_i, v in enumerate(row):
            c = t.rows[r_i].cells[c_i]
            if r_i in hl_rows: shade(c, "EBF3FB")
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run(p, str(v), size=8.3, bold=(r_i in hl_rows))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def ms(res_list, key):
    a = np.array([r[key] for r in res_list]); return a.mean(), a.std()


def s3(name, lab, key):
    m, s = ms(S3[name]['per_seed'][lab], key); return f"{m:.1f}±{s:.1f}"


def s3d(name, lab, key):
    b = np.array([r[key] for r in S3['v6 (기준)']['per_seed'][lab]]); x = np.array([r[key] for r in S3[name]['per_seed'][lab]])
    d = x - b; return f"{d.mean():+.1f}", f"{np.mean(d > 0) * 100:.0f}%"


def build():
    n_v7 = "v7a: 중복2개 제거(12D)"
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7); s.left_margin = s.right_margin = Inches(0.75)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "vidio-health-trainer\n바벨 스티프 데드리프트 sd1(무릎 고정 위반)\nv6 구성요소 전수 절제(Ablation) 및 v7 최적화 / 성능 비교 보고서", bold=True, size=16, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "작성일: 2026-10-03 | 데이터: AI-Hub 피험자 1~9 전수, 유효 668시퀀스 | 검증: 피험자 격리(LOSO / 5-Fold)", size=8.5, color=RGBColor(100, 100, 100))

    a_f1, b_f1 = ms(S3['v6 (기준)']['per_seed']['clean'], 'f1d'), ms(S3[n_v7]['per_seed']['clean'], 'f1d')
    callout(doc, "핵심 결론",
            "v6의 구성요소 44개 설정을 전수 절제한 결과, 일관되게 이득이 확인된 변경은 '상관 r≥0.9 중복 정규화 피처 2개 제거(14D→12D)' 하나뿐이었습니다. "
            "이를 적용한 v7은 시드 10개 쌍체 검증에서 AUC +1.3%p, 특이도 +2.0%p, 결함 정밀도 +3.8%p, 결함 F1 +3.1%p(정제 라벨 기준)로 개선되었고, 원본 AI-Hub 라벨 기준과 다른 분할 방식에서도 같은 방향이었습니다. "
            "다만 개선 폭은 작고(AUC 약 +1%p), 시드 잡음(F1 ±3~4%p)과 비슷한 수준이므로 '확정적 도약'이 아니라 '복잡도를 줄이면서 일관되게 소폭 개선'으로 해석해야 합니다.")
    callout(doc, "반드시 알아야 할 평가 한계 (v6 보고 수치 정정)",
            "① v6의 3-State 결함 F1 57.95%는 분할 시드 42의 유리한 분할 한 번의 값입니다. 분할 시드 10회 평균은 49.8±5.3%입니다(9명 피험자로는 서로 다른 분할이 사실상 2가지뿐이라 58.0 또는 46.3 두 값만 나옴). "
            "② 정제 라벨로 채점하면 정제 규칙이 같은 피처로 테스트 라벨을 바꾸는 순환이 있어 낙관적입니다. 원본 AI-Hub 라벨로 채점하면 v6 AUC는 55.6%, 결함 F1은 36.6%입니다. "
            "③ 668개 시퀀스는 실제로 고유 텐서 203개입니다(시점 1·2·3·5가 동일 텐서). 시점별 성능이 4개 시점에서 동일하게 나온 이유입니다.", alert=True)

    heading(doc, "1. 조사 방법 (과적합을 피하기 위한 설계)")
    for t in [
        "• 평가 분할: LOSO(피험자 1명씩 제외, 9-Fold, 결정적 분할) + 모델 시드 반복. 기존 5-Fold는 분할 자체가 운에 좌우되어 보조 확인용으로만 사용.",
        "• 이중 채점: 원본 AI-Hub 라벨 / v6 정제 라벨 양쪽에서 모두 악화되지 않는 변경만 인정.",
        "• 선택 지표: F1은 3-State 밴드 때문에 '결함을 더 많이 부르는' 방향으로 쉽게 부풀려집니다(예: iters=100은 F1 +11.9%p이나 특이도 66.7%로 하락). 따라서 순위 품질인 AUC를 주 지표로, 특이도·정밀도·보류율을 방어선으로 사용.",
        "• 선택 편향 통제: 결과를 보고 고르는 방식은 203개 고유 표본에서 과적합 위험이 크므로, 상관 r≥0.9 같이 '결과를 보기 전에 정할 수 있는 원칙'으로 후보를 사전 지정. 탐색적 후보는 별도 표기.",
        "• 확인: 선택에 쓰지 않은 새 모델 시드 10개(100~109)로 v6와 쌍체 비교(차이 평균, 이긴 시드 비율).",
    ]:
        para(doc, t, after=2)

    heading(doc, "2. 전수 절제 결과 (44개 설정, v6=14D 대비 Δ, 시드 3개 평균)")
    base2 = S2["BASE v6 (14D/v6정제/CB d4)"]['res']

    def d2(name, lab, key):
        b = np.mean([r[key] for r in base2[lab]]); x = np.mean([r[key] for r in S2[name]['res'][lab]]); return x - b

    sel = [
        ("-정규화중복(rel,norm_y)", "중복 정규화 2개 제거", "채택(→v7a)"),
        ("-신규v6(kadr,excess,zrom)", "신규 v6 피처 5개 전체 제거", "유지(제거 시 AUC -4.6)"),
        ("-수직변위계(3,7)", "무릎 수직변위 계열 제거", "유지(AUC -2.7)"),
        ("+pli", "pli 재투입", "기각(악화)"),
        ("+hip_hinge_isolation", "hip_hinge 재투입", "기각(악화)"),
        ("+femur_shank_ratio", "femur_shank_ratio 재투입", "기각(악화)"),
        ("19D 전체", "19D 전체 사용", "기각(F1 -6.4, AUC -2.0)"),
        ("정제=none", "라벨 정제 없음", "기각(정제라벨 점수 급락, 원본 F1만 상승)"),
        ("정제=strict", "정제 strict", "v7a와 결합 시 이득 없음"),
        ("정제=loose", "정제 loose", "기각(악화)"),
        ("CB depth=5", "트리 깊이 5", "기각(분할 방식에 따라 F1 악화)"),
        ("CB iters=100,lr=0.04", "반복 100", "기각(F1 상승이나 특이도 66.7%)"),
        ("CB class_weight=None", "클래스 가중치 제거", "기각(F1 -16.5)"),
    ]
    rows = [[lab, f"{d2(n,'orig','auc'):+.1f}", f"{d2(n,'clean','auc'):+.1f}", f"{d2(n,'orig','f1d'):+.1f}", f"{d2(n,'clean','f1d'):+.1f}", f"{d2(n,'orig','spec'):+.1f}", verdict] for n, lab, verdict in sel]
    table(doc, ["변경", "AUC Δ(원본)", "AUC Δ(정제)", "F1 Δ(원본)", "F1 Δ(정제)", "특이도 Δ", "판정"], rows,
          widths=[1.7, 0.7, 0.7, 0.7, 0.7, 0.6, 1.9], hl_rows=(1,))
    para(doc, "※ 피처 1개씩 제거(LOFO 14개) 결과는 대부분 ±2%p로 시드 잡음 범위 안이어서 단독 제거는 근거가 약했습니다. "
              "유일하게 잡음을 넘은 신호는 '서로 겹치는 두 피처(rel_bottom_knee_delta↔bottom_knee_delta_mean r=0.938, norm_knee_disp_y↔knee_disp_y_rom r=0.976)를 함께 빼는 것'이었습니다. "
              "한 쌍의 한쪽만 빼면 다른 쪽이 정보를 대신하므로 단독 LOFO에서는 보이지 않았던 것입니다.", size=8.8)

    heading(doc, "3. v7 구성")
    table(doc, ["항목", "v6", "v7", "비고"], [
        ["피처", "14D", "12D", "rel_bottom_knee_delta, norm_knee_disp_y 제거"],
        ["라벨 정제", "Asymmetric CleanLab 103건", "동일", "strict 등 변경은 이득 없음"],
        ["모델", "CatBoost d4, 200iter", "동일", "하이퍼파라미터 변경은 모두 기각"],
        ["3-State", "Conformal 보류율 ≤10%", "동일", "tau 0.4153~0.5847"],
        ["신규 피처 추가", "-", "없음", "추가는 모두 악화(과적합)"],
    ], widths=[1.3, 1.8, 1.4, 2.6])

    heading(doc, "4. v6 vs v7 성능 비교 (세 가지 평가 방식)")
    para(doc, "4-1. LOSO 9-Fold, 새 모델 시드 10개 (평균±표준편차, Δ=v7−v6 쌍체, 승=v7이 더 좋은 시드 비율)", bold=True)
    rows = []
    for lab, lname in [('clean', '정제 라벨'), ('orig', '원본 라벨')]:
        for key, kn in [('f1d', '3-State 결함 F1'), ('prec', '결함 정밀도'), ('spec', '특이도'), ('auc', 'ROC-AUC')]:
            d, w = s3d(n_v7, lab, key)
            rows.append([f"{kn} ({lname})", s3('v6 (기준)', lab, key), s3(n_v7, lab, key), d, w])
    rows.append(["평가 보류율", f"{ms(S3['v6 (기준)']['per_seed']['clean'],'abst')[0]:.1f}", f"{ms(S3[n_v7]['per_seed']['clean'],'abst')[0]:.1f}", "-", "-"])
    table(doc, ["지표", "v6", "v7", "Δ", "v7 승률"], rows, widths=[2.2, 1.2, 1.2, 0.8, 0.9], hl_rows=(4, 8))

    para(doc, "4-2. 다른 분할 방식: 반복 StratifiedGroupKFold(분할 시드 10회, 모델 시드 42)", bold=True)
    table(doc, ["지표", "v6", "v7", "Δ", "v7 승률"], [
        ["3-State 결함 F1 (정제)", "49.8", "51.8", "+2.0", "70%"], ["3-State 결함 F1 (원본)", "36.6", "37.4", "+0.8", "70%"],
        ["ROC-AUC (정제)", "79.6", "81.3", "+1.7", "70%"], ["ROC-AUC (원본)", "55.6", "55.9", "+0.3", "70%"],
        ["특이도 (원본)", "78.8", "80.7", "+1.9", "-"]], widths=[2.2, 1.2, 1.2, 0.8, 0.9])

    para(doc, "4-3. 기존 보고 방식 재현: 5-Fold 분할 시드 42 + 모델 시드 42 (단일 분할)", bold=True)
    a, b = V6['tier2_3state_metrics'], V7['tier2_3state_metrics']
    table(doc, ["지표", "v6", "v7", "Δ"], [
        ["3-State 결함 F1", f"{a['selective_defect_f1']*100:.2f}%", f"{b['selective_defect_f1']*100:.2f}%", f"{(b['selective_defect_f1']-a['selective_defect_f1'])*100:+.2f}%p"],
        ["3-State 정상 F1", f"{a['selective_clean_f1']*100:.2f}%", f"{b['selective_clean_f1']*100:.2f}%", f"{(b['selective_clean_f1']-a['selective_clean_f1'])*100:+.2f}%p"],
        ["특이도", f"{a['selective_specificity']*100:.2f}%", f"{b['selective_specificity']*100:.2f}%", f"{(b['selective_specificity']-a['selective_specificity'])*100:+.2f}%p"],
        ["평가 보류율", f"{a['abstention_rate']:.2f}%", f"{b['abstention_rate']:.2f}%", f"{b['abstention_rate']-a['abstention_rate']:+.2f}%p"],
        ["ROC-AUC", f"{V6['tier2_2state_metrics']['roc_auc']*100:.2f}%", f"{V7['tier2_2state_metrics']['roc_auc']*100:.2f}%", f"{(V7['tier2_2state_metrics']['roc_auc']-V6['tier2_2state_metrics']['roc_auc'])*100:+.2f}%p"]],
        widths=[2.2, 1.4, 1.4, 1.2])
    callout(doc, "해석 주의",
            "이 단일 분할에서는 v7의 결함 F1이 v6보다 3.69%p 낮게 나옵니다. 그러나 이 분할은 v6가 49.8% 평균보다 8%p 높게 나온 유리한 분할(분포의 상위 쪽)이며, "
            "같은 분할에서 v7의 특이도는 +2.69%p 높습니다. 분할 10회·시드 10회 어느 쪽 평균에서도 v7이 우세하므로, 한 번의 분할 결과로 v6가 더 낫다고 결론 내리지 않습니다.", alert=True)

    heading(doc, "5. 탐색 후보 및 기각 사유")
    rows = []
    for n in ["v7b: v7a + 정제 strict", "v7c: v7a + depth5", "v7d: v7a + strict + depth5", "탐색: v7a - kadr_max", "탐색: v7a - address_knee_angle", "탐색: v7a - bottom_knee_delta_mean"]:
        d, w = s3d(n, 'orig', 'auc'); d2_, _ = s3d(n, 'clean', 'f1d')
        sa = S3[n]['subj_auc']['orig']; ba = S3['v6 (기준)']['subj_auc']['orig']
        imp = sum(1 for k, v in sa.items() if v > ba.get(k, -1))
        rows.append([n, d, d2_, f"{imp}/{len(sa)}"])
    table(doc, ["후보", "AUC Δ(원본)", "F1 Δ(정제)", "AUC 개선 피험자"], rows, widths=[3.0, 1.0, 1.0, 1.4])
    para(doc, "• 정제 strict·depth5: v7a 위에 얹어도 일관된 이득이 없거나 분할 방식에 따라 F1이 악화되어 기각.\n"
              "• 탐색 후보(결과를 보고 고른 것): address_knee_angle 제거는 9명 중 7명에서 AUC가 개선되어 유망하나, v6 기준 절제(Stage 2)에서는 F1이 -1.8%p였고 v7a 위에서는 +1.0%p로 방향이 일관되지 않아 "
              "우연일 가능성을 배제할 수 없습니다. 새 데이터가 확보된 뒤 재검증할 후보로만 기록합니다.", size=8.8)

    heading(doc, "6. 한계와 권고")
    for t in [
        "• 이 모델의 근본 한계는 피처/모델이 아니라 데이터입니다. 모든 피처의 단독 AUC가 50~57%, 원본 라벨 기준 모델 AUC가 약 57%로, 원본 라벨에 대해서는 피처 선택·튜닝으로 올릴 수 있는 폭이 작습니다.",
        "• 고유 시퀀스가 203개(피험자 9명)뿐이고 시점 4개가 동일 텐서입니다. 시점별 성능 표는 시점 간 비교 의미가 없으며, 시점 B·C·E 등이 서로 다른 영상에서 추출되도록 전처리 단계를 점검할 것을 권고합니다(전처리 파이프라인 변경이므로 승인 후 진행).",
        "• 성능 보고는 앞으로 '단일 분할 수치'가 아니라 LOSO 시드 평균(±표준편차)과 원본 라벨 기준 값을 함께 표기할 것을 권고합니다.",
    ]:
        para(doc, t, after=2)

    heading(doc, "7. 배포 상태")
    para(doc, "• 생성: weights/sd1_v7_bundle.pkl, sd1_v7_classifier.py, train_and_eval_stiff_knee_v7_pipeline.py, test_sd1_v7_classifier.py (단위 테스트 5개 통과)\n"
              "• 공식 프로덕션(sd1_classifier.py, sd1_bundle.pkl)은 v6 그대로 유지했습니다. v7 승격 여부는 사용자 승인 후 진행합니다.\n"
              "• 실험 재현 스크립트: v7_ablation_core.py, v7_stage1_noise_floor.py, v7_stage2_ablation.py, v7_stage3_candidates.py")
    doc.save(DOCX_PATH); print("[+] DOCX:", DOCX_PATH)
    try:
        from docx2pdf import convert
        convert(DOCX_PATH, PDF_PATH); print("[+] PDF:", PDF_PATH)
    except Exception as e:
        print("[-] PDF 변환 실패:", e)


if __name__ == '__main__':
    build()
