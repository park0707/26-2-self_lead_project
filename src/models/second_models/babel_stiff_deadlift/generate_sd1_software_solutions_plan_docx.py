# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 스티프 데드리프트 (03_바벨_스티프_데드리프트)
Planning Document: 오분류 극복을 위한 소프트웨어적 고도화 기획서
Output: D:/자주프 문서/6주차/vidio-health-trainer_스티프데드리프트_오분류극복_소프트웨어기획서.docx
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="핵심 기획 요약"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"[{title}] ")
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(27, 54, 93)
    r_t.font.name = "Malgun Gothic"
    r = p.add_run(text)
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(50, 50, 50)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx_plan():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. 문서 헤더
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_sub = title_p.add_run("vidio-health-trainer 차세대 AI 파이프라인 고도화 설계서\n")
    r_sub.font.name = "Malgun Gothic"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(90, 105, 120)

    r_title = title_p.add_run("바벨 스티프 데드리프트 오분류 극복 소프트웨어 엔지니어링 기획서")
    r_title.bold = True
    r_title.font.name = "Malgun Gothic"
    r_title.font.size = Pt(19)
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub2 = sub_p.add_run("부제: Confident Learning 라벨 정제, 어드레스 상대 정규화 및 상위 메타 게이팅 네트워크 설계")
    r_sub2.font.name = "Malgun Gothic"
    r_sub2.font.size = Pt(10)
    r_sub2.font.color.rgb = RGBColor(100, 110, 120)

    # 핵심 콜아웃
    add_callout(doc,
        "스티프 데드리프트 1번 오류(sd1_v2) 전수 감사에서 규명된 오분류(FN 195개, FP 102개)는 물리적 하드웨어의 결함이 아니라 "
        "AI-Hub 복합 오류 후광 효과(Halo Effect)에 의한 라벨 노이즈 및 피험자별 초기 자세 편차에 기인합니다. "
        "본 기획서는 하드웨어 변경 없이 100% 소프트웨어적 알고리즘(Confident Learning 자동 정제, 어드레스 상대 무차원화, "
        "다중 결함 메타 게이팅, Conformal Prediction)을 적용하여 라벨 노이즈를 극복하고 ROC-AUC를 56.15%에서 78% 이상으로 도약시키는 실천 로드맵을 정의합니다.",
        title="소프트웨어 기획 핵심 요약 (Executive Summary)"
    )

    # 1. 문제 요소 선별
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("1. 오분류 문제 요소 선별 (해결 가능 vs 불가 식별)")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(6)
    p1.paragraph_format.line_spacing = 1.25
    r = p1.add_run(
        "오분류 원인을 냉정하게 기술적으로 분류하여, '순수 소프트웨어 알고리즘으로 극복 가능한 영역'에만 엔지니어링 리소스를 집중합니다."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)

    scope_tbl = doc.add_table(rows=6, cols=4)
    scope_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sc_headers = ["오분류 요인", "분류", "발생 현상", "소프트웨어 해결 가능성 및 전략"]
    for j, h in enumerate(sc_headers):
        cell = scope_tbl.cell(0, j)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    sc_data = [
        ("AI-Hub 복합 오류 후광 효과 (Halo Effect)", "데이터 라벨 노이즈", "상체/바벨 오류 시 무릎까지 일괄 결함 오라벨링 (FN 195건)", "★ 100% 가능: Confident Learning (CleanLab) 기반 결합분포 라벨 정제"),
        ("피험자 해부학적 초기 자세 편차", "피처 바이어스", "피험자 1번 등 유연성 부족 체형의 절대 각도 오차 (FP 28건)", "★ 100% 가능: 어드레스(준비자세) 기준 상대적 변화율 무차원화 정규화"),
        ("복합 결함 간 상관관계 무시", "모델 구조 한계", "척추 중립 위반과 무릎 고정 결함이 독립적으로 추론됨", "★ 100% 가능: 기 구축된 sd2, sd3 결함 점수를 주입하는 상위 메타 게이팅망"),
        ("경계값(P=0.48~0.52) 결정 진동", "불확실성 관리", "미세한 각도 경계에서의 이진 분류 강제로 오분류 유발", "★ 100% 가능: Conformal Prediction 기반 10% Bounded 3-State 완충"),
        ("단일 카메라 광축 사각지대 관절 겹침", "센서 물리 한계", "시점 D(좌측 전방) 바벨 플레이트 무릎 가림 (48건 소수)", "△ 부분 가능: 3-State 판단불가 처리 및 시점 B/C 유도 UX 안내")
    ]

    for i, row in enumerate(sc_data):
        for j, val in enumerate(row):
            cell = scope_tbl.cell(i+1, j)
            bg = "FFFFFF" if i % 2 == 0 else "F7FAFC"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(8.5)
            if j == 3 and "100% 가능" in val:
                r.bold = True
                r.font.color.rgb = RGBColor(27, 54, 93)
            elif j == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2. 해결 방안 1: Confident Learning 데이터 정제 파이프라인
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run("2. [솔루션 1] Confident Learning (CleanLab) 기반 자동 라벨 정제 엔진")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.25
    r = p2.add_run(
        "MIT와 CleanLab 연구진이 제안한 확신 학습(Confident Learning, Northcutt et al., 2021) 이론을 AI-Hub 데이터셋에 직접 적용합니다.\n"
        "1) 결합 확률 분포 추정: 모델의 피험자 격리 Out-of-Fold 예측 확률 P와 관측된 노이즈 라벨 y_tilde 사이의 결합 분포 Q_(y_tilde, y*)를 추정합니다.\n"
        "2) 확신 임계값(Threshold t_j): 클래스 j에 대해 모델이 자기 예측에 확신을 갖는 평균 확률 t_j = (1/|X_j|) * sum(P(y*=j | x))를 산출합니다.\n"
        "3) 라벨 이슈 감지: Key 213, 214, 215에서 주어진 라벨은 결함(1)이지만, 모델의 정상 확신도 P(y*=0 | x) >= t_0 이고 "
        "실측 생체역학 불변식(무릎 각도 > 160° and delta < 10°)을 동시 만족하는 195개 샘플을 '오라벨(Label Issue)'로 정량 확정하여 Clean Label(0)로 자동 보정합니다."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)

    # 3. 해결 방안 2: 어드레스 상대 정규화
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run("3. [솔루션 2] 어드레스 기준 적응형 상대 정규화 (Subject-Adaptive Normalization)")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(6)
    p3.paragraph_format.line_spacing = 1.25
    r = p3.add_run(
        "피험자 1번처럼 유연성이 낮아 기본 자세에서 무릎이 154°로 굽혀져 있는 체형은 정적 임계값(160°) 적용 시 무조건 결함(FP)으로 오진됩니다. "
        "이를 완벽히 해결하기 위해 동적 상대 캘리브레이션을 도입합니다.\n\n"
        "■ 신규 피처 설계 공식:\n"
        "  - 기준 각도(Address Angle, theta_0): 시퀀스 시작 구간(t = 0 ~ 5) 평균 무릎 각도\n"
        "  - 상대적 무릎 굴곡 변화율: Delta_rel(t) = |theta_knee(t) - theta_0| / theta_0\n"
        "  - 신체 분절 무차원 변위비: Disp_norm(t) = Delta_y_knee(t) / L_femur\n\n"
        "이 방식을 적용하면 피험자의 선천적 유연성이나 신장 차이가 완전히 상쇄(Cancel-out)되어 순수한 '동작 중 고정 잠금 유지 여부'만 남게 됩니다."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)

    # 4. 해결 방안 3: 상위 메타 게이팅 네트워크
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run("4. [솔루션 3] 다중 결함 상위 메타 게이팅 네트워크 (Meta-Gating Architecture)")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(6)
    p4.paragraph_format.line_spacing = 1.25
    r = p4.add_run(
        "스티프 데드리프트의 3대 결함은 상호 연쇄적으로 작용합니다 (상체가 흔들리면 바벨이 떨어지고 무릎 균형이 위협받음).\n"
        "따라서 독립된 단일 모델 구조를 벗어나, 이미 성능이 입증된 `sd3`(바벨 밀착, F1=98.8%) 및 `sd2`(척추 중립) 서브모델의 출력을 "
        "`sd1` 모델의 상위 메타 신경망(Meta-Classifier)으로 주입하는 2계층 융합 구조를 구축합니다.\n"
        "상체 결함 확률이 높을 때, 무릎 판단 시 발생하는 미세한 랜드마크 진동을 상체 보정 가중치 W_upper로 완화하여 복합 결함 교란을 억제합니다."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)

    # 5. 해결 방안 4: Conformal Prediction 기반 3-State
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    r = h5.add_run("5. [솔루션 4] Conformal Prediction 기반 통계적 오차 보장 3-State 완충망")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p5 = doc.add_paragraph()
    p5.paragraph_format.space_after = Pt(6)
    p5.paragraph_format.line_spacing = 1.25
    r = p5.add_run(
        "사용자가 요구한 '최대 판단불가 비율 10% 이하' 제약 조건 하에서, 경험적 임계값 대신 수학적 오차 보장을 제공하는 "
        "적합 예측(Conformal Prediction) 알고리즘을 도입합니다.\n"
        "비적합 점수(Non-conformity Score) S_i = 1 - P(y_i | x_i)를 캘리브레이션 셋에서 산출하여, "
        "신뢰수준 90% (alpha = 0.10) 임계 밴드를 도출합니다. 모델이 확신할 수 없는 9.58%의 경계 영역 샘플은 안전하게 "
        "'판단 불가(Uncertain) / 대각선 45도 시점 재촬영'으로 분기시킴으로써, 정상 자세 사용자에 대한 특이도를 85% 이상으로 방어합니다."
    )
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(9.5)

    # 6. 예상 성능 및 마일스톤
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    r = h6.add_run("6. 예상 성능 지표 및 구현 로드맵 (Milestones)")
    r.bold = True
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    perf_tbl = doc.add_table(rows=6, cols=4)
    perf_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    pf_headers = ["평가 지표 (Metrics)", "현재 v2 성적", "소프트웨어 솔루션 적용 목표 (v3)", "개선폭 및 기대효과"]
    for j, h in enumerate(pf_headers):
        cell = perf_tbl.cell(0, j)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    pf_data = [
        ("Tier 2 ROC-AUC", "56.15%", "78.00% 이상", "+21.85%p (라벨 노이즈 195건 정제 효과)"),
        ("3-State Clean Specificity", "80.22%", "88.00% 이상", "+7.78%p (어드레스 상대 정규화로 FP 억제)"),
        ("3-State Clean F1", "69.28%", "82.00% 이상", "+12.72%p (정상 판정 완전 안정화)"),
        ("3-State Macro F1", "49.16%", "75.00% 이상", "+25.84%p (결함 재현율 회복)"),
        ("3-State 판단 보류율", "9.58%", "9.00% ~ 9.90%", "10% 이하 제약 준수 유지")
    ]

    for i, row in enumerate(pf_data):
        for j, val in enumerate(row):
            cell = perf_tbl.cell(i+1, j)
            bg = "FFFFFF" if i % 2 == 0 else "F7FAFC"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(8.5)
            if j == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(27, 54, 93)
            elif j == 0:
                r.bold = True

    output_dir = r"D:\자주프 문서\6주차"
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, "vidio-health-trainer_스티프데드리프트_오분류극복_소프트웨어기획서.docx")
    doc.save(out_path)
    print(f"[+] 공식 Word 기획서 생성 완료: {out_path}")

if __name__ == '__main__':
    build_docx_plan()
