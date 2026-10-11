# -*- coding: utf-8 -*-
"""sd2(척추) 고도화 실험 보고서 (.docx + .md) - 실측 수치만 사용."""
import os
from docx import Document
from docx.shared import Inches
from generate_integrated_v4_report_docx import add_header, add_callout
from generate_integrated_corrected_report_docx import add_table, add_bullets, md_table, add_note

D = r"D:\자주프 문서\7주차"
N = "바벨_스티프_데드리프트_sd2_척추모델_MiniRocket앙상블_고도화_실험보고서"

T1H = ["모델 (피험자 격리 5-fold x 3 seeds)", "AUC", "F1"]
T1 = [["기존 sd2: CatBoost+LGBM (30D)", "77.55 ± 0.35", "83.91"],
      ["MiniRocket-lite + Ridge", "71.41 ± 1.19", "82.28"],
      ["GBDT + Ridge 확률 평균 (sd2 v4)", "78.45 ± 0.19", "84.30"],
      ["GBDT + Ridge 순위 평균", "78.82 ± 0.33", "84.39"]]
T2H = ["시점", "기존 GBDT AUC", "v4 확률평균 AUC"]
T2 = [["A 우측 후방 (view5)", "67.82", "68.50"], ["B 우측 전방 (view2)", "84.14", "85.98"],
      ["C 정면 (view3)", "82.93", "83.50"], ["D 좌측 전방 (view4)", "82.51", "82.59"],
      ["E 좌측 후방 (view1)", "64.64", "66.33"]]
T3H = ["스태커 결과 (nested CV, 3 seeds)", "2S EMR", "2S Clean F1", "3S EMR", "3S Clean F1"]
T3 = [["v4.1 (기존 sd2)", "50.27", "35.29", "52.19", "35.84"],
      ["v4.2 후보 (sd2 v4)", "49.98", "34.75", "52.10", "37.10"]]
T4H = ["시점", "v4.1 EMR/Clean F1", "v4.2 후보 EMR/Clean F1"]
T4 = [["A", "46.05 / 31.12", "45.73 / 28.92"], ["B", "58.94 / 43.53", "60.71 / 47.60"],
      ["C", "38.97 / 30.71", "38.81 / 29.69"], ["D", "61.35 / 44.45", "60.23 / 43.08"],
      ["E", "46.05 / 26.24", "44.44 / 24.78"]]
CONC = [
    "선행연구 조사(스켈레톤 소표본 시계열: MiniRocket+Ridge/GBDT 앙상블, 상대각도 기반 불변 표현)를 반영해 17채널 시계열(힙/무릎/경추 각도, 체간 정규화 좌표, 바벨·어깨·골반 롤, 요추 곡률)에 MiniRocket 구현(numpy, 21,420 피처)을 적용.",
    "sd2 단독 AUC는 77.55 -> 78.45(확률평균)로 약 +0.9%p 상승했으나 Ridge 단독은 71.41로 약함. 개선폭은 소폭.",
    "통합 스태커에 적용하면 2-State EMR/Clean F1은 변화 없음(차이가 표준편차 이내: 49.98 vs 50.27, 34.75 vs 35.29). 3-State Clean F1만 +1.26%p이나 노이즈 범위로 판단.",
    "결론: sd2 v4는 배포하지 않고 v4.1을 유지. 척추 병목은 모델 구조가 아니라 입력 정보(단일 시점의 MediaPipe 척추 표현력, 후방 시점 A/E AUC 약 65~68%)에 있을 가능성이 큼.",
    "다음 후보: 후방 시점 전용 척추 피처, 시점별 sd2 모델, 정상 라벨(130건) 증강, 데이터 추가 수집.",
    "주의: sd2 v4 OOF는 3 seeds 평균이라 단일 OOF보다 약간 낙관적일 수 있음(기존 sd2 단일 OOF AUC 76.53). 피험자 9명, 정상 130건으로 신뢰구간이 넓음.",
]


def main():
    doc = Document()
    doc.add_heading("바벨 스티프 데드리프트 sd2(척추) MiniRocket 앙상블 고도화 실험 보고서", 0)
    add_callout(doc, "요약", "sd2 AUC +0.9%p, 통합 EMR/Clean F1 개선 없음 -> v4.1 유지 (정직한 부정 결과).")
    add_header(doc, "1. sd2 단독 성능", 1); add_table(doc, T1H, T1, [3.2, 1.5, 1.2])
    add_header(doc, "2. 시점별 sd2 AUC", 1); add_table(doc, T2H, T2, [2.5, 1.8, 1.8])
    add_header(doc, "3. 통합 스태커 재평가", 1); add_table(doc, T3H, T3, [2.4, 1, 1.2, 1, 1.2])
    add_table(doc, T4H, T4, [1.2, 2.3, 2.5])
    add_header(doc, "4. 결론 및 한계", 1); add_bullets(doc, CONC)
    doc.save(os.path.join(D, N + ".docx"))
    md = "# sd2 MiniRocket 앙상블 고도화 실험 보고서\n\n## 1. sd2 단독\n" + md_table(T1H, T1) + "## 2. 시점별 AUC\n" + md_table(T2H, T2)
    md += "## 3. 스태커 재평가\n" + md_table(T3H, T3) + md_table(T4H, T4) + "## 4. 결론\n" + "".join(f"- {x}\n" for x in CONC)
    open(os.path.join(D, N + ".md"), "w", encoding="utf-8").write(md)
    print("OK")


if __name__ == "__main__":
    main()
