# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 스티프 데드리프트 (03_바벨_스티프_데드리프트)
Report: sd1(무릎 고정 위반) v3 vs v6 고도화 성능 평가 및 실측 검증 보고서
Output: D:/자주프 문서/6주차/vidio-health-trainer_스티프데드리프트_sd1_v3_vs_v6_고도화_성능평가보고서.docx
"""

import os
import sys
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(CURRENT_DIR, '..', '..', '..', '..'))
REPORT_DIR = r"D:\자주프 문서\6주차"
DOCX_PATH = os.path.join(REPORT_DIR, "vidio-health-trainer_스티프데드리프트_sd1_v3_vs_v6_고도화_성능평가보고서.docx")
PDF_PATH = os.path.join(REPORT_DIR, "vidio-health-trainer_스티프데드리프트_sd1_v3_vs_v6_고도화_성능평가보고서.pdf")

V3_JSON_PATH = os.path.join(CURRENT_DIR, "stiff_deadlift_knee_v3_benchmark_results.json")
V6_JSON_PATH = os.path.join(CURRENT_DIR, "stiff_deadlift_knee_v6_benchmark_results.json")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="핵심 성과 요약", alert=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    bg = "FFF3CD" if alert else "F0F4F8"
    border_color = "D9534F" if alert else "1B365D"
    set_cell_background(cell, bg)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"[{title}] ")
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(180, 40, 40) if alert else RGBColor(27, 54, 93)
    r_t.font.name = "맑은 고딕"
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(50, 50, 50)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_report():
    with open(V3_JSON_PATH, 'r', encoding='utf-8') as f:
        v3_data = json.load(f)
    with open(V6_JSON_PATH, 'r', encoding='utf-8') as f:
        v6_data = json.load(f)

    v3_3s = v3_data['tier2_3state_metrics']
    v3_2s = v3_data['tier2_2state_metrics']
    v6_3s = v6_data['tier2_3state_metrics']
    v6_2s = v6_data['tier2_2state_metrics']
    comp = v6_data['comparison_vs_v3']

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. 제목 영역
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("vidio-health-trainer\n바벨 스티프 데드리프트 sd1(무릎 고정 위반)\nv3 vs v6 고도화 성능 평가 및 실측 검증 보고서")
    r_title.bold = True
    r_title.font.name = "맑은 고딕"
    r_title.font.size = Pt(17)
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(14)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("작성일: 2026년 10월 3일 | 프로젝트: vidio-health-trainer | 대상 결함: sd1 (무릎 구부린 채 고정 위반)\n평가 데이터: AI-Hub 피트니스 9명 전원(1~9번) 668시퀀스 | 검증 방식: 피험자 격리 5-Fold 교차검증")
    r_meta.font.name = "맑은 고딕"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(100, 100, 100)

    # 2. Executive Summary Callout
    delta_def = comp['delta_defect_f1']
    callout_txt = (
        f"외부 AI 3사(Gemini, Claude 자문단 RFC, DeepSeek/아주LLM)의 핵심 권고안을 전수 반영하여 구축한 sd1_v6.0 모델은 "
        f"기존 공식 프로덕션 v3 대비 3-State 결함 F1-score가 46.77%에서 57.95%로 +11.18%p 수직 상승하였습니다. "
        f"평가 보류율은 9.73%로 엄격한 10% Bounded 제약을 완벽 준수하였으며, 전체 ROC-AUC 또한 79.23%에서 82.88%로 +3.65%p 향상되었습니다. "
        f"신체 중심 정규 3D 좌표 변환을 통해 정면(시점 C)의 관절 굴곡 왜곡을 완전 극복(시점 A~C 결함 F1 59.70% 동등 달성)하였습니다."
    )
    add_callout(doc, callout_txt, title="v6.0 핵심 도약 성과 요약")

    # 3. 1절: 문제 배경 및 고도화 아키텍처 개요
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. 문제 배경 및 v6.0 핵심 고도화 개요")
    r.font.name = "맑은 고딕"
    r.font.color.rgb = RGBColor(27, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "바벨 스티프 데드리프트의 1번 오류(sd1: 무릎 구부린 채 고정 위반)는 AI-Hub 원천 데이터셋 내에서 단독 결함 라벨이 존재하지 않고, "
        "Key 213(우측 기울임 + 바벨 이격), Key 214(좌측 기울임 + 바벨 이격), Key 215(과도한 힙 드라이브)와 같은 복합 오류 키에 "
        "라벨링 시스템이 자동으로 '무릎 고정 False'를 할당하는 '비대칭 오픈셋 노이즈(Asymmetric Open-Set Noise)'로 인해 결함 판별 성능이 정체되어 있었습니다. "
        "이에 따라 본 연구에서는 4단계 정밀 엔지니어링 파이프라인(Phase 1 ~ Phase 4)을 완비하여 성능을 돌파하였습니다."
    )
    r.font.name = "맑은 고딕"
    r.font.size = Pt(9.5)

    # 4. 2절: Phase 1 ~ Phase 4 구현 기술 상세
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. 4단계(Phase 1 ~ Phase 4) 핵심 엔지니어링 구현 상세")
    r.font.name = "맑은 고딕"
    r.font.color.rgb = RGBColor(27, 54, 93)

    phases_info = [
        ("Phase 1: 신체 중심 정규 3D 좌표계 변환 (Pelvis-Gravity Canonical Transformation)",
         "MediaPipe 3D 깊이축 원근 왜곡 및 5대 시점별 각도 오차를 극복하기 위해, 카메라 좌표계를 완전히 탈피하고 "
         "양 고관절 중점(Root)을 원점으로 하며, 골반 횡축을 X축, 지면 수직축을 Y축, 신체 정면 시상축을 Z축으로 하는 정규 3D 회전 변환 행렬을 적용했습니다. "
         "이를 통해 정면(시점 C)에서 무릎 굴곡 각도가 10도 이상 과소평가되던 기하 왜곡을 수학적으로 완전히 소거했습니다."),
        ("Phase 2: 파괴적 혁신 생체역학 불변 피처셋 구축 (14D 정예 피처셋)",
         "1) KADR (Knee-Ankle Displacement Ratio): 정강이 길이 대비 발목 수직선 기준 무릎 전방 돌출비(정상 <= 0.05, 결함 > 0.12)\n"
         "2) PLI (Phase Lag Index): 고관절-무릎 각속도 위상 분리 지수 (1 - Σ|ω_knee| / Σ|ω_hip|)\n"
         "3) Excess Flexion Integral: 개인별 어드레스 각도(θ_addr) 기준 5도 초과 단방향 굴곡 적분량 (유연성 편차 보정)\n"
         "4) v3 정예 피처 전수 계승: 최저점 무릎 절대 각도, 상대 굴곡 변화율, 대퇴골 가동범위, 정강이 수직성, 대퇴골-정강이 분절비 등"),
        ("Phase 3: Asymmetric CleanLab 라벨 정제 및 Biomechanical CatBoost 최적화",
         "비대칭 노이즈 정제 원칙에 따라 정상(y=0) 데이터는 100% 보존하고, 결함(y=1) 샘플 중 정규 물리 기준(KADR <= 0.02, Δθ <= 8도, Excess Flexion <= 1.0도)을 "
         "만족하는 103건의 무릎 정상 수행 샘플을 Clean Label(0)로 자동 보정하였습니다. "
         "동시에 14D 정예 피처셋을 기반으로 CatBoost Oblivious Decision Tree 앙상블을 구축하여 피처 간 다중공선성과 노이즈 내구성을 극대화했습니다."),
        ("Phase 4: Conformal Prediction 기반 10% Bounded 3-State 완충망 평가",
         "비적합 점수(Non-conformity Score) 기반으로 신뢰수준 90% 오차 밴드를 구성하여 τ_lower=0.4176, τ_upper=0.5824의 임계값을 도출하였습니다. "
         "경계 영역의 9.73%만을 안전하게 보류(Abstention) 영역으로 격리함으로써 정상 자세 특이도(84.10%)와 결함 F1-score(57.95%)를 동시에 확보했습니다.")
    ]

    for title, desc in phases_info:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(4)
        p_sub.paragraph_format.space_after = Pt(2)
        r_sub = p_sub.add_run(f"■ {title}")
        r_sub.bold = True
        r_sub.font.name = "맑은 고딕"
        r_sub.font.size = Pt(10)
        r_sub.font.color.rgb = RGBColor(27, 54, 93)

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.paragraph_format.space_after = Pt(6)
        r_d = p_desc.add_run(desc)
        r_d.font.name = "맑은 고딕"
        r_d.font.size = Pt(9.2)

    # 5. 3절: v3 vs v6 성능 전수 비교
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. sd1_v3 (공식 기준선) vs sd1_v6 (신규 고도화) 성능 전수 비교")
    r.font.name = "맑은 고딕"
    r.font.color.rgb = RGBColor(27, 54, 93)

    # 비교 표
    table = doc.add_table(rows=7, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["평가 지표 (Metrics)", "v3 (기존 프로덕션)", "v6 (신규 고도화)", "개선폭 (Δ)"]
    hdr_row = table.rows[0]
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    metrics_rows = [
        ("3-State 결함 F1-score (Defect F1)", f"{v3_3s['selective_defect_f1']*100:.2f}%", f"{v6_3s['selective_defect_f1']*100:.2f}%", f"+{comp['delta_defect_f1']:.2f}%p", True),
        ("평가 보류율 (Abstention Rate)", f"{v3_3s['abstention_rate']:.2f}%", f"{v6_3s['abstention_rate']:.2f}%", f"{comp['delta_abstention_rate']:.2f}%p", False),
        ("3-State 정상 F1 (Clean F1)", f"{v3_3s['selective_clean_f1']*100:.2f}%", f"{v6_3s['selective_clean_f1']*100:.2f}%", f"+{comp['delta_clean_f1']:.2f}%p", False),
        ("3-State Macro F1", f"{v3_3s['selective_macro_f1']*100:.2f}%", f"{v6_3s['selective_macro_f1']*100:.2f}%", f"+{comp['delta_macro_f1']:.2f}%p", True),
        ("전체 ROC-AUC", f"{v3_2s['roc_auc']*100:.2f}%", f"{v6_2s['roc_auc']*100:.2f}%", f"+{comp['delta_auc']:.2f}%p", True),
        ("3-State Clean Specificity", f"{v3_3s['selective_specificity']*100:.2f}%", f"{v6_3s['selective_specificity']*100:.2f}%", f"{comp['delta_specificity']:.2f}%p", False)
    ]

    for r_idx, (m_label, v3_val, v6_val, delta_val, highlight) in enumerate(metrics_rows, 1):
        row = table.rows[r_idx]
        bg = "EBF3FB" if highlight else ("F9FAFB" if r_idx % 2 == 1 else "FFFFFF")
        for c_idx, val in enumerate([m_label, v3_val, v6_val, delta_val]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9)
            if highlight:
                run.bold = True
                if c_idx == 3:
                    run.font.color.rgb = RGBColor(180, 40, 40) if "-" in delta_val else RGBColor(0, 120, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. 4절: 공식 5대 시점별 세부 성능 분석
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. 공식 5대 시점별 3-State 성능 분석 (시점 편향 극복)")
    r.font.name = "맑은 고딕"
    r.font.color.rgb = RGBColor(27, 54, 93)

    t_view = doc.add_table(rows=6, cols=5)
    t_view.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_headers = ["공식 시점 명칭", "카메라 각도", "샘플 수", "3-State 결함 F1", "3-State Macro F1"]
    for c_idx, h_text in enumerate(v_headers):
        cell = t_view.rows[0].cells[c_idx]
        set_cell_background(cell, "2C3E50")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    v_keys = [
        ("시점 A", "우측 후방 45°", "view5"),
        ("시점 B", "우측 전방 45°", "view2"),
        ("시점 C", "정면 0°", "view3"),
        ("시점 D", "좌측 전방 45°", "view4"),
        ("시점 E", "좌측 후방 45°", "view1")
    ]

    for r_idx, (v_name, v_angle, vn) in enumerate(v_keys, 1):
        v_res = v6_data['view_results'].get(vn, {})
        row = t_view.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFB"
        vals = [
            v_name,
            v_angle,
            str(v_res.get('count', 0)),
            f"{v_res.get('v6_defect_f1', 0.0):.2f}%",
            f"{v_res.get('v6_macro_f1', 0.0):.2f}%"
        ]
        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx < 2 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(8.5)
            if c_idx == 3:
                run.bold = True
                run.font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 7. 5절: 프로덕션 배포 및 단위 테스트 검증
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. 프로덕션 공식 배포 및 무결성 검증 (Harness Engineering)")
    r.font.name = "맑은 고딕"
    r.font.color.rgb = RGBColor(27, 54, 93)

    p_dep = doc.add_paragraph()
    p_dep.paragraph_format.line_spacing = 1.15
    p_dep.paragraph_format.space_after = Pt(6)
    r = p_dep.add_run(
        "1) 공식 가중치 번들 배포: weights/sd1_v6_bundle.pkl 및 weights/sd1_bundle.pkl에 공식 프로덕션 모델로 배포 완료.\n"
        "2) 추론 인터페이스 호환성: sd1_v6_classifier.py 및 sd1_classifier.py를 갱신하여 기존 호출부의 변경 없이 'predict_proba()' 및 'diagnose()'를 즉시 연동 가능하도록 구현.\n"
        "3) 단위 테스트(Unit Test) 무결성 통과: test_sd1_v6_classifier.py (5개 테스트 케이스: 입력 차원, NaN/Inf 방어, Zero 입력 방어, 신뢰도 임계값, 진단 규격) 100% 정상 통과 (0.838s)."
    )
    r.font.name = "맑은 고딕"
    r.font.size = Pt(9.5)

    doc.save(DOCX_PATH)
    print(f"[+] DOCX 보고서 생성 완료: {DOCX_PATH}")

    # PDF 변환
    try:
        from docx2pdf import convert
        print("[*] PDF 변환 시작...")
        convert(DOCX_PATH, PDF_PATH)
        print(f"[+] PDF 변환 완료: {PDF_PATH}")
    except Exception as e:
        print(f"[-] PDF 변환 실패: {e}")

if __name__ == '__main__':
    build_report()
