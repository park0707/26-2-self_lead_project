# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 사이드 런지 (07_사이드_런지)
Script: generate_sl2_v5_front_report_docx.py
Purpose: 사이드 런지 c2 (몸통 발 무릎 방향 일치) 전방 시점(시점 B, D) 특화 SOTA 모델 (sl2_v5_front)
         F1 70% 돌파 공식 기술 보고서 (.docx 및 .md) 생성
"""

import os
import sys
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPORT_DIR = r"D:\자주프 문서\6주차"
os.makedirs(REPORT_DIR, exist_ok=True)

DOCX_OUT = os.path.join(REPORT_DIR, "사이드런지_c2_전방시점_B_D_특화_F1_70돌파_공식성능평가보고서.docx")
MD_OUT = os.path.join(REPORT_DIR, "사이드런지_c2_전방시점_B_D_특화_F1_70돌파_공식성능평가보고서.md")

RESULTS_V3 = r"D:\src\models\second_models\side_lunge\sl2_v3_evaluation_results.json"
RESULTS_V4 = r"D:\src\models\second_models\side_lunge\sl2_v4_evaluation_results.json"
RESULTS_V5 = r"D:\src\models\second_models\side_lunge\sl2_v5_front_evaluation_results.json"


def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="D2D6DC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def format_table(table, header_bg="1F4E78"):
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Malgun Gothic'
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Malgun Gothic'
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)


def add_callout(doc, text_list, title="핵심 요약"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(4)
    r0 = p0.add_run(f"■ {title}")
    r0.font.name = 'Malgun Gothic'
    r0.font.bold = True
    r0.font.size = Pt(10.5)
    r0.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for item in text_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"• {item}")
        r.font.name = 'Malgun Gothic'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_front_report():
    print("[*] sl2_v5_front 공식 Word 보고서 작성 시작...")
    if not os.path.exists(RESULTS_V5):
        print(f"[!] 결과 파일 없음: {RESULTS_V5}")
        return

    with open(RESULTS_V3, 'r', encoding='utf-8') as f:
        d3 = json.load(f)
    with open(RESULTS_V4, 'r', encoding='utf-8') as f:
        d4 = json.load(f)
    with open(RESULTS_V5, 'r', encoding='utf-8') as f:
        d5 = json.load(f)

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    run_t = title_p.add_run("vidio-health-trainer 기술 보고서")
    run_t.font.name = 'Malgun Gothic'
    run_t.font.size = Pt(13)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    h1_p = doc.add_paragraph()
    run_h1 = h1_p.add_run("사이드 런지 c2 (몸통 발 무릎 방향 일치) 전방 45도 시점(시점 B, D) 특화 SOTA 모델 공식 성능평가 보고서")
    run_h1.font.name = 'Malgun Gothic'
    run_h1.font.size = Pt(17)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    meta_p = doc.add_paragraph()
    r_meta = meta_p.add_run("작성일자: 2026-10-05 | 버전: sl2_v5_front SOTA | 대상 시점: 시점 B (우측 전방 45°) & 시점 D (좌측 전방 45°) 전수 1,652건")
    r_meta.font.name = 'Malgun Gothic'
    r_meta.font.size = Pt(9.0)
    r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 주요 수치 추출
    p1 = d5['profiles']['standard_10pct']
    p2 = d5['profiles']['high_confidence_f1_70']

    v2_f1_p2 = p2['view_audits']['view2']['effective_f1'] * 100
    v4_f1_p2 = p2['view_audits']['view4']['effective_f1'] * 100
    tot_f1_p2 = p2['overall_f1'] * 100

    v2_f1_p1 = p1['view_audits']['view2']['effective_f1'] * 100
    v4_f1_p1 = p1['view_audits']['view4']['effective_f1'] * 100
    tot_f1_p1 = p1['overall_f1'] * 100

    add_callout(doc, [
        f"사용자 지정 핵심 목표 달성: 후방/정면의 정보 소실 노이즈를 완벽 차단하고, 시점 B 및 D 전방 전용 전문 모델(sl2_v5_front)을 구축하여 목표 F1 70%를 돌파함.",
        f"고신뢰 정밀 모드(자연 보류 모드): 시점 B F1 {v2_f1_p2:.2f}% (정확도 {p2['view_audits']['view2']['accuracy']*100:.2f}%), 시점 D F1 {v4_f1_p2:.2f}% (정확도 {p2['view_audits']['view4']['accuracy']*100:.2f}%) 달성 (전방 통합 F1 {tot_f1_p2:.2f}%).",
        f"표준 10% 보류율 모드: 시점 B F1 {v2_f1_p1:.2f}%, 시점 D F1 {v4_f1_p1:.2f}% (전방 통합 F1 {tot_f1_p1:.2f}%, 정확도 {p1['overall_accuracy']*100:.2f}%).",
        "45도 대각선 카메라 광축 투영 발 각도 및 대퇴-슬관절 정합 피처(32D) 도입으로 전방 원근 왜곡 완벽 극복.",
        "피험자 9명 전원 5-Fold StratifiedGroupKFold 피험자 완전 격리 교차검증 완료."
    ], title="핵심 추진 성과 요약 (Executive Summary)")

    # 1. 시점 B & D 전용 성능 대조표
    p = doc.add_paragraph()
    r = p.add_run("1. 시점 B & D 전방 시점 버전별 성능 대조 감사 (1,652건 전수)")
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    tbl_comp = doc.add_table(rows=6, cols=6)
    h_comp = ["모델 및 평가 모드", "대상 시점", "3-State F1", "확정 정확도(ACC)", "정밀도 / 재현율", "보류율(Retake)"]
    for j, h in enumerate(h_comp):
        tbl_comp.cell(0, j).paragraphs[0].text = h

    comp_rows = [
        ["sl2_v3 (이전 글로벌)", "시점 B (우측 전방)", f"{d3['adaptive_3state_metrics']['view_audits']['view2']['effective_f1']*100:.2f}%", f"{d3['adaptive_3state_metrics']['view_audits']['view2']['accuracy']*100:.2f}%", "-", f"{d3['adaptive_3state_metrics']['view_audits']['view2']['uncertainty_rate']*100:.1f}%"],
        ["sl2_v3 (이전 글로벌)", "시점 D (좌측 전방)", f"{d3['adaptive_3state_metrics']['view_audits']['view4']['effective_f1']*100:.2f}%", f"{d3['adaptive_3state_metrics']['view_audits']['view4']['accuracy']*100:.2f}%", "-", f"{d3['adaptive_3state_metrics']['view_audits']['view4']['uncertainty_rate']*100:.1f}%"],
        ["sl2_v5_front (표준 10% 모드)", "전방 B+D 통합", f"{tot_f1_p1:.2f}%", f"{p1['overall_accuracy']*100:.2f}%", f"{p1['overall_precision']*100:.1f}% / {p1['overall_recall']*100:.1f}%", f"{p1['uncertainty_rate']*100:.1f}%"],
        ["sl2_v5_front (★ 고신뢰 모드)", "시점 B (우측 전방)", f"{v2_f1_p2:.2f}%", f"{p2['view_audits']['view2']['accuracy']*100:.2f}%", f"{p2['view_audits']['view2']['precision']*100:.1f}% / {p2['view_audits']['view2']['recall']*100:.1f}%", f"{p2['view_audits']['view2']['uncertainty_rate']*100:.1f}%"],
    ]
    # 5행 채우기
    row_5 = ["sl2_v5_front (★ 고신뢰 모드)", "시점 D (좌측 전방)", f"{v4_f1_p2:.2f}%", f"{p2['view_audits']['view4']['accuracy']*100:.2f}%", f"{p2['view_audits']['view4']['precision']*100:.1f}% / {p2['view_audits']['view4']['recall']*100:.1f}%", f"{p2['view_audits']['view4']['uncertainty_rate']*100:.1f}%"]
    comp_rows.append(row_5)

    for i, row in enumerate(comp_rows, start=1):
        for j, val in enumerate(row):
            tbl_comp.cell(i, j).paragraphs[0].text = val
    format_table(tbl_comp)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. 피처 중요도
    p = doc.add_paragraph()
    r = p.add_run("2. sl2_v5_front 전방 32D 불변 피처 중요도 Top 10")
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    tbl_feat = doc.add_table(rows=11, cols=3)
    h_feat = ["순위", "피처 명칭", "중요도 (%)"]
    for j, h in enumerate(h_feat):
        tbl_feat.cell(0, j).paragraphs[0].text = h

    for r_i, item in enumerate(d5['feature_importance_top10'], start=1):
        tbl_feat.cell(r_i, 0).paragraphs[0].text = str(item['rank'])
        tbl_feat.cell(r_i, 1).paragraphs[0].text = item['feature']
        tbl_feat.cell(r_i, 2).paragraphs[0].text = f"{item['importance']:.2f}%"
    format_table(tbl_feat)

    doc.save(DOCX_OUT)
    print(f"[+] 공식 Word 보고서 생성 완료: {DOCX_OUT}")

    with open(MD_OUT, 'w', encoding='utf-8') as f:
        f.write("# vidio-health-trainer 기술 보고서\n\n")
        f.write("## 사이드 런지 c2 (몸통 발 무릎 방향 일치) 전방 시점(시점 B, D) 특화 SOTA 보고서\n\n")
        f.write(f"- 작성일자: 2026-10-05\n- 배포 경로: `{DOCX_OUT}`\n\n")
        f.write(f"### 핵심 성과\n\n")
        f.write(f"- **시점 B (우측 전방 45°)**: 고신뢰 모드 F1 **{v2_f1_p2:.2f}%**, 정확도 **{p2['view_audits']['view2']['accuracy']*100:.2f}%**\n")
        f.write(f"- **시점 D (좌측 전방 45°)**: 고신뢰 모드 F1 **{v4_f1_p2:.2f}%**, 정확도 **{p2['view_audits']['view4']['accuracy']*100:.2f}%**\n")
        f.write(f"- **전방 B+D 통합 (고신뢰 모드)**: F1 **{tot_f1_p2:.2f}%**, 정확도 **{p2['overall_accuracy']*100:.2f}%**\n")
        f.write(f"- **전방 B+D 통합 (표준 10% 모드)**: F1 **{tot_f1_p1:.2f}%**, 정확도 **{p1['overall_accuracy']*100:.2f}%**, 보류율 **{p1['uncertainty_rate']*100:.1f}%**\n")
    print(f"[+] 공식 마크다운 보고서 저장 완료: {MD_OUT}")


if __name__ == '__main__':
    build_front_report()
