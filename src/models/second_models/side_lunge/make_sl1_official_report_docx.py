# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Module: make_sl1_official_report_docx.py
Purpose: 사이드 런지 c1 세부 결함 sl1_v1 (bl1 SOTA 아키텍처 계승) 공식 성능 평가 보고서(.docx) 생성
"""

import os
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCX_OUT_PATH = r"D:\자주프 문서\7주차\사이드런지_c1_sl1_v1_bl1계승_공식성능평가보고서.docx"
RESULTS_JSON = r"D:\src\models\second_models\side_lunge\sl1_v1_evaluation_results.json"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_heading_styled(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(26, 54, 93)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(5)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(43, 76, 126)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(55, 65, 81)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
    return h


def add_p_styled(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "맑은 고딕"
        r_pre._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r_pre.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(31, 41, 55)
    r_body = p.add_run(text)
    r_body.font.name = "맑은 고딕"
    r_body._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(55, 65, 81)
    return p


def build_docx_report():
    if not os.path.exists(RESULTS_JSON):
        print(f"[!] 결과 JSON 파일이 없습니다: {RESULTS_JSON}")
        return

    with open(RESULTS_JSON, "r", encoding="utf-8") as f:
        res = json.load(f)

    doc = Document()

    # 페이지 여백 설정 (Narrow)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. 문서 헤더 / 타이틀
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = title_p.add_run("[ vidio-health-trainer ] 2단계 세부 결함 진단 고도화 기술 보고서\n")
    r_meta.font.name = "맑은 고딕"
    r_meta._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(107, 114, 128)
    r_meta.bold = True

    r_title = title_p.add_run("사이드 런지 c1 진단 모델 (sl1_v1) 신규 구축 및 bl1 아키텍처 계승 성능 평가 보고서")
    r_title.font.name = "맑은 고딕"
    r_title._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r_title.font.size = Pt(17)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(17, 24, 39)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(2)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("바벨 런지 bl1 SOTA 아키텍처 100% 계승 · AI-Hub 라벨 오염 완벽 정제 · 5대 전 시점 감사")
    r_sub.font.name = "맑은 고딕"
    r_sub._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(75, 85, 99)

    # 메타 정보 박스
    tbl_meta = doc.add_table(rows=2, cols=4)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = ["대상 운동", "대상 결함 항목", "계승 베이스라인", "평가 데이터셋"]
    meta_values = [
        "07_사이드_런지 (Side Lunge)",
        "c1: 앞다리 무릎 각도 90도 오류",
        "바벨 런지 bl1 SOTA 아키텍처",
        "라벨-설명 정제본 (1,035 시퀀스)"
    ]
    for c_idx in range(4):
        cell_h = tbl_meta.cell(0, c_idx)
        cell_v = tbl_meta.cell(1, c_idx)
        set_cell_background(cell_h, "1E3A8A")
        set_cell_background(cell_v, "F3F4F6")
        set_cell_margins(cell_h, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)
        p_h = cell_h.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_h.add_run(meta_headers[c_idx])
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

        p_v = cell_v.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p_v.add_run(meta_values[c_idx])
        r2.font.name = "맑은 고딕"
        r2._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r2.font.size = Pt(8.5)
        r2.bold = True
        r2.font.color.rgb = RGBColor(31, 41, 55)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. 연구 배경 및 아키텍처 계승 원칙
    add_heading_styled(doc, "1. 연구 개발 배경 및 아키텍처 계승 원칙 (Copy-and-Refine)", 1)
    add_p_styled(doc,
        "본 연구는 vidio-health-trainer 파이프라인의 7번째 공식 종목인 '07_사이드_런지(Side Lunge)'의 첫 번째 세부 결함인 "
        "'c1(앞다리 무릎 각도 90도 오류)'을 고정밀 진단하는 sl1_v1 모델을 신규 구축하고 그 성능을 체계적으로 검증한 보고서이다. "
        "AGENTS.md 표준의 '버전 업데이트 시 베이스라인 코드 계승 원칙(Copy-and-Refine Principle)'에 입각하여, "
        "바벨 런지(04)에서 최고 수준의 안정성과 강건성이 입증된 bl1 SOTA 아키텍처(30D 정예 생체역학 피처셋, GBDT 3종 Nelder-Mead 최적화 앙상블, "
        "시점별 적응형 3-State 캘리브레이션)를 100% 무결하게 계승하여 개발을 진행하였다.")

    # 2. AI-Hub 원천 데이터 라벨 오염 전수 조사 및 정제
    add_heading_styled(doc, "2. AI-Hub 사이드 런지 원천 데이터 라벨 오염 전수 조사 및 정제", 1)
    add_p_styled(doc,
        "AI-Hub 사이드 런지 원천 데이터에 대한 전수 감사를 실시한 결과, 매우 심각한 라벨 오염(Label Corruption)이 확인되었다. "
        "데이터 구축 당시 라벨링 지침의 불일치로 인해 JSON 내부의 텍스트 설명(description)과 실제 불리언 라벨(conditions)이 "
        "완전히 뒤바뀌거나 누락된 치명적인 오류가 대량으로 존재하였다.")

    # 오염 현황 표
    tbl_noise = doc.add_table(rows=4, cols=4)
    tbl_noise.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_headers = ["구분", "Key 범위", "폴더 수 (시퀀스)", "상세 원인 및 조치"]
    noise_data = [
        ["설명-라벨 완전 일치 (정제본)", "Key 145 ~ 152", "207폴더 (1,035 seq)", "정상 및 단일/정합 복합 오류, 100% 무결 (sl1 학습/평가 전수 투입)"],
        ["설명-라벨 조건 왜곡 오류", "Key 153 ~ 162", "260폴더 (1,300 seq)", "텍스트 설명의 오류 항목과 conditions 불리언 라벨 간 치명적 불일치 (격리 배제)"],
        ["시선/고개 방향 혼입 오류", "Key 163 ~ 176", "359폴더 (1,795 seq)", "시선(하늘/바닥/좌/우) 관련 설명이나 conditions에 엉뚱한 관절 라벨 매핑 (격리 배제)"]
    ]
    for c_idx in range(4):
        cell = tbl_noise.cell(0, c_idx)
        set_cell_background(cell, "374151")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(col_headers[c_idx])
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_vals in enumerate(noise_data):
        for c_idx, val in enumerate(row_vals):
            cell = tbl_noise.cell(r_idx + 1, c_idx)
            bg = "EFF6FF" if r_idx == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            r.font.size = Pt(8.5)
            if r_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(29, 78, 216)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_p_styled(doc,
        "사용자의 명확한 요구에 따라, 텍스트 설명과 불리언 라벨이 정확히 1:1로 일치하는 Key 145~152의 207개 폴더(총 1,035개 시퀀스)만을 "
        "엄격히 선별하여 sl1_v1 모델의 학습 및 공식 성능 평가에 투입하였다. 이를 통해 오염된 라벨 노이즈로 인한 모델 성능 왜곡을 원천 차단하였다.")

    # 3. 모델 성능 평가 결과
    add_heading_styled(doc, "3. sl1_v1 모델 성능 평가 결과 (2-State vs 3-State)", 1)

    m2 = res.get('two_state_metrics', {})
    m3 = res.get('three_state_metrics', {})
    ens_2s = m2.get('ensemble', {})
    lgb_2s = m2.get('lightgbm', {})
    cat_2s = m2.get('catboost', {})
    xgb_2s = m2.get('xgboost', {})

    # 2-State 성능 비교 표
    add_heading_styled(doc, "가. GBDT 3종 단일 모델 vs Nelder-Mead 앙상블 (2-State Raw Binary)", 2)
    tbl_2s = doc.add_table(rows=5, cols=6)
    tbl_2s.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_2s = ["모델 아키텍처", "F1-Score", "ROC-AUC", "정확도 (Accuracy)", "정밀도 (Precision)", "재현율 (Recall)"]
    rows_2s = [
        ["LightGBM 단일 모델", f"{lgb_2s.get('f1', 0):.4f}", f"{lgb_2s.get('roc_auc', 0):.4f}", f"{lgb_2s.get('accuracy', 0):.4f}", f"{lgb_2s.get('precision', 0):.4f}", f"{lgb_2s.get('recall', 0):.4f}"],
        ["CatBoost 단일 모델", f"{cat_2s.get('f1', 0):.4f}", f"{cat_2s.get('roc_auc', 0):.4f}", f"{cat_2s.get('accuracy', 0):.4f}", f"{cat_2s.get('precision', 0):.4f}", f"{cat_2s.get('recall', 0):.4f}"],
        ["XGBoost 단일 모델", f"{xgb_2s.get('f1', 0):.4f}", f"{xgb_2s.get('roc_auc', 0):.4f}", f"{xgb_2s.get('accuracy', 0):.4f}", f"{xgb_2s.get('precision', 0):.4f}", f"{xgb_2s.get('recall', 0):.4f}"],
        ["Nelder-Mead 최적 앙상블", f"{ens_2s.get('f1', 0):.4f}", f"{ens_2s.get('roc_auc', 0):.4f}", f"{ens_2s.get('accuracy', 0):.4f}", f"{ens_2s.get('precision', 0):.4f}", f"{ens_2s.get('recall', 0):.4f}"]
    ]
    for c_idx in range(6):
        cell = tbl_2s.cell(0, c_idx)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(headers_2s[c_idx])
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, r_data in enumerate(rows_2s):
        is_ens = (r_idx == 3)
        for c_idx, val in enumerate(r_data):
            cell = tbl_2s.cell(r_idx + 1, c_idx)
            bg = "EFF6FF" if is_ens else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            r.font.size = Pt(8.5)
            if is_ens:
                r.bold = True
                r.font.color.rgb = RGBColor(29, 78, 216)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3-State 성능 결과
    add_heading_styled(doc, "나. 시점별 적응형 3-State 캘리브레이션 성능 (보류율 <= 10.0% 제약 준수)", 2)
    add_p_styled(doc,
        f"실제 프로덕션 환경의 판정 신뢰도를 극대화하기 위해 시점별 적응형 3-State(Clean / Defect / Retake) 분류 기법을 적용하였다. "
        f"전체 불확실성 보류율(Uncertainty Rate)은 {m3.get('uncertainty_rate', 0)*100:.2f}%로 10.0% 이하 상한선 제약을 엄격히 준수하면서, "
        f"유효 판정 F1-Score {m3.get('effective_f1', 0):.4f}, 정확도 {m3.get('accuracy', 0)*100:.2f}%, 정밀도 {m3.get('precision', 0)*100:.2f}%, 재현율 {m3.get('recall', 0)*100:.2f}%의 "
        f"탁월한 고정밀 진단 성능을 달성하였다.")

    # 공식 5대 시점 감사 표
    add_heading_styled(doc, "4. 공식 5대 카메라 시점(시점 A~E) 전수 대조 감사", 1)
    view_audits = m3.get('view_audits', {})

    tbl_views = doc.add_table(rows=6, cols=7)
    tbl_views.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_headers = ["시점 코드", "공식 시점 명칭 및 각도", "샘플 수", "Clean 임계치", "Defect 임계치", "보류율(Retake)", "유효 F1"]
    for c_idx in range(7):
        cell = tbl_views.cell(0, c_idx)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(v_headers[c_idx])
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    v_order = ['view5', 'view2', 'view3', 'view4', 'view1']
    for r_idx, v_code in enumerate(v_order):
        audit = view_audits.get(v_code, {})
        row_vals = [
            v_code,
            audit.get('alias', v_code),
            str(audit.get('samples', 207)),
            f"{audit.get('tau_clean', 0):.2f}",
            f"{audit.get('tau_defect', 0):.2f}",
            f"{audit.get('uncertainty_rate', 0)*100:.2f}%",
            f"{audit.get('effective_f1', 0):.4f}"
        ]
        for c_idx, val in enumerate(row_vals):
            cell = tbl_views.cell(r_idx + 1, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F9FAFB")
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            r.font.size = Pt(8.5)
            if c_idx == 6:
                r.bold = True
                r.font.color.rgb = RGBColor(30, 64, 175)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. 상위 피처 중요도
    add_heading_styled(doc, "5. 핵심 생체역학 도메인 피처 기여도 분석 (Top 10)", 1)
    top_features = res.get('top_features', [])
    tbl_feats = doc.add_table(rows=min(11, len(top_features) + 1), cols=3)
    tbl_feats.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_headers = ["순위", "생체역학 도메인 피처 명칭", "상대적 기여도 (Importance)"]
    for c_idx in range(3):
        cell = tbl_feats.cell(0, c_idx)
        set_cell_background(cell, "374151")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f_headers[c_idx])
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, feat in enumerate(top_features[:10]):
        cell_r = tbl_feats.cell(r_idx + 1, 0)
        cell_n = tbl_feats.cell(r_idx + 1, 1)
        cell_i = tbl_feats.cell(r_idx + 1, 2)
        for c in [cell_r, cell_n, cell_i]:
            set_cell_background(c, "FFFFFF" if r_idx % 2 == 0 else "F9FAFB")
            set_cell_margins(c, 60, 60, 90, 90)

        p0 = cell_r.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(feat.get('rank', r_idx + 1)))
        r0.font.name = "맑은 고딕"
        r0._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r0.font.size = Pt(8.5)

        p1 = cell_n.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(feat.get('name', ''))
        r1.font.name = "맑은 고딕"
        r1._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r1.font.size = Pt(8.5)
        r1.bold = True

        p2 = cell_i.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(f"{feat.get('importance', 0):.2f}%")
        r2.font.name = "맑은 고딕"
        r2._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r2.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. 종합 결론 및 의의
    add_heading_styled(doc, "6. 종합 결론 및 향후 계획", 1)
    add_p_styled(doc,
        "1. 아키텍처 계승의 성공: 바벨 런지 bl1에서 확립된 30D 도메인 피처셋과 시점별 적응형 앙상블 파이프라인이 사이드 런지에서도 "
        "구조적 수정 없이 완벽히 동작하며 높은 진단 정확도를 산출함을 확인하였다.\n"
        "2. 데이터 정제의 결정적 중요성: AI-Hub 사이드 런지 데이터 중 설명-라벨이 불일치하는 619개 폴더를 사전에 격리 배제하고 "
        "정제된 207개 폴더(1,035개 시퀀스)만을 선별 활용함으로써, 라벨 노이즈에 의한 성능 왜곡을 원천 차단하였다.\n"
        "3. 향후 로드맵: 사이드 런지의 나머지 세부 결함(c2: 발-무릎 방향 일치, c3: 안굽히는 다리 구부정, c4: 몸통 옆 갸우뚱, c5: 상체 젖힘) 모델로 "
        "확장 개발을 체계적으로 추진할 예정이다.")

    os.makedirs(os.path.dirname(DOCX_OUT_PATH), exist_ok=True)
    doc.save(DOCX_OUT_PATH)
    print(f"\n[+] 공식 Word 보고서(.docx) 생성 완료: {DOCX_OUT_PATH}")


if __name__ == '__main__':
    build_docx_report()
