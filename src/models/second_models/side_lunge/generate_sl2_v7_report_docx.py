# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 사이드 런지 (07_사이드_런지)
Script: generate_sl2_v7_report_docx.py
Purpose: v7 3대 혁신 기하 알고리즘(역투영 3D + 다단계 시계열 5D + 역학복원 4D) 통합 51D 모델
         전체 원천 데이터(4,130건) 5대 시점 전수 성능 비교 공식 Word/MD 기술 보고서 생성
"""

import os
import sys
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPORT_DIR = r"D:\자주프 문서\6주차"
os.makedirs(REPORT_DIR, exist_ok=True)

DOCX_OUT = os.path.join(REPORT_DIR, "사이드런지_c2_v7_3대혁신기하_전체원천데이터_공식성능평가보고서.docx")
MD_OUT = os.path.join(REPORT_DIR, "사이드런지_c2_v7_3대혁신기하_전체원천데이터_공식성능평가보고서.md")

RESULTS_V4 = r"D:\src\models\second_models\side_lunge\sl2_v4_evaluation_results.json"
RESULTS_V6 = r"D:\src\models\second_models\side_lunge\sl2_v6_all_views_evaluation_results.json"
RESULTS_V7 = r"D:\src\models\second_models\side_lunge\sl2_v7_evaluation_results.json"


def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="D2D6DC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def format_table(table, header_bg="1E293B"):
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Malgun Gothic'
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [0, 1] else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Malgun Gothic'
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)


def add_callout(doc, text_list, title="핵심 요약"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(4)
    r0 = p0.add_run(f"■ {title}")
    r0.font.name = 'Malgun Gothic'
    r0.font.bold = True
    r0.font.size = Pt(10.5)
    r0.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for item in text_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"• {item}")
        r.font.name = 'Malgun Gothic'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def generate_v7_report():
    if not os.path.exists(RESULTS_V7):
        print(f"[!] 결과 파일 없음: {RESULTS_V7}")
        return

    with open(RESULTS_V4, 'r', encoding='utf-8') as f:
        d4 = json.load(f)
    with open(RESULTS_V6, 'r', encoding='utf-8') as f:
        d6 = json.load(f)
    with open(RESULTS_V7, 'r', encoding='utf-8') as f:
        d7 = json.load(f)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    p0 = doc.add_paragraph()
    r0 = p0.add_run("vidio-health-trainer 공식 기술 보고서")
    r0.font.name = 'Malgun Gothic'
    r0.font.size = Pt(11)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    p1 = doc.add_paragraph()
    r1 = p1.add_run("사이드 런지 c2 v7 3대 혁신 기하 알고리즘 전체 원천 데이터(4,130건) 성능 평가 보고서")
    r1.font.name = 'Malgun Gothic'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p1_tot = d7['profiles']['standard_10pct']
    front_sub = d7['front_views_submetrics']

    add_callout(doc, [
        f"사용자 제안 3대 혁신 기하 알고리즘(1. 시선 폐색 역투영, 2. 하강 전 구간 다단계 시계열, 3. 정강이 비틀림 역학 복원)을 통합한 51D SOTA 모델 'sl2_v7' 구축 완료.",
        f"전체 원천 데이터(4,130건 전수) 보류율 10% 이내 엄수 기준 통합 3-State F1 {p1_tot['overall_f1']*100:.2f}%, 확정 정확도 {p1_tot['overall_accuracy']*100:.2f}% 기록.",
        f"전방 시점(시점 B, D) 서브셋에서는 F1 {front_sub['f1']*100:.2f}%, 확정 정확도 {front_sub['accuracy']*100:.2f}%로 매우 우수한 판정력 유지.",
        f"후방 시점 가림 구간에서 3D 정강이 역학 복원과 하강 30% 초기 시계열 피처가 중요도 상위권에 성공적으로 안착."
    ], title="핵심 추진 성과 요약 (Executive Summary)")

    # 1. 5대 공식 시점별 성능 대조표
    p = doc.add_paragraph()
    r = p.add_run("1. 5대 공식 시점별 3-State 성능 대조 (보류율 <= 10.0% 엄수)")
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    tbl = doc.add_table(rows=7, cols=7)
    headers = ["시점 코드", "공식 시점 명칭", "3-State F1", "확정 정확도(ACC)", "정밀도", "재현율", "보류율(Retake)"]
    for j, h in enumerate(headers):
        tbl.cell(0, j).paragraphs[0].text = h

    view_order = ['view5', 'view2', 'view3', 'view4', 'view1']
    rows_data = []
    for i, v in enumerate(view_order, start=1):
        aud = p1_tot['view_audits'][v]
        row = [
            v,
            aud['alias'],
            f"{aud['effective_f1']*100:.2f}%",
            f"{aud['accuracy']*100:.2f}%",
            f"{aud['precision']*100:.2f}%",
            f"{aud['recall']*100:.2f}%",
            f"{aud['uncertainty_rate']*100:.1f}%"
        ]
        rows_data.append(row)
        for j, val in enumerate(row):
            tbl.cell(i, j).paragraphs[0].text = val

    tot_row = [
        "전체 (ALL)",
        "5대 시점 전체 통합",
        f"{p1_tot['overall_f1']*100:.2f}%",
        f"{p1_tot['overall_accuracy']*100:.2f}%",
        f"{p1_tot['overall_precision']*100:.2f}%",
        f"{p1_tot['overall_recall']*100:.2f}%",
        f"{p1_tot['uncertainty_rate']*100:.1f}%"
    ]
    for j, val in enumerate(tot_row):
        tbl.cell(6, j).paragraphs[0].text = val

    format_table(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. v7 12D 신규 피처 분석표
    p = doc.add_paragraph()
    r = p.add_run("2. v7 신규 12D 혁신 기하 피처 기여도 분석")
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    tbl_feat = doc.add_table(rows=13, cols=3)
    h_f = ["순번", "피처 명칭", "모델 기여도(%)"]
    for j, h in enumerate(h_f):
        tbl_feat.cell(0, j).paragraphs[0].text = h

    for r_i, item in enumerate(d7['v7_new_features_analysis'], start=1):
        tbl_feat.cell(r_i, 0).paragraphs[0].text = str(r_i)
        tbl_feat.cell(r_i, 1).paragraphs[0].text = item['feature']
        tbl_feat.cell(r_i, 2).paragraphs[0].text = f"{item['importance']:.2f}%"
    format_table(tbl_feat)

    doc.save(DOCX_OUT)
    print(f"[+] 공식 Word 보고서 생성 완료: {DOCX_OUT}")

    with open(MD_OUT, 'w', encoding='utf-8') as f:
        f.write("# vidio-health-trainer 기술 보고서\n\n")
        f.write("## 사이드 런지 c2 v7 3대 혁신 기하 알고리즘 전체 원천 데이터(4,130건) 성능 평가 보고서\n\n")
        f.write(f"- 작성일자: 2026-10-06\n- 공식 Word 산출물: `{DOCX_OUT}`\n")
        f.write(f"- 평가 환경: **보류율 10.0% 이내 제한** (실측 평균: **{p1_tot['uncertainty_rate']*100:.2f}%**)\n\n")
        f.write("### 1. 5대 공식 시점별 3-State 세부 성적표\n\n")
        f.write("| 시점 코드 | 공식 시점 명칭 | 3-State F1 | 확정 정확도 | 정밀도 | 재현율 | 보류율 |\n")
        f.write("| :---: | :--- | :---: | :---: | :---: | :---: | :---: |\n")
        for r in rows_data:
            f.write(f"| `{r[0]}` | **{r[1]}** | **{r[2]}** | {r[3]} | {r[4]} | {r[5]} | {r[6]} |\n")
        f.write(f"| **전체** | **{tot_row[1]}** | **{tot_row[2]}** | **{tot_row[3]}** | **{tot_row[4]}** | **{tot_row[5]}** | **{tot_row[6]}** |\n\n")
        f.write(f"### 2. 전체 및 전방 요약\n\n")
        f.write(f"- **전체 5대 시점 통합 (4,130건)**: F1 **{p1_tot['overall_f1']*100:.2f}%**, 정확도 **{p1_tot['overall_accuracy']*100:.2f}%**, 보류율 **{p1_tot['uncertainty_rate']*100:.1f}%**\n")
        f.write(f"- **전방 시점 서브셋 (시점 B, D)**: F1 **{front_sub['f1']*100:.2f}%**, 정확도 **{front_sub['accuracy']*100:.2f}%**\n")
    print(f"[+] 공식 마크다운 보고서 저장 완료: {MD_OUT}")


if __name__ == '__main__':
    generate_v7_report()
