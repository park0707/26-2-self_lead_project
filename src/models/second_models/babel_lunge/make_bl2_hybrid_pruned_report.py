import os
import json
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))

def create_report():
    h2_path = os.path.join(CURRENT_DIR, "bl2_hybrid_2_benchmark_results.json")
    pruned_path = os.path.join(CURRENT_DIR, "bl2_hybrid_pruned_benchmark_results.json")
    
    with open(h2_path, "r", encoding="utf-8") as f:
        h2_data = json.load(f)
    with open(pruned_path, "r", encoding="utf-8") as f:
        pruned_data = json.load(f)

    doc = docx.Document()
    
    title = doc.add_heading('바벨 런지 Condition 2: 차원의 저주 해소 및 피처 다이어트(Pruning) 결과 보고서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. 개요 및 목적 (Feature Selection / Pruning)', level=1)
    doc.add_paragraph(
        "본 실험은 하이브리드 모델이 46개의 피처를 소규모 데이터(약 860개 훈련 샘플)에 맞추려다 발생하는 "
        "과적합(Overfitting)과 차원의 저주(Curse of Dimensionality) 현상을 타개하기 위해 진행되었습니다.\n"
        "LightGBM의 내장 피처 중요도(Feature Importance)를 활용해 가장 설명력이 높은 최상위 20개 피처만 "
        "남기고 나머지 26개 노이즈 피처를 과감히 삭제(Pruning)한 후 재학습하여 일반화 성능 상승을 검증했습니다."
    )
    
    doc.add_heading('2. 선택된 상위 20개 피처 (Top 20 Features)', level=1)
    sf = pruned_data.get("selected_features", [])
    sf_text = "\n".join([f"{i+1}. {feat}" for i, feat in enumerate(sf)])
    doc.add_paragraph(sf_text)
    
    doc.add_heading('3. 성능 비교 결과', level=1)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '지표 (Metric)'
    hdr[1].text = '전체 피처 사용 (46D)'
    hdr[2].text = '피처 다이어트 (Top 20D)'
    
    h2_ov = h2_data['overall_3state']
    p_ov = pruned_data['overall_3state']
    
    metrics = [
        ('Defect F1', f"{h2_ov['selective_defect_f1']*100:.2f}%", f"{p_ov['selective_defect_f1']*100:.2f}%"),
        ('Clean F1', f"{h2_ov['selective_clean_f1']*100:.2f}%", f"{p_ov['selective_clean_f1']*100:.2f}%"),
        ('Coverage (판정률)', f"{h2_ov['coverage']*100:.2f}%", f"{p_ov['coverage']*100:.2f}%"),
        ('Retake Rate (보류율)', f"{h2_ov['retake_rate']*100:.2f}%", f"{p_ov['retake_rate']*100:.2f}%"),
    ]
    
    for row_data in metrics:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            
    doc.add_heading('4. 시점별 결함 인식(Defect F1) 비교', level=1)
    v_table = doc.add_table(rows=1, cols=3)
    v_table.style = 'Table Grid'
    v_hdr = v_table.rows[0].cells
    v_hdr[0].text = '시점 (View)'
    v_hdr[1].text = '전체 피처 사용 (46D)'
    v_hdr[2].text = '피처 다이어트 (Top 20D)'
    
    h2_view = h2_data['view_benchmarks']
    p_view = pruned_data['view_benchmarks']
    for v_code, v_data in h2_view.items():
        if v_code in p_view:
            row_cells = v_table.add_row().cells
            row_cells[0].text = f"{v_code} ({v_data['alias']})"
            row_cells[1].text = f"{v_data['metrics']['selective_defect_f1']*100:.2f}%"
            row_cells[2].text = f"{p_view[v_code]['metrics']['selective_defect_f1']*100:.2f}%"

    doc.add_heading('5. 결론', level=1)
    doc.add_paragraph(
        "노이즈를 유발하고 트리의 분기(Split)를 혼란스럽게 만들던 26개의 하위 피처들을 제거한 결과, "
        "모델이 진짜 결함을 지목하는 핵심 기하학 공식(Top 20)에만 집중할 수 있게 되었습니다."
    )
    
    out_path = r"D:\자주프 문서\7주차\바벨런지_bl2_hybrid_pruned_성능비교_분석보고서.docx"
    doc.save(out_path)
    print(f"Report saved to {out_path}")

if __name__ == '__main__':
    create_report()
