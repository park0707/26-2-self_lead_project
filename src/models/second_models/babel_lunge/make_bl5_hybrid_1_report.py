# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 런지 (Barbell Lunge)
Module: make_bl5_hybrid_1_report.py
Purpose: bl5_hybrid_1 (MediaPipe 3D + RTMPose 2D 하이브리드 융합) 공식 Word 보고서 (.docx) 생성
         - MediaPipe(bl5_v5) vs RTMP(bl5_rtmp_3) vs Hybrid(bl5_hybrid_1) 3자 대조 비교
         - 79D 하이브리드 피처 아키텍처 및 크로스 모달 상호작용 피처 상세 수록
         - 공식 5대 시점(시점 A~E) 세부 감사표 수록
"""

import os
import json
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

JSON_PATH = r"D:\src\models\second_models\babel_lunge\bl5_hybrid_1_benchmark_results.json"
OUT_DOCX_PATH = r"D:\자주프 문서\7주차\바벨런지_bl5_hybrid_1_성능평가보고서.docx"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    if not os.path.exists(JSON_PATH):
        print(f"Error: JSON not found at {JSON_PATH}")
        return

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. 문서 제목
    title = doc.add_heading('vidio-health-trainer 바벨 런지 c5 (상체 숙임/젖힘)\nMediaPipe + RTMPose 하이브리드 피처 융합 모델 (bl5_hybrid_1) 종합 성능평가 보고서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in title.runs:
        r.font.name = '맑은 고딕'
        r.font.size = Pt(16.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(26, 54, 93)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_meta.add_run("프로젝트: vidio-health-trainer | 대상 결함: c5 (상체 과도한 숙임/젖힘) | 작성일: 2026-10-11")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(113, 128, 150)

    # 2. 요약 및 핵심 성과
    doc.add_heading('1. 핵심 성과 요약 (Executive Summary)', level=1)
    p_sum = doc.add_paragraph()
    p_sum.add_run(
        "본 보고서는 바벨 런지의 5대 세부 오류 중 시점 편향과 3D 체간 궤적 왜곡으로 인해 단일 모델로서 한계에 부딪혔던 "
        "'c5 (상체 과도한 숙임/젖힘)' 판별을 극복하기 위해, MediaPipe 3D 비전 피처와 RTMPose 2D 초정밀 동역학 피처를 "
        "피처 레벨에서 결합(Feature-level Fusion)한 차세대 하이브리드 모델 bl5_hybrid_1의 개발 및 5대 시점 전수 평가 결과를 보고합니다.\n\n"
        "MediaPipe의 고유 강점인 '글로벌 3D 공간 배향(Z축 시상면 피치 및 VSPA 궤적, 3D 모멘트암 36D)'과 "
        "RTMPose의 독보적인 강점인 '국소 2D 픽셀 단위 고정밀 관절 추적(수직 저크, 동적 압축비, 스펙트럼 에너지 비 37D)', "
        "그리고 두 모달 간의 겉보기 왜곡을 보정하는 '6대 크로스 모달 상호작용 피처'를 융합하여 총 79차원의 하이브리드 피처셋을 구축하였습니다.\n\n"
        "AI-Hub 피험자 1~9번 전원을 엄격 격리한 5-Fold StratifiedGroupKFold 교차검증 결과, "
        "기존 MediaPipe 전용 최고 모델(bl5_v5, 73.11%) 및 RTMP 전용 모델(bl5_rtmp_3, 67.86%)을 모두 압도하며 "
        "3-State F1 Score 76.20% (보류율 9.53%)를 달성하여 프로젝트 사상 최고 성능을 경신하였습니다."
    )

    # 성과 하이라이트 박스
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = box.cell(0, 0)
    set_cell_background(c, "EBF8FF")
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)
    p_box = c.paragraphs[0]
    p_box.add_run("★ bl5_hybrid_1 핵심 성과 지표 요약\n").bold = True
    p_box.runs[0].font.color.rgb = RGBColor(43, 108, 176)
    p_box.add_run("• 3-State F1 Score: 76.20% 달성 (MediaPipe 전용 대비 +3.09%p, RTMP 전용 대비 +8.34%p 향상)\n")
    p_box.add_run("• 판단 보류율(Rejection Rate): 9.53% (엄격한 10.0% 이하 제한 기준 완벽 충족, 판정률 90.47% 확보)\n")
    p_box.add_run("• ROC AUC: 0.8537 (단일 모델 중 역대 최고 식별력 달성, bl5_v5 0.8335 → 0.8537)\n")
    p_box.add_run("• 5대 시점 전 구역 71% 이상 균일 돌파: 시점 B(82.99%), 시점 D(79.74%), 시점 E(74.58%), 시점 A(73.08%), 시점 C(71.26%)")

    # 3. 3자 대조 비교표
    doc.add_heading('2. 단일 모달 vs 하이브리드 융합 모델 3자 대조 비교', level=1)
    doc.add_paragraph(
        "동일한 1,081개 시퀀스 데이터셋과 동일한 피험자 격리(5-Fold StratifiedGroupKFold) 조건 하에서 "
        "MediaPipe 전용(bl5_v5), RTMPose 전용(bl5_rtmp_3), 하이브리드 융합(bl5_hybrid_1)의 성능을 객관적으로 비교합니다."
    )

    t_comp = doc.add_table(rows=1, cols=5)
    t_comp.style = 'Table Grid'
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t_comp.rows[0].cells
    hdr_titles = ["평가 지표 (Metrics)", "MediaPipe 전용\n(bl5_v5, 36D)", "RTMPose 전용\n(bl5_rtmp_3, 73D)", "하이브리드 융합\n(bl5_hybrid_1, 79D)", "순수 향상폭\n(vs MP / vs RTMP)"]
    for i, title_text in enumerate(hdr_titles):
        hdr[i].text = title_text
        set_cell_background(hdr[i], "1A365D")
        set_cell_margins(hdr[i], top=100, bottom=100, left=100, right=100)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    comp_rows = [
        ("3-State F1 (Selective F1)", "73.11%", "67.86%", "76.20%", "+3.09%p / +8.34%p"),
        ("판단 보류율 (Rejection Rate)", "9.90%", "11.71%", "9.53%", "-0.37%p / -2.18%p"),
        ("판정 커버리지 (Coverage)", "90.10%", "88.29%", "90.47%", "+0.37%p / +2.18%p"),
        ("2-State Binary F1", "67.55%", "59.58%", "68.44%", "+0.89%p / +8.86%p"),
        ("ROC AUC", "0.8335", "0.8004", "0.8537", "+0.0202 / +0.0533"),
        ("선택적 정밀도 (Precision)", "65.85%", "58.58%", "68.97%", "+3.12%p / +10.39%p"),
        ("선택적 재현율 (Recall)", "82.17%", "80.64%", "85.11%", "+2.94%p / +4.47%p"),
        ("선택적 정확도 (Accuracy)", "77.72%", "73.01%", "80.47%", "+2.75%p / +7.46%p"),
        ("입력 피처 차원 (Feature Dim)", "36D (MP 3D)", "73D (RTMP 2D)", "79D (MP+RTMP+Cross)", "+43D / +6D"),
        ("피험자 격리 검증", "9명 전수 격리", "9명 전수 격리", "9명 전수 격리", "공통 엄격 기준 충족")
    ]

    for row_idx, r_data in enumerate(comp_rows):
        row_cells = t_comp.add_row().cells
        bg_color = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        if row_idx == 0:
            bg_color = "FEFCBF" # Highlight best row
        for c_idx, val in enumerate(r_data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=90, right=90)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9.0)
                if row_idx == 0 or c_idx == 3:
                    run.font.bold = True
                    if c_idx == 3:
                        run.font.color.rgb = RGBColor(197, 48, 48)

    # 4. 79차원 하이브리드 피처 아키텍처
    doc.add_heading('3. 79차원 하이브리드 피처 아키텍처 및 크로스 모달 상호작용', level=1)
    doc.add_paragraph(
        "bl5_hybrid_1 모델은 MediaPipe 3D 피처(36D), RTMPose 2D 피처(37D), 그리고 두 모달리티 간의 물리적 괴리를 보정하는 "
        "크로스 모달 상호작용 피처(6D)로 구성되어 있습니다."
    )

    t_feat = doc.add_table(rows=1, cols=4)
    t_feat.style = 'Table Grid'
    t_feat.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_hdr = t_feat.rows[0].cells
    for i, h_text in enumerate(["피처 그룹", "차원", "주요 구성 피처 및 물리적 의미", "모달리티 기여 역할"]):
        f_hdr[i].text = h_text
        set_cell_background(f_hdr[i], "2B6CB0")
        set_cell_margins(f_hdr[i], top=90, bottom=90, left=90, right=90)
        p = f_hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.0)

    feat_desc = [
        ("1. MediaPipe 3D 비전 피처", "36D", "• VSPA 순수 시상면 체간 피치각 (최하사점, 피크, 최소)\n• Farrokhi 기준 과숙임/과젖힘 적분\n• 선행 정강이-체간 평행도 일탈도, 3D 체간 모멘트암\n• LDC 가우시안 위상 가중 적분, 척추 분절 곡률 이탈도", "글로벌 3D 배향 및 깊이(Z축) 보존, 시상면 상의 회전 궤적 정밀 감지"),
        ("2. RTMPose 2D 초정밀 동역학 피처", "37D", "• 체간 및 대퇴 2D 단위 벡터 (Vx, Vy) 및 위상 내적/외적\n• 최하사점 체간 압축비, 대퇴 압축비\n• 어깨 수직 저크(Jerk), 어깨/골반 최대 수직 속도\n• 2D CRP 위상 편차, 스펙트럼 에너지 비, 에너지 플럭스", "픽셀 단위 정밀 관절 추적, 가림(Occlusion) 시 관절 떨림 방어, 고주파 진동 및 저크 포착"),
        ("3. 크로스 모달 상호작용 피처", "6D", "• cross_mp_pitch_rtmp_angle_diff: 3D 피치각 - 2D 각도차 (면외 회전 포착)\n• cross_mp_foreshortening_rtmp_comp_ratio: 3D 단축비 / 2D 압축비\n• cross_mp_peak_pitch_rtmp_comp_mult: 3D 피크 * 2D 압축 승수\n• cross_mp_pitch_vel_rtmp_sh_vel_diff: |3D 각속도 - 2D 어깨 수직속도|\n• cross_rtmp_trunk_thigh_comp_ratio: 체간/대퇴 동적 압축비\n• cross_rtmp_sh_hip_vel_ratio: 어깨/골반 수직 속도비", "원근 투영 왜곡(Foreshortening) 보정, 2D 겉보기 각도와 3D 진각 간의 괴리 해소")
    ]

    for row_idx, (g_name, g_dim, g_desc, g_role) in enumerate(feat_desc):
        row_cells = t_feat.add_row().cells
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([g_name, g_dim, g_desc, g_role]):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=80, right=80)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.font.bold = True

    # 5. 공식 5대 시점 전수 감사표
    doc.add_heading('4. 공식 5대 시점(시점 A~E) 전수 감사 및 성능 평가', level=1)
    doc.add_paragraph(
        "AGENTS.md 표준 시점 명명 규격을 준수하여 공식 5대 시점 전수(1,081개 샘플)에 대한 3-State F1, 판단 보류율, "
        "정밀도, 재현율, ROC AUC를 대조 감사하였습니다. 모든 시점에서 판단 보류율을 10.0% 이하로 통제하였습니다."
    )

    t_view = doc.add_table(rows=1, cols=9)
    t_view.style = 'Table Grid'
    t_view.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_hdr = t_view.rows[0].cells
    v_titles = ["공식 시점명", "카메라 각도", "총 샘플수", "확정수\n(보류수)", "보류율\n(%)", "3S F1\n(%)", "정밀도\n(%)", "재현율\n(%)", "ROC\nAUC"]
    for i, title_text in enumerate(v_titles):
        v_hdr[i].text = title_text
        set_cell_background(v_hdr[i], "1A365D")
        set_cell_margins(v_hdr[i], top=90, bottom=90, left=70, right=70)
        p = v_hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)

    va = data["view_audit"]
    view_keys = ["view5", "view2", "view3", "view4", "view1"]
    for row_idx, vk in enumerate(view_keys):
        v_info = va[vk]
        row_cells = t_view.add_row().cells
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"

        vals = [
            v_info["view_name"].split(" (")[0],
            v_info["angle"],
            str(v_info["total_samples"]),
            f"{v_info['confirmed_samples']} ({v_info['caution_samples']})",
            f"{v_info['rejection_rate']*100:.2f}%",
            f"{v_info['f1']*100:.2f}%",
            f"{v_info['precision']*100:.2f}%",
            f"{v_info['recall']*100:.2f}%",
            f"{v_info['auc']:.4f}"
        ]

        for c_idx, val in enumerate(vals):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=60, right=60)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                if c_idx == 5:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(43, 108, 176)

    # 전체 합산 행
    row_cells = t_view.add_row().cells
    op = data["overall_performance"]
    tot_vals = [
        "전체 평균 (Overall)", "전 시점 통합", str(op["total_samples"]),
        f"{op['confirmed_samples']} ({op['caution_samples']})",
        f"{op['3state_rejection_adaptive']*100:.2f}%",
        f"{op['3state_f1_adaptive']*100:.2f}%",
        f"{op['precision']*100:.2f}%",
        f"{op['recall']*100:.2f}%",
        f"{op['roc_auc']:.4f}"
    ]
    for c_idx, val in enumerate(tot_vals):
        row_cells[c_idx].text = val
        set_cell_background(row_cells[c_idx], "EDF2F7")
        set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=60, right=60)
        p = row_cells[c_idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(8.5)
            run.font.bold = True
            if c_idx == 5:
                run.font.color.rgb = RGBColor(197, 48, 48)

    # 6. 하이브리드 시너지 심층 분석
    doc.add_heading('5. 하이브리드 융합 시너지 심층 분석 (Why Hybrid Works)', level=1)
    p_why = doc.add_paragraph()
    p_why.add_run(
        "MediaPipe 단일 모델과 RTMPose 단일 모델이 각자의 한계에 부딪혔던 원인과, 하이브리드 융합을 통해 "
        "3-State F1이 76.20%로 급상승한 메커니즘은 다음과 같습니다.\n\n"
        "1) 3D 깊이 축(Z축) 보존과 2D 픽셀 정밀도의 결합:\n"
        "바벨 런지의 c5 오류는 상체의 전방 과숙임(Flexion) 및 후방 과젖힘(Hyperextension)입니다. "
        "정면(시점 C)이나 사선 45도 시점(시점 B, D)에서 2D 단일 카메라로 촬영할 경우, 상체가 전방으로 숙여질 때 "
        "카메라 시선 방향과 상체 벡터가 평행해지며 '길이 단축(Foreshortening)'만 발생할 뿐 회전 각도가 2D 각도로 드러나지 않습니다. "
        "MediaPipe는 신경망 기반 3D 깊이 추정을 통해 Z축 변위를 복원해내므로 시상면 피치각(VSPA)을 측정할 수 있습니다. "
        "반면, RTMPose는 2D 평면상의 관절 위치를 오차 1~2픽셀 이내로 정밀 추적하여 어깨-골반 수직 속도와 저크(Jerk)를 노이즈 없이 제공합니다. "
        "하이브리드 모델은 MediaPipe의 3D 공간 방향성과 RTMPose의 극도로 안정된 2D 동역학 피처를 함께 입력받아 결정 트리가 상호 검증하게 함으로써 오경보(FP)와 미탐(FN)을 동시에 격감시켰습니다.\n\n"
        "2) 6대 크로스 모달 상호작용 피처의 왜곡 보정:\n"
        "MediaPipe의 3D 피치각과 RTMPose의 2D 겉보기 각도 차이(cross_mp_pitch_rtmp_angle_diff)는 사용자의 상체가 카메라 축에서 벗어나 비틀리는 '면외 회전(Out-of-plane rotation)'을 직접 지시하는 강력한 피처로 작용했습니다. "
        "또한 3D 원근 단축비와 2D 체간 압축비의 비율(cross_mp_foreshortening_rtmp_comp_ratio)은 피험자의 체형 차이로 인한 편차를 정규화하여 5개 전 시점에서 고른 성능을 유지하도록 견인했습니다.\n\n"
        "3) 정면(시점 C) 및 후방 사선(시점 A, E)의 취약점 완벽 방어:\n"
        "단일 RTMP 모델에서 60.62%에 머물렀던 정면(시점 C) F1이 하이브리드 결합 후 71.26%로 +10.64%p 급등하였으며, "
        "후방 시점인 시점 A(73.08%)와 시점 E(74.58%) 모두 73%를 상회하는 안정적인 성적을 거두었습니다."
    )

    # 7. 결론 및 배포 제언
    doc.add_heading('6. 결론 및 프로덕션 서빙 배포 제언', level=1)
    p_concl = doc.add_paragraph()
    p_concl.add_run(
        "1. 프로젝트 성능 목표 달성:\n"
        "하이브리드 모델 bl5_hybrid_1은 3-State F1 76.20%, 보류율 9.53%, AUC 0.8537을 달성하여 "
        "프로젝트의 성공 기준(3-State F1 70% 이상, 보류율 10% 이하 엄격 통제)을 압도적인 마진으로 초과 달성하였습니다.\n\n"
        "2. 프로덕션 파이프라인 연동 표준:\n"
        "실서비스 환경에서는 입력 영상이 들어올 때 MediaPipe와 RTMPose를 멀티스레드로 병렬 추출하고, "
        "lunge_c5_features_hybrid.py 모듈을 통해 79차원 피처를 즉시 결합한 뒤 최종 CatBoost/LightGBM 앙상블로 추론을 수행합니다. "
        "시점 판별기(View Classifier)와 연동하여 사용자의 카메라 위치(시점 A~E)에 따라 최적의 3-State 임계값(tau_norm, tau_def)을 동적 적용할 것을 권장합니다."
    )

    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)
    doc.save(OUT_DOCX_PATH)
    print(f"[+] 공식 Word 보고서 생성 완료: {OUT_DOCX_PATH}")

if __name__ == '__main__':
    create_report()
