# -*- coding: utf-8 -*-
"""bl4_v4b vs bl4_v3 비교 보고서(.docx/.md) 생성."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"D:\자주프 문서\6주차\바벨_런지_c4_bl4_v4b_vs_v3_노이즈강건_성능비교보고서"

SECTIONS = [
    ("1. 목적과 결론", [
        "목적: bl4_v3의 성능 향상 방법을 조사해 v4b로 구현하고 v3와 비교한다.",
        "결론: 라벨 노이즈 강건 학습(v4b-cl)은 순위 변별력(AUC)을 +1.63%p 올렸다(피험자 부트스트랩 95% CI [+0.07, +2.11]%p). "
        "그러나 3-State F1은 커버리지 95% 기준 -0.10%p로 사실상 동일하고 2-State F1은 -0.68%p 낮다. 'F1이 오른 개선'은 아니다.",
        "v3 공식 수치(F1 77.04%, AUC 82.15%)는 TCN 에포크를 검증 폴드 F1로 선택한 값이라 낙관적이다. "
        "동일 프로토콜(고정 에포크)로 재측정한 v3(ref)는 AUC 81.30%, 3-State F1(커버리지 95%) 76.54%이다.",
    ]),
    ("2. 조사 기반 적용 기법", [
        "Confident Learning (Northcutt et al., arXiv:1911.00068): 학습 폴드 내부 피험자-그룹 교차예측으로 라벨 오류 의심 샘플을 탐지. "
        "제외하지 않고 가중치 0.25로 하향(AGENTS.md 전수 사용 원칙 준수). cleanlab 라이브러리는 쓰지 않고 직접 구현(라이선스 이슈 없음).",
        "Symmetric Cross Entropy (Wang et al., ICCV 2019, arXiv:1908.06112): TCN 손실에 RCE 항(beta=0.3, A=4) 추가.",
        "피험자 균형 가중(1/sqrt(n)): 피험자1~3의 지배 완화. 효과 미미하여 최종안에서 제외.",
        "미확인 사항: 2023~24년 표/포즈 라벨 노이즈 논문은 검색 시간 초과로 확인하지 못해 인용하지 않았다.",
    ]),
    ("3. 공정 비교 프로토콜", [
        "동일 5-Fold StratifiedGroupKFold(피험자 9명 격리, seed 42), 동일 24D 피처, 동일 GBDT 하이퍼파라미터, 앙상블 가중치 0.25 균등.",
        "TCN은 모든 변형에서 30 에포크 고정, 마지막 가중치 사용(검증 누수 제거).",
        "평가는 원본 라벨 787건 전수. 라벨 오류 의심 샘플도 평가에서 제외하지 않았다.",
        "3-State 임계치는 모든 변형에 동일한 격자·동일 커버리지 하한으로 OOF에서 탐색(양쪽 모두 낙관 편향 동일).",
    ]),
]

HDR = ["변형", "AUC", "2-State F1", "고정 0.44/0.51 (cov / F1)", "cov>=95% F1", "cov>=90% F1"]
ROWS = [
    ["ref (v3, 동일 프로토콜)", "81.30%", "75.38%", "95.9% / 76.21%", "76.54%", "77.38%"],
    ["★ cl (v4b: CL+SCE)", "82.93%", "74.70%", "97.2% / 75.35%", "76.44%", "78.01%"],
    ["sb (피험자 균형)", "81.14%", "75.04%", "96.2% / 75.92%", "76.95%", "77.56%"],
    ["cl_sb (둘 다)", "82.57%", "73.57%", "96.3% / 74.92%", "75.56%", "77.15%"],
    ["참고: v3 원본 공식", "82.15%", "75.72%", "96.06% / 77.04%", "-", "-"],
]
HDR2 = ["서브모델", "ref AUC", "cl AUC", "차이"]
ROWS2 = [["TCN", "75.43%", "80.61%", "+5.18%p"], ["LightGBM", "80.28%", "81.25%", "+0.97%p"],
         ["XGBoost", "80.37%", "81.27%", "+0.90%p"], ["CatBoost", "80.79%", "81.68%", "+0.89%p"]]
HDR3 = ["시점", "ref AUC / 2sF1", "cl AUC / 2sF1"]
ROWS3 = [["시점 A (우측 후방 45°)", "76.7 / 69.1", "78.2 / 67.6"], ["시점 B (우측 전방 45°)", "83.9 / 78.2", "85.5 / 78.2"],
         ["시점 C (정면 0°)", "83.8 / 80.3", "86.0 / 79.3"], ["시점 D (좌측 전방 45°)", "80.6 / 73.1", "82.1 / 72.6"],
         ["시점 E (좌측 후방 45°)", "79.1 / 76.8", "81.5 / 76.6"]]
TAIL = [
    ("5. 해석", [
        "AUC 향상은 5개 시점 모두에서 나타났고(+1.5~2.4%p), TCN이 +5.18%p로 가장 크게 개선되었다. 라벨 노이즈 가중이 TCN의 과적합을 줄인 것으로 보인다(추정).",
        "F1이 오르지 않은 이유: 평가 라벨 자체에 노이즈가 있어 순위가 좋아져도 라벨 기준 F1은 한계에 막힌다. "
        "피험자2의 AUC는 60~62%로 변하지 않았고, 피험자1/3도 79~85%에 머문다. 병목은 모델이 아니라 라벨 품질로 보인다.",
        "2-State F1(-0.68%p)은 임계 0.5 고정 지표라 확률 분포 이동에 민감하다. 3-State는 임계치를 재탐색하므로 이 영향이 줄어든다.",
    ]),
    ("6. 산출물과 한계", [
        "bl4_v4b_weights.pkl (cl 모드, 임계치 tau_clean=0.37, tau_def=0.50, 커버리지 95% 기준 OOF 탐색), bl4_v4b.py, test_bl4_v4b.py 단위 테스트 통과.",
        "bl4_v4b.py는 bl4_v3.py를 복제해 이름·번들 경로만 변경했다(추론 구조 동일). v3 파일은 수정하지 않았다.",
        "임계치는 OOF에서 튜닝한 값이라 낙관 편향이 있다. 판정 전 이중 검증 세트가 필요하다.",
        "결정 필요: v3 유지 vs v4b 채택. F1 이득이 없으므로 AUC 향상(특히 TCN)이 중요할 때만 v4b를 권한다.",
        "참고: 이 학습은 다른 파이썬 작업과 CPU를 공유해 느렸다. 결과에는 영향 없음.",
    ]),
]


def build():
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(10)
    t = doc.add_paragraph(); r = t.add_run("vidio-health-trainer | 바벨 런지 c4 bl4_v4b vs bl4_v3 성능 비교")
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    def table(hdr, rows):
        tb = doc.add_table(rows=1, cols=len(hdr)); tb.style = 'Light Grid Accent 1'
        for i, h in enumerate(hdr):
            tb.rows[0].cells[i].text = h
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = v
        doc.add_paragraph()

    def sec(title, items):
        h = doc.add_paragraph(); rr = h.add_run(title); rr.font.bold = True; rr.font.size = Pt(13)
        for it in items:
            doc.add_paragraph(it, style='List Bullet')

    for s in SECTIONS:
        sec(*s)
    h = doc.add_paragraph(); rr = h.add_run("4. 결과 (OOF, N=787, 피험자 격리 5-Fold)"); rr.font.bold = True; rr.font.size = Pt(13)
    table(HDR, ROWS); table(HDR2, ROWS2); table(HDR3, ROWS3)
    for s in TAIL:
        sec(*s)
    doc.save(OUT + ".docx")

    md = ["# bl4_v4b vs bl4_v3 성능 비교\n"]
    def mdt(hdr, rows):
        md.append("| " + " | ".join(hdr) + " |\n|" + "---|" * len(hdr) + "\n" + "".join("| " + " | ".join(r) + " |\n" for r in rows))
    for title, items in SECTIONS:
        md.append(f"## {title}\n" + "".join(f"- {i}\n" for i in items))
    md.append("## 4. 결과\n"); mdt(HDR, ROWS); mdt(HDR2, ROWS2); mdt(HDR3, ROWS3)
    for title, items in TAIL:
        md.append(f"## {title}\n" + "".join(f"- {i}\n" for i in items))
    open(OUT + ".md", 'w', encoding='utf-8').write("\n".join(md))
    print("saved", OUT)


if __name__ == '__main__':
    build()
