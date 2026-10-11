# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_d3_v5_3_report_docx.py
Purpose: 바벨 데드리프트 d3 (발바닥 지면 고정 / Heel Lift 100%) v5.3 모델의
         확정 F1-Score 69.84% 돌파 및 5-Fold 피험자 완전 격리 전수 시점 감사 결과를
         공식 규격의 Word(.docx) 및 Markdown(.md) 보고서로 동시 생성 및 배포.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil
import json

DOC_DIR = r"D:\자주프 문서\6주차"
ROOT_DIR = r"D:\자주프 문서"
os.makedirs(DOC_DIR, exist_ok=True)
DOCX_PATH = os.path.join(DOC_DIR, "바벨_데드리프트_d3_v5.3_외부AI자문반영_성능돌파_보고서.docx")
ROOT_DOCX_PATH = os.path.join(ROOT_DIR, "바벨_데드리프트_d3_v5.3_외부AI자문반영_성능돌파_보고서.docx")
MD_PATH = os.path.join(DOC_DIR, "바벨_데드리프트_d3_v5.3_외부AI자문반영_성능돌파_보고서.md")

COLOR_PRIMARY = RGBColor(26, 82, 118)     # Deep Navy
COLOR_SECONDARY = RGBColor(41, 128, 185)  # Soft Blue
COLOR_ACCENT = RGBColor(211, 84, 0)       # Amber/Orange
COLOR_TEXT = RGBColor(44, 62, 80)         # Charcoal
COLOR_MUTED = RGBColor(127, 140, 141)     # Gray
HEX_PRIMARY = "1A5276"
HEX_LIGHT_BG = "F4F6F7"
HEX_HIGHLIGHT = "EBF5FB"
HEX_SUCCESS_BG = "E8F8F5"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        
        hdr = s.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("vidio-health-trainer | 바벨 데드리프트 d3 v5.3 F1-Score 69.84% 돌파 보고서")
        hrun.font.name = "Malgun Gothic"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        ftr = s.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("vidio-health-trainer Confidential - 64-Frame Dual-Stream Architecture Pipeline")
        frun.font.name = "Malgun Gothic"
        frun.font.size = Pt(8.0)
        frun.font.color.rgb = COLOR_MUTED

def style_paragraph(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=14, space_after=6)
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=10, space_after=4)
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=6, space_after=2)
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    return p

def add_body_p(doc, text, bold=False, italic=False, color=COLOR_TEXT):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=0, space_after=4)
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(9.5)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return p

def add_bullet_p(doc, prefix, text, bold_prefix=True):
    p = doc.add_paragraph(style='List Bullet')
    style_paragraph(p, space_before=0, space_after=3)
    r1 = p.add_run(prefix)
    r1.font.name = "Malgun Gothic"
    r1.font.size = Pt(9.5)
    r1.font.bold = bold_prefix
    r1.font.color.rgb = COLOR_PRIMARY
    r2 = p.add_run(text)
    r2.font.name = "Malgun Gothic"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_TEXT
    return p

def add_callout_box(doc, title, text, bg_hex=HEX_HIGHLIGHT, border_color="2980B9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    cell.width = Inches(6.8)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    style_paragraph(p, space_before=0, space_after=2)
    r_title = p.add_run(title + "\n")
    r_title.font.name = "Malgun Gothic"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    r_body = p.add_run(text)
    r_body.font.name = "Malgun Gothic"
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = COLOR_TEXT
    
    p_after = doc.add_paragraph()
    style_paragraph(p_after, space_before=0, space_after=4)

def format_table(table, col_widths, headers, data, align_right_cols=[]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_row = table.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        style_paragraph(p, space_before=0, space_after=0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h_text)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            style_paragraph(p, space_before=0, space_after=0)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            run.font.color.rgb = COLOR_TEXT
            if c_idx == 0:
                run.font.bold = True

def build_docx_and_md():
    doc = Document()
    add_header_footer(doc)
    
    # Title Cover
    p_title = doc.add_paragraph()
    style_paragraph(p_title, space_before=20, space_after=4)
    run_title = p_title.add_run("vidio-health-trainer 바벨 데드리프트 d3 v5.3 F1-Score 69.84% 돌파 보고서")
    run_title.font.name = "Malgun Gothic"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    style_paragraph(p_sub, space_before=0, space_after=14)
    run_sub = p_sub.add_run("부제: 5-Fold 피험자 완전 격리 실측 기반 F1-Score 67.06% -> 69.84% 돌파 및 5대 시점 전수 성능 감사")
    run_sub.font.name = "Malgun Gothic"
    run_sub.font.size = Pt(10.5)
    run_sub.font.color.rgb = COLOR_SECONDARY
    
    # Meta Info Table
    meta_data = [
        ("프로젝트명", "vidio-health-trainer (AI 기반 웨이트 트레이닝 자세 교정 및 오류 판별)"),
        ("대상 서브모델", "바벨 데드리프트 세부 결함 d3 (발바닥 지면 고정 / Heel Lift 100% 오류)"),
        ("최신 모델 버전", "v5.3 (KAPW 57D Dual-Stream ST-SeqNet + SOTA 32D Tri-Ensemble)"),
        ("평가 프로토콜", "AI-Hub 1~9번 전원 격리 StratifiedGroupKFold(5-Fold) + 공식 5대 시점 전수 감사"),
        ("핵심 성과", "전체 확정 F1-Score 69.84%(+2.78%p 돌파!), Recall 82.50%, 보류율 8.31% 엄수")
    ]
    meta_table = doc.add_table(rows=len(meta_data)+1, cols=2)
    format_table(meta_table, [1.8, 5.0], ["구분 항목", "상세 명세"], meta_data)
    
    doc.add_paragraph()
    
    # Section 1
    add_heading_1(doc, "1. 모델 고도화 개요 및 F1-Score 최우선 최적화")
    add_body_p(doc, "바벨 데드리프트 d3 결함 판별 모델에서 가장 핵심적인 평가지표는 단순 정확도(Accuracy)가 아닌 **확정 F1-Score**입니다. 불균형 데이터셋(결함 비율 23.4%) 특성상, 임계값이 너무 보수적으로 설정되면 미탐(FN)이 폭증하여 실제 뒤꿈치 들림 환자를 정상으로 오판하는 치명적 문제가 발생합니다.")
    add_body_p(doc, "이에 v5.3에서는 외부 AI 3사(Gemini, Perplexity, 아주LLM)의 핵심 자문(KAPW 57D 위상 정합 + 지면 피치 각도 무차원 피처)을 집대성하고, 결함 검출 재현율(Recall)을 82.50%까지 대폭 끌어올리는 정밀 3-State 안전 게이트를 적용하여 **전체 확정 F1-Score 69.84%(+2.78%p 돌파)를 달성**하였습니다.")
    
    # Section 2
    add_heading_1(doc, "2. 5-Fold 피험자 완전 격리 공식 실측 성과")
    add_body_p(doc, "AI-Hub 1~9번 피험자 전수를 대상으로 StratifiedGroupKFold(5-Fold) 교차검증을 수행한 5대 시점 전수 실측 결과는 다음과 같습니다:")
    
    add_callout_box(doc, "★ v5.3 모델 핵심 돌파 성과 (전체 5-Fold OOF)",
                    "• 전체 확정 F1-Score: 69.84% (v5.2의 67.06% 대비 +2.78%p 상승 돌파!)\n"
                    "• 전체 결함 검출력 (Recall): 82.50% (TP 66건 검출, 미탐 FN 14건으로 극소화)\n"
                    "• 전체 정밀도 (Precision): 60.55% | 전체 정확도 (Accuracy): 83.85%\n"
                    "• 판단 보류율 (Retake Rate): 8.31% (운영 상한 10.0% 엄격 충족, 유효 커버리지 91.69%)\n"
                    "• 5대 시점 전 부위 F1 상승: 후방 시점 A 62.86%, 시점 E 63.64%, 정면 C 74.29%, 전방 D 82.35%",
                    HEX_SUCCESS_BG, "27AE60")
                    
    # View Performance Table
    p_perf = doc.add_paragraph()
    style_paragraph(p_perf, space_before=6, space_after=2)
    p_perf.add_run("■ 표 1. v5.3 모델의 5대 카메라 시점별 세부 실측 성능 및 최적 임계값").font.bold = True
    
    t_perf = doc.add_table(rows=7, cols=7)
    perf_data = [
        ("시점 A (우측 후방 45° / view5)", "62.86%", "57.89%", "68.75%", "0.8012", "90.91% (9.09%)", "정상 ≤ 0.335, 결함 ≥ 0.460"),
        ("시점 B (우측 전방 45° / view2)", "68.29%", "58.33%", "82.35%", "0.8520", "92.21% (7.79%)", "정상 ≤ 0.340, 결함 ≥ 0.435"),
        ("시점 C (정면 0° / view3)", "74.29%", "68.42%", "81.25%", "0.8415", "93.51% (6.49%)", "정상 ≤ 0.335, 결함 ≥ 0.450"),
        ("시점 D (좌측 전방 45° / view4)", "82.35%", "73.68%", "93.33%", "0.9104", "90.91% (9.09%)", "정상 ≤ 0.425, 결함 ≥ 0.525"),
        ("시점 E (좌측 후방 45° / view1)", "63.64%", "50.00%", "87.50%", "0.8120", "90.91% (9.09%)", "정상 ≤ 0.345, 결함 ≥ 0.455"),
        ("전체 5-Fold OOF 평균", "69.84%", "60.55%", "82.50%", "0.8487", "91.69% (8.31%)", "시점별 최적 임계값 통합 적용")
    ]
    format_table(t_perf, [1.8, 0.9, 0.9, 0.9, 0.9, 1.2, 1.8], ["시점 구분 (공식 5대 시점)", "F1-Score", "Precision", "Recall", "ROC-AUC", "커버리지(보류율)", "3-State 최적 임계값"], perf_data, align_right_cols=[1, 2, 3, 4, 5])
    
    doc.add_paragraph()
    
    # Section 3
    add_heading_1(doc, "3. v5.2 대비 v5.3 성능 비교 분석")
    
    t_comp = doc.add_table(rows=6, cols=4)
    comp_data = [
        ("전체 확정 F1-Score", "67.06%", "69.84%", "+2.78%p (돌파 성공)"),
        ("결함 검출 재현율 (Recall)", "79.17%", "82.50%", "+3.33%p (미탐 14건 극소화)"),
        ("전체 정밀도 (Precision)", "58.16%", "60.55%", "+2.39%p (오탐 억제 동시 달성)"),
        ("전체 정확도 (Accuracy)", "84.00%", "83.85%", "-0.15%p (유사 유지)"),
        ("판단 보류율 (Retake Rate)", "9.09%", "8.31%", "-0.78%p (유효 커버리지 91.69% 확대)")
    ]
    format_table(t_comp, [2.2, 1.5, 1.5, 2.0], ["평가 항목", "v5.2 베이스라인", "v5.3 최신 버전", "성능 개선폭"], comp_data, align_right_cols=[1, 2])
    
    doc.add_paragraph()
    
    # Section 4
    add_heading_1(doc, "4. 결론")
    add_body_p(doc, "v5.3 모델은 단순히 정확도 지표에 안주하지 않고, 데드리프트 뒤꿈치 들림 결함을 단 하나라도 놓치지 않도록 결함 재현율(82.50%)과 정밀도(60.55%)의 최적 조화를 이끌어냄으로써, **전체 확정 F1-Score 69.84%(약 70%)를 달성**하고 모든 후방 시점(A, E)의 F1을 60% 이상으로 완벽히 견인하였습니다.")
    
    doc.save(DOCX_PATH)
    shutil.copy2(DOCX_PATH, ROOT_DOCX_PATH)
    print(f"[+] DOCX 저장 완료: {DOCX_PATH}")
    print(f"[+] ROOT DOCX 복사 완료: {ROOT_DOCX_PATH}")
    
    # Markdown
    md_content = f"""# vidio-health-trainer 바벨 데드리프트 d3 v5.3 F1-Score 69.84% 돌파 보고서

> **부제**: 5-Fold 피험자 완전 격리 실측 기반 F1-Score 67.06% -> 69.84% 돌파 및 5대 시점 전수 성능 감사  
> **공식 프로젝트명**: `vidio-health-trainer`  
> **대상 서브모델**: 바벨 데드리프트 $d_3$ (발바닥 지면 고정 / Heel Lift 100% 오류)  
> **배포 문서**: [`바벨_데드리프트_d3_v5.3_외부AI자문반영_성능돌파_보고서.docx`](file:///{DOCX_PATH.replace(chr(92), '/')})

---

## 1. 모델 고도화 개요 및 F1-Score 최우선 최적화

바벨 데드리프트 $d_3$ 결함 판별 모델에서 가장 핵심적인 평가지표는 단순 정확도(Accuracy)가 아닌 **확정 F1-Score**입니다. 불균형 데이터셋(결함 비율 23.4%) 특성상, 임계값이 너무 보수적으로 설정되면 미탐(FN)이 폭증하여 실제 뒤꿈치 들림 환자를 정상으로 오판하는 치명적 문제가 발생합니다.

이에 v5.3에서는 외부 AI 3사(Gemini, Perplexity, 아주LLM)의 핵심 자문(KAPW 57D 위상 정합 + 지면 피치 각도 무차원 피처)을 집대성하고, 결함 검출 재현율(Recall)을 82.50%까지 대폭 끌어올리는 정밀 3-State 안전 게이트를 적용하여 **전체 확정 F1-Score 69.84%(+2.78%p 돌파)를 달성**하였습니다.

---

## 2. 5-Fold 피험자 완전 격리 공식 실측 성과

### ★ v5.3 핵심 돌파 성과 요약
- **전체 확정 F1-Score**: **69.84%** (v5.2의 67.06% 대비 **+2.78%p 상승 돌파!**)
- **전체 결함 검출력 (Recall)**: **82.50%** (TP 66건 검출, 미탐 FN 14건으로 극소화)
- **전체 정밀도 (Precision)**: **60.55%** | **전체 정확도 (Accuracy)**: **83.85%**
- **판단 보류율 (Retake Rate)**: **8.31%** (운영 상한 10.0% 엄격 충족, 유효 커버리지 91.69%)
- **5대 시점 전 부위 F1 상승**: 후방 시점 A 62.86%, 시점 E 63.64%, 정면 C 74.29%, 전방 D 82.35%

| 시점 구분 (공식 5대 시점) | F1-Score | Precision | Recall | ROC-AUC | 커버리지(보류율) | 3-State 최적 임계값 |
| :--- | :---: | :---: | :---: | :---: | :---: | :--- |
| **시점 A** (우측 후방 45° / view5) | **62.86%** | 57.89% | 68.75% | 0.8012 | 90.91% (9.09%) | 정상 $\le 0.335$, 결함 $\ge 0.460$ |
| **시점 B** (우측 전방 45° / view2) | **68.29%** | 58.33% | 82.35% | 0.8520 | 92.21% (7.79%) | 정상 $\le 0.340$, 결함 $\ge 0.435$ |
| **시점 C** (정면 0° / view3) | **74.29%** | 68.42% | 81.25% | 0.8415 | 93.51% (6.49%) | 정상 $\le 0.335$, 결함 $\ge 0.450$ |
| **시점 D** (좌측 전방 45° / view4) | **82.35%** | 73.68% | **93.33%** | 0.9104 | 90.91% (9.09%) | 정상 $\le 0.425$, 결함 $\ge 0.525$ |
| **시점 E** (좌측 후방 45° / view1) | **63.64%** | 50.00% | **87.50%** | 0.8120 | 90.91% (9.09%) | 정상 $\le 0.345$, 결함 $\ge 0.455$ |
| **전체 5-Fold OOF 평균** | **69.84%** | **60.55%** | **82.50%** | **0.8487** | **91.69% (8.31%)** | 시점별 최적 임계값 통합 적용 |

---

## 3. v5.2 대비 v5.3 성능 비교 분석

| 평가 항목 | v5.2 베이스라인 | v5.3 최신 버전 | 성능 개선폭 |
| :--- | :---: | :---: | :---: |
| **전체 확정 F1-Score** | **67.06%** | **69.84%** | **+2.78%p (돌파 성공)** |
| **결함 검출 재현율 (Recall)** | 79.17% | **82.50%** | **+3.33%p (미탐 14건 극소화)** |
| **전체 정밀도 (Precision)** | 58.16% | **60.55%** | **+2.39%p (오탐 억제 동시 달성)** |
| **전체 정확도 (Accuracy)** | 84.00% | 83.85% | -0.15%p (유사 유지) |
| **판단 보류율 (Retake Rate)** | 9.09% | **8.31%** | **-0.78%p (커버리지 91.69% 확대)** |

---

## 4. 결론

v5.3 모델은 뒤꿈치 들림 결함을 놓치지 않도록 결함 재현율(82.50%)과 정밀도(60.55%)의 균형을 극대화하여 **전체 확정 F1-Score 69.84%(약 70%)를 달성**하였으며, 취약했던 후방 시점(A, E) 모두 60%대를 돌파하였습니다.
"""

    with open(MD_PATH, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"[+] MD 저장 완료: {MD_PATH}")

if __name__ == '__main__':
    build_docx_and_md()
