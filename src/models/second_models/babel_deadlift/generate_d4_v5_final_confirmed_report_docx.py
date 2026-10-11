# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_d4_v5_final_confirmed_report_docx.py
Purpose: 바벨 데드리프트 d4 v5.0 최종 확정 성능 보고서 (.docx) 생성
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d4_v5.0_최종확정_성능보고서.docx"
MD_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d4_v5.0_최종확정_성능보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("바벨 데드리프트 결함 d4 v5.0 최종 확정 기술 보고서")
    run_title.font.name = "맑은 고딕"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(24, 43, 73)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("vidio-health-trainer | Two-Stream Meta Fusion (43D Tri-GBDT + 64F Kinematic 1D-TCN) 최종 확정")
    run_sub.font.name = "맑은 고딕"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 110, 120)

    # 1. 모델 버전별 성능 변천사
    h1 = doc.add_heading("1. 바벨 데드리프트 d4 모델 버전별 종합 성능 변천사", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    headers = ["평가 지표", "v1.0 (초기 24D)", "v2.0 (체형 정규화)", "v3.0 (CRP+MVC)", "v4.0 (Two-Stream)", "★ v5.0 (최종 확정)", "v1 대비 개선폭"]
    rows = [
        ["F1-Score", "22.38%", "26.49%", "33.87%", "47.80%", "57.03%", "+34.65%p (2.5배 도약)"],
        ["재현율 (Recall)", "18.52%", "23.70%", "31.85%", "67.44%", "52.59%", "+34.07%p (오류 적발 급증)"],
        ["정밀도 (Precision)", "28.41%", "30.00%", "36.13%", "37.02%", "62.24%", "+33.83%p (오탐 억제 2배)"],
        ["ROC-AUC", "0.5820", "0.6210", "0.6970", "0.7560", "0.8049", "+0.2229 (AUC 0.80 벽 돌파)"],
        ["정확도 (Accuracy)", "76.10%", "77.40%", "78.44%", "72.47%", "85.84%", "+9.74%p 안정적 향상"],
        ["3-State F1", "-", "-", "40.14%", "47.80%", "53.04%", "Coverage 92% 엄수"]
    ]

    t1 = doc.add_table(rows=len(rows)+1, cols=len(headers))
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER

    for j, h in enumerate(headers):
        cell = t1.cell(0, j)
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for i, r in enumerate(rows):
        bg = "F2F5F9" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(r):
            cell = t1.cell(i+1, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in [0, 6] else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(8.5)
            if j == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(180, 40, 40)

    # 2. 공식 5대 시점 전수 감사
    h2 = doc.add_heading("2. 공식 5대 시점 전수 감사 (Official 5-View Audit)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    v_headers = ["공식 시점 명칭", "내부 코드", "2-State F1", "Tuned F1 (τ=0.70)", "ROC-AUC", "시점별 판정 강건성"]
    v_rows = [
        ["시점 D (좌측 전방 45°)", "view4", "66.67%", "69.57%", "0.8524", "★ 전방 대각선 뷰 최고의 신전 동기화 포착 (F1 70%)"],
        ["시점 B (우측 전방 45°)", "view2", "62.86%", "65.52%", "0.8410", "★ 우측 대각선 뷰 First Pull 각속도 비 정밀 탐지"],
        ["시점 E (좌측 후방 45°)", "view1", "52.78%", "56.00%", "0.8120", "후방 각도에서도 1D-TCN 시계열 변별력 견고히 유지"],
        ["시점 C (정면 0°)", "view3", "38.36%", "42.11%", "0.7420", "정면 깊이(Z축) 왜곡 한계 극복하고 40%대 수복"],
        ["시점 A (우측 후방 45°)", "view5", "38.24%", "41.67%", "0.7180", "후방 골반 상승 속도차로 힙 슈트업 안정적 변별"]
    ]

    t2 = doc.add_table(rows=len(v_rows)+1, cols=len(v_headers))
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER

    for j, h in enumerate(v_headers):
        cell = t2.cell(0, j)
        set_cell_background(cell, "2E5B82")
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for i, r in enumerate(v_rows):
        bg = "F9FAFC" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(r):
            cell = t2.cell(i+1, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in [0, 5] else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(8.5)
            if i in [0, 1] and j == 3:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 120, 50)

    # 3. 아키텍처 및 배포 정보
    h3 = doc.add_heading("3. 확정 아키텍처 및 실서비스 배포 명세", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)

    p_body = doc.add_paragraph()
    p_body.paragraph_format.line_spacing = 1.25
    r_body = p_body.add_run(
        "1. 공식 모델 소스 파일: d:\\src\\models\\second_models\\babel_deadlift\\d4_model_v5.py\n"
        "2. 최종 가중치 패키징: d:\\src\\models\\second_models\\babel_deadlift\\weights_v5\\babel_deadlift_d4_v5_final.pkl\n"
        "3. 엔터프라이즈 인터페이스:\n"
        "   - model.predict_proba_from_seq(seq_64f): p_final, p_gbdt, p_tcn, binary_pred 반환\n"
        "   - model.predict_3state_from_seq(seq_64f): NORMAL(0), DEFECT(1), RETAKE(-1) 3상태 자동 판정\n"
        "4. 권장 배포 임계치: tau_binary=0.70 (F1 57.03%), tau_normal=0.20, tau_defect=0.62 (Coverage 92%)"
    )
    r_body.font.name = "맑은 고딕"
    r_body.font.size = Pt(9.5)

    doc.save(DOCX_PATH)
    print(f"[+] Word 보고서 저장 완료: {DOCX_PATH}")

    # 마크다운 요약본 저장
    md_content = f"""# 바벨 데드리프트 결함 d4 v5.0 최종 확정 기술 보고서

> **프로젝트**: vidio-health-trainer | AI-Hub 바벨 데드리프트 결함 d4 모델  
> **보고서 파일**: `{DOCX_PATH}`  
> **모델 파일**: `d:\\src\\models\\second_models\\babel_deadlift\\d4_model_v5.py`  
> **가중치 파일**: `d:\\src\\models\\second_models\\babel_deadlift\\weights_v5\\babel_deadlift_d4_v5_final.pkl`  
> **작성 일자**: 2026-09-30  

---

## 1. 버전별 종합 성능 변천사 (v1.0 ~ v5.0)

| 평가 지표 | v1.0 (초기 24D) | v2.0 (체형 정규화) | v3.0 (CRP+MVC) | v4.0 (Two-Stream) | **★ v5.0 (최종 확정)** | v1 대비 개선폭 |
| :--- | :---: | :---: | :---: | :---: | :---: | :--- |
| **F1-Score** | 22.38% | 26.49% | 33.87% | 47.80% | <span style="color:red">**57.03%**</span> | **+34.65%p (2.5배 도약)** |
| **재현율 (Recall)** | 18.52% | 23.70% | 31.85% | 67.44% | **52.59%** | **+34.07%p (오류 적발 급증)** |
| **정밀도 (Precision)**| 28.41% | 30.00% | 36.13% | 37.02% | <span style="color:red">**62.24%**</span> | **+33.83%p (오탐 억제 2배)** |
| **ROC-AUC** | 0.5820 | 0.6210 | 0.6970 | 0.7560 | <span style="color:red">**0.8049**</span> | **+0.2229 (0.80 벽 돌파)** |
| **정확도 (Accuracy)** | 76.10% | 77.40% | 78.44% | 72.47% | **85.84%** | **+9.74%p 안정적 향상** |
| **3-State F1** | - | - | 40.14% | 47.80% | **53.04%** | Coverage 92% 엄수 |

---

## 2. 공식 5대 시점 전수 감사 (Official 5-View Audit)

| 공식 시점 명칭 | 내부 코드 | 2-State F1 | **Tuned F1 (τ=0.70)** | **ROC-AUC** | 시점별 판정 강건성 |
| :--- | :---: | :---: | :---: | :---: | :--- |
| **시점 D (좌측 전방 45°)** | view4 | **66.67%** | <span style="color:red">**69.57%**</span> | **0.8524** | ★ 전방 대각선 뷰 최고의 신전 동기화 포착 (F1 70%) |
| **시점 B (우측 전방 45°)** | view2 | **62.86%** | <span style="color:red">**65.52%**</span> | **0.8410** | ★ 우측 대각선 뷰 First Pull 각속도 비 정밀 탐지 |
| **시점 E (좌측 후방 45°)** | view1 | **52.78%** | **56.00%** | **0.8120** | 후방 각도에서도 1D-TCN 시계열 변별력 견고히 유지 |
| **시점 C (정면 0°)** | view3 | 38.36% | **42.11%** | 0.7420 | 정면 깊이(Z축) 왜곡 한계 극복하고 40%대 수복 |
| **시점 A (우측 후방 45°)** | view5 | 38.24% | **41.67%** | 0.7180 | 후방 골반 상승 속도차로 힙 슈트업 변별 |
"""
    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[+] 마크다운 요약본 저장 완료: {MD_PATH}")

if __name__ == "__main__":
    create_report()
