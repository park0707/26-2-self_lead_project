# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 데드리프트 (Barbell Deadlift)
Defect: d1. 척추의 중립 위반 (Spine Neutrality)
Script: generate_d1_v5_final_spec_docx.py
Purpose: 바벨 데드리프트 d1 최종 v5.0 확정 및 배포명세서 공식 Word 및 Markdown 문서 생성
"""

import os
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_OUTPUT_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d1_최종_v5.0_확정_및_배포명세서.docx"
MD_OUTPUT_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d1_최종_v5.0_확정_및_배포명세서.md"


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_callout(doc, text, title="공식 최종 확정 요약 (Confirmed Summary)", border_hex="1F497D", bg_hex="F2F5F9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}\n")
    run_t.bold = True
    run_t.font.name = "Malgun Gothic"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    run_b = p.add_run(text)
    run_b.font.name = "Malgun Gothic"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def build_final_spec():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. 헤더
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_sub = p_title.add_run("vidio-health-trainer AI Core 배포 사양서\n")
    r_sub.font.name = "Malgun Gothic"
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    r_main = p_title.add_run("바벨 데드리프트 d1 최종 v5.0 공식 확정 및 운영 배포명세서")
    r_main.bold = True
    r_main.font.name = "Malgun Gothic"
    r_main.font.size = Pt(20)
    r_main.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(2)
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta = p_meta.add_run("공식 프로젝트명: vidio-health-trainer | 운동 종목: 바벨 데드리프트 | 결함 항목: d1. 척추의 중립 위반\n배포 모델 버전: Deadlift d1 v5.0 (2-Tier Ambiguity Pruned & Calibrated 3-State)")
    r_meta.font.name = "Malgun Gothic"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

    # 2. Executive Summary
    summary_text = (
        "사용자의 명시적 승인에 따라 '바벨 데드리프트 d1 최종 3-State 선택 판정 (표준 윈도우 [0.16, 0.43])' 모델을 "
        "vidio-health-trainer 운영 서빙 엔진의 공식 최종 버전(v5.0)으로 확정 배포합니다.\n\n"
        "• 공식 확정 지표: 3-State Selective F1-Score 81.61% ★ (81% 돌파), 결함 검출 재현율 93.81% ★, 정확도 87.50% ★\n"
        "• 서비스 안정성 제약: 채점 커버리지 91.88% (보류율 8.12%, 상용 기준 10% 이하 제약 완벽 준수)\n"
        "• 핵심 아키텍처: Level 4 오염(7.27%) 폐기 정제 학습 + 25D 다이어트 피처 + 24D Dilated TCN + GBDT Tri-Ensemble"
    )
    add_callout(doc, summary_text)

    # 3. 확정 사양 테이블
    h1 = doc.add_heading("1. d1 v5.0 공식 최종 배포 규격 및 성능 지표", level=1)
    tbl_spec = doc.add_table(rows=8, cols=3)
    tbl_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_spec)

    spec_headers = ["항목", "공식 확정 사양", "상세 설명 및 비고"]
    for c_idx, h in enumerate(spec_headers):
        cell = tbl_spec.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_background(cell, "1F497D")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

    spec_rows = [
        ["배포 모델명", "Deadlift d1 Model v5.0 Confirmed", "d1_model_v5_confirmed.py"],
        ["3-State 판정 윈도우", "[0.16, 0.43]", "P <= 0.16 (정상 확정), P >= 0.43 (결함 확정), 그 사이 (보류)"],
        ["Selective F1-Score", "81.61% ★", "이전 v4.3(80.37%) 대비 +1.25%p, 2진(77.87%) 대비 +3.74%p"],
        ["결함 검출 재현율", "93.81% ★", "위험한 등굽음 결함을 94% 수준으로 확실하게 검출"],
        ["판정 정확도", "87.50% ★", "이전 v4.3(86.65%) 대비 +0.85%p 상승"],
        ["채점 커버리지 / 보류율", "91.88% / 8.12% (29건)", "상용 안정성 제약(보류율 10% 이하) 완벽 준수"],
        ["가중치 번들 경로", "weights_d1_v5/d1_v5_clean_gbdt_models.pkl\nweights_d1_v5/d1_v5_clean_tcn.pth", "LightGBM + CatBoost + Dilated TCN 가중치"]
    ]

    for r_idx, row in enumerate(spec_rows):
        for c_idx, val in enumerate(row):
            cell = tbl_spec.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if "81.61%" in val:
                cell.paragraphs[0].runs[0].bold = True
                set_cell_background(cell, "EBF5FB")
            else:
                set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F9FBFD")
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

    # 4. 차기 작업 안내
    h2 = doc.add_heading("2. 향후 순차 개발 로드맵", level=1)
    p_road = doc.add_paragraph()
    p_road.add_run(
        "d1의 2계층 모호성 정제 및 3-State F1 81.61% 최종 확정에 이어, "
        "사용자의 'd1부터 하나씩 진행' 방침에 따라 2단계인 **d2(무릎-발목 모임 / 무릎 외반)**의 "
        "4단계 모호성 분류 및 2계층 구조 도입으로 즉시 착수합니다."
    )

    doc.save(DOCX_OUTPUT_PATH)
    print(f"[+] 공식 Word 배포명세서 생성 완료: {DOCX_OUTPUT_PATH}")

    # 마크다운 배포명세서
    md_content = f"""# [vidio-health-trainer] 바벨 데드리프트 d1 최종 v5.0 공식 확정 및 운영 배포명세서

- **공식 프로젝트명**: `vidio-health-trainer`
- **배포 모델명**: Deadlift d1 Model v5.0 Confirmed (2-Tier Ambiguity Pruned & 3-State Calibrated)
- **배포 상태**: **Officially Confirmed & Production Deployed (공식 최종 확정 및 운영 서빙 배포)**

---

## 1. 공식 확정 성능 지표

| 항목 | 공식 확정 사양 | 상세 내용 및 비고 |
| :--- | :---: | :--- |
| **3-State 판정 윈도우** | **$[0.16, 0.43]$** | $P \le 0.16$ 정상 확정 / $P \ge 0.43$ 결함 확정 / 그 사이 보류 |
| **Selective F1-Score** | **81.61% ★** | **81% 돌파 (이전 v4.3 대비 +1.25%p, 2진 대비 +3.74%p)** |
| **결함 검출 재현율** | **93.81% ★** | **위험 결함의 94% 확실한 검출 보장** |
| **진단 정확도** | **87.50% ★** | **이전 대비 +0.85%p 향상** |
| **채점 커버리지** | **91.88%** | **보류율 8.12% (상용 제약 10% 이하 완벽 준수)** |
| **강제 2진 전수 F1** | **77.87%** | **2진 강제 분류 시에도 전수 77.87% 달성** |

---

## 2. 배포 산출물 및 API 규격

1. **운영 서빙 모듈**:
   [`d1_model_v5_confirmed.py`](file:///D:/src/models/second_models/babel_deadlift/d1_model_v5_confirmed.py)
2. **공식 가중치 번들**:
   - GBDT: [`weights_d1_v5/d1_v5_clean_gbdt_models.pkl`](file:///D:/src/models/second_models/babel_deadlift/weights_d1_v5/d1_v5_clean_gbdt_models.pkl)
   - TCN: [`weights_d1_v5/d1_v5_clean_tcn.pth`](file:///D:/src/models/second_models/babel_deadlift/weights_d1_v5/d1_v5_clean_tcn.pth)
3. **공식 Word 배포명세서**:
   [`바벨_데드리프트_d1_최종_v5.0_확정_및_배포명세서.docx`](file:///D:/자주프%20문서/6주차/바벨_데드리프트_d1_최종_v5.0_확정_및_배포명세서.docx)
"""
    with open(MD_OUTPUT_PATH, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[+] 공식 마크다운 배포명세서 생성 완료: {MD_OUTPUT_PATH}")


if __name__ == '__main__':
    build_final_spec()
