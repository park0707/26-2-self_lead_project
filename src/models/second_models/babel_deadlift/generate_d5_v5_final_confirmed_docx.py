# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 데드리프트 (Barbell Deadlift)
Defect: d5. 바벨 궤적과 몸 밀착 (Barbell Path Proximity to Body)
Script: generate_d5_v5_final_confirmed_docx.py
Purpose: 사용자가 승인/확정한 v5.0 최종 배포 규격
         (±0.06 신뢰도 마진 [0.44, 0.56], 커버리지 91.69%, F1 82.07%, Acc 81.30%)
         공식 배포 명세서 (.docx 및 .md) 생성
"""

import os
import sys
import json
from datetime import datetime

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUT_DOCX_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v5.0_확정_및_배포명세서.docx"
OUT_MD_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v5.0_확정_및_배포명세서.md"

COLOR_PRIMARY = RGBColor(26, 54, 93)      # Deep Navy
COLOR_SECONDARY = RGBColor(43, 108, 176)  # Blue
COLOR_HIGHLIGHT = RGBColor(197, 48, 48)   # Crimson Red
COLOR_TEXT_DARK = RGBColor(45, 55, 72)    # Slate Dark
COLOR_MUTED = RGBColor(113, 128, 150)     # Gray


def set_cell_background(cell, fill_hex: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout(doc, text: str, title: str = "사용자 최종 승인 규격", fill_hex="EBF8FF", border_hex="3182CE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, fill_hex)
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)

    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"★ {title}\n")
    run_t.font.name = "Malgun Gothic"
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_SECONDARY

    run_b = p.add_run(text)
    run_b.font.name = "Malgun Gothic"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = COLOR_TEXT_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_sub = p_title.add_run("vidio-health-trainer | 최종 배포 명세서\n")
    r_sub.font.name = "Malgun Gothic"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    r_main = p_title.add_run("바벨 데드리프트 d5 (바벨 궤적과 몸 밀착)\n최종 v5.0 모델 공식 확정 및 운영 배포 명세서")
    r_main.font.name = "Malgun Gothic"
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = COLOR_PRIMARY

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run(f"확정일자: {datetime.now().strftime('%Y년 %m월 %d일')} | 모델 버전: v5.0 Final Confirmed | 검증: 5-Fold StratifiedGroupKFold 전수")
    r_meta.font.name = "Malgun Gothic"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = COLOR_MUTED

    # Callout
    add_callout(
        doc,
        "사용자의 최종 지시에 따라, 바벨 데드리프트 d5(바벨 궤적과 몸 밀착) 모델의 최종 운영 배포 규격이 "
        "신뢰도 마진 ±0.06 (불확실성 보류 구간 [0.44, 0.56], 커버리지 91.69%, 706건 확정 판정, F1-Score 82.07%, 정확도 81.30%)으로 "
        "공식 확정(Sign-off)되었습니다.\n"
        "해당 배포 가중치는 'd5_model_v5_final.pkl'로 고정되었으며, 전용 추론 API(infer_d5_v5_sequence)가 즉시 연동 가능하도록 구축 완료되었습니다.",
        title="v5.0 최종 모델 공식 확정 요약 (F1 82.07% / 커버리지 91.69%)"
    )

    # Section 1
    h1 = doc.add_heading("1. 확정된 최종 v5.0 배포 규격 명세", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    tbl = doc.add_table(rows=8, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    spec_headers = ["항목 구분", "확정 파라미터 / 사양", "상세 설명 및 엔지니어링 근거"]
    for i, h in enumerate(spec_headers):
        c = tbl.cell(0, i)
        set_cell_background(c, "2B6CB0")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    specs = [
        ("공식 모델 버전", "v5.0 Final Confirmed", "외부 AI 3사 종합 자문 반영 및 시점별 동적 캘리브레이션 완성형"),
        ("입력 피처 규격", "40D View-Adaptive Gated", "32D 기하 단축 보정 + 8D 무릎-팔꿈치 상대 기하학(SEK 각도, 돌출도 등)"),
        ("의사결정 정책", "3-State Uncertainty Aware", "판정 확신 구간(확정보통) vs 신뢰도 부족 구간(보류/재촬영 권고)"),
        ("신뢰도 마진 (Margin)", "±0.06", "사후확률 0.50 기준 ±0.06 오차 대역 설정"),
        ("보류(Uncertain) 구간", "[0.44, 0.56]", "애매한 64건(8.31%)을 사용자 재촬영으로 유도하여 오분류 원천 차단"),
        ("실서비스 커버리지", "91.69% (706건 / 770건)", "10명 중 9명 이상에게 즉시 확정 피드백 제공 (초고가용성)"),
        ("최종 달성 성능", "F1 82.07% | Acc 81.30%", "★ 사용자 요구치(80.0%+)를 2.07%p 초과 달성 (ROC-AUC 0.8811)")
    ]

    for row_idx, (cat, val, desc) in enumerate(specs, start=1):
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "EDF2F7"
        if row_idx in [4, 5, 7]:
            bg_col = "FEFCBF"
        for col_idx, text in enumerate([cat, val, desc]):
            c = tbl.cell(row_idx, col_idx)
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if row_idx in [4, 5, 7] and col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = COLOR_HIGHLIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 2: 시점별 임계값
    h2 = doc.add_heading("2. 공식 5대 시점별 캘리브레이션 임계값 및 성능", level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)

    tbl_v = doc.add_table(rows=6, cols=6)
    tbl_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_headers = ["카메라 시점", "시점 명칭 및 각도", "동적 결정 임계값", "F1-Score", "정확도(Acc)", "재현율(Recall)"]
    for i, h in enumerate(v_headers):
        c = tbl_v.cell(0, i)
        set_cell_background(c, "1A365D")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    v_rows = [
        ("시점 A", "우측 후방 (45도)", "0.48", "81.58%", "81.82%", "80.52%"),
        ("시점 B", "우측 전방 (45도)", "0.50", "86.49%", "87.01%", "83.12%"),
        ("시점 C", "정면 (0도) [원근 축소 극복]", "0.32", "66.30%", "59.74%", "79.22%"),
        ("시점 D", "좌측 전방 (45도)", "0.50", "96.00%", "96.10%", "93.51%"),
        ("시점 E", "좌측 후방 (45도)", "0.50", "77.30%", "75.97%", "81.82%")
    ]

    for row_idx, r_data in enumerate(v_rows, start=1):
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "EDF2F7"
        if row_idx == 3:
            bg_col = "FEFCBF"
        for col_idx, text in enumerate(r_data):
            c = tbl_v.cell(row_idx, col_idx)
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            if col_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if row_idx == 3 and col_idx in [2, 3, 5]:
                run.font.bold = True
                run.font.color.rgb = COLOR_HIGHLIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: 파일 아티팩트
    h3 = doc.add_heading("3. 배포 아티팩트 및 추론 API 사용 가이드", level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_after = Pt(4)
    r_code = p_code.add_run(
        "■ 최종 배포 파일 목록:\n"
        "  1. 최종 모델 가중치: D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_v5_final.pkl\n"
        "  2. 배포 메타데이터: D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_v5_final_meta.json\n"
        "  3. 핵심 추론 모듈: D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v5_calibrated.py\n\n"
        "■ 표준 파이프라인 연동 예시 (Python):\n"
        "   from src.models.second_models.babel_deadlift.d5_model_v5_calibrated import infer_d5_v5_sequence\n\n"
        "   # kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 키포인트 배열\n"
        "   result = infer_d5_v5_sequence(kps_64f, view_name='view3', uncertainty_margin=0.06)\n"
        "   # 반환값: {'defect_prob': 0.72, 'prediction': 1, 'is_decided': True, 'decision_label': '결함(바벨 몸 이탈 감지)'}"
    )
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = COLOR_PRIMARY

    doc.save(OUT_DOCX_PATH)
    print(f"[+] DOCX 최종 확정 배포 완료: {OUT_DOCX_PATH}")

    # Markdown
    md_content = f"""# vidio-health-trainer | 최종 배포 명세서
# 바벨 데드리프트 d5 (바벨 궤적과 몸 밀착) 최종 v5.0 모델 공식 확정 및 운영 배포 명세서

- **확정일자:** {datetime.now().strftime('%Y-%m-%d')}
- **모델 버전:** `v5.0 Final Confirmed`
- **의사결정 정책:** 3-State View-Calibrated Uncertainty Gating
- **사용자 확정 사양:**
  - **신뢰도 마진:** $\pm 0.06$ (보류 구간: $[0.44, 0.56]$)
  - **커버리지(채택율):** **91.69%** (706건 / 770건 채택)
  - **최종 F1-Score:** **82.07%** (기본 목표치 80.0%+ 초과 달성)
  - **최종 정확도(Acc):** **81.30%** (ROC-AUC **0.8811**)

---

## 1. 확정된 최종 v5.0 배포 규격 명세

| 항목 구분 | 확정 파라미터 / 사양 | 상세 설명 및 엔지니어링 근거 |
| :--- | :---: | :--- |
| **공식 모델 버전** | `v5.0 Final Confirmed` | 외부 AI 3사 종합 자문 반영 및 시점별 동적 캘리브레이션 완성형 |
| **입력 피처 규격** | 40D View-Adaptive Gated | 32D 기하 단축 보정 + 8D 무릎-팔꿈치 상대 기하학(SEK 각도, 돌출도 등) |
| **의사결정 정책** | 3-State Uncertainty Aware | 판정 확신 구간(확정) vs 신뢰도 부족 구간(보류/재촬영 권고) |
| **신뢰도 마진 (Margin)** | **$\pm 0.06$** | 사후확률 0.50 기준 $\pm 0.06$ 오차 대역 설정 |
| **보류(Uncertain) 구간** | **$[0.44, 0.56]$** | 애매한 64건(8.31%)을 사용자 재촬영으로 유도하여 오분류 원천 차단 |
| **실서비스 커버리지** | **91.69% (706건 / 770건)** | 10명 중 9명 이상에게 즉시 확정 피드백 제공 (초고가용성) |
| **최종 달성 성능** | **F1 82.07% \| Acc 81.30%** | **★ 사용자 요구치(80.0%+)를 2.07%p 초과 달성 (ROC-AUC 0.8811)** |

---

## 2. 공식 5대 시점별 캘리브레이션 임계값 및 성능

| 카메라 시점 | 시점 명칭 및 각도 | 동적 결정 임계값 | F1-Score | 정확도(Acc) | 재현율(Recall) |
| :--- | :--- | :---: | :---: | :---: | :---: |
| **시점 A** | 우측 후방 (45도) | 0.48 | **81.58%** | 81.82% | 80.52% |
| **시점 B** | 우측 전방 (45도) | 0.50 | **86.49%** | 87.01% | 83.12% |
| **시점 C** | **정면 (0도) [원근 축소 극복]** | **0.32** | **66.30%** | **59.74%** | **79.22%** |
| **시점 D** | 좌측 전방 (45도) | 0.50 | **96.00%** | 96.10% | 93.51% |
| **시점 E** | 좌측 후방 (45도) | 0.50 | **77.30%** | 75.97% | 81.82% |

---

## 3. 배포 아티팩트 및 추론 API

```python
from src.models.second_models.babel_deadlift.d5_model_v5_calibrated import infer_d5_v5_sequence

# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 키포인트
result = infer_d5_v5_sequence(kps_64f, view_name="view3", uncertainty_margin=0.06)
print(result)
# {{'defect_prob': 0.72, 'prediction': 1, 'is_decided': True, 'decision_label': '결함(바벨 몸 이탈 감지)'}}
```
"""

    with open(OUT_MD_PATH, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"[+] MD 최종 확정 배포 완료: {OUT_MD_PATH}")


if __name__ == "__main__":
    main()
