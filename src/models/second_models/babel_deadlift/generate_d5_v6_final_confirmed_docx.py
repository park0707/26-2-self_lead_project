# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_d5_v6_final_confirmed_docx.py
Purpose: 바벨 데드리프트 d5 최종 v6.0 Zero-Calibrated 확정 및 운영 배포 명세서
         Word(.docx) 및 Markdown(.md) 동시 생성
Target Paths:
  - D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v6.0_선정_및_운영배포명세서.docx
  - D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v6.0_선정_및_운영배포명세서.md
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DOCX_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v6.0_선정_및_운영배포명세서.docx"
OUT_MD_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v6.0_선정_및_운영배포명세서.md"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_header_box(doc, title, subtitle):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "1E3A8A")  # Navy
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p.add_run(title + "\n")
    run1.font.name = "맑은 고딕"
    run1.font.size = Pt(17)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(255, 255, 255)

    run2 = p.add_run(subtitle)
    run2.font.name = "맑은 고딕"
    run2.font.size = Pt(10.5)
    run2.font.color.rgb = RGBColor(224, 231, 255)

    doc.add_paragraph()


def add_section_title(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(30, 58, 138)
    elif level == 2:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(37, 99, 235)
    else:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(71, 85, 105)


def add_bullet_point(doc, title, desc, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))

    run_t = p.add_run(title + ": ")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 41, 59)

    run_d = p.add_run(desc)
    run_d.font.name = "맑은 고딕"
    run_d.font.size = Pt(9.5)
    run_d.font.color.rgb = RGBColor(71, 85, 105)


def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "맑은 고딕"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    doc.add_paragraph()


def add_callout_box(doc, text, title="핵심 안내"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")  # Blue-50
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    r_title = p.add_run(f"[{title}] ")
    r_title.font.name = "맑은 고딕"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    r_body = p.add_run(text)
    r_body.font.name = "맑은 고딕"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph()


def generate_documents():
    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 헤더 박스
    add_header_box(
        doc,
        "vidio-health-trainer | 최종 모델 선정 및 운영 배포 명세서",
        "바벨 데드리프트 d5(바벨 궤적과 몸 밀착) 최종 v6.0 Zero-Calibrated 공식 선정 및 배포 명세서\n"
        "확정일자: 2026-10-02 | 배포 상태: Official Production Release (최종 확정)"
    )

    # 핵심 요약 박스
    add_callout_box(
        doc,
        "사용자 승인 지침에 따라 바벨 데드리프트 d5 세부 결함(바벨 궤적 밀착)의 최종 프로덕션 모델로 "
        "'v6.0 Zero-Calibrated (37D, 마진 ±0.06)' 사양을 공식 최종 확정합니다.\n"
        "• 실서비스 가용 커버리지: 91.30% (703건 / 770건 확정 판정)\n"
        "• 3-State 최종 확정 F1: ★ 82.34% (정확도 80.65%)\n"
        "• 시점 D(좌측 전방 45°) F1: ★ 96.64% (정밀도 100.00%, FP 0건 무결점)\n"
        "• 체형 극복: 상체 긴 피험자(4번) FN 50% 급감 및 F1 +11.83%p 폭등 달성",
        title="최종 모델 공식 채택 선언"
    )

    # 1. 확정된 최종 운영 규격 명세
    add_section_title(doc, "1. 확정된 최종 운영 규격 명세서")

    headers_spec = ["규격 분류 항목", "최종 확정 사양 / 수치", "엔지니어링 근거 및 상세 설명"]
    data_spec = [
        ["공식 모델 명칭", "v6.0 Zero-Calibrated (37D)", "32D 최적 피처 + 어드레스 영점화 및 체형비율 5D 결합"],
        ["입력 피처 규격", "37D Zero-Calibrated Set", "개인별 셋업 기준선 차감(Δ)을 통해 체형 왜곡 원천 보정"],
        ["의사결정 정책", "3-State View-Calibrated Gating", "5대 시점별 비대칭 임계값 + 불확실성 보류(Reject Option) 결합"],
        ["신뢰도 마진 대역", "±0.06 대역 ([0.44, 0.56])", "모호한 67건(8.70%)을 재촬영 유도하여 오분류 위험 사전 차단"],
        ["실서비스 커버리지", "91.30% (703건 / 770건)", "10명 중 9명 이상에게 즉시 확정 코칭 피드백 제공 (초고가용성)"],
        ["기본 OOF 성능", "F1 80.76% (Acc 79.09%)", "ROC-AUC 0.8797, 결함 재현율 87.79%로 결함 미탐 극소화"],
        ["3-State 최종 F1", "★ 82.34% (Acc 80.65%)", "확정 판정 시 F1 82.3% 이상의 초고신뢰도 운동 피드백 보장"]
    ]
    add_styled_table(doc, headers_spec, data_spec, [1.5, 1.8, 3.2])

    # 2. 공식 5대 시점별 최종 운영 결정 임계값 및 성능
    add_section_title(doc, "2. 공식 5대 시점별 최종 운영 결정 임계값 및 성능")

    headers_view = ["카메라 시점", "공식 명칭 및 각도", "최종 결정 임계값", "F1-Score", "정확도(Acc)", "결함 재현율(Rec)"]
    data_view = [
        ["시점 A", "우측 후방 (45도)", "0.35", "80.24%", "78.57%", "87.01%"],
        ["시점 B", "우측 전방 (45도)", "0.49", "85.71%", "86.36%", "81.82%"],
        ["시점 C", "정면 (0도) [원근 보정]", "0.20", "67.92%", "55.84%", "93.51% ★"],
        ["시점 D", "좌측 전방 (45도) [최적 권장]", "0.50", "★ 96.64%", "★ 96.75%", "93.51% (Prec 100%)"],
        ["시점 E", "좌측 후방 (45도)", "0.49", "79.01%", "77.92%", "83.12%"]
    ]
    add_styled_table(doc, headers_view, data_view, [1.0, 1.8, 1.1, 1.1, 1.1, 1.4])

    # 3. 프로덕션 실서비스 추론 파이프라인 연동 코드
    add_section_title(doc, "3. 프로덕션 실서비스 추론 파이프라인 표준 연동 가이드")

    p = doc.add_paragraph()
    p.add_run(
        "실서비스 추론 엔진 연동 시 아래 표준 인터페이스 함수를 호출하면, 64프레임 키포인트로부터 37D 영점화 피처 추출, "
        "가중 앙상블 추론, 시점별 확률 캘리브레이션 및 3-State 신뢰도 판정까지 원스톱으로 처리됩니다."
    )

    code_text = (
        "from src.models.second_models.babel_deadlift.d5_model_v6 import infer_d5_v6_sequence\n\n"
        "# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 관절 좌표 배열\n"
        "result = infer_d5_v6_sequence(\n"
        "    kps_64f=kps_64f,\n"
        "    view_name='view4',  # 시점 D (좌측 전방 45°)\n"
        "    uncertainty_margin=0.06  # 확정 사양: ±0.06 ([0.44, 0.56])\n"
        ")\n\n"
        "print(result)\n"
        "# 출력 예시:\n"
        "# {\n"
        "#   'defect_prob': 0.071,           # 모델 원시 예측 확률\n"
        "#   'calibrated_prob': 0.081,       # 시점 보정 확률\n"
        "#   'prediction': 0,                # 0: 정상, 1: 결함, -1: 판정 보류\n"
        "#   'is_decided': True,             # 확정 여부 (커버리지 91.3% 내 진입)\n"
        "#   'decision_label': '정상 (바벨 밀착 궤적 양호)',\n"
        "#   'view_name': 'view4',\n"
        "#   'features_dim': 37\n"
        "# }"
    )

    table_code = doc.add_table(rows=1, cols=1)
    table_code.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_code.autofit = False
    c_code = table_code.cell(0, 0)
    c_code.width = Inches(6.5)
    set_cell_background(c_code, "0F172A")  # Slate-900
    set_cell_margins(c_code, top=120, bottom=120, left=160, right=160)
    p_c = c_code.paragraphs[0]
    r_c = p_c.add_run(code_text)
    r_c.font.name = "Consolas"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(226, 232, 240)
    doc.add_paragraph()

    # 4. 배포 산출물 및 무결성 검증 내역
    add_section_title(doc, "4. 프로덕션 배포 산출물 및 무결성 검증 내역")
    add_bullet_point(doc, "모델 가중치 파일", "D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_v6_best.pkl")
    add_bullet_point(doc, "5-Fold 가중치 일체", "D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_v6_fold1~5.pkl")
    add_bullet_point(doc, "최종 메타데이터", "D:\\src\\models\\second_models\\babel_deadlift\\d5_model_final_v6_meta.json")
    add_bullet_point(doc, "피처 추출 엔진", "D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v6_features.py (37D)")
    add_bullet_point(doc, "추론 엔진 모듈", "D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v6.py (Tri-Ensemble)")
    add_bullet_point(doc, "단일 시퀀스 검증", "test_d5_v6_inference.py 5대 시점 전수 100% 통과 (Shape, Type, Range 무결성 확인)")

    doc.save(OUT_DOCX_PATH)
    print(f"[+] DOCX Release Spec Created: {OUT_DOCX_PATH}")

    # 마크다운 파일 동시 생성
    md_content = f"""# vidio-health-trainer | 최종 모델 선정 및 운영 배포 명세서
# 바벨 데드리프트 d5 (바벨 궤적과 몸 밀착) 최종 v6.0 Zero-Calibrated 모델 공식 선정 및 운영 배포 명세서

- **최종선정일:** 2026-10-02
- **공식 모델 명칭:** `v6.0 Zero-Calibrated (37D)`
- **배포 상태:** **Official Production Release (최종 확정)**
- **확정 사양:** **신뢰도 마진 $\pm 0.06$ 대역 [$0.44, 0.56$]** | 커버리지 **91.30% (703건 / 770건)** | 3-State F1 **★ 82.34%** (정확도 **80.65%**)

---

## 1. 확정된 최종 운영 규격 명세서

| 규격 분류 항목 | 최종 확정 사양 / 수치 | 엔지니어링 근거 및 상세 설명 |
| :--- | :---: | :--- |
| **공식 모델 명칭** | `v6.0 Zero-Calibrated (37D)` | 32D 최적 피처 + 어드레스 영점화 및 체형비율 5D 결합 모델 |
| **입력 피처 규격** | **37D Zero-Calibrated Feature Set** | 개인별 셋업 기준선 차감($\Delta$)을 통해 체형 왜곡 원천 보정 |
| **의사결정 정책** | **3-State View-Calibrated Gating** | 5대 시점별 비대칭 임계값 + 불확실성 보류(Reject Option) 결합 |
| **신뢰도 마진 대역** | **$\pm 0.06$ [$0.44, 0.56$]** | 모호한 67건(8.70%)을 재촬영 유도하여 오분류 위험 사전 차단 |
| **실서비스 커버리지** | **91.30% (703건 / 770건)** | 10명 중 9명 이상에게 즉시 확정 코칭 피드백 제공 (초고가용성) |
| **기본 OOF 성능** | **F1 80.76% (Acc 79.09%)** | ROC-AUC **0.8797**, 결함 재현율 **87.79%**로 결함 미탐 극소화 |
| **3-State 최종 F1** | **★ 82.34% (Acc 80.65%)** | 확정 판정 시 F1 82.3% 이상의 초고신뢰도 운동 피드백 보장 |

---

## 2. 공식 5대 시점별 최종 운영 결정 임계값 및 성능

| 카메라 시점 | 공식 명칭 및 각도 | 최종 결정 임계값 | F1-Score | 정확도(Acc) | 결함 재현율(Rec) |
| :--- | :--- | :---: | :---: | :---: | :---: |
| **시점 A** | 우측 후방 (45도) | 0.35 | **80.24%** | 78.57% | **87.01%** |
| **시점 B** | 우측 전방 (45도) | 0.49 | **85.71%** | 86.36% | **81.82%** |
| **시점 C** | 정면 (0도) [원근 보정] | 0.20 | **67.92%** | 55.84% | **93.51% ★** |
| **시점 D** | **좌측 전방 (45도) [최적 권장]** | **0.50** | **★ 96.64%** | **★ 96.75%** | **93.51% (Prec 100%)** |
| **시점 E** | 좌측 후방 (45도) | 0.49 | **79.01%** | 77.92% | **83.12%** |

---

## 3. 프로덕션 실서비스 추론 파이프라인 연동 코드

```python
from src.models.second_models.babel_deadlift.d5_model_v6 import infer_d5_v6_sequence

# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 관절 좌표 배열
result = infer_d5_v6_sequence(
    kps_64f=kps_64f,
    view_name="view4",  # 시점 D (좌측 전방 45°)
    uncertainty_margin=0.06  # 최종 확정 사양: ±0.06 대역
)

print(result)
# 출력 예시:
# {{
#   'defect_prob': 0.071,
#   'calibrated_prob': 0.081,
#   'prediction': 0,                # 0: 정상, 1: 결함, -1: 판정 보류
#   'is_decided': True,             # 커버리지 91.3% 내 확정 판정
#   'decision_label': '정상 (바벨 밀착 궤적 양호)',
#   'view_name': 'view4',
#   'features_dim': 37
# }}
```

---

## 4. 프로덕션 배포 산출물 및 무결성 검증

1. **최종 가중치 파일**: `D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_v6_best.pkl`
2. **최종 메타데이터**: `D:\\src\\models\\second_models\\babel_deadlift\\d5_model_final_v6_meta.json`
3. **피처 추출 엔진**: `D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v6_features.py` (37D)
4. **추론 엔진 모듈**: `D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v6.py` (Tri-Ensemble)
5. **무결성 검증**: `test_d5_v6_inference.py` 5대 시점 전수 100% 통과 (Shape, Type, Range 무결성 확인)
"""

    with open(OUT_MD_PATH, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[+] MD Release Spec Created: {OUT_MD_PATH}")


if __name__ == "__main__":
    generate_documents()
