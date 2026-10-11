# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: generate_v6_official_report_docx.py
Purpose:
  바벨 데드리프트 통합 진단 엔진 v6.0 (Two-Stream Hybrid: GBDT Meta + Kinematic 1D-TCN)
  공식 성능 평가 보고서 생성 (.docx & .md)
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_DIR_6TH = r"D:\자주프 문서\6주차"
DOCX_DIR_ROOT = r"D:\자주프 문서"
os.makedirs(DOCX_DIR_6TH, exist_ok=True)
os.makedirs(DOCX_DIR_ROOT, exist_ok=True)

TARGET_DOCX_6TH = os.path.join(DOCX_DIR_6TH, "바벨_데드리프트_통합진단엔진_v6.0_TwoStream_하이브리드_공식_성능보고서.docx")
TARGET_DOCX_ROOT = os.path.join(DOCX_DIR_ROOT, "바벨_데드리프트_통합진단엔진_v6.0_TwoStream_하이브리드_공식_성능보고서.docx")
TARGET_MD_6TH = os.path.join(DOCX_DIR_6TH, "바벨_데드리프트_통합진단엔진_v6.0_TwoStream_하이브리드_공식_성능보고서.md")
TARGET_MD_ROOT = os.path.join(DOCX_DIR_ROOT, "바벨_데드리프트_통합진단엔진_v6.0_TwoStream_하이브리드_공식_성능보고서.md")

doc = Document()

for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

NAVY = RGBColor(26, 54, 93)
BLUE = RGBColor(43, 108, 176)
GRAY = RGBColor(113, 128, 150)
BLACK = RGBColor(45, 55, 72)
DARK_GREEN = RGBColor(39, 103, 73)
CRIMSON = RGBColor(155, 44, 44)

def style_heading(p, text, level=1):
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = NAVY
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.color.rgb = BLUE
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = BLACK

def add_callout_box(doc, text_list, title="", border_color="2B6CB0", bg_color="F7FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.8)
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    if title:
        run_t = p.add_run(f"■ {title}\n")
        run_t.font.name = "맑은 고딕"
        run_t.font.size = Pt(10.5)
        run_t.bold = True
        run_t.font.color.rgb = BLUE
    for line in text_list:
        run = p.add_run(line + "\n")
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9.5)
        run.font.color.rgb = BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def set_cell_properties(cell, bg_color=None, bold=False, color=BLACK, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    if bg_color:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            r.font.name = "맑은 고딕"
            r.font.size = Pt(font_size)
            r.bold = bold
            r.font.color.rgb = color

# ================= 본문 =================
p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(2)
r_meta = p_meta.add_run("vidio-health-trainer | 인공지능 헬스 코칭 시스템 공식 기술 보고서")
r_meta.font.name = "맑은 고딕"
r_meta.font.size = Pt(9)
r_meta.font.color.rgb = GRAY

p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(4)
p_title.paragraph_format.space_after = Pt(4)
r_title = p_title.add_run("바벨 데드리프트 통합 진단 엔진 v6.0\nTwo-Stream Hybrid 공식 성능 평가 보고서")
r_title.font.name = "맑은 고딕"
r_title.font.size = Pt(18)
r_title.bold = True
r_title.font.color.rgb = NAVY

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(12)
r_sub = p_sub.add_run("전문가 도메인 메타 모델(GBDT)과 시계열 동역학 신경망(1D-TCN)의 융합을 통한 EMR 35.19% 및 3-State 고신뢰도 진단 달성")
r_sub.font.name = "맑은 고딕"
r_sub.font.size = Pt(11)
r_sub.font.color.rgb = BLUE

add_callout_box(
    doc,
    [
        "1. 아키텍처 혁신: 도메인 통계 모델(Stream 1)과 64프레임 20채널 생체역학 1D-TCN(Stream 2)을 5:5로 결합하고, 생체역학 법칙(스트리퍼 풀 d4 발생 시 바벨 이격 d5 100% 필연성)을 연동한 Two-Stream Hybrid 구조를 공식 확정함.",
        "2. 100% Full-Coverage (전수 채점, 판정 보류 없음) 성능:",
        "   - 전수 EMR: 29.87% (v5.0) -> 35.19% (v6.0 표준) / 37.01% (v6.0 보정) 달성 (+5.32%p~+7.14%p 도약, 271~285건 완전 적중)",
        "   - 정상 특이도: 62.50% -> 84.43% (+21.93%p 급상승, 정상 오탐 극적 억제)",
        "   - Clean F1-Score: 58.48% -> 63.27% (+4.79%p 상승, 정밀도 55.4%, 재현율 73.8%)",
        "   - 5대 전 결함 동시 향상: 특히 d4(동시성) F1 49.43% -> 58.78% (+9.35%p 대폭 향상)",
        "3. 3-State 선택적 판정 (Selective Classification, 판정 보류 적용) 성능:",
        "   - 최소 마진 신뢰도(Min-Margin) 기반 판정 보류 도입 시, 보류율 47.8%(커버리지 52.2%)에서 EMR 45.52%, Clean F1 71.81% 달성",
        "   - 고신뢰도 모드(보류율 73.0%, 커버리지 27.0%)에서는 EMR 57.69%, Clean F1 79.82% 달성"
    ],
    title="v6.0 핵심 성능 요약 (Executive Summary)",
    border_color="1A365D",
    bg_color="F0F4F8"
)

# 1. 버전별 종합 성능 진화 추이
style_heading(doc.add_paragraph(), "1. 데드리프트 엔진 버전별 성능 진화 추이 (v1 ~ v6)", level=1)

tbl_evo = doc.add_table(rows=6, cols=6)
tbl_evo.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_evo.autofit = False
for col, w in zip(tbl_evo.columns, [Inches(1.5), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.0), Inches(1.0)]):
    col.width = w

headers_evo = ["엔진 버전", "전수 EMR", "Clean Spec", "Clean F1", "d4 F1", "완전 적중수"]
for j, h in enumerate(headers_evo):
    bg = "1A365D" if j == 1 else "2B6CB0"
    set_cell_properties(tbl_evo.cell(0, j), bg_color=bg, bold=True, color=RGBColor(255, 255, 255), font_size=9)
    tbl_evo.cell(0, j).paragraphs[0].text = h

evo_data = [
    ["독립 베이스라인", "24.03%", "50.00%", "53.16%", "42.63%", "185 / 770"],
    ["통합 v3.0 (Two-Track)", "27.01%", "66.88%", "57.07%", "48.20%", "208 / 770"],
    ["통합 v4.0 (View-Cond)", "29.22%", "66.88%", "57.99%", "50.56%", "225 / 770"],
    ["통합 v5.0 (Fail-Safe)", "29.87%", "62.50%", "58.48%", "49.43%", "230 / 770"],
    ["통합 v6.0 (Two-Stream)", "35.19% ★", "84.43% ★", "63.27% ★", "58.78% ★", "271 / 770 (+41건)"]
]

for i, row in enumerate(evo_data, 1):
    bg = "EBF8FF" if i == 5 else ("F7FAFC" if i % 2 == 1 else "FFFFFF")
    for j, val in enumerate(row):
        cell = tbl_evo.cell(i, j)
        is_hl = (i == 5 and "★" in val)
        set_cell_properties(cell, bg_color=bg, bold=(j==0 or is_hl), color=DARK_GREEN if is_hl else BLACK, font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER)
        cell.paragraphs[0].text = val

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 2. 5대 카메라 시점별 세부 성능
style_heading(doc.add_paragraph(), "2. 5대 카메라 시점별 세부 성능 (v6.0 전수 감사)", level=1)

tbl_view = doc.add_table(rows=6, cols=5)
tbl_view.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_view.autofit = False
for col, w in zip(tbl_view.columns, [Inches(1.8), Inches(1.2), Inches(1.2), Inches(1.3), Inches(1.3)]):
    col.width = w

headers_view = ["카메라 시점 규격", "샘플 수", "v5.0 EMR", "v6.0 EMR", "v6.0 Clean F1"]
for j, h in enumerate(headers_view):
    set_cell_properties(tbl_view.cell(0, j), bg_color="2B6CB0", bold=True, color=RGBColor(255, 255, 255), font_size=9)
    tbl_view.cell(0, j).paragraphs[0].text = h

view_data = [
    ["시점 D (좌측 전방 45°)", "154건", "37.66%", "46.75% ★", "74.42%"],
    ["시점 B (우측 전방 45°)", "154건", "32.47%", "38.96% ★", "70.42%"],
    ["시점 E (좌측 후방 45°)", "154건", "24.03%", "34.42% ★", "51.35%"],
    ["시점 C (정면 0°)", "154건", "32.47%", "31.17%", "64.71%"],
    ["시점 A (우측 후방 45°)", "154건", "22.73%", "24.68% ★", "54.05%"]
]

for i, row in enumerate(view_data, 1):
    bg = "F7FAFC" if i % 2 == 1 else "FFFFFF"
    for j, val in enumerate(row):
        cell = tbl_view.cell(i, j)
        is_hl = ("★" in val)
        set_cell_properties(cell, bg_color=bg, bold=(j==0 or is_hl), color=DARK_GREEN if is_hl else BLACK, font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER)
        cell.paragraphs[0].text = val

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 3. 3-State 선택적 판정 트레이드오프 분석
style_heading(doc.add_paragraph(), "3. 3-State 선택적 판정 (판정 보류 적용) 트레이드오프 곡선", level=1)

p_3state = doc.add_paragraph()
p_3state.paragraph_format.space_after = Pt(4)
p_3state.add_run(
    "사용자의 질문에 대한 핵심 답변: 상기 1절의 EMR 35.19%는 '판정 보류를 전혀 적용하지 않은 100% Full-Coverage' 수치입니다.\n"
    "여기에 3-State 선택적 판정(판정 보류, Selective Prediction)을 활성화하면, "
    "최소 마진 신뢰도(Min-Margin Confidence)가 임계치 미만인 모호한 샘플을 보류 처리하여 판정 성공률을 극적으로 높일 수 있습니다:"
)

tbl_3s = doc.add_table(rows=7, cols=5)
tbl_3s.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_3s.autofit = False
for col, w in zip(tbl_3s.columns, [Inches(1.5), Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.4)]):
    col.width = w

headers_3s = ["신뢰도 임계치 (tau)", "커버리지 (판정율)", "보류율 (Abstain)", "선택 EMR", "선택 Clean F1"]
for j, h in enumerate(headers_3s):
    set_cell_properties(tbl_3s.cell(0, j), bg_color="2B6CB0", bold=True, color=RGBColor(255, 255, 255), font_size=9)
    tbl_3s.cell(0, j).paragraphs[0].text = h

data_3s = [
    ["tau = 0.00 (전수 판정)", "100.0% (770/770)", "0.0%", "35.19%", "63.27%"],
    ["tau = 0.05", "52.2% (402/770)", "47.8%", "45.52%", "71.81% ★"],
    ["tau = 0.10", "27.0% (208/770)", "73.0%", "57.69%", "79.82% ★"],
    ["tau = 0.15", "13.1% (101/770)", "86.9%", "66.34%", "85.07% ★"],
    ["tau = 0.20", "6.0% (46/770)", "94.0%", "69.57%", "82.35% ★"],
    ["tau = 0.25 (초고신뢰)", "1.6% (12/770)", "98.4%", "75.00%", "84.21% ★"]
]

for i, row in enumerate(data_3s, 1):
    bg = "EBF8FF" if i in [2, 3] else ("F7FAFC" if i % 2 == 1 else "FFFFFF")
    for j, val in enumerate(row):
        cell = tbl_3s.cell(i, j)
        is_hl = ("★" in val)
        set_cell_properties(cell, bg_color=bg, bold=(j==0 or is_hl), color=DARK_GREEN if is_hl else BLACK, font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER)
        cell.paragraphs[0].text = val

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 4. 배포 산출물 및 결론
style_heading(doc.add_paragraph(), "4. 모델 가중치 번들 및 배포 명세", level=1)

p_dep = doc.add_paragraph()
p_dep.paragraph_format.space_after = Pt(6)
p_dep.add_run(
    "v6.0 Two-Stream Hybrid 파이프라인의 모델 가중치와 검증 데이터가 공식 배포 완료되었습니다:\n\n"
    "1. 파이프라인 스크립트: D:\\src\\models\\second_models\\babel_deadlift\\run_deadlift_integrated_v6_pipeline.py\n"
    "2. 모델 가중치 번들: D:\\src\\models\\second_models\\babel_deadlift\\weights_integrated_v6\\deadlift_integrated_v6_bundle.pkl\n"
    "   - 5-Fold 1D-TCN 신경망 State Dicts (PyTorch)\n"
    "   - 5-Fold 결합 가중치 및 생체역학 연동 규칙\n"
    "3. 벤치마크 JSON 데이터: D:\\src\\models\\second_models\\babel_deadlift\\deadlift_integrated_v6_benchmark_results.json\n\n"
    "본 v6.0 엔진은 100% Full-Coverage 모드에서도 35.19% EMR을 보장하며, "
    "실제 사용자 서비스에서는 3-State Selective 판정을 활성화하여 Clean F1 70% 이상의 초고신뢰 코칭을 제공할 수 있습니다."
)

doc.save(TARGET_DOCX_6TH)
doc.save(TARGET_DOCX_ROOT)
print(f"[*] v6.0 Word 보고서 생성 완료:\n  - {TARGET_DOCX_6TH}\n  - {TARGET_DOCX_ROOT}")

md_content = f"""# [vidio-health-trainer] 바벨 데드리프트 통합 진단 엔진 v6.0 공식 보고서

## 1. 개요 및 핵심 질문 답변
- **질문 1**: "v5.0 + temp 1 하이브리드 융합 결과를 v6으로 만들어 줘"
  -> **반영 완료**: `run_deadlift_integrated_v6_pipeline.py`를 공식 Two-Stream Hybrid 아키텍처로 구현 및 가중치 번들(`deadlift_integrated_v6_bundle.pkl`) 저장 완료.
- **질문 2**: "그리고 3-state 적용한 건가?"
  -> **답변**: **아닙니다!** 앞서 보고드린 **EMR 35.19%(271건 적중)는 판정 보류를 전혀 하지 않은 100% Full-Coverage(전수 채점)** 수치입니다.
  -> 3-State Selective Classification(판정 보류)을 적용하면 모호한 샘플을 보류 처리하여 **EMR 45.52% ~ 75.00%, Clean F1 71.81% ~ 85.07%**까지 올라갑니다.

## 2. v6.0 전수 성능 대조표 (100% Full Coverage 기준)
| 평가 지표 | v5.0 (이전 채택) | v6.0 Two-Stream (공식 채택) | 개선폭 |
| :--- | :---: | :---: | :---: |
| **전수 EMR (완전 일치율)** | 29.87% (230건) | **35.19% (271건) ★** | **+5.32%p (+41건 적중)** |
| **정상 특이도 (Clean Spec)** | 62.50% | **84.43% ★** | **+21.93%p (정상 오탐 억제)** |
| **Clean F1-Score** | 58.48% | **63.27% ★** | **+4.79%p** |
| **Hamming Loss** | 0.2532 | **0.2088 ★** | **17.5% 오차 감소** |
| d1 F1 (척추 중립 위반) | 56.85% | **61.25%** | +4.40%p |
| d2 F1 (무릎-발목 모임) | 48.88% | **53.02%** | +4.14%p |
| d3 F1 (발바닥 접지 불안) | 60.47% | **66.79%** | +6.32%p |
| d4 F1 (무릎-골반 동시성) | 49.43% | **58.78% ★** | **+9.35%p 대폭 개선** |
| d5 F1 (바벨 몸 이격) | 76.38% | **78.54%** | +2.16%p |

## 3. 3-State 선택적 판정 (판정 보류 적용 시) 성능
| 신뢰도 임계치 (tau) | 판정 커버리지 | 보류율 | 선택 EMR | 선택 Clean F1 |
| :--- | :---: | :---: | :---: | :---: |
| tau = 0.00 (전수 판정) | 100.0% | 0.0% | **35.19%** | 63.27% |
| tau = 0.05 | 52.2% | 47.8% | **45.52%** | **71.81%** |
| tau = 0.10 | 27.0% | 73.0% | **57.69%** | **79.82%** |
| tau = 0.15 | 13.1% | 86.9% | **66.34%** | **85.07%** |
| tau = 0.20 | 6.0% | 94.0% | **69.57%** | **82.35%** |
| tau = 0.25 | 1.6% | 98.4% | **75.00%** | **84.21%** |
"""

with open(TARGET_MD_6TH, "w", encoding="utf-8") as f:
    f.write(md_content)
with open(TARGET_MD_ROOT, "w", encoding="utf-8") as f:
    f.write(md_content)

print(f"[*] v6.0 Markdown 보고서 생성 완료:\n  - {TARGET_MD_6TH}\n  - {TARGET_MD_ROOT}")
