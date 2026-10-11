# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 바벨 데드리프트 (Barbell Deadlift)
Defect: d5. 바벨 궤적과 몸 밀착 (Barbell Path Proximity to Body)
Script: generate_d5_v5_1_final_release_docx.py
Purpose: 사용자가 선정한 v5.1 Pruned (32D 최적화) 최종 모델에 대한
         공식 운영 배포 명세서 (.docx 및 .md) 생성
"""

import os
import sys
import json
from datetime import datetime

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUT_DOCX_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v5.1_선정_및_운영배포명세서.docx"
OUT_MD_PATH = r"D:\자주프 문서\6주차\바벨_데드리프트_d5_최종_v5.1_선정_및_운영배포명세서.md"
JSON_META_PATH = r"d:\src\models\second_models\babel_deadlift\weights\d5_model_final_meta.json"

COLOR_PRIMARY = RGBColor(26, 54, 93)      # Deep Navy
COLOR_SECONDARY = RGBColor(43, 108, 176)  # Blue
COLOR_HIGHLIGHT = RGBColor(197, 48, 48)   # Crimson Red
COLOR_TEXT_DARK = RGBColor(45, 55, 72)    # Slate Dark
COLOR_MUTED = RGBColor(113, 128, 150)     # Gray


def set_cell_background(cell, fill_hex: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout(doc, text: str, title: str = "사용자 공식 선정 요약", fill_hex="EBF8FF", border_hex="3182CE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, fill_hex)
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)

    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"★ {title}\n")
    run_t.font.name = "Malgun Gothic"
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_SECONDARY

    run_b = p.add_run(text)
    run_b.font.name = "Malgun Gothic"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = COLOR_TEXT_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    os.makedirs(os.path.dirname(OUT_DOCX_PATH), exist_ok=True)
    with open(JSON_META_PATH, 'r', encoding='utf-8') as f:
        meta = json.load(f)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_sub = p_title.add_run("vidio-health-trainer | 최종 모델 선정 및 프로덕션 배포 명세서\n")
    r_sub.font.name = "Malgun Gothic"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    r_main = p_title.add_run("바벨 데드리프트 d5 (바벨 궤적과 몸 밀착)\n최종 v5.1 Pruned 모델 공식 선정 및 운영 배포 명세서")
    r_main.font.name = "Malgun Gothic"
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = COLOR_PRIMARY

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run(f"최종선정일: {datetime.now().strftime('%Y년 %m월 %d일')} | 모델 버전: v5.1 Pruned (32D) | 상태: Official Production Release")
    r_meta.font.name = "Malgun Gothic"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = COLOR_MUTED

    # Callout
    add_callout(
        doc,
        "사용자의 최종 결정에 따라, 바벨 데드리프트 d5(바벨 궤적과 몸 밀착) 결함 분류의 최종 프로덕션 모델로 "
        "'v5.1 Pruned (32D 최적화 모델)'이 공식 선정(Final Sign-off)되었습니다.\n"
        "8개 저해 및 중복 피처를 영구 제거하여 20% 경량화(40D -> 32D)를 달성함과 동시에, "
        "기본 OOF F1 81.60% (정확도 80.26%, 재현율 87.53%, ROC-AUC 0.8841), 정면 뷰(시점 C) F1 70.00% 돌파, "
        "3-State 신뢰도 판정(±0.06 마진) F1 82.57% (커버리지 91.43%)의 최고 사양으로 운영 가중치(d5_model_final.pkl)가 영구 고정 배포되었습니다.",
        title="v5.1 Pruned 최종 모델 공식 선정 요약"
    )

    # Section 1
    h1 = doc.add_heading("1. 확정된 최종 운영 규격 명세서", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    tbl = doc.add_table(rows=8, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    spec_headers = ["규격 분류 항목", "최종 확정 사양 / 수치", "엔지니어링 근거 및 상세 설명"]
    for i, h in enumerate(spec_headers):
        c = tbl.cell(0, i)
        set_cell_background(c, "2B6CB0")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    specs = [
        ("공식 모델 명칭", "v5.1 Pruned (32D Optimized)", "40D 감사 기반 8개 저해/중복 피처 소거 완료 최적화 모델"),
        ("입력 피처 규격", "32D Optimized Feature Set", "노이즈 및 |r| >= 0.85 다중공선성 원천 차단 (20% 경량화)"),
        ("의사결정 정책", "3-State View-Calibrated Gating", "시점별 최적 결정선 캘리브레이션 + 불확실성 보류 모드"),
        ("신뢰도 마진 대역", "±0.06 [0.44, 0.56]", "애매한 66건(8.57%)을 재촬영 유도하여 오분류 원천 방어"),
        ("실서비스 커버리지", "91.43% (704건 / 770건)", "10명 중 9명 이상에게 즉시 확정 피드백 제공 (초고가용성)"),
        ("100% 기본 OOF F1", "81.60% (Acc 80.26%)", "ROC-AUC 0.8841 (결함 재현율 87.53%로 결함 미탐 최소화)"),
        ("3-State 최종 F1", "82.57% (Acc 81.11%)", "★ 확정 판정 시 F1 82.5% 이상의 초고신뢰도 운동 코칭 보장")
    ]

    for row_idx, (cat, val, desc) in enumerate(specs, start=1):
        bg = "F7FAFC" if row_idx % 2 == 1 else "EDF2F7"
        if row_idx in [2, 6, 7]:
            bg = "FEFCBF"
        for col_idx, text in enumerate([cat, val, desc]):
            c = tbl.cell(row_idx, col_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if row_idx in [2, 6, 7] and col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = COLOR_HIGHLIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 2: 시점별 임계값
    h2 = doc.add_heading("2. 공식 5대 시점별 최종 운영 결정 임계값 및 성능", level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)

    tbl_v = doc.add_table(rows=6, cols=6)
    tbl_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_headers = ["카메라 시점", "공식 명칭 및 각도", "최종 결정 임계값", "F1-Score", "정확도(Acc)", "결함 재현율(Rec)"]
    for i, h in enumerate(v_headers):
        c = tbl_v.cell(0, i)
        set_cell_background(c, "1A365D")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    v_rows = [
        ("시점 A", "우측 후방 (45도)", "0.32", "81.66%", "79.87%", "89.61%"),
        ("시점 B", "우측 전방 (45도)", "0.54", "85.71%", "86.36%", "81.82%"),
        ("시점 C", "정면 (0도) [원근 소실 극복]", "0.26", "70.00% ★", "61.04%", "90.91% ★"),
        ("시점 D", "좌측 전방 (45도)", "0.44", "96.00%", "96.10%", "93.51%"),
        ("시점 E", "좌측 후방 (45도)", "0.52", "78.75%", "77.92%", "81.82%")
    ]

    for row_idx, r_data in enumerate(v_rows, start=1):
        bg = "F7FAFC" if row_idx % 2 == 1 else "EDF2F7"
        if row_idx == 3:
            bg = "FEFCBF"
        for col_idx, text in enumerate(r_data):
            c = tbl_v.cell(row_idx, col_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            if col_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if row_idx == 3 and col_idx in [2, 3, 5]:
                run.font.bold = True
                run.font.color.rgb = COLOR_HIGHLIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: 배포 아티팩트
    h3 = doc.add_heading("3. 프로덕션 배포 아티팩트 및 추론 API 엔드포인트", level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_after = Pt(4)
    r_code = p_code.add_run(
        "■ 최종 운영 배포 파일 목록:\n"
        "  1. 최종 운영 가중치: D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_final.pkl\n"
        "  2. 배포 메타데이터: D:\\src\\models\\second_models\\babel_deadlift\\weights\\d5_model_final_meta.json\n"
        "  3. 핵심 추론 모듈: D:\\src\\models\\second_models\\babel_deadlift\\d5_model_v5_1.py\n\n"
        "■ 프로덕션 실서비스 추론 파이프라인 연동 코드:\n"
        "   from src.models.second_models.babel_deadlift.d5_model_v5_1 import infer_d5_v5_1_sequence\n\n"
        "   # kps_64f: (64, 33, 3) 전처리된 MediaPipe 관절 키포인트 시퀀스\n"
        "   # view_name: 'view1' ~ 'view5' (공식 시점 코드)\n"
        "   result = infer_d5_v5_1_sequence(kps_64f, view_name='view3', uncertainty_margin=0.06)\n\n"
        "   print(result)\n"
        "   # 출력: {'defect_prob': 0.572, 'prediction': 1, 'is_decided': True, 'decision_label': '결함(바벨 몸 이탈 감지)', 'features_dim': 32}"
    )
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = COLOR_PRIMARY

    doc.save(OUT_DOCX_PATH)
    print(f"[+] DOCX 최종 선정 및 운영 배포 명세서 생성 완료: {OUT_DOCX_PATH}")

    # Markdown
    md_content = f"""# vidio-health-trainer | 최종 모델 선정 및 프로덕션 배포 명세서
# 바벨 데드리프트 d5 (바벨 궤적과 몸 밀착) 최종 v5.1 Pruned 모델 공식 선정 및 운영 배포 명세서

- **최종선정일:** {datetime.now().strftime('%Y-%m-%d')}
- **공식 모델 명칭:** `v5.1 Pruned (32D Optimized)`
- **배포 상태:** **Official Production Release (최종 확정)**
- **주요 성능:** 기본 F1 **81.60%** (정면 시점 C **70.00%** 돌파) | 3-State F1 **82.57%** (커버리지 **91.43%**)

---

## 1. 확정된 최종 운영 규격 명세서

| 규격 분류 항목 | 최종 확정 사양 / 수치 | 엔지니어링 근거 및 상세 설명 |
| :--- | :---: | :--- |
| **공식 모델 명칭** | `v5.1 Pruned (32D Optimized)` | 40D 감사 기반 8개 저해/중복 피처 소거 완료 최적화 모델 |
| **입력 피처 규격** | **32D Optimized Feature Set** | 노이즈 및 \|r\| >= 0.85 다중공선성 원천 차단 (20% 경량화) |
| **의사결정 정책** | 3-State View-Calibrated Gating | 시점별 최적 결정선 캘리브레이션 + 불확실성 보류 모드 |
| **신뢰도 마진 대역** | **$\pm 0.06$ [$0.44, 0.56$]** | 애매한 66건(8.57%)을 재촬영 유도하여 오분류 원천 방어 |
| **실서비스 커버리지** | **91.43% (704건 / 770건)** | 10명 중 9명 이상에게 즉시 확정 피드백 제공 (초고가용성) |
| **100% 기본 OOF F1** | **81.60% (Acc 80.26%)** | ROC-AUC **0.8841** (결함 재현율 **87.53%**로 결함 미탐 최소화) |
| **3-State 최종 F1** | **★ 82.57% (Acc 81.11%)** | 확정 판정 시 F1 82.5% 이상의 초고신뢰도 운동 코칭 보장 |

---

## 2. 공식 5대 시점별 최종 운영 결정 임계값 및 성능

| 카메라 시점 | 공식 명칭 및 각도 | 최종 결정 임계값 | F1-Score | 정확도(Acc) | 결함 재현율(Rec) |
| :--- | :--- | :---: | :---: | :---: | :---: |
| **시점 A** | 우측 후방 (45도) | 0.32 | **81.66%** | 79.87% | **89.61%** |
| **시점 B** | 우측 전방 (45도) | 0.54 | **85.71%** | 86.36% | 81.82% |
| **시점 C** | **정면 (0도) [원근 소실 극복]** | **0.26** | **70.00% ★** | 61.04% | **90.91% ★** |
| **시점 D** | 좌측 전방 (45도) | 0.44 | **96.00%** | 96.10% | 93.51% |
| **시점 E** | 좌측 후방 (45도) | 0.52 | **78.75%** | 77.92% | **81.82%** |

---

## 3. 프로덕션 실서비스 추론 파이프라인 연동 코드

```python
from src.models.second_models.babel_deadlift.d5_model_v5_1 import infer_d5_v5_1_sequence

# kps_64f: (64, 33, 3) 3차 스플라인 보간된 MediaPipe 키포인트
result = infer_d5_v5_1_sequence(kps_64f, view_name="view3", uncertainty_margin=0.06)
print(result)
# {{
#   'defect_prob': 0.572,
#   'prediction': 1,
#   'is_decided': True,
#   'decision_label': '결함(바벨 몸 이탈 감지)',
#   'view_name': 'view3',
#   'uncertainty_band': [0.44, 0.56],
#   'features_dim': 32
# }}
```
"""

    with open(OUT_MD_PATH, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"[+] MD 최종 선정 및 운영 배포 명세서 생성 완료: {OUT_MD_PATH}")


if __name__ == "__main__":
    main()
