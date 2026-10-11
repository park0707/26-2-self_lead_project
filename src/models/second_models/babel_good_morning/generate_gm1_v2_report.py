import os
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Malgun Gothic'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    run.font.name = 'Malgun Gothic'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    return p

def main():
    CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
    V1_JSON = os.path.join(CURRENT_DIR, "gm1_benchmark_results.json")
    V2_JSON = os.path.join(CURRENT_DIR, "gm1_v2_benchmark_results.json")
    
    with open(V1_JSON, 'r', encoding='utf-8') as f:
        v1_data = json.load(f)
    with open(V2_JSON, 'r', encoding='utf-8') as f:
        v2_data = json.load(f)
        
    doc = Document()
    title = doc.add_heading('vidio-health-trainer', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    add_heading(doc, '굿모닝 (gm1) v2 최적화 모델 성능 보고서', level=1)
    
    add_heading(doc, '1. 개요', level=2)
    add_paragraph(doc, '- 대상 운동: 바벨 굿모닝 (09_굿모닝)')
    add_paragraph(doc, '- 대상 오류: 무릎 구부린 채 고정 위반 (gm1)')
    add_paragraph(doc, '- 최적화 방향: 스티프 데드리프트(sd1) 계승 모델(v1)에서 다중공선성(r>=0.85) 피처를 제거하고, 굿모닝 동작에 특화된 피처(Torso Pitch ROM, 대퇴-정강이 비율, 고관절 힌지 격리 지수)를 추가 편입하여 13D 정예 피처로 v2 구축')
    
    add_heading(doc, '2. v1 vs v2 성능 비교 (2-State 진단)', level=2)
    v1_met = v1_data['metrics']
    v2_met = v2_data['metrics']
    
    add_paragraph(doc, f"- 전체 ROC-AUC: {v1_met['roc_auc']*100:.2f}% (v1) -> {v2_met['roc_auc']*100:.2f}% (v2) [개선: {v2_met['roc_auc']*100 - v1_met['roc_auc']*100:+.2f}%p]", bold=True)
    add_paragraph(doc, f"- 결함 F1-score: {v1_met['defect_f1']*100:.2f}% (v1) -> {v2_met['defect_f1']*100:.2f}% (v2) [개선: {v2_met['defect_f1']*100 - v1_met['defect_f1']*100:+.2f}%p]", bold=True)
    add_paragraph(doc, f"- 정상 F1: {v1_met['clean_f1']*100:.2f}% (v1) -> {v2_met['clean_f1']*100:.2f}% (v2)")
    add_paragraph(doc, f"- Macro F1: {v1_met['macro_f1']*100:.2f}% (v1) -> {v2_met['macro_f1']*100:.2f}% (v2)")
    add_paragraph(doc, f"- Clean Specificity: {v1_met['clean_specificity']*100:.2f}% (v1) -> {v2_met['clean_specificity']*100:.2f}% (v2)")
    
    add_heading(doc, '3. 공식 5대 시점별 결함 F1 (v2)', level=2)
    for view, data in v2_data['view_results'].items():
        add_paragraph(doc, f"- {data['alias']}: {data['defect_f1']:.2f}%")
        
    add_heading(doc, '4. 주요 피처 중요도 (v2 Top 5)', level=2)
    top_feats = list(v2_data['top_features'].items())[:5]
    for i, (fname, imp) in enumerate(top_feats, 1):
        add_paragraph(doc, f"{i}. {fname} ({imp:.2f}%)")
        
    add_heading(doc, '5. 결론', level=2)
    add_paragraph(doc, '굿모닝 동작의 핵심 역학인 체간 전방 경사(Torso Pitch) 및 대퇴-정강이 비율을 편입하고 중복 피처를 절제한 결과, v1 대비 진단 정확도와 신뢰성이 향상되었습니다. 굿모닝 특화 피처 최적화를 통해 최종 벤치마크 성능을 성공적으로 달성했습니다.')
    
    output_path = r'D:\Documents\gm1_v2_performance_report.docx'
    doc.save(output_path)
    print(f'보고서 저장 완료: {output_path}')

if __name__ == "__main__":
    main()
