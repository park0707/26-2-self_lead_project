import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Malgun Gothic'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    run.font.name = 'Malgun Gothic'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    return p

def main():
    doc = Document()
    
    # Title
    title = doc.add_heading('vidio-health-trainer', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    add_heading(doc, '굿모닝 (gm1) 오류 모델 성능 보고서', level=1)
    
    add_heading(doc, '1. 개요', level=2)
    add_paragraph(doc, '- 대상 운동: 바벨 굿모닝 (09_굿모닝)')
    add_paragraph(doc, '- 대상 오류: 무릎 구부린 채 고정 위반 (gm1)')
    add_paragraph(doc, '- 적용 기법: 스티프 데드리프트 (sd1) 오류 모델의 정규 3D 변환, KADR, Excess Flexion 등의 생체역학 피처 및 CatBoost 기반 5-Fold 파이프라인 구조 계승')
    
    add_heading(doc, '2. 데이터셋 구성', level=2)
    add_paragraph(doc, '- 총 샘플 수: 678개')
    add_paragraph(doc, '- 클래스 분포: 정상(Clean) 339개, 결함(Defect) 339개')
    
    add_heading(doc, '3. 성능 평가 결과 (2-State 진단)', level=2)
    add_paragraph(doc, '- 전체 ROC-AUC: 84.79%', bold=True)
    add_paragraph(doc, '- 결함 F1-score: 76.29%', bold=True)
    add_paragraph(doc, '- 정상 F1: 76.22%')
    add_paragraph(doc, '- Macro F1: 76.25%')
    add_paragraph(doc, '- Clean Specificity: 76.11%')
    
    add_heading(doc, '4. 공식 5대 시점별 성능 (결함 F1)', level=2)
    add_paragraph(doc, '- 시점 A (우측 후방 45°): 75.95%')
    add_paragraph(doc, '- 시점 B (우측 전방 45°): 75.95%')
    add_paragraph(doc, '- 시점 C (정면 0°): 75.95%')
    add_paragraph(doc, '- 시점 D (좌측 전방 45°): 80.85%')
    add_paragraph(doc, '- 시점 E (좌측 후방 45°): 75.95%')
    
    add_heading(doc, '5. 주요 피처 중요도', level=2)
    add_paragraph(doc, '1. bottom_knee_angle (28.03%)')
    add_paragraph(doc, '2. shank_svii_mean (12.65%)')
    add_paragraph(doc, '3. excess_flex_integral (9.97%)')
    add_paragraph(doc, '4. knee_disp_y_rom (7.82%)')
    add_paragraph(doc, '5. knee_disp_z_canonical_rom (7.66%)')
    
    add_heading(doc, '6. 결론', level=2)
    add_paragraph(doc, '스티프 데드리프트의 sd1 (무릎 고정 위반) 구조를 굿모닝 gm1에 성공적으로 이식하여 초기 기준선(Baseline) 성능을 확인했습니다. 전체 ROC-AUC 약 84.8% 및 F1-Score 76.3%로 준수한 벤치마크를 기록했으며, 특히 무릎 절대 각도(bottom_knee_angle)와 정강이 수직 불변성 지수(shank_svii_mean)가 가장 중요한 분류 기준으로 작용함을 확인했습니다.')
    
    output_path = r'D:\Documents\gm1_performance_report.docx'
    doc.save(output_path)
    print(f'보고서 저장 완료: {output_path}')

if __name__ == "__main__":
    main()
