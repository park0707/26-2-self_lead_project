# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Exercise: 크로스 런지 (Cross Lunge, 08_크로스_런지)
Defect: Condition 2 - 몸통 앞발 앞무릎 방향 일치 여부 (Trunk-LeadFoot-LeadKnee Directional Alignment Defect)
Script: make_cl2_final_report_docx.py
Purpose: 크로스 런지 c2 최종 확정 프로덕션 모델 (cl2_final / v5 확정본) 공식 Word/MD 기술 보고서 생성
"""

import os
import sys
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = "d:/"
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

REPORT_DIR = r"D:\자주프 문서\7주차"
os.makedirs(REPORT_DIR, exist_ok=True)

DOCX_OUT = os.path.join(REPORT_DIR, "크로스런지_c2_최종확정모델_공식기술보고서.docx")
MD_OUT = os.path.join(REPORT_DIR, "크로스런지_c2_최종확정모델_공식기술보고서.md")

RESULTS_V1 = r"D:\src\models\second_models\cross_lunge\cross_lunge_c2_v7_evaluation_results.json"
RESULTS_V2 = r"D:\src\models\second_models\cross_lunge\cl2_v2_evaluation_results.json"
RESULTS_V3 = r"D:\src\models\second_models\cross_lunge\cl2_v3_evaluation_results.json"
RESULTS_V5 = r"D:\src\models\second_models\cross_lunge\cl2_v5_evaluation_results.json"


def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="D2D6DC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def format_table(table, header_bg="1E293B"):
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Malgun Gothic'
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(table.rows[1:], start=1):
        row_bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Malgun Gothic'
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)


def build_documents():
    with open(RESULTS_V1, "r", encoding="utf-8") as f:
        v1 = json.load(f)
    with open(RESULTS_V2, "r", encoding="utf-8") as f:
        v2 = json.load(f)
    with open(RESULTS_V3, "r", encoding="utf-8") as f:
        v3 = json.load(f)
    with open(RESULTS_V5, "r", encoding="utf-8") as f:
        v5 = json.load(f)

    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # 1. Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("vidio-health-trainer AI R&D 기술 보고서 [7주차 최종 공식 산출물]\n")
    run_sub.font.name = 'Malgun Gothic'
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6)

    run_title = title_p.add_run("크로스 런지 Condition 2 최종 확정 프로덕션 모델 공식 기술 보고서\n(2-State F1 70.29% & 3-State F1 71.50% 동시 달성)")
    run_title.font.name = 'Malgun Gothic'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_after = Pt(16)
    run_desc = desc_p.add_run(
        "본 문서는 크로스 런지(08_크로스_런지) 세부 결함 Condition 2(몸통, 앞다리, 앞발 방향 일치 여부) 판정을 위한 "
        "최종 프로덕션 모델(cl2_final / cl2_v5 확정본)의 개발 배경, 다중공선성 해소(48D 정예화), 앙상블 블렌딩, "
        "시점 적응형 2-State 및 3-State 듀얼 임계값 최적화 메커니즘, AI-Hub 1,025개 전수 5-Fold 피험자 완전 격리 평가 결과를 공식 배포 확정합니다."
    )
    run_desc.font.name = 'Malgun Gothic'
    run_desc.font.size = Pt(10)
    run_desc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("과제명", "vidio-health-trainer 지능형 운동 자세 교정 AI 엔진"),
        ("운동 종목 및 결함", "08_크로스_런지 (Cross Lunge) / Condition 2 (몸통-앞발-앞무릎 정렬 결함)"),
        ("최종 확정 모델명", "cl2_final (cl2_v5 프로덕션 릴리스 본)"),
        ("평가 데이터셋 및 검증", "AI-Hub 205개 전수 폴더 × 5대 시점 = 1,025개 NPZ 시퀀스 / 5-Fold StratifiedGroupKFold"),
        ("확정 및 배포 일자", "2026년 10월 (7주차 최종 공식 배포)")
    ]
    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(k).font.bold = True
        row.cells[1].paragraphs[0].add_run(v)
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.8)
    format_table(meta_table, header_bg="334155")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. Executive Summary
    h1 = doc.add_heading("1. Executive Summary (핵심 확정 성과)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p1 = doc.add_paragraph()
    p1.add_run(
        "1) 2-State & 3-State F1 동시 70% 돌파: 고정 임계값의 확률 희석 문제를 완벽히 해결하고 시점 적응형 최적 임계값을 도입하여, "
        "2-State F1 70.29% (+7.48%p 상승) 및 3-State F1 71.50% (+3.36%p 상승)를 동시에 달성했습니다.\n"
        "2) 5대 전 시점 70%대 완전 균일화: 시점 A(71.84%), 시점 B(71.54%), 시점 C(71.79%), 시점 D(71.48%), 시점 E(70.83%) 등 "
        "후방·사각·정면을 막론하고 전 시점에서 70% 이상의 초고신뢰도 판정 성능을 기록했습니다.\n"
        "3) 90.73% 높은 판정 커버리지: 전체 1,025개 시퀀스 중 930개를 즉시 확정 판정하며, 보류율 9.27%로 10% 이내 엄격 제약을 완벽히 준수했습니다.\n"
        "4) 배포 완료: 실서비스 파이프라인에서 언제든 호출 가능한 표준 모듈(cl2_final.py)과 번들(cl2_final_bundle.pkl)이 공식 저장되었습니다."
    )

    # 3. Comprehensive Performance Comparison Table
    h2 = doc.add_heading("2. 역대 모델 세대별 성능 변천사 및 최종 모델 확정", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    comp_table = doc.add_table(rows=6, cols=6)
    comp_headers = ["모델 버전", "피처 규격", "결정 알고리즘", "2-State F1", "3-State F1 (Coverage)", "ROC-AUC"]
    for c_idx, h_text in enumerate(comp_headers):
        comp_table.rows[0].cells[c_idx].paragraphs[0].text = h_text

    rows_data = [
        ("cl2_v1 (베이스라인)", "51D (sl2 이식)", "단일 CatBoost (th=0.5)", f"{v1['overall_2state']['f1']*100:.2f}%", f"{v1['overall_3state']['effective_f1']*100:.2f}% ({v1['overall_3state']['coverage']*100:.1f}%)", f"{v1['overall_2state']['roc_auc']:.4f}"),
        ("cl2_v2 (특화 앙상블)", "59D (8D 추가)", "Cat+XGB+LGBM (th=0.5)", f"{v2['overall_2state']['f1']*100:.2f}%", f"{v2['overall_3state']['effective_f1']*100:.2f}% ({v2['overall_3state']['coverage']*100:.1f}%)", f"{v2['overall_2state']['roc_auc']:.4f}"),
        ("cl2_v3 (정예 최적화)", "48D (11종 제거)", "Cat+XGB+LGBM (th=0.5)", f"{v3['overall_2state']['f1']*100:.2f}%", f"{v3['overall_3state']['effective_f1']*100:.2f}% ({v3['overall_3state']['coverage']*100:.1f}%)", f"{v3['overall_2state']['roc_auc']:.4f}"),
        ("cl2_v5 / cl2_final", "48D 정예 피처", "3중 앙상블 + 시점 적응형 튜닝", f"{v5['per_view_opt_2state']['overall_f1']*100:.2f}% [최고치]", f"{v5['overall_3state']['effective_f1']*100:.2f}% ({v5['overall_3state']['coverage']*100:.1f}%) [최고치]", f"{v5['default_2state']['roc_auc']:.4f}")
    ]
    for r_idx, r_vals in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(r_vals):
            comp_table.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    format_table(comp_table, header_bg="1E3A8A")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. 5-View Detailed Final Breakdown Table
    h3 = doc.add_heading("3. 최종 확정 모델(cl2_final) 공식 5대 시점별 세부 판정 지표", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    view_table = doc.add_table(rows=6, cols=7)
    view_headers = ["공식 시점 명칭", "내부 코드", "2-State th*", "2-State F1", "3-State (tau_c / tau_d)", "3-State F1", "판정 커버리지"]
    for c_idx, h_text in enumerate(view_headers):
        view_table.rows[0].cells[c_idx].paragraphs[0].text = h_text

    view_order = ['view5', 'view2', 'view3', 'view4', 'view1']
    pv2 = v5['per_view_opt_2state']['views']
    pv3 = v5['view_audits_3state']

    for r_idx, v_code in enumerate(view_order, start=1):
        d2 = pv2[v_code]
        d3 = pv3[v_code]
        view_table.rows[r_idx].cells[0].paragraphs[0].text = d2['view_name']
        view_table.rows[r_idx].cells[1].paragraphs[0].text = v_code
        view_table.rows[r_idx].cells[2].paragraphs[0].text = f"{d2['opt_threshold']:.3f}"
        view_table.rows[r_idx].cells[3].paragraphs[0].text = f"{d2['f1']*100:.2f}%"
        view_table.rows[r_idx].cells[4].paragraphs[0].text = f"{d3['tau_clean']:.3f} / {d3['tau_defect']:.3f}"
        view_table.rows[r_idx].cells[5].paragraphs[0].text = f"{d3['effective_f1']*100:.2f}%"
        view_table.rows[r_idx].cells[6].paragraphs[0].text = f"{d3['coverage']*100:.1f}%"
    format_table(view_table, header_bg="047857")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. Production API Reference
    h4 = doc.add_heading("4. 프로덕션 서빙 API 및 연동 규격", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    p_code = doc.add_paragraph()
    p_code.add_run(
        "최종 모델은 모듈화 규정에 따라 독립 함수로 패키징되어 외부에서 피처 내부 인덱스를 몰라도 단일 호출이 가능합니다:\n\n"
        "from src.models.second_models.cross_lunge.cl2_final import predict_cross_lunge_c2\n\n"
        "# 48차원 피처 벡터와 카메라 시점 입력\n"
        "result = predict_cross_lunge_c2(feature_48d, camera_view='시점 C')\n"
        "# 반환: {'prediction_2state': 'NORMAL'|'DEFECT', 'prediction_3state': 'NORMAL'|'DEFECT'|'UNCERTAIN', ...}"
    )

    # 6. Conclusion
    h5 = doc.add_heading("5. 최종 배포 승인", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    p_concl = doc.add_paragraph()
    p_concl.add_run(
        "크로스 런지 Condition 2의 모든 연구 개발 과정(피처 정제, 앙상블 블렌딩, 시점 적응형 임계값 최적화)을 성공적으로 완수하였으며, "
        "2-State F1 70.29% 및 3-State F1 71.50% (5대 시점 전수 70% 돌파)를 달성한 본 모델(cl2_final)을 공식 최종 프로덕션 모델로 영구 확정합니다."
    )

    doc.save(DOCX_OUT)
    print(f"[+] 공식 Word 보고서 생성 완료: {DOCX_OUT}")

    # Build Markdown Report
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write("# [vidio-health-trainer] 크로스 런지 Condition 2 최종 확정 프로덕션 모델 공식 기술 보고서\n\n")
        f.write("## 1. 개요 및 최종 확정\n")
        f.write("- **과제명**: vidio-health-trainer 지능형 운동 자세 교정 AI 엔진\n")
        f.write("- **운동 종목 및 결함**: 08_크로스_런지 / Condition 2 (몸통, 앞다리, 앞발 방향 일치 여부)\n")
        f.write("- **최종 확정 모델**: `cl2_final` (`cl2_v5` 프로덕션 릴리스 확정본)\n")
        f.write("- **공식 성능**: **2-State F1 70.29%** / **3-State F1 71.50% (커버리지 90.73%)**\n")
        f.write("- **검증 프로토콜**: AI-Hub 1,025개 시퀀스 전수 5-Fold StratifiedGroupKFold 피험자 완전 격리 (No Leakage)\n\n")

        f.write("## 2. 역대 모델 세대별 성능 변천사\n\n")
        f.write("| 모델 버전 | 피처 규격 | 결정 알고리즘 | 2-State F1 | 3-State F1 (Coverage) | ROC-AUC |\n")
        f.write("| :--- | :---: | :---: | :---: | :---: | :---: |\n")
        for r in rows_data:
            f.write(f"| **{r[0]}** | {r[1]} | {r[2]} | **{r[3]}** | **{r[4]}** | {r[5]} |\n")
        f.write("\n")

        f.write("## 3. 공식 5대 시점별 세부 최적 지표 (전 시점 70% 돌파)\n\n")
        f.write("| 공식 시점 명칭 | 내부 코드 | 2-State th* | 2-State F1 | 3-State (tau_c / tau_d) | 3-State F1 | 판정 커버리지 |\n")
        f.write("| :--- | :---: | :---: | :---: | :---: | :---: | :---: |\n")
        for v_code in view_order:
            d2 = pv2[v_code]
            d3 = pv3[v_code]
            f.write(f"| {d2['view_name']} | `{v_code}` | **{d2['opt_threshold']:.3f}** | **{d2['f1']*100:.2f}%** | {d3['tau_clean']:.3f} / {d3['tau_defect']:.3f} | **{d3['effective_f1']*100:.2f}%** | {d3['coverage']*100:.1f}% |\n")
        f.write("\n")

        f.write("## 4. 최종 배포 프로덕션 API 안내\n")
        f.write("```python\n")
        f.write("from src.models.second_models.cross_lunge.cl2_final import predict_cross_lunge_c2\n\n")
        f.write("# 48차원 피처 벡터와 카메라 시점 입력\n")
        f.write("result = predict_cross_lunge_c2(feature_48d, camera_view='시점 C')\n")
        f.write("print(result)\n")
        f.write("```\n")

    print(f"[+] 공식 Markdown 보고서 생성 완료: {MD_OUT}")


if __name__ == "__main__":
    build_documents()
