import os
import json
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

RESULTS_JSON = r"D:\src\models\second_models\good_morning\gm2_v1_evaluation_results.json"
REPORT_DOCX = r"D:\gm2_v1_performance_report.docx"

def create_report():
    if not os.path.exists(RESULTS_JSON):
        print("결과 JSON 파일이 없습니다.")
        return

    with open(RESULTS_JSON, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = docx.Document()
    
    # Title
    title = doc.add_heading('vidio-health-trainer 성능 보고서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. 개요', level=1)
    p = doc.add_paragraph()
    p.add_run('• 모델명: ').bold = True
    p.add_run(f"{data.get('model_name', 'gm2_v1')}\n")
    p.add_run('• 운동 종목: ').bold = True
    p.add_run(f"{data.get('exercise', '09_굿모닝')}\n")
    p.add_run('• 대상 결함: ').bold = True
    p.add_run(f"{data.get('defect', 'c2 (시선 방향 유지 오류)')}\n")
    p.add_run('• 특징: ').bold = True
    p.add_run('얼굴 피쳐가 가려지는 후면/측후면 시점에서도 동작하도록, 어깨(Shoulder), 귀(Ear), 골반(Hip) 기반의 간접 측정 생체역학 피처를 구축하여 학습했습니다.')

    doc.add_heading('2. 시점별 적응형 3-State 캘리브레이션 결과', level=1)
    tsm = data.get('three_state_metrics', {})
    
    p = doc.add_paragraph()
    p.add_run(f"• 유효 의사결정 F1: {tsm.get('effective_f1', 0):.4f}\n").bold = True
    p.add_run(f"• 커버리지 (결정 비율): {tsm.get('coverage', 0)*100:.2f}%\n")
    p.add_run(f"• 불확실성 보류율: {tsm.get('uncertainty_rate', 0)*100:.2f}%\n")
    p.add_run(f"• 정확도 (Accuracy): {tsm.get('accuracy', 0):.4f}\n")
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '시점 (View)'
    hdr_cells[1].text = 'F1 Score'
    hdr_cells[2].text = '정확도 (Acc)'
    hdr_cells[3].text = '보류율 (Retake %)'
    
    for cell in hdr_cells:
        for p_table in cell.paragraphs:
            for r in p_table.runs:
                r.bold = True
    
    view_audits = tsm.get('view_audits', {})
    for v_code, audit in view_audits.items():
        row_cells = table.add_row().cells
        row_cells[0].text = audit.get('alias', v_code)
        row_cells[1].text = f"{audit.get('effective_f1', 0):.4f}"
        row_cells[2].text = f"{audit.get('accuracy', 0):.4f}"
        row_cells[3].text = f"{audit.get('uncertainty_rate', 0)*100:.2f}%"

    doc.add_heading('3. 핵심 간접 측정 피처 중요도 (Top 10)', level=1)
    p = doc.add_paragraph("시선 방향 유지를 판단하기 위해 새롭게 설계된 간접 측정 피처들의 중요도 기여도입니다.")
    
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = 'Table Grid'
    hdr2 = table2.rows[0].cells
    hdr2[0].text = '순위'
    hdr2[1].text = '피처명'
    hdr2[2].text = '중요도 (%)'
    
    for cell in hdr2:
        for p_table in cell.paragraphs:
            for r in p_table.runs:
                r.bold = True
                
    top_features = data.get('top_features', [])
    for feat in top_features:
        row_cells = table2.add_row().cells
        row_cells[0].text = str(feat.get('rank', ''))
        row_cells[1].text = feat.get('name', '')
        row_cells[2].text = f"{feat.get('importance', 0):.2f}%"

    doc.save(REPORT_DOCX)
    print(f"[+] 보고서 생성 완료: {REPORT_DOCX}")

if __name__ == '__main__':
    create_report()
