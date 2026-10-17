# -*- coding: utf-8 -*-
r"""
vidio-health-trainer
Script: src/reports/generate_week8_4_defects_pass_fail_report.py (v2, 정정판)
Purpose: 바벨 스쿼트 4대 세부 오류 모델 지표별 합격/불합격 감사 정정 보고서
  - 입력: D:\data\일반화\squat_4_defects_audit_v2.json (audit_squat_4defects_v2.py 산출)
  - 산출: D:\자주프 문서\8주차\8주차_바벨스쿼트_4대_세부오류모델_전수_성능측정_및_통과심사보고서.docx
          D:\자주프 문서\8주차\md\...동일명.md
"""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

AUDIT = r"D:\data\일반화\squat_4_defects_audit_v2.json"
DOCX_OUT = r"D:\자주프 문서\8주차\8주차_바벨스쿼트_4대_세부오류모델_전수_성능측정_및_통과심사보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_바벨스쿼트_4대_세부오류모델_전수_성능측정_및_통과심사보고서.md"
LABEL = {"balanced_acc": "균형 정확도", "recall": "재현율", "specificity": "특이도",
         "precision": "정밀도", "roc_auc": "ROC-AUC", "coverage": "판정 도달률"}
TAGS = {"in_sample_production": "(A) 현재 프로덕션 임계치 (in-sample)",
        "leave_one_view_out": "(B) 시점 제외 검증 (Leave-One-View-Out)"}


def shade(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hexcolor}"/>'))


def put(cell, text, bold=False, size=9, bg=None, color=None):
    cell.text = text
    r = cell.paragraphs[0].runs[0]
    r.font.name = "맑은 고딕"; r.font.size = Pt(size); r.font.bold = bold
    if color: r.font.color.rgb = color
    if bg: shade(cell, bg)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        put(t.cell(0, i), h, True, 9, "1A365D", RGBColor(255, 255, 255))
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row):
            bad = isinstance(v, str) and ("불합격" in v or "미충족" in v)
            put(t.cell(ri, ci), str(v), ci == 0, 9, "FDE8E8" if bad else ("F7FAFC" if ri % 2 else None))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    return t


def para(doc, text, size=10, bold=False):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.name = "맑은 고딕"; r.font.size = Pt(size); r.font.bold = bold
    return p


def head(doc, text, lv):
    h = doc.add_heading(text, level=lv)
    for r in h.runs:
        r.font.name = "맑은 고딕"; r.font.color.rgb = RGBColor(26, 54, 93)


def main():
    a = json.load(open(AUDIT, encoding="utf-8"))
    M = a["models"]; C = a["criteria"]
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)
    md = []
    doc = Document()

    head(doc, "[vidio-health-trainer] 8주차 공식 기술 보고서 (정정판)", 0)
    para(doc, "바벨 스쿼트 4대 세부 오류 모델 지표별 합격/불합격 감사", 12, True)
    md += ["# [vidio-health-trainer] 8주차 공식 기술 보고서 (정정판)",
           "## 바벨 스쿼트 4대 세부 오류 모델 지표별 합격/불합격 감사", ""]

    head(doc, "0. 정정 사항", 1)
    corr = [
        "이전 보고서의 '시스템 종합 우수 통과(87.0점)' 판정을 철회한다.",
        "이전 채점은 '최소 합격선 통과 시 배점 80% 부여'인데 'High Pass' 기준선도 80점이어서, 모든 지표가 최소선만 넘어도 '우수 통과'가 되는 구조 결함이 있었다. 합산 점수 체계를 폐기하고 지표별 합격/불합격으로 대체한다.",
        "합격 기준이 결과 확인 이후에 정해졌다. 본 정정판은 기준을 사전 고정하고 코드(audit_squat_4defects_v2.py)에 명시한다.",
        "시점별 임계치가 동일 75클립으로 튜닝되어 in-sample 평가였다. 시점 제외 검증(LOVO)을 병행한다.",
        "모델 명칭 오기재 정정: 메타데이터 기준 c1=척추 중립, c2=고개 정면/시선, c3=무릎-발 방향 일치, c4=발바닥 지면 고정. (이전 보고서의 '무릎 모임/무릎 돌출/가동범위 부족' 표기는 오류)",
    ]
    for c in corr: para(doc, "• " + c, 9.5)
    md += ["### 0. 정정 사항"] + ["- " + c for c in corr] + [""]

    head(doc, "1. 합격 기준 (사전 고정)", 1)
    rows = [[LABEL[k], f">= {v}%"] for k, v in C.items()]
    table(doc, ["지표", "합격선"], rows, [2.5, 1.5])
    para(doc, "판정: 6개 전부 충족=통과 / 4~5개=조건부 / 3개 이하=불합격. 최종 판정은 (A), (B) 중 낮은 쪽을 따른다.", 9.5)
    md += ["### 1. 합격 기준 (사전 고정)", "", "| 지표 | 합격선 |", "|---|---|"] + [f"| {r[0]} | {r[1]} |" for r in rows]
    md += ["", "판정: 6개 전부 충족=통과 / 4~5개=조건부 / 3개 이하=불합격. 최종 판정은 (A), (B) 중 낮은 쪽.", ""]

    head(doc, "2. 최종 판정 총괄", 1)
    rows = []
    for c, r in M.items():
        A = r["in_sample_production"]; B = r["leave_one_view_out"]
        rows.append([f"{c.upper()} {r['name']}", f"{A['verdict']} ({A['passed']}/6)", f"{B['verdict']} ({B['passed']}/6)", r["final_verdict"]])
    table(doc, ["모델", "(A) in-sample", "(B) LOVO", "최종 판정"], rows, [2.3, 1.5, 1.5, 1.2])
    md += ["### 2. 최종 판정 총괄", "", "| 모델 | (A) in-sample | (B) LOVO | 최종 판정 |", "|---|---|---|---|"]
    md += [f"| {r[0]} | {r[1]} | {r[2]} | **{r[3]}** |" for r in rows]
    md += [""]

    head(doc, "3. 모델별 상세", 1)
    md += ["### 3. 모델별 상세", ""]
    for c, r in M.items():
        head(doc, f"{c.upper()} {r['name']} — 최종 {r['final_verdict']}", 2)
        md += [f"#### {c.upper()} {r['name']} — 최종 {r['final_verdict']}", ""]
        for tag, title in TAGS.items():
            t = r[tag]; m = t["metrics"]
            para(doc, f"{title}: {t['verdict']} ({t['passed']}/6) | TP {m['tp']} FP {m['fp']} FN {m['fn']} TN {m['tn']} | F1 {m['f1']}% | 균형정확도 95% CI {t['ci95_balanced_acc']}", 9.5, True)
            rows = [[LABEL[k], f"{v['value']}%", f">= {v['min']}%", "충족" if v["pass"] else "미충족"] for k, v in t["checks"].items()]
            table(doc, ["지표", "실측", "합격선", "결과"], rows, [1.8, 1.2, 1.2, 1.0])
            md += [f"**{title}**: {t['verdict']} ({t['passed']}/6) — TP {m['tp']} FP {m['fp']} FN {m['fn']} TN {m['tn']}, F1 {m['f1']}%, 균형정확도 95% CI {t['ci95_balanced_acc']}", "",
                   "| 지표 | 실측 | 합격선 | 결과 |", "|---|---|---|---|"] + [f"| {x[0]} | {x[1]} | {x[2]} | {x[3]} |" for x in rows] + [""]

    head(doc, "4. 결론", 1)
    concl = [
        "c4(발바닥 지면 고정)만 (A),(B) 모두 6/6 충족하여 통과한다.",
        "c3(무릎-발 방향)은 재현율은 높으나 오탐이 많아(정밀도 41%/27%) 불합격이다. in-sample에서는 조건부였다.",
        "c1(척추 중립), c2(고개/시선)는 불합격이다. 재현율·ROC-AUC 미달.",
        "따라서 4개 모델 전체 통과가 아니며 '시스템 종합 우수 통과'는 성립하지 않는다. 표본(75클립, 양성 15)이 작아 신뢰구간이 넓으므로, 임계치 튜닝에 쓰지 않은 별도 홀드아웃 영상 평가가 필요하다.",
        "한계: LOVO는 단순화한 근사(단일 임계치, 간격 0.06 고정)이며 프로덕션 튜닝 절차와 동일하지 않다.",
    ]
    for c in concl: para(doc, "• " + c, 10)
    md += ["### 4. 결론"] + ["- " + c for c in concl]

    doc.save(DOCX_OUT)
    open(MD_OUT, "w", encoding="utf-8").write("\n".join(md))
    print("OK", DOCX_OUT)


if __name__ == "__main__":
    main()
