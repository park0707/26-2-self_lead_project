# -*- coding: utf-8 -*-
r"""
vidio-health-trainer
Script: src/reports/generate_week8_5_dimensions_generalization_report.py (v2, 정정판)
Purpose: 바벨 스쿼트 5대 일반화 차원 실측 보고서 (정정판) - 실측치만 기재하고 임의 합격 문구를 쓰지 않음.
  - 입력: D:\data\일반화\squat_75reps_5_dimensions_generalization_metrics.json
  - 산출: D:\자주프 문서\8주차\8주차_바벨스쿼트_5대_핵심_일반화_차원_실측평가보고서.docx
          D:\자주프 문서\8주차\md\...동일명.md
"""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

SRC = r"D:\data\일반화\squat_75reps_5_dimensions_generalization_metrics.json"
DOCX_OUT = r"D:\자주프 문서\8주차\8주차_바벨스쿼트_5대_핵심_일반화_차원_실측평가보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_바벨스쿼트_5대_핵심_일반화_차원_실측평가보고서.md"
NAMES = {"c1": "척추 중립", "c2": "고개 정면/시선", "c3": "무릎-발 방향 일치", "c4": "발바닥 지면 고정"}


def shade(cell, hx):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hx}"/>'))


def put(cell, text, bold=False, bg=None, color=None):
    cell.text = str(text); r = cell.paragraphs[0].runs[0]
    r.font.name = "맑은 고딕"; r.font.size = Pt(9); r.font.bold = bold
    if color: r.font.color.rgb = color
    if bg: shade(cell, bg)


def table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header): put(t.cell(0, i), h, True, "1A365D", RGBColor(255, 255, 255))
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row): put(t.cell(ri, ci), v, ci == 0, "F7FAFC" if ri % 2 else None)


def para(doc, text, size=10, bold=False):
    r = doc.add_paragraph().add_run(text); r.font.name = "맑은 고딕"; r.font.size = Pt(size); r.font.bold = bold


def head(doc, text, lv):
    h = doc.add_heading(text, level=lv)
    for r in h.runs: r.font.name = "맑은 고딕"; r.font.color.rgb = RGBColor(26, 54, 93)


def main():
    d = json.load(open(SRC, encoding="utf-8"))
    d1 = d["dimension_1_single_defect_quality"]; d2 = d["dimension_2_multi_defect_service_quality"]
    d3 = d["dimension_3_camera_angle_robustness"]; d4 = d["dimension_4_false_alarm_suppression"]
    d5 = d["dimension_5_reliability_and_calibration"]
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)
    doc = Document(); md = []

    def sec(title, header=None, rows=None, notes=()):
        head(doc, title, 1); md.append(f"### {title}\n")
        if header:
            table(doc, header, rows)
            md.extend(["| " + " | ".join(header) + " |", "|" + "---|" * len(header)] + ["| " + " | ".join(map(str, r)) + " |" for r in rows] + [""])
        for n in notes: para(doc, "• " + n, 9.5); md.append("- " + n)
        md.append("")

    head(doc, "[vidio-health-trainer] 8주차 공식 기술 보고서 (정정판)", 0)
    para(doc, "바벨 스쿼트 5대 일반화 차원 실측 평가 (75클립, SpinePose 3D 결합)", 12, True)
    md += ["# [vidio-health-trainer] 8주차 공식 기술 보고서 (정정판)", "## 바벨 스쿼트 5대 일반화 차원 실측 평가 (75클립, SpinePose 3D 결합)", ""]

    sec("0. 정정 사항 및 평가 전제", notes=[
        "이전 판의 '상용 서비스 기준 완벽 충족', 'ECE 우수', '시점 안정성 양호', '11.33% 라벨 오류율' 등 문구는 실측과 달라 철회한다. (ECE=24.44%, 시점 간 F1 표준편차=37.15%p, Hamming Loss=0.1833)",
        "모델 명칭 정정: c1=척추 중립, c2=고개 정면/시선, c3=무릎-발 방향 일치, c4=발바닥 지면 고정.",
        "전제: 시점별 임계치는 동일 75클립으로 튜닝되었으므로 아래 수치는 in-sample이며 낙관적이다. 시점 제외 검증 결과는 별도 4대 모델 정정 보고서 참조.",
        "표본: 75클립(양성 결함당 15). 신뢰구간이 넓어 소수점 단위 비교는 의미가 작다.",
        "별도 합격선은 사전 합의되지 않았으므로 본 보고서는 합격/불합격을 선언하지 않고 실측치만 제시한다."])

    sec("1. 차원 1: 단일 결함 품질",
        ["모델", "Balanced Acc", "F1", "ROC-AUC", "PR-AUC"],
        [[f"{c.upper()} {NAMES[c]}", f"{v['balanced_accuracy']}%", f"{v['f1_score']}%", f"{v['roc_auc']}%", f"{v['pr_auc']}%"] for c, v in d1["by_defect"].items()]
        + [["Macro 평균", f"{d1['macro_balanced_accuracy']}%", f"{d1['macro_f1']}%", f"{d1['macro_roc_auc']}%", f"{d1['macro_pr_auc']}%"]])

    sec("2. 차원 2: 종합 서비스 품질",
        ["지표", "값"],
        [["EMR (4결함 완전 일치)", f"{d2['emr_pct']}% ({d2['exact_matches']}/{d2['total_reps']})"],
         ["Hamming Loss", f"{d2['hamming_loss']:.4f} (라벨 오류율 {d2['hamming_loss']*100:.2f}%)"],
         ["Subset 오류율", f"{d2['subset_error_rate']}%"]],
        notes=["EMR은 보류(retake)를 '결함 아님'으로 처리한 값이다. 이전 EMR 18.67%는 임계치 재튜닝 전 값이며, 49.33%는 in-sample 튜닝 후 값이다."])

    rows = []
    for v, i in d3["views"].items():
        f = i["f1_scores"]
        rows.append([f"{v} ({i['view_code']})", f"{i['view_emr_pct']}% ({i['exact_matches']}/{i['reps_count']})", f"{f['c1']}%", f"{f['c2']}%", f"{f['c3']}%", f"{f['c4']}%"])
    sec("3. 차원 3: 카메라 시점 강건성",
        ["시점", "EMR", "c1 F1", "c2 F1", "c3 F1", "c4 F1"], rows,
        notes=[f"시점 간 F1 표준편차 평균 {d3['mean_view_f1_std']}%p, EMR 표준편차 {d3['view_emr_std']}%p → 시점 편차가 매우 크다(안정적이지 않음).",
               "시점 C(정면)는 15/15 완전 일치이나, 시점별 임계치를 해당 시점 15클립으로 튜닝한 영향이 커 과대평가 가능성이 높다.",
               "시점 D, E(좌측)는 성능이 낮다."])

    sec("4. 차원 4: 허위 경보 억제력",
        ["모델", "FP", "FPR", "특이도", "정밀도"],
        [[f"{c.upper()} {NAMES[c]}", v["fp"], f"{v['fpr_pct']}%", f"{v['specificity_pct']}%", f"{v['precision_pct']}%"] for c, v in d4["by_defect"].items()]
        + [["Macro", "-", f"{d4['macro_fpr_pct']}%", f"{d4['macro_specificity_pct']}%", f"{d4['macro_precision_pct']}%"]],
        notes=["c2, c4는 FP 0건이나 c1(FP 14), c3(FP 17)은 오경보가 많다."])

    sec("5. 차원 5: 3-State 신뢰도 및 보정",
        ["모델", "보류 건수", "Coverage", "Selective Precision", "ECE"],
        [[f"{c.upper()} {NAMES[c]}", v["holds_count"], f"{v['coverage_pct']}%", f"{v['selective_precision']}%", f"{v['ece_pct']}%"] for c, v in d5["by_defect"].items()]
        + [["Macro", "-", f"{d5['macro_coverage_pct']}%", "-", f"{d5['macro_ece_pct']}%"]],
        notes=["ECE 평균 24.4%로 확률 보정이 좋지 않다(일반적으로 10% 이하를 양호로 본다). 확률 캘리브레이션(Platt/Isotonic 등) 개선이 필요하다.",
               "Selective Precision은 확정 판정(보류 제외) 중 결함 판정의 정밀도이며, 현재 구조에서는 차원 4의 정밀도와 같은 값이다."])

    sec("6. 결론", notes=[
        "c4는 전 차원에서 가장 안정적이다. c1, c3는 오경보가 많고 c1, c2는 재현율/AUC가 낮다.",
        "시점 편차(σ 37%p)와 ECE(24%)가 크므로 '일반화 성능이 확인되었다'고 결론 내릴 수 없다.",
        "다음 단계: 튜닝에 쓰지 않은 별도 촬영 홀드아웃 평가, 확률 보정, 좌측 시점 보강."])

    doc.save(DOCX_OUT); open(MD_OUT, "w", encoding="utf-8").write("\n".join(md)); print("OK")


if __name__ == "__main__":
    main()
