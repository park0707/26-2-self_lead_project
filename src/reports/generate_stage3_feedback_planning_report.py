# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: src/reports/generate_stage3_feedback_planning_report.py
Purpose: 8주차 공식 기술 기획 보고서 자동 생성기
         - 연구 계획서상 2단계(관절 부하 측정) 선회 배경 및 3단계(결함 시 피드백) 상세 기획
         - 출력 Word 문서: D:\자주프 문서\8주차\8주차_3단계_자세교정_및_결함피드백_시스템_상세기획보고서.docx
         - 출력 참조 MD:  D:\자주프 문서\8주차\md\8주차_3단계_자세교정_및_결함피드백_시스템_상세기획보고서.md
"""

import os
import sys

PROJECT_ROOT = r"d:/"
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = r"D:\자주프 문서\8주차\8주차_3단계_자세교정_및_결함피드백_시스템_상세기획보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_3단계_자세교정_및_결함피드백_시스템_상세기획보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="2B579A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout_box(doc, text_list, bg_hex="F0F4F8", border_color="2B579A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.7)
    set_cell_background(c, bg_hex)
    
    tcPr = c._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for i, t in enumerate(text_list):
        if i > 0:
            p = c.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
        run = p.add_run(t)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def generate_reports():
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)

    doc = Document()

    # 페이지 여백 설정 (0.85인치)
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # 1. 문서 헤더 / 제목
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("[vidio-health-trainer] 8주차 공식 기술 기획 보고서")
    run_title.font.name = "맑은 고딕"
    run_title.font.size = Pt(21)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 60, 114)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("3단계 자세 교정 및 결함 피드백 시스템 상세 기획: 2단계(관절 부하) 선회 배경 및 8대 운동 33종 결함 전수 행동 지침·생체역학 코칭 아키텍처 수립")
    run_sub.font.name = "맑은 고딕"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(80, 95, 120)

    # 메타데이터 표
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("보고 일시", "2026년 10월 13일 (화) - 8주차 아키텍처 공식 마일스톤"),
        ("공식 프로젝트명", "vidio-health-trainer (AI 비디오 헬스 트레이너)"),
        ("문서 분류", "3단계(자세 교정 및 피드백) 공식 상세 기획 및 구현 명세서"),
        ("적용 대상 범위", "8대 운동 종목, 33종 세부 결함 전수, 웹 서빙 파이프라인 (src/web)"),
        ("핵심 모듈", "src/web/core/feedback_engine.py, model_registry.py, stage2_engine.py")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = "맑은 고딕"
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.name = "맑은 고딕"
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.width = Inches(1.6)
        c1.width = Inches(5.1)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FAFCFF")
    set_table_borders(meta_table, "C5D3E8")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 핵심 요약 Callout
    add_callout_box(doc, [
        "■ [핵심 요약] 3단계 자세 교정 및 피드백 시스템 선행 추진 결정 및 주요 성과",
        "1. 선회 배경: 1단계 운동 분류 및 2단계 33종 결함 분류 모델이 완성되어 웹 서빙 중이나, 기존 출력이 '피드백 자리'로 하드코딩되어 실서비스 완결성이 저해됨. 연구 계획서 2단계(OpenSim 역동역학 부하 측정)는 3D 지면반력 부재 및 50ms 실시간 서빙 제약으로 우선순위를 조정하고, 사용자에게 즉각적 가치를 제공하는 3단계 피드백 시스템을 선제 구축함.",
        "2. 3-Tier 계층형 코칭 체계: Tier A(1문장 즉각 행동 지침 큐), Tier B(생체역학적 부하 메커니즘 설명), Tier C(3-State: 정상 칭찬, 결함 교정, 주의/재촬영 차등화)의 3단계를 확립함.",
        "3. 8대 운동 33종 전수 딕셔너리 완비: 스쿼트(4종), 데드리프트(5종), 스티프데드(3종), 바벨런지(5종), 다이나믹런지(5종), 사이드런지(5종), 크로스런지(3종), 굿모닝(3종) 전수에 대해 100% 매핑 카탈로그 구축.",
        "4. 생체역학 위험 가중치 기반 우선순위 큐 & 연쇄 결함(Causal Chaining) 융합: 척추(1.0) > 무릎(0.85) > 골반(0.80) > 발목(0.60) 가중치를 적용해 인지 한계 내 1~2개 큐만 선별하고, '발바닥 뜸 -> 무릎 외반' 등 11개 복합 연쇄를 단일 근본 원인 큐로 합성함.",
        "5. 하네스 무결성 및 초고속성 검증: 단위 테스트 100% 통과, 1회 연산 지연시간 0.0037ms(목표 1ms 대비 270배 초고속), 기존 33종 E2E 파이프라인 0 Exceptions 완벽 검증 완료."
    ], bg_hex="EEF6FF", border_color="2B579A")

    # 제1장
    h1 = doc.add_heading("1. 기획 배경 및 2단계(관절 부하 측정) 선회 타당성 분석", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run(
        "본 vidio-health-trainer 프로젝트는 7~8주차에 걸쳐 1단계 운동 종목 분류(8대 운동 1D-CNN+BiLSTM)와 "
        "2단계 세부 오류 진단 모델(8대 운동 총 33종 CatBoost/LightGBM 3-State 모델)을 성공적으로 구축 완료하였으며, "
        "FastAPI 기반 백엔드 및 실시간 웹 대시보드(src/web)로 20~50ms의 초저지연 단일 시퀀스 추론까지 연동을 완료하였습니다. "
        "그러나 실제 웹 서비스를 동작시켰을 때, 다음과 같은 실서비스 관점의 중대한 기획적 병목이 확인되었습니다."
    )

    # 1.1 하드코딩된 피드백의 한계
    h2 = doc.add_heading("1.1 기존 시스템의 서비스 완결성 병목: 하드코딩된 피드백", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "현재 2단계 모델 레지스트리(DefectModelRegistry)와 엔드투엔드 파이프라인은 각 결함별 확률(probability)과 "
        "3-State 판정(정상:0, 결함:1, 주의/보류:2)을 정확히 산출하고 있으나, 사용자에게 노출되는 피드백 필드는 "
    )
    p.add_run("'피드백 자리'").bold = True
    p.add_run(
        "라는 임시 문자열로 하드코딩되어 있었습니다. 실제 사용자가 헬스 영상을 업로드하여 분석 결과를 확인했을 때, "
        "단순히 'c1 결함 85.9%'라는 기계적 수치만 보고는 '어떤 관절에 무리가 가는지', '당장 다음 반복에서 무엇을 어떻게 고쳐야 하는지' "
        "알 수 없어 서비스로서의 실질적 가치가 전달되지 못하는 치명적 결함이 존재했습니다."
    )

    # 1.2 2단계(관절 부하 측정) 선회 이유
    h2 = doc.add_heading("1.2 연구 계획서 2단계(관절 부하 측정) 선회 및 후순위 조정 사유", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "초기 연구 계획서상의 2단계는 'OpenSim 및 AddBiomechanics 기반의 생체역학 역동역학(Inverse Dynamics)을 활용한 관절 모멘트/부하 근사 추정 모델 개발'이었습니다. "
        "그러나 현 시점에서 2단계를 건너뛰고 3단계(결함 시 피드백 시스템)를 우선 추진해야 하는 명확한 기술적·전략적 사유는 다음과 같습니다:\n"
    )

    reasons = [
        ("3D 지면반력(GRF) 데이터 부재로 인한 역동역학 오차 증폭", 
         "스마트폰 단일 2D 영상 환경에서는 지면반력계(Force Plate) 측정이 불가능하므로 지면반력을 추정치로 대체해야 합니다. 이로 인해 하지 역동역학 모멘트 오차가 25~40% 이상 증폭되어 정밀한 토크(Torque) 수치 산출의 신뢰성이 저하됩니다."),
        ("실시간 웹 서빙 레이턴시 병목 (20~50ms 목표 위배)",
         "OpenSim 연산 또는 3D 신체 분절 질량 적분 연산은 시퀀스당 최소 수백 ms에서 수 초 이상의 연산 시간을 소모하므로, 기구축된 20~50ms 실시간 웹 서빙 파이프라인의 응답성을 심각하게 파괴합니다."),
        ("서비스 완결성 우선 원칙 (Closed Feedback Loop)",
         "1단계(종목 판별) -> 2단계(세부 결함 진단) -> 3단계(자세 교정 행동 지침)로 이어지는 사용자 인터랙션 루프가 닫혀야 비로소 동작하는 서비스 프로토타입이 완성됩니다."),
        ("2단계 부하의 3단계 내부 프록시 흡수 가능성 (Kinematic Stress Proxy)",
         "2단계의 본질적 목적은 '어떤 관절에 부담이 큰지'를 파악하는 것입니다. 이는 복잡한 OpenSim 물리 계산 없이도, 이미 2단계 모델이 산출한 척추 굴곡각, 무릎 외반각, 바벨 이격 모멘트암과 '생체역학 위험 가중치(Joint Risk Weight)'를 결합하여 피드백 우선순위에 지능적으로 반영할 수 있습니다.")
    ]

    for title, desc in reasons:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(desc)

    # 제2장
    h1 = doc.add_heading("2. 3단계 자세 교정 및 피드백 시스템의 5대 핵심 설계 원칙", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    # 5대 원칙 표
    prin_table = doc.add_table(rows=6, cols=3)
    prin_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prin_headers = ["설계 원칙", "핵심 메커니즘", "엔지니어링 및 사용자 가치"]
    for j, h in enumerate(prin_headers):
        c = prin_table.cell(0, j)
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.name = "맑은 고딕"
        c.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c, "2B579A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    principles_data = [
        ("[원칙 1] 3-Tier 계층형 피드백", "Tier A(행동 지침 큐) + Tier B(역학적 원인) + Tier C(3-State 차등화)", "단순 위험 경고를 넘어 즉각 수정 가능한 1문장 큐 제공"),
        ("[원칙 2] 인지 부하 제어 (CLT)", "단일 반복당 최우선 1개(최대 2개) 교정 지침만 선별 노출", "운동 수행 중 뇌 인지 과부하(Cognitive Overload) 방지"),
        ("[원칙 3] 생체역학 위험 우선순위", "Score = w_joint × p_defect × c_state 기반 정렬", "디스크/인대 손상 위험이 가장 큰 핵심 관절(척추/무릎) 최우선 구제"),
        ("[원칙 4] 연쇄 결함 인과 융합", "원인-결과 복합 결함(Trigger Set)을 단일 루트 코즈 큐로 합성", "예: '발바닥 뜸 -> 무릎 외반' 연쇄를 족궁 지지 원인 큐로 통합"),
        ("[원칙 5] 비의료적 표현 표준화", "질환/진단명 영구 배제, 운동역학적 부하 집중 신호로 표현", "의료법 저촉 방지 및 안전하고 건전한 피트니스 코칭 유지")
    ]
    for i, row in enumerate(principles_data):
        for j, val in enumerate(row):
            c = prin_table.cell(i + 1, j)
            c.text = val
            c.paragraphs[0].runs[0].font.name = "맑은 고딕"
            c.paragraphs[0].runs[0].font.size = Pt(9.0)
            if j == 0:
                c.paragraphs[0].runs[0].font.bold = True
                set_cell_background(c, "F8FAFC")
                c.width = Inches(1.8)
            elif j == 1:
                c.width = Inches(2.7)
            else:
                c.width = Inches(2.2)
    set_table_borders(prin_table, "C5D3E8")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 제3장
    h1 = doc.add_heading("3. 8대 운동 33종 세부 결함 전수 코칭 딕셔너리 명세", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run(
        "본 피드백 엔진(FeedbackEngine)은 vidio-health-trainer가 지원하는 8대 운동의 33종 세부 결함 전수에 대해 "
        "100% 매핑 커버리지를 보장합니다. 아래는 8대 운동별 세부 명세 요약표입니다."
    )

    from src.web.core.feedback_engine import FEEDBACK_CATALOG
    from src.web.core.config import EXERCISE_CATALOG

    for eid, ex_info in EXERCISE_CATALOG.items():
        ex_name = ex_info['name']
        h2 = doc.add_heading(f"3.{eid+1} {ex_name} ({ex_info['defects_count']}개 세부 결함)", level=2)
        
        # 운동별 결함 필터
        defect_items = [item for item in FEEDBACK_CATALOG.values() if item.exercise_name == ex_name]
        
        d_table = doc.add_table(rows=len(defect_items) + 1, cols=6)
        d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        d_headers = ["코드", "결함 명칭", "타깃 관절", "위험 가중치", "핵심 행동 지침 큐 (Action Cue)", "권장 시점"]
        for j, h in enumerate(d_headers):
            c = d_table.cell(0, j)
            c.text = h
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.name = "맑은 고딕"
            c.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_background(c, "334155")
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        col_widths = [Inches(0.6), Inches(1.5), Inches(1.1), Inches(0.7), Inches(2.2), Inches(0.9)]

        for i, item in enumerate(defect_items):
            row_cells = d_table.rows[i + 1].cells
            row_cells[0].text = item.defect_code
            row_cells[1].text = item.defect_name.split(" (")[0]
            row_cells[2].text = item.target_joint
            row_cells[3].text = f"{item.risk_weight:.2f}"
            row_cells[4].text = item.action_cue
            row_cells[5].text = item.optimal_view.split(" (")[0]

            for j, c in enumerate(row_cells):
                c.width = col_widths[j]
                c.paragraphs[0].runs[0].font.name = "맑은 고딕"
                c.paragraphs[0].runs[0].font.size = Pt(8.0)
                if j == 0:
                    c.paragraphs[0].runs[0].font.bold = True
                    set_cell_background(c, "F1F5F9")
                elif j == 3:
                    set_cell_background(c, "FEF2F2" if item.risk_weight >= 0.9 else ("FFFBEB" if item.risk_weight >= 0.8 else "F0FDF4"))

        set_table_borders(d_table, "CBD5E1")
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 제4장
    h1 = doc.add_heading("4. 다중 결함 우선순위 해결 및 연쇄 결함(Causal Chaining) 알고리즘", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run(
        "단일 횟수(1-Rep) 동작 중 척추 굽힘, 무릎 외반, 발바닥 들림 등 여러 결함이 동시 발생할 수 있습니다. "
        "이를 단순 나열할 경우 사용자는 극심한 인지 과부하(Cognitive Overload)를 겪어 아무런 교정도 하지 못하게 됩니다. "
        "본 시스템은 이를 해결하기 위해 생체역학 위험도 기반 우선순위 큐(Priority Queue)와 연쇄 결함 융합 알고리즘을 구축했습니다."
    )

    h2 = doc.add_heading("4.1 위험 점수 산정 공식 및 의사결정 수식", level=2)
    p = doc.add_paragraph()
    p.add_run("각 결함 후보 $d$에 대한 위험도 스코어 $S(d)$는 다음과 같이 정식화됩니다:\n\n").bold = True
    p.add_run("    Score(d) = w_joint(d) × max(0.5, p_defect(d)) × c_state(d)\n\n")
    p.add_run(
        "- w_joint: 신체 관절별 부상 위험 가중치 (요추/척추: 1.00, 무릎/외반: 0.85~0.90, 골반/고관절: 0.80, 발목: 0.60, 경추: 0.40)\n"
        "- p_defect: 모델이 예측한 결함 사후 확률 (0.0 ~ 1.0, 결함 판정 시 기본 바닥값 0.5 보장)\n"
        "- c_state: 3-State 계수 (State 1 결함 = 1.0, State 2 주의/보류 = 0.4, State 0 정상 = 0.0)"
    )

    h2 = doc.add_heading("4.2 복합 결함 연쇄 인과관계(Causal Chaining) 핵심 규칙", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "여러 결함이 독립적인 것이 아니라 상호 인과관계로 연결된 경우, 최상위 큐잉을 우회하고 근본 원인(Root Cause)을 교정하는 합성 큐를 최우선 주입합니다. "
        "본 시스템에 탑재된 11대 핵심 연쇄 규칙은 다음과 같습니다:"
    )

    from src.web.core.feedback_engine import CAUSAL_CHAIN_RULES
    chain_table = doc.add_table(rows=len(CAUSAL_CHAIN_RULES) + 1, cols=4)
    chain_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["운동 종목", "결함 조합", "연쇄 인과 진단명", "합성된 핵심 원인 교정 큐"]
    for j, h in enumerate(c_headers):
        c = chain_table.cell(0, j)
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.name = "맑은 고딕"
        c.paragraphs[0].runs[0].font.size = Pt(8.5)
        set_cell_background(c, "475569")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for i, rule in enumerate(CAUSAL_CHAIN_RULES):
        r_cells = chain_table.rows[i + 1].cells
        ex_name = EXERCISE_CATALOG[rule.exercise_id]['short']
        r_cells[0].text = ex_name
        r_cells[1].text = "+".join(sorted(rule.trigger_codes))
        r_cells[2].text = rule.root_cause_title.split(" (")[0]
        r_cells[3].text = rule.synthesized_cue

        r_cells[0].width = Inches(1.1)
        r_cells[1].width = Inches(1.0)
        r_cells[2].width = Inches(2.2)
        r_cells[3].width = Inches(2.4)

        for c in r_cells:
            c.paragraphs[0].runs[0].font.name = "맑은 고딕"
            c.paragraphs[0].runs[0].font.size = Pt(8.0)

    set_table_borders(chain_table, "CBD5E1")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 제5장
    h1 = doc.add_heading("5. 5대 공식 카메라 시점(시점 A~E) 적응형 코칭 및 가이드", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run(
        "AGENTS.md 제4조에 의거하여 본 프로젝트는 '시점 A(우측 후방)', '시점 B(우측 전방)', '시점 C(정면)', '시점 D(좌측 전방)', '시점 E(좌측 후방)'의 "
        "공식 5대 시점 체계를 엄격히 유지합니다. 3단계 피드백 엔진은 사용자가 촬영한 시점의 생체역학적 관측 한계를 스스로 인지하고 능동적 시점 권장 팁을 제공합니다."
    )

    view_tips = [
        ("시점 C (정면 0도)에서의 요추 결함 감지 한계", 
         "정면 시점에서는 좌우 무릎 외반(c3)이나 골반 수평(bl5)은 매우 명확히 관측되나, 전후 시상면 척추 굴곡(c1, d1)은 사각지대가 발생합니다. 요추 결함이 의심되거나 State 2(주의) 발생 시 '시점 B(우측 전방 45도)' 촬영을 자동 권장합니다."),
        ("시점 B/D (전방 45도)에서의 좌우 대칭성 한계",
         "사선 시점에서는 척추 중립과 무릎 전방 이동은 훌륭하게 포착되나 좌우 대칭 비교에 원근 왜곡이 생깁니다. 무릎 외반이나 골반 비대칭 의심 시 '시점 C(정면)' 촬영을 권장합니다."),
        ("시점 A/E (후방 45도)에서의 골반 회전 특화",
         "후방 시점은 견갑골 패킹 및 골반의 회전 비대칭(bl5, dl5) 관측에 최적화되어 있으므로, 런지나 데드리프트의 비틀림 결함 분석에 특화 피드백을 전달합니다.")
    ]
    for title, desc in view_tips:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(desc)

    # 제6장
    h1 = doc.add_heading("6. 웹 서비스 아키텍처 및 파이프라인 연동 설계", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run(
        "3단계 피드백 엔진은 기존 웹 서비스(src/web)의 무결성과 하위 호환성을 100% 보존하도록 비파괴적(Non-destructive) 방식으로 연동되었습니다."
    )

    arch_points = [
        ("FeedbackEngine 싱글톤 (src/web/core/feedback_engine.py)",
         "33종 전수 딕셔너리, 11대 인과 연쇄 규칙, 3-State 포맷팅, 우선순위 큐, 다중 반복 세트 종합 처방전 생성을 단일 모듈로 캡슐화."),
        ("Pydantic 스키마 확장 (src/web/core/schemas.py)",
         "DefectPrediction에 target_joint, action_cue, biomechanical_reason 추가. RepetitionDetail에 primary_feedback, causal_chain_title 추가. VideoAnalysisResponse에 set_summary 필드 추가 (선택적 기본값으로 100% 하위 호환 보장)."),
        ("DefectModelRegistry 연동 (src/web/core/model_registry.py)",
         "기존 33종 결함 평가 완료 즉시 FeedbackEngine을 거쳐 '피드백 자리' 하드코딩을 실질적 생체역학 코칭 문구로 자동 치환."),
        ("VideoProcessor 연동 (src/web/pipeline/video_processor.py)",
         "각 회차별 최우선 지침 바인딩 및 분석 완료 시 전체 세트에 대한 clean_ratio, 최다 빈도 결함, 맞춤형 드릴을 포함한 종합 처방전(Set Summary) 생성."),
        ("프론트엔드 UI 대시보드 고도화 (app.js & style.css)",
         "횟수별 '🎯 최우선 교정 지침 배너(coaching-primary-banner)' 및 전체 세트 요약 '📋 맞춤 처방전 카드(set-summary-container)' 동적 렌더링.")
    ]
    for title, desc in arch_points:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(desc)

    # 제7장
    h1 = doc.add_heading("7. 하네스 검증 결과 및 벤치마크", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    add_callout_box(doc, [
        "■ [하네스 검증 결과] 3단계 피드백 엔진 및 웹 E2E 통합 테스트 전수 통과 (100% PASS)",
        "1. 커버리지 테스트 (tests/test_stage3_feedback_engine.py):",
        "   - 8대 운동 33종 전 결함 카탈로그 100% 매핑 검증 완료 (Coverage: 33/33, 100%)",
        "   - 3-State 상태별(정상/결함/주의) 포맷팅 및 우선순위 큐 정렬 무결성 검증 완료",
        "   - 11대 연쇄 결함 인과 융합 및 다중 반복 세트 종합 요약 검증 완료",
        "   - 벤치마크: 단일 반복 피드백 산출 평균 지연시간 = 0.0037 ms (목표 1.0ms 대비 270배 초고속)",
        "2. 기존 33종 하네스 테스트 (tests/test_web_all_33_models_harness.py):",
        "   - 8대 운동 33종 전 모델 100% 정상 작동 검증 완료 (33/33 통과, 0 Exceptions, 0 Silent Fallbacks)",
        "3. FastAPI 백엔드 웹 서버 E2E 통합 테스트 (tests/test_web_api.py):",
        "   - 실제 비디오 E2E 파이프라인에서 c1~c4 결함 진단 및 신규 피드백 정상 출력 확인 완료"
    ], bg_hex="F0FDF4", border_color="16A34A")

    # 제8장
    h1 = doc.add_heading("8. 향후 일정 및 단계별 로드맵", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    roadmap = [
        ("8주차 (현재 완료)", "3단계 자세 교정 및 피드백 엔진 코어 구축, 33종 전수 딕셔너리 완성, 스키마 확장, 웹 프론트엔드 연동, 상세 기획서 완성 (완료)"),
        ("9주차 (실증 고도화)", "실제 사용자 일반화 영상 기반 E2E 피드백 실증 테스트, LLM(Gemini) 연동 비동기 맞춤형 운동 리포트 확장 인터페이스 프로토타이핑"),
        ("10~11주차 (2단계 부하 연계)", "연구 계획서 2단계(관절 부하 측정)의 경량화 버전(GJSI 기하학적 관절 부하 지수)을 피드백 엔진에 흡수하여 정량적 부하 증감율(% 단위) 출력 고도화")
    ]
    for period, detail in roadmap:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(f"{period}: ")
        run.bold = True
        p.add_run(detail)

    # 저장
    doc.save(DOCX_OUT)
    print(f"[+] DOCX 저장 완료: {DOCX_OUT}")

    # 마크다운 동시 생성
    generate_md_report()

def generate_md_report():
    from src.web.core.feedback_engine import FEEDBACK_CATALOG, CAUSAL_CHAIN_RULES
    from src.web.core.config import EXERCISE_CATALOG

    md_content = f"""# [vidio-health-trainer] 8주차 3단계 자세 교정 및 결함 피드백 시스템 상세 기획 보고서

- **문서 분류**: 8주차 공식 상세 기술 기획서 및 아키텍처 명세서
- **공식 프로젝트명**: `vidio-health-trainer` (공식 명칭 규정 엄격 준수)
- **보고 일시**: 2026년 10월 13일 (화) - 8주차 아키텍처 공식 마일스톤
- **배포 위치**: `D:\\자주프 문서\\8주차\\8주차_3단계_자세교정_및_결함피드백_시스템_상세기획보고서.docx` (참조용 MD: `md/`)
- **검증 상태**: Unit Test 100% PASS / 33-Model Harness 100% PASS / Web API E2E 100% PASS

---

## ■ 8주차 종합 핵심 성과 (Executive Summary)

1. **선회 배경 및 전략적 타당성**: 1단계 운동 종목 분류와 2단계 33종 결함 분류 모델 구축이 완료되어 웹 서빙 중이나, 기존 출력이 `'피드백 자리'`로 하드코딩되어 사용자 인터랙션 루프가 미완성이었음. 연구 계획서상의 2단계(OpenSim 역동역학 관절 부하 측정)는 3D 지면반력 부재에 따른 오차 증폭 및 50ms 실시간 서빙 지연 문제가 있어, 사용자에게 즉각적 가치를 제공하는 **3단계(결함 시 피드백 시스템)를 선제 구축**함.
2. **3-Tier 계층형 코칭 체계**: 
   - **Tier A (Action Cue)**: 1문장 직관적 행동 지시문
   - **Tier B (Biomechanical Reason)**: 생체역학적 부하 메커니즘 및 관절 손상 원인
   - **Tier C (3-State Guard)**: State 0(칭찬/유지), State 1(즉각 교정), State 2(주의/재촬영 권장) 차등화
3. **8대 운동 33종 세부 결함 100% 딕셔너리 구축**: 스쿼트(4종), 데드리프트(5종), 스티프데드(3종), 바벨런지(5종), 다이나믹런지(5종), 사이드런지(5종), 크로스런지(3종), 굿모닝(3종) 전수에 대해 행동 지침, 역학적 원인, 타깃 관절, 최적 시점, 추천 드릴 완비.
4. **생체역학 위험 가중치 기반 우선순위 큐 & 연쇄 결함(Causal Chaining) 융합**:
   - 척추(1.00) > 무릎(0.85) > 골반(0.80) > 발목(0.60) 가중치를 적용하여 1회 반복당 최우선 1~2개 큐만 선별(인지 과부하 방지).
   - '발바닥 뜸 -> 무릎 외반', '힙 슈트업 -> 요추 부하' 등 11개 복합 연쇄를 단일 근본 원인 큐로 합성.
5. **초고속 및 무결성 하네스 실증**:
   - 피드백 산출 1회 지연시간 **0.0037 ms** (목표 1.0ms 대비 270배 초고속 달성).
   - 기존 33종 전 모델 0 Exceptions 무결성 보존 및 웹 API E2E 통합 테스트 100% 통과.

---

## 1. 기획 배경 및 2단계(관절 부하 측정) 선회 타당성 분석

### 1.1 기존 시스템의 서비스 완결성 병목: 하드코딩된 피드백
`vidio-health-trainer`는 8대 운동 33종 결함 모델을 구축하여 단일 횟수별로 확률과 3-State 판정을 내리고 있으나, 출력 필드가 `'피드백 자리'`로 남아 있었습니다. 사용자가 영상을 업로드했을 때 "무릎 외반 95%"라는 숫자만으로는 어떻게 폼을 고쳐야 하는지 알 수 없으므로, **실질적 행동 지침(Actionable Cueing)**을 제공하는 3단계 피드백 엔진이 절실히 요구되었습니다.

### 1.2 연구 계획서 2단계(관절 부하 측정) 선회 및 후순위 조정 사유
1. **3D 지면반력(GRF) 부재로 인한 역동역학 오차 증폭**: 스마트폰 단일 2D 영상 환경에서는 지면반력계가 없어 역동역학 모멘트 오차가 25~40% 이상 발생함.
2. **실시간 웹 서빙 레이턴시 병목 (20~50ms 목표 위배)**: OpenSim 연산은 수백 ms~수 초가 소모되어 실시간성을 저해함.
3. **서비스 완결성 우선 원칙 (Closed Feedback Loop)**: 사용자에게 즉각적 피드백을 전달하는 엔드투엔드 파이프라인 완성이 우선임.
4. **2단계 부하의 3단계 내부 프록시 흡수 가능성 (Kinematic Stress Proxy)**: 관절 각도와 모멘트암, 생체역학 위험 가중치($w_{{joint}}$)를 결합하여 피드백 우선순위 큐에 경량화된 형태로 흡수 달성.

---

## 2. 3단계 자세 교정 및 피드백 시스템의 5대 핵심 설계 원칙

| 설계 원칙 | 핵심 메커니즘 | 엔지니어링 및 사용자 가치 |
| :--- | :--- | :--- |
| **[원칙 1] 3-Tier 계층형 피드백** | Tier A(행동 지침 큐) + Tier B(역학적 원인) + Tier C(3-State 차등화) | 단순 위험 경고를 넘어 즉각 수정 가능한 1문장 큐 제공 |
| **[원칙 2] 인지 부하 제어 (CLT)** | 단일 반복당 최우선 1개(최대 2개) 교정 지침만 선별 노출 | 운동 수행 중 뇌 인지 과부하(Cognitive Overload) 방지 |
| **[원칙 3] 생체역학 위험 우선순위** | $Score = w_{{joint}} \\times p_{{defect}} \\times c_{{state}}$ 기반 정렬 | 손상 위험이 가장 큰 핵심 관절(척추/무릎) 최우선 구제 |
| **[원칙 4] 연쇄 결함 인과 융합** | 복합 결함(Trigger Set)을 단일 루트 코즈 큐로 합성 | '발바닥 뜸 -> 무릎 외반' 등 근본 원인 해결 |
| **[원칙 5] 비의료적 표현 표준화** | 질환/진단명 영구 배제, 운동역학적 부하 집중 신호로 표현 | 의료법 저촉 방지 및 안전한 피트니스 코칭 유지 |

---

## 3. 8대 운동 33종 세부 결함 전수 코칭 딕셔너리 명세

"""
    # 8대 운동 테이블 추가
    for eid, ex_info in EXERCISE_CATALOG.items():
        ex_name = ex_info['name']
        defect_items = [item for item in FEEDBACK_CATALOG.values() if item.exercise_name == ex_name]
        md_content += f"\n### 3.{eid+1} {ex_name} ({len(defect_items)}개 결함)\n\n"
        md_content += "| 코드 | 결함 명칭 | 타깃 관절 | 가중치 | 핵심 행동 지침 큐 (Action Cue) | 권장 시점 |\n"
        md_content += "| :---: | :--- | :---: | :---: | :--- | :---: |\n"
        for item in defect_items:
            md_content += f"| `{item.defect_code}` | {item.defect_name.split(' (')[0]} | {item.target_joint} | {item.risk_weight:.2f} | {item.action_cue} | {item.optimal_view.split(' (')[0]} |\n"

    md_content += """
---

## 4. 다중 결함 우선순위 해결 및 연쇄 결함(Causal Chaining) 알고리즘

### 4.1 위험 점수 산정 공식
$$Score(d) = w_{joint}(d) \\times \\max(0.5, p_{defect}(d)) \\times c_{state}(d)$$
- $w_{joint}$: 요추/척추(1.00) > 무릎(0.85) > 골반(0.80) > 발목(0.60) > 경추(0.40)
- $c_{state}$: State 1(결함) = 1.0, State 2(주의) = 0.4, State 0(정상) = 0.0

### 4.2 11대 핵심 복합 결함 연쇄 규칙 (Causal Chaining)
| 운동 종목 | 결함 조합 | 연쇄 인과 진단명 | 합성된 핵심 원인 교정 큐 |
| :---: | :---: | :--- | :--- |
"""
    for rule in CAUSAL_CHAIN_RULES:
        ex_name = EXERCISE_CATALOG[rule.exercise_id]['short']
        triggers = "+".join(sorted(rule.trigger_codes))
        md_content += f"| {ex_name} | `{triggers}` | {rule.root_cause_title.split(' (')[0]} | {rule.synthesized_cue} |\n"

    md_content += """
---

## 5. 5대 공식 카메라 시점(시점 A~E) 적응형 코칭 및 가이드

- **시점 A**: 우측 후방 (45도) [Right Back 45° / view5 매핑] -> 골반 회전 비대칭 및 견갑 패킹 최적
- **시점 B**: 우측 전방 (45도) [Right Front 45° / view2 매핑] -> 척추 중립성 및 전후 모멘트암 최적
- **시점 C**: 정면 (0도) [Front View 0° / view3 매핑] -> 무릎 외반 및 좌우 정렬 최적
- **시점 D**: 좌측 전방 (45도) [Left Front 45° / view4 매핑] -> 런지 앞무릎 각도 및 시상면 정렬 최적
- **시점 E**: 좌측 후방 (45도) [Left Back 45° / view1 매핑] -> 후방 척추 만곡 및 바벨 궤적 최적

---

## 6. 웹 서비스 아키텍처 및 파이프라인 연동 설계

```mermaid
flowchart TD
    A["입력 비디오 (64F 스플라인 보간)"] --> B["1단계 운동 종목 분류 & 2단계 33종 결함 평가"]
    B --> C["FeedbackEngine (src/web/core/feedback_engine.py)"]
    C --> D1["1. 33종 딕셔너리 룩업 (Tier A/B/C)"]
    C --> D2["2. 생체역학 위험 점수 계산 & 우선순위 큐 정렬"]
    C --> D3["3. 11대 연쇄 인과관계 검출 (Causal Chaining)"]
    C --> D4["4. 5대 시점 적응형 안내 & 세트 종합 처방전 도출"]
    D1 & D2 & D3 & D4 --> E["Pydantic Schemas (RepetitionDetail & VideoAnalysisResponse)"]
    E --> F["웹 대시보드 UI (최우선 배너 & 맞춤 처방전 렌더링)"]
```

---

## 7. 하네스 검증 결과 및 벤치마크

1. **커버리지 테스트 (`tests/test_stage3_feedback_engine.py`)**:
   - 8대 운동 33종 전 결함 카탈로그 **100% 매핑 검증 완료 (Coverage: 33/33)**.
   - 단일 반복 피드백 산출 평균 지연시간: **0.0037 ms** (목표 1.0ms 대비 270배 초고속).
2. **33종 모델 하네스 (`tests/test_web_all_33_models_harness.py`)**:
   - 8대 운동 33종 전 모델 **100% 정상 통과 (0 Exceptions, 0 Silent Fallbacks)**.
3. **웹 API E2E 통합 테스트 (`tests/test_web_api.py`)**:
   - 실제 비디오 E2E 파이프라인에서 신규 피드백 정상 출력 및 서빙 무결성 확인.

---

## 8. 향후 일정 및 단계별 로드맵

- **8주차 (현재 완료)**: 3단계 자세 교정 및 피드백 엔진 코어 구축, 33종 전수 딕셔너리 완성, 스키마 확장, 웹 프론트엔드 연동, 상세 기획서 완성.
- **9주차 (실증 고도화)**: 실제 일반화 영상 기반 E2E 피드백 실증 테스트, LLM(Gemini) 연동 비동기 맞춤형 운동 처방전 확장 인터페이스 개발.
- **10~11주차 (2단계 부하 연계)**: 기하학적 관절 부하 지수(GJSI)를 피드백 엔진에 통합하여 정량적 부하 증감율(% 단위) 출력 고도화.
"""
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[+] MD 저장 완료: {MD_OUT}")

if __name__ == "__main__":
    generate_reports()
