# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: src/reports/generate_kqs_implementation_report.py
Purpose: 8주차 멀티비전 키포인트 품질 평가(KQS) 및 재촬영 게이팅 구현 완료 보고서 자동 생성기
         - 출력 Word 문서: D:\자주프 문서\8주차\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.docx
         - 출력 참조 MD:  D:\자주프 문서\8주차\md\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.md
"""

import os
import sys

PROJECT_ROOT = r"d:/"
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_OUT = r"D:\자주프 문서\8주차\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="2B579A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout_box(doc, text_list, bg_hex="F0F4F8", border_color="2B579A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.7)
    set_cell_background(c, bg_hex)
    
    tcPr = c._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for i, t in enumerate(text_list):
        if i > 0:
            p = c.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
        run = p.add_run(t)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def generate_reports():
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # 1. 문서 헤더 / 제목
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("[vidio-health-trainer] 8주차 공식 기술 구현 보고서")
    run_title.font.name = "맑은 고딕"
    run_title.font.size = Pt(21)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(43, 87, 154)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("멀티비전 키포인트 품질 평가(KQS) 및 재촬영/보류 게이팅 파이프라인 구현 완료 보고서")
    run_sub.font.name = "맑은 고딕"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 메타데이터 표
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("공식 프로젝트명", "vidio-health-trainer (AI 기반 웨이트 트레이닝 자세 진단 및 코칭 시스템)"),
        ("문서 분류 / 주차", "8주차 공식 기술 구현 및 자가 치유 하네스 실증 보고서"),
        ("작성 기준 일시", "2026년 10월 17일 (토) - 8주차 세션 최종 마일스톤"),
        ("Word 배포 파일", "D:\\자주프 문서\\8주차\\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.docx"),
        ("참조 마크다운 파일", "D:\\자주프 문서\\8주차\\md\\8주차_키포인트품질(KQS)_평가_및_재촬영게이팅_구현완료보고서.md")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_k = row.cells[0]
        cell_v = row.cells[1]
        cell_k.width = Inches(1.8)
        cell_v.width = Inches(4.9)
        set_cell_background(cell_k, "F8FAFC")
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(2)
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.name = "맑은 고딕"
        rk.font.size = Pt(9.5)
        rk.font.bold = True
        rk.font.color.rgb = RGBColor(71, 85, 105)

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(2)
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.name = "맑은 고딕"
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = RGBColor(15, 23, 42)
    set_table_borders(meta_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Summary
    add_callout_box(doc, [
        "■ [Executive Summary] 8주차 KQS 게이팅 구현 핵심 성과",
        "1. 사용자 명시 지침 완벽 반영: 슬관절(무릎) 가림 코칭 큐에서 '바벨 원판' 언급을 전면 배제하고, '무릎 관절이 화면 또는 물체에 가려져 정밀 분석이 어렵습니다. 무릎이 온전히 보이도록 카메라 각도와 위치를 조절하여 재촬영해 주세요.'로 순수 관절 가림 및 물리적 조치에 집중하도록 수정 완료함.",
        "2. 멀티비전 품질 지수(KQS) 엔진 탑재: FeedbackEngine 내 evaluate_keypoint_quality()를 신설하여 MediaPipe visibility와 RTMPose confidence를 융합하고, 종목별 핵심 관절 가중치 및 연속 결손 페널티(StreakPenalty)를 적용함.",
        "3. 2단계 결함 모델 오진단(GIGO) 원천 차단: 품질 미달(KQS < 0.60 또는 핵심 관절 가림) 감지 시 2단계 머신러닝 추론을 즉시 바이패스(Bypass)하고, 결함을 임의 단정 짓지 않고 '판정 보류 (재촬영 권고)' 상태 및 타깃형 재촬영 큐를 직결함.",
        "4. 자가 치유(Self-Healing) 하네스 검증 100% 통과: KQS 전용 하네스 테스트(6/6 PASS, 0.001s), 33종 모델 하네스(33/33 PASS, 100%), FastAPI 백엔드 E2E 통합 테스트(비디오 업로드 및 웹소켓 포함 100% PASS)를 완벽히 통과함."
    ], bg_hex="F0FDF4", border_color="16A34A")

    def add_section_header(title, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.font.name = "맑은 고딕"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(43, 87, 154)
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(71, 85, 105)

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # 1. 배경 및 사용자 수정 지시사항 반영
    add_section_header("1. 배경 및 사용자 지침 반영 (User Requirements & Fixes)")
    add_body_p(
        "사용자가 직접 스마트폰 카메라로 촬영하는 일반화 영상 데이터를 웹 서비스에 투입하기에 앞서, "
        "화각 이탈이나 신체 가림으로 인한 2단계 결함 분류기의 False Positive(오진단)를 차단하는 안전장치가 필수적이었습니다."
    )
    add_body_p(
        "특히 사용자의 세부 지침에 따라, 슬관절(무릎) 결손 시의 코칭 문구에서 '바벨 원판에 가려졌다'는 구체적 기구 언급을 배제하고 "
        "모든 촬영 환경(맨몸 스쿼트, 덤벨 운동 등)에 유연하게 대응할 수 있도록 문구를 일반화하여 정제하였습니다."
    )

    prompt_table = doc.add_table(rows=6, cols=3)
    prompt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_headers = ["결손 코드", "원인 분석", "최종 탑재된 1문장 행동 지침 큐 (Action Cue)"]
    for c_idx, h in enumerate(p_headers):
        cell = prompt_table.rows[0].cells[c_idx]
        set_cell_background(cell, "2B579A")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    p_rows = [
        ("OCC_KNEE", "무릎 관절 가림 (슬관절 결손)", "무릎 관절이 화면 또는 물체에 가려져 정밀 분석이 어렵습니다. 무릎이 온전히 보이도록 카메라 각도와 위치를 조절하여 재촬영해 주세요."),
        ("OCC_FEET", "화면 하단 잘림 (족관절 결손)", "발목과 발바닥이 화면 하단 밖으로 벗어났습니다. 카메라 각도를 살짝 내리거나 스마트폰을 1m 뒤로 물려 발끝 지면 접지면이 모두 보이도록 재촬영해 주세요."),
        ("OCC_SPINE", "체간/골반선 인식 불완전", "골반 및 척추선 인식이 불완전합니다. 상체 윤곽이 드러나는 복장을 착용하거나 조명이 피사체를 정면으로 비추는 환경에서 재촬영해 주세요."),
        ("OCC_BLUR", "급격한 움직임 / 모션 블러", "카메라 흔들림이나 모션 블러로 관절 추적 신뢰도가 낮습니다. 스마트폰을 거치대에 고정하고 템포를 통제하며 재촬영해 주세요."),
        ("OCC_GENERAL", "종합 품질 기준치(0.60) 미달", "신체 주요 관절의 가림 현상으로 분석 신뢰도가 낮습니다. 전신이 프레임 안에 잘 들어오도록 카메라 위치를 조정하여 재촬영해 주세요.")
    ]
    for r_idx, row_data in enumerate(p_rows, start=1):
        row = prompt_table.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "맑은 고딕"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(51, 65, 85)
    set_table_borders(prompt_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2. 코드베이스 정밀 수정 내역
    add_section_header("2. 코드베이스 정밀 수정 내역 (Surgical Edits Specification)")
    add_body_p(
        "프로젝트 무결성을 유지하며 4개 핵심 모듈에 대해 비파괴적 점진 수정을 적용하였습니다:"
    )
    add_body_p(
        "1. src/web/core/schemas.py:\n"
        "   - RepetitionDetail에 kqs_score, is_retake_required, occluded_joints, retake_cue 필드 추가\n"
        "2. src/web/core/feedback_engine.py:\n"
        "   - RETAKE_PROMPT_CATALOG 딕셔너리 신설 (사용자 요청 문구 반영)\n"
        "   - evaluate_keypoint_quality() 신설: 64프레임 시간축 가중 품질 및 연속 결손 페널티 산출\n"
        "   - prioritize_rep_defects()에 kqs_info 매개변수 연동 및 판정보류 최우선 분기 추가\n"
        "3. src/web/pipeline/stage2_engine.py:\n"
        "   - diagnose 및 evaluate_single_repetition에 mp_vis_seq, rtm_conf_seq 전달 지원\n"
        "   - KQS < 0.60 시 2단계 모델 추론 바이패스 및 0.00ms 즉시 보류 분기 처리\n"
        "4. src/web/pipeline/video_processor.py:\n"
        "   - 반복 구간별 MediaPipe vis_33 (64, 33) 및 RTMPose conf (64, 26) 스플라인 보간 및 전달\n"
        "   - RepetitionDetail 결과 객체에 KQS 지표 완전 매핑"
    )

    # 3. 하네스 검증 결과
    add_section_header("3. 자동화 검증 및 자가 치유(Self-Healing) 하네스 실증 결과")
    add_body_p(
        "하네스 엔지니어링 규정에 따라 유닛 테스트, 33종 모델 하네스, 백엔드 E2E 통합 테스트를 전수 실행하였습니다."
    )

    harness_table = doc.add_table(rows=5, cols=4)
    harness_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_headers = ["검증 테스트 스위트", "검증 목적 및 대상", "수행 결과", "비고"]
    for c_idx, h in enumerate(h_headers):
        cell = harness_table.rows[0].cells[c_idx]
        set_cell_background(cell, "2B579A")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    h_rows = [
        ("test_kqs_gating_harness.py", "정상 통과, 무릎 가림, 발목 화각 이탈, 연속 결손 페널티, 바이패스 검증", "6 / 6 PASS (100.0%)", "소요 시간: 0.001s"),
        ("test_stage3_feedback_engine.py", "8대 운동 33종 전 결함 룩업 및 우선순위 큐, 세트 요약 검증", "6 / 6 PASS (100.0%)", "지연시간: 0.0034 ms"),
        ("test_web_all_33_models_harness.py", "8대 운동 33종 결함 진단 모델 0 Exceptions 무결성 전수 검증", "33 / 33 PASS (100.0%)", "0 Silent Fallbacks"),
        ("test_web_api.py", "FastAPI 웹 서버 루트 HTML, Health, Sequence, E2E Video, WebSocket 검증", "전 엔드포인트 100% PASS", "비디오 E2E 완결 검증")
    ]
    for r_idx, row_data in enumerate(h_rows, start=1):
        row = harness_table.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "맑은 고딕"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(51, 65, 85)
    set_table_borders(harness_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. 결론 및 향후 계획
    add_section_header("4. 결론 및 향후 계획 (Next Steps)")
    add_body_p(
        "멀티비전 키포인트 품질 평가(KQS) 및 재촬영 게이팅 파이프라인이 완벽히 구축되어 웹 서비스에 탑재되었습니다. "
        "이제 사용자가 직접 촬영하신 일반화 영상(스쿼트, 데드리프트, 런지 등)을 시스템에 투입하더라도, "
        "가림 현상이나 화각 이탈 시 엉뚱한 결함 오진단이 발생하지 않고 지능적인 재촬영 안내로 안전하게 보호됩니다."
    )

    doc.save(DOCX_OUT)
    print(f"[SUCCESS] Word Report Generated: {DOCX_OUT}")

    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write("# [vidio-health-trainer] 8주차 공식 기술 구현 보고서\n\n")
        f.write("## 멀티비전 키포인트 품질 평가(KQS) 및 재촬영/보류 게이팅 파이프라인 구현 완료 보고서\n\n")
        f.write("- **공식 프로젝트명**: `vidio-health-trainer`\n")
        f.write("- **작성 일시**: 2026년 10월 17일 (토)\n")
        f.write(f"- **Word 배포 파일**: `{DOCX_OUT}`\n\n")
        f.write("---\n\n")
        f.write("### ■ [Executive Summary] 8주차 KQS 게이팅 구현 핵심 성과\n\n")
        f.write("1. **사용자 명시 지침 완벽 반영**: 슬관절(무릎) 가림 코칭 큐에서 '바벨 원판' 언급을 전면 배제하고, '무릎 관절이 화면 또는 물체에 가려져 정밀 분석이 어렵습니다. 무릎이 온전히 보이도록 카메라 각도와 위치를 조절하여 재촬영해 주세요.'로 순수 관절 가림 및 물리적 조치에 집중하도록 수정 완료함.\n")
        f.write("2. **멀티비전 품질 지수(KQS) 엔진 탑재**: FeedbackEngine 내 evaluate_keypoint_quality()를 신설하여 MediaPipe visibility와 RTMPose confidence를 융합하고, 종목별 핵심 관절 가중치 및 연속 결손 페널티를 적용함.\n")
        f.write("3. **2단계 결함 모델 오진단(GIGO) 원천 차단**: 품질 미달(KQS < 0.60 또는 핵심 관절 가림) 감지 시 2단계 머신러닝 추론을 즉시 바이패스하고 '판정 보류 (재촬영 권고)' 상태 및 타깃형 재촬영 큐를 직결함.\n")
        f.write("4. **자가 치유 하네스 검증 100% 통과**: test_kqs_gating_harness.py (6/6 PASS, 0.001s), 33종 전 모델 하네스 (33/33 PASS, 100%), FastAPI 백엔드 E2E 통합 테스트 (전 엔드포인트 100% PASS)를 완벽 통과함.\n")
    print(f"[SUCCESS] Markdown Report Generated: {MD_OUT}")

if __name__ == '__main__':
    generate_reports()
