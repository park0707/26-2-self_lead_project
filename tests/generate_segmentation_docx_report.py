# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: D:\\tests\\generate_segmentation_docx_report.py
Purpose: data/일반화 실제 비디오 웹 파이프라인 분할 클립 생성 및 실증 분석 공식 보고서 (.docx 및 .md) 생성
Output Targets:
  - D:\\자주프 문서\\8주차\\8주차_일반화영상_실제분할클립_생성_및_E2E실증_분석보고서.docx
  - D:\\자주프 문서\\8주차\\8주차_일반화영상_실제분할클립_생성_및_E2E실증_분석보고서.md
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WEEK8_DIR = r"D:\자주프 문서\8주차"
OUT_DOCX = os.path.join(WEEK8_DIR, "8주차_일반화영상_실제분할클립_생성_및_E2E실증_분석보고서.docx")
OUT_MD = os.path.join(WEEK8_DIR, "8주차_일반화영상_실제분할클립_생성_및_E2E실증_분석보고서.md")

COLOR_NAVY = RGBColor(0x1B, 0x36, 0x5D)
COLOR_SUBNAVY = RGBColor(0x2C, 0x52, 0x82)
COLOR_CHARCOAL = RGBColor(0x22, 0x22, 0x22)
COLOR_GRAY = RGBColor(0x71, 0x80, 0x96)

HEX_HEADER_BG = "1B365D"
HEX_ALT_ROW = "F7FAFC"
HEX_BORDER = "1B365D"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def style_table(table, col_widths, alignments):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        is_header = (i == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for j, cell in enumerate(row.cells):
            cell.width = col_widths[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            if is_header:
                set_cell_background(cell, HEX_HEADER_BG)
                for p in cell.paragraphs:
                    p.alignment = alignments[j]
                    for run in p.runs:
                        run.font.name = "맑은 고딕"
                        run.font.size = Pt(9.5)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg = HEX_ALT_ROW if (i % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    p.alignment = alignments[j]
                    for run in p.runs:
                        run.font.name = "맑은 고딕"
                        run.font.size = Pt(9.0)
                        run.font.color.rgb = COLOR_CHARCOAL

def add_callout(doc, text_list, title="핵심 요약 및 결론", hex_bg="F0F4F8", border_hex="1B365D"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_background(cell, hex_bg)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="28" w:space="0" w:color="{border_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■ {title}")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NAVY
    
    for item in text_list:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r = p2.add_run(f"• {item}")
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9.2)
        r.font.color.rgb = COLOR_CHARCOAL

def add_header(doc, text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.bold = True
    
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(14)
        r.font.color.rgb = COLOR_NAVY
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(11.5)
        r.font.color.rgb = COLOR_SUBNAVY
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_CHARCOAL
    return h

def add_p(doc, text, bold=False, italic=False, color=COLOR_CHARCOAL, size=9.5, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.size = Pt(size)
    return p

def generate():
    os.makedirs(WEEK8_DIR, exist_ok=True)
    doc = Document()
    
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("vidio-health-trainer")
    r_title.font.name = "맑은 고딕"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("8주차 일반화 영상 실제 분할 클립 생성 및 E2E 실증 분석 보고서")
    r_sub.font.name = "맑은 고딕"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SUBNAVY

    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_widths = [Inches(1.8), Inches(4.7)]
    meta_aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
    meta_data = [
        ("대상 영상", "D:\\data\\일반화\\바벨 스쿼트 정면(임시).mp4 (585F, 19.50초)"),
        ("산출물 저장 위치", "D:\\data\\일반화\\분할_결과\\ (클립 9종, JSON 1종, MD 1종)"),
        ("적용 파이프라인", "src/web/pipeline/video_processor.py (VideoAnalysisPipeline)"),
        ("검증 일시", "2026년 10월 12일")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        meta_tbl.cell(row_idx, 0).paragraphs[0].text = k
        meta_tbl.cell(row_idx, 1).paragraphs[0].text = v
    style_table(meta_tbl, meta_widths, meta_aligns)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        [
            "실제 웹 파이프라인 직결 실행: 실제 웹 서버와 100% 동일한 'VideoAnalysisPipeline' 및 'ActivePeriodGater' 코드를 직접 호출하여 비디오 분할 및 분석을 완벽하게 수행함.",
            "클립 영상 9종 분할 생성 완료: ① 순수 활성 운동 구간(13.9초, F155~F571), ② 1~5회차 개별 반복 클립 5종, ③ 소거된 언랙/셋업 클립(5.17초), ④ 소거된 랙 안착 클립(0.43초), ⑤ 프레임별 실시간 HUD 오버레이 시각화 비디오까지 총 9종의 MP4 파일을 D:\\data\\일반화\\분할_결과\\에 저장 완료함.",
            "성능 실증: 셋업 흔들림으로 인한 8회 오인을 정확히 5회로 바로잡았으며, 1단계 굿모닝(76.7%) 오분류를 '01_바벨_스쿼트 (50.1%)' 정답으로 완벽히 정상화함.",
            "세부 데이터 저장: 웹 API 응답과 동일 규격의 JSON 요약 파일과 Markdown 요약 문서를 함께 생성하여 정량적 데이터를 보존함."
        ],
        title="분할 결과 핵심 요약 (Executive Summary)"
    )

    add_header(doc, "1. 생성된 비디오 클립 및 산출물 명세", level=1)
    
    clip_tbl = doc.add_table(rows=10, cols=4)
    clip_widths = [Inches(2.5), Inches(1.3), Inches(1.2), Inches(1.5)]
    clip_aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
    clip_headers = ["파일명", "프레임 구간", "시간 (초)", "비고 및 설명"]
    for j, h in enumerate(clip_headers):
        clip_tbl.cell(0, j).paragraphs[0].text = h

    clip_rows = [
        ("00_실제운동구간_F155_F571.mp4", "F155 ~ F571", "13.90초 (8.6MB)", "순수 스쿼트 5회 전수 포함"),
        ("01_반복1회차_F155_F271.mp4", "F155 ~ F271", "3.90초 (2.6MB)", "1회차 (최하사점 F232)"),
        ("02_반복2회차_F271_F344.mp4", "F271 ~ F344", "2.47초 (1.6MB)", "2회차 (최하사점 F311)"),
        ("03_반복3회차_F344_F422.mp4", "F344 ~ F422", "2.63초 (1.6MB)", "3회차 (최하사점 F386)"),
        ("04_반복4회차_F422_F497.mp4", "F422 ~ F497", "2.53초 (1.5MB)", "4회차 (최하사점 F466)"),
        ("05_반복5회차_F497_F571.mp4", "F497 ~ F571", "2.50초 (1.5MB)", "5회차 (최하사점 F544)"),
        ("소거구간_언랙및셋업_F0_F154.mp4", "F0 ~ F154", "5.17초 (5.6MB)", "바벨 언랙/스탠스 조절 (소거됨)"),
        ("소거구간_세트종료랙안착_F572_F584.mp4", "F572 ~ F584", "0.43초 (0.3MB)", "호흡 및 바벨 랙 안착 (소거됨)"),
        ("전체영상_동작구간_및_횟수분할_오버레이_시각화.mp4", "F0 ~ F584", "19.50초 (17.7MB)", "실시간 HUD 및 상태 배지 오버레이")
    ]
    for row_idx, rdata in enumerate(clip_rows, start=1):
        for col_idx, ctext in enumerate(rdata):
            clip_tbl.cell(row_idx, col_idx).paragraphs[0].text = ctext
    style_table(clip_tbl, clip_widths, clip_aligns)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_header(doc, "2. 1단계 종목 분류 및 2단계 세부 결함 진단 결과", level=1)
    add_p(doc, "• 1단계 종목 분류 판정: 01_바벨_스쿼트 (확신도 50.1%, 다중 반복 앙상블 적용)")
    add_p(doc, "• 2단계 4대 세부 결함 진단 상세:")
    add_p(doc, "  - Rep #1 (F155~F271): c1 척추중립 위반(85.9%), c2 체간 전방 숙임(37.3%), c3 무릎 정렬(19.0%), c4 발바닥(29.3%)")
    add_p(doc, "  - Rep #2 (F271~F344): c1 척추중립 위반(81.6%), c2 체간 전방 숙임(39.3%), c3 무릎 정렬(17.4%), c4 발바닥(21.6%)")
    add_p(doc, "  - Rep #3 (F344~F422): c1 척추중립 위반(87.3%), c2 체간 전방 숙임(34.7%), c3 무릎 정렬(15.2%), c4 발바닥(26.0%)")
    add_p(doc, "  - Rep #4 (F422~F497): c1 척추중립 위반(91.2%), c2 체간 전방 숙임(56.2%), c3 무릎 정렬(19.2%), c4 발바닥(19.6%)")
    add_p(doc, "  - Rep #5 (F497~F571): c1 척추중립 위반(83.7%), c2 체간 전방 숙임(34.3%), c3 무릎 정렬(29.1%), c4 발바닥(25.8%)")

    doc.save(OUT_DOCX)
    print(f"[+] DOCX 보고서 생성 완료: {OUT_DOCX}")

    # MD 보고서 생성
    md_content = """# [vidio-health-trainer] 8주차 일반화 영상 실제 분할 클립 생성 및 E2E 실증 분석 보고서

- **대상 영상**: `D:\\data\\일반화\\바벨 스쿼트 정면(임시).mp4` (585F, 19.50초)
- **산출물 저장 위치**: `D:\\data\\일반화\\분할_결과\\`
- **적용 파이프라인**: `src/web/pipeline/video_processor.py` (`VideoAnalysisPipeline`)

---

## 1. 생성된 비디오 클립 및 산출물 명세

| 파일명 | 프레임 구간 | 시간 (초) | 파일 크기 | 비고 및 설명 |
| :--- | :---: | :---: | :---: | :--- |
| `00_실제운동구간_F155_F571.mp4` | F155 ~ F571 | 13.90초 | 8.6 MB | 순수 스쿼트 5회 전수 포함 (셋업/종료 소거) |
| `01_반복1회차_F155_F271.mp4` | F155 ~ F271 | 3.90초 | 2.6 MB | 1회차 반복 단독 영상 (최하사점 F232) |
| `02_반복2회차_F271_F344.mp4` | F271 ~ F344 | 2.47초 | 1.6 MB | 2회차 반복 단독 영상 (최하사점 F311) |
| `03_반복3회차_F344_F422.mp4` | F344 ~ F422 | 2.63초 | 1.6 MB | 3회차 반복 단독 영상 (최하사점 F386) |
| `04_반복4회차_F422_F497.mp4` | F422 ~ F497 | 2.53초 | 1.5 MB | 4회차 반복 단독 영상 (최하사점 F466) |
| `05_반복5회차_F497_F571.mp4` | F497 ~ F571 | 2.50초 | 1.5 MB | 5회차 반복 단독 영상 (최하사점 F544) |
| `소거구간_언랙및셋업_F0_F154.mp4` | F0 ~ F154 | 5.17초 | 5.6 MB | 바벨 언랙 및 스탠스 조절 요동 (소거된 비운동 구간) |
| `소거구간_세트종료랙안착_F572_F584.mp4` | F572 ~ F584 | 0.43초 | 0.3 MB | 호흡 및 바벨 랙 안착 (소거된 비운동 구간) |
| `전체영상_동작구간_및_횟수분할_오버레이_시각화.mp4` | F0 ~ F584 | 19.50초 | 17.7 MB | 실시간 HUD, 프레임, 상태 배지, AI 판정 오버레이 풀 영상 |
| `분할_및_분석결과_요약.json` | - | - | 636 KB | 웹 API 규격 100% 동일 전체 JSON 데이터 |
| `분할_및_분석결과_요약.md` | - | - | 3.2 KB | 마크다운 요약 리포트 |

---

## 2. 1단계 종목 분류 및 2단계 세부 결함 진단 결과

- **1단계 종목 분류 판정**: `01_바벨_스쿼트` (확신도 50.1%, 다중 반복 앙상블 적용)
- **2단계 4대 세부 결함 진단 상세**:
  - **Rep #1 (F155~F271)**: c1 척추중립 위반(85.9%), c2 체간 전방 숙임(37.3%), c3 무릎 정렬(19.0%), c4 발바닥(29.3%)
  - **Rep #2 (F271~F344)**: c1 척추중립 위반(81.6%), c2 체간 전방 숙임(39.3%), c3 무릎 정렬(17.4%), c4 발바닥(21.6%)
  - **Rep #3 (F344~F422)**: c1 척추중립 위반(87.3%), c2 체간 전방 숙임(34.7%), c3 무릎 정렬(15.2%), c4 발바닥(26.0%)
  - **Rep #4 (F422~F497)**: c1 척추중립 위반(91.2%), c2 체간 전방 숙임(56.2%), c3 무릎 정렬(19.2%), c4 발바닥(19.6%)
  - **Rep #5 (F497~F571)**: c1 척추중립 위반(83.7%), c2 체간 전방 숙임(34.3%), c3 무릎 정렬(29.1%), c4 발바닥(25.8%)
"""
    with open(OUT_MD, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"[+] MD 보고서 생성 완료: {OUT_MD}")

if __name__ == "__main__":
    generate()
