import docx
import sys
try:
    doc = docx.Document('d:\\자주프 문서\\7주차\\바벨_데드리프트_MediaPipe_vs_RTMP_성능비교보고서.docx')
    with open('d:\\report_content.txt', 'w', encoding='utf-8') as f:
        for p in doc.paragraphs:
            f.write(p.text + '\n')
        f.write('\nTABLES:\n')
        for t in doc.tables:
            for row in t.rows:
                f.write('\t'.join([c.text for c in row.cells]) + '\n')
            f.write('\n')
except Exception as e:
    with open('d:\\report_content.txt', 'w', encoding='utf-8') as f:
        f.write(str(e))
