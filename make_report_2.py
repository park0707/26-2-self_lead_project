import os
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()
title = doc.add_heading('vidio-health-trainer Barbell Deadlift Final Model Performance', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Project: vidio-health-trainer')
doc.add_heading('1. Overview', level=1)
doc.add_paragraph('Performance comparison of Barbell Deadlift MediaPipe vs RTMPose final model.')

doc.add_heading('2. Performance Table', level=1)
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Metric'
hdr_cells[1].text = 'MediaPipe Final Model'
hdr_cells[2].text = 'RTMPose Final Model'
hdr_cells[3].text = 'Delta'

records = [
    ('원천 데이터 3-State F1', '71.88%', '73.98%', '+2.11%p'),
    ('보류율 (Rejection Rate)', '11.7%', '11.9%', '+0.2%p'),
    ('Coverage', '88.31%', '88.05%', '-0.26%p'),
    ('3-State Precision', '67.65%', '72.80%', '+5.15%p'),
    ('3-State Recall', '76.67%', '75.21%', '-1.46%p'),
    ('3-State Accuracy', '78.82%', '81.12%', '+2.30%p'),
    ('5-Defect EMR', '24.03%', '40.00%', '+15.97%p'),
]

for item in records:
    row_cells = table.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]
    row_cells[3].text = item[3]

doc.add_heading('3. Conclusion', level=1)
doc.add_paragraph('RTMPose final model outperformed MediaPipe. Recommended for production.')

doc.save(r'd:\7주차\보고서\바벨_데드리프트_성능_비교_보고서.docx')
print('DONE')
