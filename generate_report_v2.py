import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_background(cell, fill, color=None, val=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    tcPr.append(tcBorders)

    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

target_dir = r'D:\7주차\보고서'
os.makedirs(target_dir, exist_ok=True)
doc_path = os.path.join(target_dir, '바벨_스쿼트_성능_비교_보고서.docx')

doc = Document()

title = doc.add_heading('바벨 스쿼트 모델 성능 비교 보고서', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

proj_para = doc.add_paragraph('프로젝트명: vidio-health-trainer', style='Intense Quote')
proj_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('1. 개요', level=1)
doc.add_paragraph(
    '본 보고서는 vidio-health-trainer 프로젝트의 일환으로, 바벨 스쿼트 동작 분석을 위한 '
    'MediaPipe 최종 모델과 RTMPose 최종 최적화 모델의 세부 오류 4종(c1~c4)에 대한 성능을 비교하기 위해 작성되었습니다. '
    '평가는 원천 데이터(AI-Hub 바벨 스쿼트 전수 2,605건, 64-Frame Cubic Spline 보간 데이터)를 기반으로 진행되었으며, '
    '3-State F1 Score와 실제 서비스 환경에서의 안정성을 가늠하는 보류율(Hold Rate)을 중점적으로 비교하였습니다.'
)

doc.add_heading('2. 세부 오류 4종 성능 비교표', level=1)
table = doc.add_table(rows=5, cols=5)
table.style = 'Table Grid'
table.autofit = False

widths = (Inches(1.2), Inches(1.5), Inches(1.2), Inches(1.2), Inches(1.2))
for row in table.rows:
    for idx, width in enumerate(widths):
        row.cells[idx].width = width

headers = ['세부 오류 (c1~c4)', '원천 데이터', 'MediaPipe 3-State F1', 'RTMPose 3-State F1 (최종)', 'RTMPose 보류율']
for idx, text in enumerate(headers):
    cell = table.cell(0, idx)
    cell.text = text
    set_cell_background(cell, 'D9D9D9')
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.runs[0]
    run.font.bold = True

data = [
    ['c1. 척추선 중립 위반', 'AI-Hub 바벨스쿼트 (64-Frame)', '82.81%', '81.33%', '17.27%'],
    ['c2. 고개 정면 및 시선', 'AI-Hub 바벨스쿼트 (64-Frame)', '80.92%', '80.00%', '30.44%'],
    ['c3. 발-무릎 방향 일치', 'AI-Hub 바벨스쿼트 (64-Frame)', '87.46%', '94.82%', '4.11%'],
    ['c4. 발바닥 지면 고정', 'AI-Hub 바벨스쿼트 (64-Frame)', '82.80%', '88.12%', '11.98%']
]

for i, row_data in enumerate(data):
    for j, text in enumerate(row_data):
        cell = table.cell(i+1, j)
        cell.text = text
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('3. 분석 및 결론', level=1)
p = doc.add_paragraph()
p.add_run('1. 3-State F1 Score 관점:\n').bold = True
p.add_run('c3(발-무릎 방향 일치)와 c4(발바닥 지면 고정) 등 하체와 관련된 세부 오류에서 RTMPose 최종 모델이 MediaPipe 대비 각각 '
          '+7.36%p, +5.32%p 가량 향상된 우수한 정확도를 보였습니다. 척추선과 고개 등 상체 오류(c1, c2)는 큰 차이 없이 안정적인 80%대 성능을 유지했습니다.\n\n')

p.add_run('2. 보류율(Hold Rate) 및 3D 관절 정보 활용성:\n').bold = True
p.add_run('RTMPose 모델은 자연적 모호 구간에 대해 보류율을 산출(c2의 경우 최대 30.44%)하여 판단을 유보하고 안정성을 확보할 수 있도록 설계되었습니다. '
          '이를 통해 4-Defect EMR(완전 일치율)이 43.61%로 MediaPipe 대비 개선되었고, Clean F1-Score(정상 구간) 또한 50.19%로 대폭 상승하였습니다.\n\n')

p.add_run('3. 결론:\n').bold = True
p.add_run('RTMPose 최종 최적화 모델은 3D 정보를 적극 활용하여 하체 관련 판단 오류를 줄였으며, 보류율 통제를 통해 전체 모호 구간에서의 '
          '안정성을 크게 향상시켰습니다. 따라서 실서비스 환경에서는 RTMPose 모델을 주력으로 채택하는 것이 적합합니다.')

doc.save(doc_path)
