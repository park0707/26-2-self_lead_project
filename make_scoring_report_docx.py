# -*- coding: utf-8 -*-
"""
vidio-health-trainer
8주차_일반화_모델예측_채점표_및_다중결함_검증체계_구축보고서 생성 스크립트 (.docx 및 .md)
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_OUT = r"D:\자주프 문서\8주차\8주차_일반화_모델예측_채점표_및_다중결함_검증체계_구축보고서.docx"
MD_OUT = r"D:\자주프 문서\8주차\md\8주차_일반화_모델예측_채점표_및_다중결함_검증체계_구축보고서.md"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_table_header(row, col_widths=None):
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=130, bottom=130, left=130, right=130)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Malgun Gothic"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if col_widths and idx < len(col_widths):
            cell.width = Inches(col_widths[idx])

def style_table_rows(table, col_widths=None, align_cols=None):
    for r_idx, row in enumerate(table.rows[1:]):
        bg = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            for p in cell.paragraphs:
                if align_cols and c_idx < len(align_cols):
                    p.alignment = align_cols[c_idx]
                for r in p.runs:
                    r.font.name = "Malgun Gothic"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])

def build_docx():
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    os.makedirs(os.path.dirname(MD_OUT), exist_ok=True)
    
    doc = Document()
    
    # 여백
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # 문서 제목
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("vidio-health-trainer 일반화 데이터셋 웹 모델 예측 결과\n및 사용자 직접 채점표 검증체계 구축 보고서")
    trun.font.name = "Malgun Gothic"
    trun.font.size = Pt(18)
    trun.font.bold = True
    trun.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    # 메타데이터 표
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("문서 번호", "8주차 공식 기술 보고서 (2026-10-16)"),
        ("프로젝트 명", "vidio-health-trainer"),
        ("대상 데이터셋", "D:\\data\\일반화 (3대 운동 5대 시점 72개 정밀 분할 세트 클립)"),
        ("산출물 경로", "D:\\data\\일반화\\ (통합 및 종목별 채점표 CSV 4종, 가이드 2종)")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        c0 = meta_table.rows[r_idx].cells[0]
        c1 = meta_table.rows[r_idx].cells[1]
        c0.text = k
        c1.text = v
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = "Malgun Gothic"
        c0.paragraphs[0].runs[0].font.size = Pt(9)
        c1.paragraphs[0].runs[0].font.name = "Malgun Gothic"
        c1.paragraphs[0].runs[0].font.size = Pt(9)
        c0.width = Inches(1.5)
        c1.width = Inches(5.0)

    doc.add_paragraph() # Spacing

    # Section 1
    h1 = doc.add_paragraph()
    hrun1 = h1.add_run("1. 구축 배경 및 사용자 직접 채점 체계의 필요성")
    hrun1.font.name = "Malgun Gothic"
    hrun1.font.size = Pt(13)
    hrun1.font.bold = True
    hrun1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p1 = doc.add_paragraph()
    p1.add_run(
        "본 프로젝트의 8주차 2단계 웹 서빙 추론 파이프라인(Stage2InferenceEngine)을 실제 사용자가 직접 촬영한 72개 일반화 세트 클립에 적용한 결과, "
        "단일 타깃 결함에 대한 모델의 재현율(Recall)은 최대 100%(스쿼트 c4 100%, 스티프 데드 sd1/sd2 100%, 스쿼트 c2 80%)에 달하는 극대화된 감지 성능을 확인하였습니다.\n\n"
        "그러나 실제 인체의 생체역학적(Biomechanical) 키네마틱스 연동 구조상, 사용자가 특정 1개의 결함(예: 시선 이탈)을 의도하여 시연하더라도 "
        "경추 굴곡에 따른 흉추 후만(등 굽힘) 및 바벨 궤적 이탈 등의 연쇄 결함(Co-occurrence)이 불가피하게 동반 발생합니다. "
        "따라서 사용자가 사전에 정의한 '단일 의도 라벨'만을 정답으로 간주하는 단순 이진 채점 방식은 모델이 실제 발생한 동반 결함을 정확히 잡아냈음에도 "
        "불구하고 오탐(False Positive)으로 과소평가되는 한계를 드러냈습니다.\n\n"
        "이에 따라 사용자가 72개 전체 세트 영상을 직접 시청하면서 모델의 예측 확률 및 경고/결함 판정이 실제 영상 동작과 얼마나 일치하는지를 정밀하게 평가하고, "
        "진성 다중 라벨 Ground Truth를 구축할 수 있도록 전용 CSV 채점표 및 원클릭 비디오 재생 인터페이스를 전면 구축하였습니다."
    )
    for r in p1.runs:
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9.5)

    # Section 2
    h2 = doc.add_paragraph()
    hrun2 = h2.add_run("2. 배포 산출물 및 디렉터리 구성 체계")
    hrun2.font.name = "Malgun Gothic"
    hrun2.font.size = Pt(13)
    hrun2.font.bold = True
    hrun2.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    t2 = doc.add_table(rows=5, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2_headers = ["파일명", "대상 클립 수", "주요 구성 컬럼", "파일 용도"]
    for i, title_text in enumerate(t2_headers):
        t2.rows[0].cells[i].text = title_text
    style_table_header(t2.rows[0], [2.2, 1.0, 2.0, 1.8])

    t2_data = [
        ("일반화_모델예측_통합_채점표.csv", "72개 클립 (전수)", "영상재생링크, 의도라벨, 종합판정, 감지결함요약, 세부확률, 사용자판정(O/△/X), 메모", "3대 운동 전 시점 종합 감사 마스터 채점표"),
        ("일반화_채점표_바벨데드.csv", "28개 클립", "d1~d5 개별 확률/판정, [사용자확인] d1~d5 개별 O/X 체크, 종합타당성, 피드백", "바벨 데드리프트 세부 결함 정밀 검증표"),
        ("일반화_채점표_스쿼트.csv", "24개 클립", "c1~c4 개별 확률/판정, [사용자확인] c1~c4 개별 O/X 체크, 종합타당성, 피드백", "바벨 스쿼트 세부 결함 정밀 검증표"),
        ("일반화_채점표_스티프데드.csv", "20개 클립", "sd1~sd3 개별 확률/판정, [사용자확인] sd1~sd3 개별 O/X 체크, 종합타당성, 피드백", "스티프 레그 데드리프트 세부 결함 정밀 검증표")
    ]
    for r_idx, row_data in enumerate(t2_data):
        for c_idx, val in enumerate(row_data):
            t2.rows[r_idx+1].cells[c_idx].text = val
    style_table_rows(t2, [2.2, 1.0, 2.0, 1.8], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    p2 = doc.add_paragraph()
    p2.add_run(
        "\n* 이원화 격리 및 접근 편의성: 모든 채점표는 D:\\data\\일반화\\ 루트 디렉터리와 실제 영상이 저장된 D:\\data\\일반화\\분할_결과\\ 내부 양쪽에 동일하게 배포되어 "
        "사용자가 어떤 폴더를 탐색기로 열더라도 즉시 채점표를 열람할 수 있도록 하였습니다.\n"
        "* 한글 인코딩 무결성: Windows 환경의 Excel에서 UTF-8 파일 열람 시 발생하는 한글 깨짐 현상을 원천 방지하기 위해 모든 CSV 파일에 UTF-8 with BOM(utf-8-sig) 규격을 적용하였습니다."
    )
    for r in p2.runs:
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9.5)

    # Section 3
    h3 = doc.add_paragraph()
    hrun3 = h3.add_run("3. 원클릭 영상 재생 및 사용자 채점 인터페이스 설계")
    hrun3.font.name = "Malgun Gothic"
    hrun3.font.size = Pt(13)
    hrun3.font.bold = True
    hrun3.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p3 = doc.add_paragraph()
    p3.add_run(
        "1) 엑셀 내장 원클릭 비디오 재생(Hyperlink Formula) 구현:\n"
        "   - 수작업으로 72개 비디오 파일을 탐색기에서 일일이 찾아 여는 불편을 해소하기 위해, CSV 내에 =HYPERLINK(\"절대경로\", \"영상 재생 ▶\") 수식을 전면 탑재하였습니다.\n"
        "   - 사용자는 Excel에서 해당 셀을 클릭하기만 하면 윈도우 기본 미디어 플레이어(팟플레이어, 미디어 플레이어 등)로 해당 세트 클립(3~5회 반복)이 즉각 팝업되어 0.1초 만에 영상 확인이 가능합니다.\n\n"
        "2) 다중 라벨(Multi-label) Ground Truth 체크 규격:\n"
        "   - 통합 채점표 채점 기호:\n"
        "     * O (완전 타당): 모델의 결함 감지 및 상태 판정이 실제 동작과 100% 일치함.\n"
        "     * △ (부분 타당): 의도치 않았으나 실제로 동반된 연쇄 결함을 모델이 정확히 지적함 (복합 결함 확인).\n"
        "     * X (오탐): 영상에서는 해당 결함이 전혀 없으나 모델이 결함으로 오진함.\n"
        "     * FN (미탐): 영상에서 명백한 결함이 육안으로 관찰되나 모델이 감지하지 못함.\n"
        "   - 종목별 세부 채점표:\n"
        "     * 각 결함 부위별로 [사용자확인] 실제발생(O/X) 컬럼을 제공하여, 사용자가 각 부위별 발생 여부를 직관적으로 O/X 체크할 수 있도록 표준화하였습니다."
    )
    for r in p3.runs:
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9.5)

    # Section 4
    h4 = doc.add_paragraph()
    hrun4 = h4.add_run("4. 결함 코드 및 공식 5대 시점 표준 매핑 명세")
    hrun4.font.name = "Malgun Gothic"
    hrun4.font.size = Pt(13)
    hrun4.font.bold = True
    hrun4.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    t4 = doc.add_table(rows=13, cols=4)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    t4_headers = ["운동 종목", "결함 코드", "사용자 라벨", "공식 결함 명칭 및 세부 설명"]
    for i, title_text in enumerate(t4_headers):
        t4.rows[0].cells[i].text = title_text
    style_table_header(t4.rows[0], [1.5, 1.0, 1.2, 3.3])

    t4_data = [
        ("바벨 스쿼트", "c1", "결함 1", "척추의 중립 위반 (요추 후만/과신전, 벗윙크)"),
        ("바벨 스쿼트", "c2", "결함 2", "고개 정면 및 시선 이탈 (과도한 숙임 또는 과신전)"),
        ("바벨 스쿼트", "c3", "결함 3", "무릎-발 방향 일치 위반 (외반/내반 비정렬)"),
        ("바벨 스쿼트", "c4", "결함 4", "발바닥 지면 고정 위반 (뒤꿈치/앞꿈치 들림)"),
        ("바벨 데드리프트", "d1", "결함 1", "척추의 중립 위반 (요추/흉추 후만 라운드백)"),
        ("바벨 데드리프트", "d2", "결함 2", "무릎-발목 정렬 모임/외반 (Knee Medial Collapse)"),
        ("바벨 데드리프트", "d5", "결함 3", "바벨 궤적과 몸 밀착 위반 (정강이/허벅지 이탈)"),
        ("바벨 데드리프트", "d4", "결함 4", "무릎과 골반 동시 신전 위반 (힙 슈트업, 선신전)"),
        ("바벨 데드리프트", "d3", "결함 5", "발바닥 지면 고정 위반 (중심 이동 및 들림)"),
        ("스티프 레그 데드", "sd3", "결함 1", "바벨 궤적과 몸 밀착 위반 (수직 바벨 경로 이탈)"),
        ("스티프 레그 데드", "sd1", "결함 2", "무릎 구부린 채 고정 위반 (무릎 관절 가변 굴곡)"),
        ("스티프 레그 데드", "sd2", "결함 3", "척추의 중립 위반 (요추 안정성 붕괴)")
    ]
    for r_idx, row_data in enumerate(t4_data):
        for c_idx, val in enumerate(row_data):
            t4.rows[r_idx+1].cells[c_idx].text = val
    style_table_rows(t4, [1.5, 1.0, 1.2, 3.3], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    # Section 5
    h5 = doc.add_paragraph()
    hrun5 = h5.add_run("5. 향후 일정 및 기대 효과")
    hrun5.font.name = "Malgun Gothic"
    hrun5.font.size = Pt(13)
    hrun5.font.bold = True
    hrun5.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p5 = doc.add_paragraph()
    p5.add_run(
        "1) 사용자 직접 채점 데이터 피드백 수집:\n"
        "   - 사용자가 작성 완료한 CSV 채점표를 취합하여 72개 세트 클립 전체에 대한 최종 Multi-label True Matrix를 확정합니다.\n"
        "2) 시점별/부위별 임계치(Threshold) 및 가중치 사후 교정(Post-Calibration):\n"
        "   - 사용자의 실제 발생 판정(O) 대비 모델이 주의(State 2)로 주저했거나 오탐(State 1)을 낸 부위에 대해, 시점별(시점 A~E) 최적 결정 임계치를 미세 재조정합니다.\n"
        "3) 웹 서비스 자세 피드백 3단계 엔진 완성도 제고:\n"
        "   - 실제 사용자 환경에서도 고개 숙임 시 척추 굽힘 경고를 함께 제시하는 '생체역학적 원인 분석 피드백'의 정당성을 실증 확보함으로써, "
        "vidio-health-trainer 웹 서비스의 현장 적응성과 신뢰성을 극대화합니다."
    )
    for r in p5.runs:
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(9.5)

    doc.save(DOCX_OUT)
    print(f"Generated Word Report: {DOCX_OUT}")

    # Build Markdown equivalent
    md_content = f"""# vidio-health-trainer 일반화 데이터셋 웹 모델 예측 결과 및 사용자 직접 채점표 검증체계 구축 보고서

- **문서 번호**: 8주차 공식 기술 보고서 (2026-10-16)
- **프로젝트 명**: `vidio-health-trainer`
- **저장 위치**: `{DOCX_OUT}`

---

## 1. 구축 배경 및 사용자 직접 채점 체계의 필요성

본 프로젝트의 8주차 2단계 웹 서빙 추론 파이프라인(`Stage2InferenceEngine`)을 실제 사용자가 직접 촬영한 72개 일반화 세트 클립에 적용한 결과, 단일 타깃 결함에 대한 모델의 재현율(Recall)은 최대 100%(스쿼트 c4 100%, 스티프 데드 sd1/sd2 100%, 스쿼트 c2 80%)에 달하는 극대화된 감지 성능을 확인하였습니다.

그러나 실제 인체의 생체역학적(Biomechanical) 키네마틱스 연동 구조상, 사용자가 특정 1개의 결함을 의도하여 시연하더라도 경추 굴곡에 따른 흉추 후만 및 바벨 궤적 이탈 등의 연쇄 결함(Co-occurrence)이 불가피하게 동반 발생합니다. 이에 따라 단순 의도 라벨 중심의 정답 비교를 넘어, 사용자가 72개 세트 영상을 직접 시청하며 모델 예측의 타당성을 평가하고 진성 다중 라벨 Ground Truth를 구축할 수 있도록 채점 인터페이스를 전면 구축하였습니다.

---

## 2. 배포 산출물 및 디렉터리 구성 체계

모든 채점표는 `D:\\data\\일반화\\` 및 `D:\\data\\일반화\\분할_결과\\` 두 곳에 모두 배치되어 있습니다. (Excel 한글 깨짐 방지 `UTF-8 BOM` 적용)

| 파일명 | 대상 클립 수 | 주요 구성 컬럼 | 파일 용도 |
| :--- | :---: | :--- | :--- |
| **`일반화_모델예측_통합_채점표.csv`** | 72개 (전수) | 영상재생링크, 의도라벨, 종합판정, 감지결함요약, 세부확률, 사용자판정(O/△/X), 메모 | 3대 운동 전 시점 종합 감사 마스터 채점표 |
| **`일반화_채점표_바벨데드.csv`** | 28개 클립 | d1~d5 개별 확률/판정, [사용자확인] d1~d5 개별 O/X 체크, 종합타당성, 피드백 | 바벨 데드리프트 세부 결함 정밀 검증표 |
| **`일반화_채점표_스쿼트.csv`** | 24개 클립 | c1~c4 개별 확률/판정, [사용자확인] c1~c4 개별 O/X 체크, 종합타당성, 피드백 | 바벨 스쿼트 세부 결함 정밀 검증표 |
| **`일반화_채점표_스티프데드.csv`** | 20개 클립 | sd1~sd3 개별 확률/판정, [사용자확인] sd1~sd3 개별 O/X 체크, 종합타당성, 피드백 | 스티프 레그 데드리프트 세부 결함 정밀 검증표 |

---

## 3. 원클릭 영상 재생 및 사용자 채점 인터페이스 설계

1. **원클릭 비디오 재생(Excel Hyperlink) 탑재**:
   - `=HYPERLINK("절대경로", "영상 재생 ▶")` 수식이 포함되어 있어, Excel에서 셀을 클릭하기만 하면 해당 세트 영상(3~5회)이 즉시 재생됩니다.
2. **다중 라벨 평가 기준**:
   - `O` (완전 타당): 모델의 결함 감지 및 판정이 실제 영상 동작과 일치.
   - `△` (부분 타당): 의도 외에 실제로 동반된 연쇄 결함을 모델이 정탐함.
   - `X` (오탐): 정상 동작임에도 모델이 결함으로 잘못 판정함.
   - `FN` (미탐): 실제 결함이 발생했으나 모델이 감지하지 못함.
3. **종목별 개별 O/X 체크란**:
   - 각 결함 부위별로 `[사용자확인] 실제발생(O/X)` 컬럼을 제공하여 개별 부위별 다중 정답을 직접 마킹할 수 있습니다.

---

## 4. 결함 코드 및 공식 5대 시점 매핑 명세

| 운동 종목 | 결함 코드 | 사용자 라벨 | 공식 결함 명칭 및 세부 설명 |
| :--- | :---: | :---: | :--- |
| **바벨 스쿼트** | `c1` | 결함 1 | 척추의 중립 위반 (Spine Neutrality) |
| **바벨 스쿼트** | `c2` | 결함 2 | 고개 정면 및 시선 이탈 (Head & Gaze Orientation) |
| **바벨 스쿼트** | `c3` | 결함 3 | 무릎-발 방향 일치 위반 (Knee-Toe Alignment) |
| **바벨 스쿼트** | `c4` | 결함 4 | 발바닥 지면 고정 위반 (Foot Grounding) |
| **바벨 데드리프트** | `d1` | 결함 1 | 척추의 중립 위반 (Spine Neutrality) |
| **바벨 데드리프트** | `d2` | 결함 2 | 무릎-발목 정렬 모임/외반 (Knee Medial Collapse) |
| **바벨 데드리프트** | `d5` | 결함 3 | 바벨 궤적과 몸 밀착 위반 (Barbell Path Proximity) |
| **바벨 데드리프트** | `d4` | 결함 4 | 무릎과 골반 동시 신전 위반 (Knee-Hip Desynchronization) |
| **바벨 데드리프트** | `d3` | 결함 5 | 발바닥 지면 고정 위반 (Foot Ground Fixation) |
| **스티프 레그 데드** | `sd3` | 결함 1 | 바벨 궤적과 몸 밀착 위반 (Barbell Path Proximity) |
| **스티프 레그 데드** | `sd1` | 결함 2 | 무릎 구부린 채 고정 위반 (Knee Fixation Violation) |
| **스티프 레그 데드** | `sd2` | 결함 3 | 척추의 중립 위반 (Spine Neutrality) |

---

## 5. 향후 계획

1. 사용자의 직접 채점 완료 후, 다중 라벨 Ground Truth 매트릭스 확정
2. 시점별 및 결함별 최적 결정 임계치(Threshold) 미세 사후 교정(Post-Calibration)
3. 3단계 피드백 안내 생성기와의 엔드투엔드 연계 고도화
"""
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Generated Markdown Report: {MD_OUT}")

if __name__ == "__main__":
    build_docx()
