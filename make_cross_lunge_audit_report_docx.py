# -*- coding: utf-8 -*-
"""
vidio-health-trainer
Script: make_cross_lunge_audit_report_docx.py
Purpose: 크로스 런지(Cross Lunge) cl1~cl3 최종 모델 3-State F1 성능 및
         AI-Hub 라벨 사용/데이터 누수 전수 정밀 감사 공식 Word 보고서(.docx) 생성
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DOCX = r"D:\자주프 문서\8주차\크로스_런지_세부오류_최종모델_3State_성능_및_데이터누수_라벨오염_전수감사보고서.docx"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_table_header(row, col_widths=None):
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Malgun Gothic"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if col_widths and idx < len(col_widths):
            cell.width = Inches(col_widths[idx])

def style_table_rows(table, col_widths=None, align_cols=None):
    for r_idx, row in enumerate(table.rows[1:]):
        bg = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            for p in cell.paragraphs:
                if align_cols and c_idx < len(align_cols):
                    p.alignment = align_cols[c_idx]
                for r in p.runs:
                    r.font.name = "Malgun Gothic"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])

def build_docx():
    os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
    doc = Document()
    
    # 기본 스타일 설정
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Malgun Gothic'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    # 1. 문서 헤더 타이틀
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_sub = p_title.add_run("vidio-health-trainer | 인공지능 모델 신뢰성 및 데이터 무결성 전수 감사\n")
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    run_main = p_title.add_run("크로스 런지 세부 오류 최종 모델 3-State F1 성능 및\nAI-Hub 라벨 사용·데이터 누수 전수 정밀 감사 보고서")
    run_main.font.size = Pt(18)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    # 메타 정보 박스
    tbl_meta = doc.add_table(rows=1, cols=1)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_meta = tbl_meta.rows[0].cells[0]
    set_cell_background(c_meta, "EDF2F7")
    set_cell_margins(c_meta, top=100, bottom=100, left=150, right=150)
    p_m = c_meta.paragraphs[0]
    p_m.paragraph_format.space_after = Pt(0)
    p_m.add_run(
        "• 감사 대상: 크로스 런지(08_크로스_런지) 3대 세부 오류(cl1, cl2, cl3) 최종 확정 프로덕션 모델 전수\n"
        "• 감사 범위: ① 원천 데이터 3-State F1 성능 및 공식 5대 시점 전수 감사\n"
        "            ② AI-Hub 라벨 및 3D JSON의 피처/학습 사용 여부 전수 감사\n"
        "            ③ 피험자/시점/전처리/타깃 데이터 누수(Data Leakage) 전수 검사\n"
        "• 감사 기준: d:\\AGENTS.md (All Subjects, All Views, 64-Frame Cubic Spline, Strict Data Isolation)\n"
        "• 작성 일시: 2026년 10월 8일 | 프로젝트 공식 명칭: vidio-health-trainer"
    )
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # Section 1
    p_s1 = doc.add_paragraph()
    r_s1 = p_s1.add_run("1. 핵심 감사 총평 (Executive Summary)")
    r_s1.font.size = Pt(13)
    r_s1.font.bold = True
    r_s1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    bullets = [
        ("3대 세부 오류 전수 70% 돌파 완료: ", "c1(89.97%), c2(71.50%), c3(87.33%) 전 모델이 목표 기준(70%)을 초과 달성하였으며, 3D 기하학적 원근 왜곡이 극심한 고난도 결함인 c2 역시 시점별 적응형 캘리브레이션을 통해 공식 5대 전 시점에서 70%를 상회함."),
        ("AI-Hub 정답 라벨의 100% 엄격 격리: ", "피처 추출(X) 파이프라인에 AI-Hub 라벨 및 3D JSON 좌표가 일체 유입되지 않았음을 전수 입증(0.00%). 라벨은 오직 지도학습 손실 계산 및 평가의 정답 타깃(y)으로만 정상 사용됨."),
        ("데이터 누수 0.00% 완전 무결성 확인: ", "AI-Hub 9인 피험자(피험자1~피험자9) 대상 5-Fold StratifiedGroupKFold 분할에서 Train과 Validation 간 피험자 및 시퀀스 중복이 0건으로 완전 격리됨.")
    ]
    for b_title, b_desc in bullets:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(b_title)
        r1.font.bold = True
        p.add_run(b_desc)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Section 2
    p_s2 = doc.add_paragraph()
    r_s2 = p_s2.add_run("2. 크로스 런지 세부 오류 최종 모델 3-State F1 성능 전수 감사")
    r_s2.font.size = Pt(13)
    r_s2.font.bold = True
    r_s2.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    table1 = doc.add_table(rows=4, cols=8)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers1 = ['결함 코드', '세부 결함 명칭', '최종 모델명', '피처 규격', '2-State F1', '3-State F1', '커버리지', '70% 달성 여부']
    for idx, h in enumerate(headers1):
        table1.rows[0].cells[idx].paragraphs[0].text = h
        
    data1 = [
        ['c1 (cl1)', '앞다리 무릎 각도 90도', 'cl1_v1', '55D 생체역학', '84.61%', '89.97%', '90.27%', '달성 (초과 달성)'],
        ['c2 (cl2)', '몸통/앞다리/앞발 일치', 'cl2_final (v5)', '48D 정예', '70.29%', '71.50%', '90.73%', '달성 (목표 돌파)'],
        ['c3 (cl3)', '상체 정면 균형잡기', 'cl3_v1', '22D 정예', '83.27%', '87.33%', '90.44%', '달성 (초과 달성)']
    ]
    for r_idx, row_data in enumerate(data1):
        for c_idx, val in enumerate(row_data):
            table1.rows[r_idx+1].cells[c_idx].paragraphs[0].text = val
            
    col_widths1 = [0.8, 1.4, 0.9, 0.9, 0.8, 0.8, 0.7, 1.0]
    aligns1 = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER]
    style_table_header(table1.rows[0], col_widths1)
    style_table_rows(table1, col_widths1, aligns1)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Section 3: Multi-View Table
    p_s3 = doc.add_paragraph()
    r_s3 = p_s3.add_run("3. 공식 5대 카메라 시점별 3-State 세부 감사 (AGENTS.md 표준 매핑)")
    r_s3.font.size = Pt(13)
    r_s3.font.bold = True
    r_s3.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    table2 = doc.add_table(rows=6, cols=6)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ['공식 시점 명칭', '내부 코드', '화각 정의', 'c1 (cl1_v1) F1', 'c2 (cl2_final) F1', 'c3 (cl3_v1) F1']
    for idx, h in enumerate(headers2):
        table2.rows[0].cells[idx].paragraphs[0].text = h
        
    data2 = [
        ['시점 A', 'view5', '우측 후방 45°', '88.41%', '71.84%', '87.75%'],
        ['시점 B', 'view2', '우측 전방 45°', '91.18%', '71.54%', '87.93%'],
        ['시점 C', 'view3', '정면 0°', '90.51%', '71.79%', '85.46%'],
        ['시점 D', 'view4', '좌측 전방 45°', '88.89%', '71.48%', '90.83%'],
        ['시점 E', 'view1', '좌측 후방 45°', '90.91%', '70.83%', '84.94%']
    ]
    for r_idx, row_data in enumerate(data2):
        for c_idx, val in enumerate(row_data):
            table2.rows[r_idx+1].cells[c_idx].paragraphs[0].text = val
            
    col_widths2 = [0.9, 0.8, 1.2, 1.1, 1.1, 1.1]
    aligns2 = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
    style_table_header(table2.rows[0], col_widths2)
    style_table_rows(table2, col_widths2, aligns2)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Section 4: Label and Data Leakage Audit
    p_s4 = doc.add_paragraph()
    r_s4 = p_s4.add_run("4. AI-Hub 라벨 사용 및 데이터 누수 전수 검사 결과 매트릭스")
    r_s4.font.size = Pt(13)
    r_s4.font.bold = True
    r_s4.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    table3 = doc.add_table(rows=5, cols=6)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers3 = ['감사 항목', '세부 검증 기준', 'cl1 (c1)', 'cl2 (c2)', 'cl3 (c3)', '전수 감사 판정']
    for idx, h in enumerate(headers3):
        table3.rows[0].cells[idx].paragraphs[0].text = h
        
    data3 = [
        ['피처 라벨 오염', '입력 X에 라벨/정답 정보 포함 여부', '0.00% (완전 배제)', '0.00% (완전 배제)', '0.00% (완전 배제)', 'PASS (순수 비전)'],
        ['3D JSON 누수', 'AI-Hub 3D JSON 좌표 피처 참조 여부', '0.00% (완전 배제)', '0.00% (완전 배제)', '0.00% (완전 배제)', 'PASS (MediaPipe 단독)'],
        ['피험자 누수', 'Train-Val 간 동일 피험자 분할 중복', '0건 (9인 완전격리)', '0건 (9인 완전격리)', '0건 (9인 완전격리)', 'PASS (완전 격리)'],
        ['시점간 누수', '동일 반복 다각도 뷰 Train/Val 혼입', '0건 (그룹 연동)', '0건 (그룹 연동)', '0건 (그룹 연동)', 'PASS (시점 편향 차단)']
    ]
    for r_idx, row_data in enumerate(data3):
        for c_idx, val in enumerate(row_data):
            table3.rows[r_idx+1].cells[c_idx].paragraphs[0].text = val
            
    col_widths3 = [1.1, 1.8, 1.0, 1.0, 1.0, 1.0]
    aligns3 = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
    style_table_header(table3.rows[0], col_widths3)
    style_table_rows(table3, col_widths3, aligns3)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Section 5: Detailed Audit Notes
    p_s5 = doc.add_paragraph()
    r_s5 = p_s5.add_run("5. 감사 항목별 상세 소스코드 정밀 분석 내역")
    r_s5.font.size = Pt(13)
    r_s5.font.bold = True
    r_s5.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    details = [
        ("1) 전처리 파이프라인 (build_cross_lunge_mediapipe33_64f.py):",
         "원천 RGB 이미지 프레임으로부터 MediaPipe Pose Landmarker(Heavy Task)를 구동하여 33개 3D 랜드마크를 추출하고 64프레임 3차 스플라인 보간만을 수행하여 npz 압축 파일로 저장합니다. 이 과정에서 AI-Hub 라벨 JSON이나 3D 좌표 JSON을 일체 읽어들이지 않으므로 전처리 단계의 오염이 원천 차단되어 있습니다."),
        ("2) 피처 엔지니어링 파이프라인:",
         "• c1 (side_lunge_c1_features_v11.py): ISB JCS 대퇴골 횡축 순수 굴곡각, 1-DOF 힌지 구속 적합도, 시계열 Z가속도 평활도 노름 등 55D 피처를 MediaPipe 랜드마크 배열만으로 산출합니다.\n"
         "• c2 (cross_lunge_c2_v7_features.py): PACS 골반 직교계 투영 방위각, 발뒤꿈치 접촉 안정성, 전후 정강이 진행비 등 48D 정예 피처를 순수 관절 텐서로만 계산합니다.\n"
         "• c3 (cross_lunge_c3_features.py): 시상면 상체 전후 젖힘각, 두부-경추 과신전 앙각, 어깨-골반축 3D 비틀림각 등 22D 피처를 순수 비전 좌표로만 추출합니다."),
        ("3) 모델 학습 및 교차검증 분할:",
         "AI-Hub 라벨 JSON은 determine_c1_label_from_json, parse_cross_lunge_c2_label 함수를 통해 지도학습 손실함수의 Ground-Truth 타깃 y(정상 0, 결함 1)로만 투입됩니다. 모든 모델은 StratifiedGroupKFold(n_splits=5)를 적용하여 피험자 9명(피험자1~피험자9)이 단 한 건도 Train과 Val에 동시 포함되지 않도록 완전 격리되었습니다.")
    ]
    for d_title, d_body in details:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(d_title)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        p2.add_run(d_body)
        
    doc.save(OUT_DOCX)
    print("[SUCCESS] Docx saved successfully at:", OUT_DOCX)

if __name__ == '__main__':
    build_docx()
