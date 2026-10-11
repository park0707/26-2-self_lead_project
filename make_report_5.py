import os
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

report_dir = r"d:\자주프 문서\7주차\보고서"
os.makedirs(report_dir, exist_ok=True)

doc = docx.Document()
title = doc.add_heading('vidio-health-trainer 바벨 데드리프트 세부 오류 5종 모델 성능 비교 보고서', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('작성일: 2026년 10월 10일\n프로젝트: vidio-health-trainer')

doc.add_heading('1. 개요', level=1)
doc.add_paragraph(
    '본 보고서는 바벨 데드리프트의 세부 오류 모델 5가지에 대한 MediaPipe와 RTMPose의 성능을 비교합니다. '
    '평가는 원천 데이터를 기반으로 이루어졌으며, 3-State F1 지표 및 보류율(Rejection Rate)을 포함합니다.'
)

doc.add_heading('2. 세부 오류 5종 성능 비교 (원천 데이터 기준)', level=1)
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '세부 오류 항목'
hdr_cells[1].text = 'MediaPipe\n3-State F1'
hdr_cells[2].text = 'RTMPose\n3-State F1'
hdr_cells[3].text = 'F1 증감 (Δ)'
hdr_cells[4].text = '보류율\n(RTMPose 기준)'

records = [
    ('D1. 척추의 중립 위반', '62.02%', '62.76%', '+0.74%p', '13.8%'),
    ('D2. 무릎-발목 정렬 모임/외반', '57.03%', '84.56%', '+27.53%p', '10.5%'),
    ('D3. 발바닥 접지/엉덩이 선행 거상', '63.51%', '66.99%', '+3.48%p', '15.8%'),
    ('D4. 무릎-골반 동시 신전 불일치', '48.31%', '65.34%', '+17.03%p', '9.6%'),
    ('D5. 바벨 궤적 신체 밀착 실패', '80.81%', '99.87%', '+19.06%p', '0.5%')
]

for item in records:
    row_cells = table.add_row().cells
    for i in range(5):
        row_cells[i].text = item[i]

doc.add_paragraph('\n* 보류율(Rejection Rate)은 모델의 Coverage(판정률)를 기반으로 역산(100% - Coverage)된 수치입니다.')

doc.add_heading('3. 결론', level=1)
doc.add_paragraph(
    '바벨 데드리프트의 5가지 세부 오류 모델 모두에서 RTMPose 백본이 MediaPipe 대비 우수한 3-State F1 성능을 기록했습니다. '
    '특히 D2(무릎-발목 정렬)와 D5(바벨 밀착 실패)에서 매우 큰 폭의 성능 향상이 확인되었으며, 전반적인 보류율 또한 실서비스 적용 가능 범위 내에서 통제되고 있습니다.'
)

output_path = os.path.join(report_dir, '바벨_데드리프트_세부오류5종_성능_비교_보고서.docx')
doc.save(output_path)
print("SUCCESS")
